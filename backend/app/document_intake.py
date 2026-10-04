from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

from PIL import Image
from pypdf import PdfReader
import fitz
import pytesseract
from docx import Document as DocxDocument
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Document, Facility, FacilityDetail
from .settings import settings

MAX_OCR_PAGES = 20
MAX_TEXT_CHARS = 250_000
MAX_SHEET_CELLS = 12_000
MAX_DOCUMENT_BYTES = 100 * 1024 * 1024


def _storage_path(doc: Document) -> Path:
    root = Path(settings.storage_root).resolve()
    p = Path(doc.storage_path)
    if not p.is_absolute():
        p = root / p
    p = p.resolve()
    try:
        p.relative_to(root)
    except ValueError as exc:
        raise ValueError("document path is outside managed storage") from exc
    return p


def _ocr_lang() -> str:
    try:
        langs = set(pytesseract.get_languages(config=""))
    except Exception:
        return "eng"
    for candidate in ("Japanese", "jpn", "eng"):
        if candidate in langs:
            return candidate
    return next(iter(langs), "eng")


def _ocr_image(image: Image.Image) -> str:
    return pytesseract.image_to_string(image, lang=_ocr_lang())


def _extract_pdf(path: Path, force_ocr: bool) -> tuple[str, str, int, dict]:
    reader = PdfReader(str(path))
    page_count = len(reader.pages)
    parts: list[str] = []
    if not force_ocr:
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                parts.append("")
    text = "\n".join(parts).strip()
    if text and len(re.sub(r"\s+", "", text)) >= 20 and not force_ocr:
        return text[:MAX_TEXT_CHARS], "pdf_text", page_count, {"ocr_used": False}

    doc = fitz.open(str(path))
    ocr_parts: list[str] = []
    limit = min(len(doc), MAX_OCR_PAGES)
    for i in range(limit):
        pix = doc[i].get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        ocr_parts.append(_ocr_image(img))
    return "\n".join(ocr_parts)[:MAX_TEXT_CHARS], "pdf_ocr", page_count, {"ocr_used": True, "ocr_pages": limit, "ocr_page_limit": MAX_OCR_PAGES}


def _extract_image(path: Path) -> tuple[str, str, int, dict]:
    with Image.open(path) as img:
        text = _ocr_image(img.convert("RGB"))
    return text[:MAX_TEXT_CHARS], "image_ocr", 1, {"ocr_used": True, "ocr_pages": 1}


def _extract_docx(path: Path) -> tuple[str, str, int | None, dict]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)[:MAX_TEXT_CHARS], "docx_text", None, {"ocr_used": False}


def _extract_excel(path: Path) -> tuple[str, str, int | None, dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    cells = 0
    for ws in wb.worksheets:
        parts.append(f"[SHEET] {ws.title}")
        for row in ws.iter_rows():
            values = []
            for cell in row:
                if cell.value not in (None, ""):
                    values.append(str(cell.value))
                cells += 1
                if cells >= MAX_SHEET_CELLS:
                    break
            if values:
                parts.append("\t".join(values))
            if cells >= MAX_SHEET_CELLS:
                break
        if cells >= MAX_SHEET_CELLS:
            break
    return "\n".join(parts)[:MAX_TEXT_CHARS], "excel_cells", None, {"ocr_used": False, "cell_limit": MAX_SHEET_CELLS, "cells_scanned": cells}


def _extract_text_file(path: Path) -> tuple[str, str, int | None, dict]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            return raw.decode(enc)[:MAX_TEXT_CHARS], f"text_{enc}", None, {"ocr_used": False}
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")[:MAX_TEXT_CHARS], "text_fallback", None, {"ocr_used": False}


def extract_document(doc: Document, force_ocr: bool = False) -> tuple[str, str, int | None, dict]:
    path = _storage_path(doc)
    if not path.exists():
        raise FileNotFoundError(path)
    if path.stat().st_size > MAX_DOCUMENT_BYTES:
        raise ValueError("document is too large for synchronous analysis")
    ext = Path(doc.original_filename or path.name).suffix.lower()
    mime = (doc.mime_type or "").lower()
    if ext == ".pdf" or mime == "application/pdf":
        return _extract_pdf(path, force_ocr)
    if ext in {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff"} or mime.startswith("image/"):
        return _extract_image(path)
    if ext == ".docx" or "wordprocessingml" in mime:
        return _extract_docx(path)
    if ext in {".xlsx", ".xlsm"} or "spreadsheetml" in mime:
        return _extract_excel(path)
    if ext in {".txt", ".csv", ".tsv"} or mime.startswith("text/"):
        return _extract_text_file(path)
    raise ValueError(f"unsupported document type: {ext or mime or 'unknown'}")


CLASSIFIERS = {
    "equipment_inspection_report": ["消防用設備等点検結果報告", "点検結果報告書", "消防用設備等点検"],
    "fire_manager_appointment": ["防火管理者選任", "防火管理者解任", "防火管理者"],
    "fire_plan": ["消防計画作成", "消防計画変更", "消防計画"],
}


def classify_submission(text: str) -> tuple[str | None, float, dict]:
    compact = re.sub(r"[\s　]+", "", text)
    best_code = None
    best_hits: list[str] = []
    for code, keywords in CLASSIFIERS.items():
        hits = [k for k in keywords if re.sub(r"[\s　]+", "", k) in compact]
        if len(hits) > len(best_hits):
            best_code, best_hits = code, hits
    if not best_hits:
        return None, 0.0, {"matched_keywords": []}
    confidence = min(0.98, 0.72 + 0.1 * (len(best_hits) - 1))
    return best_code, confidence, {"matched_keywords": best_hits, "derivation_method": "deterministic_keyword_rule"}


def _reiwa_to_iso(match: re.Match[str]) -> str:
    year = 2018 + int(match.group(1))
    month = int(match.group(2))
    day = int(match.group(3))
    return date(year, month, day).isoformat()


def _first_group(text: str, patterns: list[str]) -> tuple[str | None, str | None]:
    for pattern in patterns:
        m = re.search(pattern, text, re.MULTILINE)
        if m:
            return m.group(1).strip(" \t　:："), pattern
    return None, None


def detect_fields(text: str, type_code: str | None) -> tuple[dict, dict]:
    fields: dict = {}
    evidence: dict = {}
    facility_name, pat = _first_group(text, [r"(?:防火対象物|建物|事業所)(?:の)?名称[\s:：]*([^\n\r]{2,100})", r"名称[\s:：]+([^\n\r]{2,100})"])
    if facility_name:
        fields["facility_name"] = facility_name
        evidence["facility_name"] = {"pattern": pat}
    address, pat = _first_group(text, [r"所在地[\s:：]*([^\n\r]{5,160})"])
    if address:
        fields["address"] = address
        evidence["address"] = {"pattern": pat}
    phone, pat = _first_group(text, [r"(?:電話番号|電話)[\s:：]*([0-9０-９\-ー()（） ]{8,30})"])
    if phone:
        fields["phone"] = phone
        evidence["phone"] = {"pattern": pat}
    if type_code == "fire_manager_appointment":
        manager, pat = _first_group(text, [r"防火管理者(?:氏名)?[\s:：]*([^\n\r]{2,50})", r"選任者(?:氏名)?[\s:：]*([^\n\r]{2,50})"])
        if manager:
            fields["manager_name"] = manager
            evidence["manager_name"] = {"pattern": pat}
    date_hits: list[str] = []
    for m in re.finditer(r"令和\s*(\d{1,2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", text):
        try:
            date_hits.append(_reiwa_to_iso(m))
        except ValueError:
            pass
    for m in re.finditer(r"(20\d{2})[./\-年](\d{1,2})[./\-月](\d{1,2})日?", text):
        try:
            date_hits.append(date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat())
        except ValueError:
            pass
    if date_hits:
        fields["date_candidates"] = list(dict.fromkeys(date_hits))[:10]
        evidence["date_candidates"] = {"derivation_method": "date_pattern"}
    return fields, evidence


def _norm(value: str | None) -> str:
    return re.sub(r"[\s　\-ー]", "", value or "").lower()


def find_facility_candidates(db: Session, text: str, fields: dict) -> list[dict]:
    compact = _norm(text)
    rows = db.scalars(select(Facility).where(Facility.status == "active")).all()
    scored: list[dict] = []
    for row in rows:
        score = 0.0
        reasons: list[str] = []
        n = _norm(row.name)
        a = _norm(row.address)
        if n and n in compact:
            score += 0.72
            reasons.append("name_in_document")
        if a and len(a) >= 6 and a in compact:
            score += 0.22
            reasons.append("address_in_document")
        fn = _norm(fields.get("facility_name"))
        if fn and n and fn == n:
            score += 0.15
            reasons.append("extracted_name_exact")
        fa = _norm(fields.get("address"))
        if fa and a and (fa == a or fa in a or a in fa):
            score += 0.12
            reasons.append("extracted_address_match")
        if score >= 0.35:
            scored.append({"building_id": row.building_id, "name": row.name, "address": row.address, "score": round(min(score, 0.99), 3), "reasons": reasons})
    return sorted(scored, key=lambda x: (-x["score"], x["name"]))[:5]


def build_difference_candidates(db: Session, building_id: str, fields: dict) -> dict:
    facility = db.get(Facility, building_id)
    if not facility:
        return {}
    diffs: dict = {}
    mapping = {
        "facility_name": ("facility.name", facility.name),
        "address": ("facility.address", facility.address),
        "phone": ("facility.phone", facility.phone),
    }
    for field, (path, current) in mapping.items():
        proposed = fields.get(field)
        if proposed and _norm(str(proposed)) != _norm(str(current or "")):
            diffs[path] = {"current": current, "proposed": proposed, "confidence": 0.65, "source_field": field}
    return diffs
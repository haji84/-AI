"""Atomic stock ledger and explicit Human service decisions; AI is never required."""
import csv,json,os
from zoneinfo import ZoneInfo
from datetime import date,datetime,timedelta,timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO,StringIO
from uuid import UUID
from fastapi import HTTPException
from pydantic import ValidationError
from sqlalchemy import select,update,func,or_
from sqlalchemy.exc import IntegrityError
from .audit import write_audit
from .authz import permission_codes
from .models import Employee,Document,Facility,now_utc
from .operations_models import Vehicle,Incident
from .operations_service import source_need
from .assets_models import OperationalAsset,AssetLocation,AssetLot,AssetBalance,AssetLoan,AssetMovement,AssetService,AssetImportPreview
from .assets_schemas import AssetInput,AssetPatch,LocationInput,LotInput,MovementInput,ServiceInput

SCHEMA_VERSION='assets-v1'
# Server-owned business timezone; no request can alter stock expiry enforcement.
BUSINESS_TIMEZONE=ZoneInfo(os.environ.get('FIRE_AI_ASSET_BUSINESS_TIMEZONE','Asia/Tokyo'))
def business_today():return now_utc().astimezone(BUSINESS_TIMEZONE).date()
DATE_FIELDS=['next_pressure_test_on','next_use_on','next_calibration_on','next_service_on']
MODELS={'registry':OperationalAsset,'locations':AssetLocation,'lots':AssetLot,'balances':AssetBalance,'movements':AssetMovement,'loans':AssetLoan,'services':AssetService}
INPUTS={'registry':AssetInput,'locations':LocationInput,'lots':LotInput,'movements':MovementInput,'services':ServiceInput}

def scalar(value):
    if isinstance(value,(datetime,date)):return value.isoformat()
    if isinstance(value,Decimal):return str(value)
    return value

def row_dict(row):return {c.name:scalar(getattr(row,c.name)) for c in row.__table__.columns}

def need(db,user,*codes):
    if not set(codes).issubset(permission_codes(db,user.user_id)):raise HTTPException(403,'required asset/source permission missing')

def get_row(db,model,key,lock=False):
    try:UUID(key)
    except (ValueError,TypeError):raise HTTPException(404,'record not found')
    pk=list(model.__table__.primary_key.columns)[0]
    stmt=select(model).where(pk==key)
    if lock:stmt=stmt.with_for_update().execution_options(populate_existing=True)
    row=db.scalar(stmt)
    if row is None:raise HTTPException(404,'record not found')
    return row

def check_version(row,expected):
    if row.version!=expected:raise HTTPException(409,'record version conflict; reload latest record')

def audit(db,user,action,row,before=None):
    pk=list(row.__table__.primary_key.columns)[0].name
    write_audit(db,user_id=user.user_id,action=action,entity_type=row.__tablename__,entity_id=getattr(row,pk),before=before,after=row_dict(row))

def flush(db):
    try:db.flush()
    except IntegrityError as exc:raise HTTPException(409,'duplicate identity or invalid relationship') from exc

def save(db):
    try:db.commit()
    except IntegrityError as exc:
        db.rollback();raise HTTPException(409,'duplicate identity or invalid relationship') from exc

def change(db,user,row,expected,values,action):
    before=row_dict(row);pk=list(row.__table__.primary_key.columns)[0]
    result=db.execute(update(type(row)).where(pk==getattr(row,pk.name),type(row).version==expected).values(**values,version=expected+1,updated_at=now_utc()))
    if result.rowcount!=1:raise HTTPException(409,'record version conflict; reload latest record')
    db.refresh(row);audit(db,user,action,row,before);return row

def document(db,user,key):
    if key:need(db,user,'document.read');get_row(db,Document,key)

def employee(db,key):
    row=get_row(db,Employee,key)
    if not row.active:raise HTTPException(422,'employee inactive')
    return row

def active(row):
    if not row.active:raise HTTPException(409,'record retired/inactive')

def location_links(db,user,data):
    if data.get('building_id'):
        need(db,user,'facility.read');f=get_row(db,Facility,data['building_id'])
        if f.deleted_at is not None:raise HTTPException(422,'facility retired')
    if data.get('vehicle_id'):
        need(db,user,'fleet.read');active(get_row(db,Vehicle,data['vehicle_id']))

def redact_documents(value):
    if isinstance(value,dict):return {k:redact_documents(v) for k,v in value.items() if k not in {'document_id','document_sha256','lot_document_id','lot_document_sha256'}}
    if isinstance(value,list):return [redact_documents(v) for v in value]
    return value

def visible(db,user,row):
    data=row_dict(row);perms=permission_codes(db,user.user_id)
    if 'document.read' not in perms:
        data.pop('document_id',None)
        if 'review_snapshot' in data:
            data['review_snapshot']=redact_documents(data['review_snapshot'])
    if 'asset.borrower.read' not in perms:
        for key in ['borrower_employee_id','handler_employee_id','loan_id']:data.pop(key,None)
        # A loan reason can contain borrower details; restricted readers get no free text.
        if isinstance(row,AssetMovement) and (row.loan_id or row.handler_employee_id):data.pop('reason',None)
    if 'facility.read' not in perms:data.pop('building_id',None)
    if 'fleet.read' not in perms:data.pop('vehicle_id',None)
    if 'incident.read' not in perms:data.pop('incident_id',None)
    elif data.get('incident_id'):
        incident=get_row(db,Incident,data['incident_id'])
        if (incident.emergency_case_id and 'emergency.case.read' not in perms) or (incident.fire_investigation_case_id and 'fire_investigation.read' not in perms):data.pop('incident_id',None)
    return data

def create_asset(db,user,payload):
    document(db,user,payload.document_id)
    row=OperationalAsset(**payload.model_dump());db.add(row);flush(db);audit(db,user,'asset.registry.create',row);return row

def patch_asset(db,user,key,payload):
    row=get_row(db,OperationalAsset,key,True);active(row)
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if any(values.get(k) is None for k in ['name','unit','reorder_threshold'] if k in values):raise HTTPException(422,'required fields cannot be empty')
    if 'unit' in values and values['unit']!=row.unit and db.scalar(select(AssetMovement.movement_id).where(AssetMovement.asset_id==key).limit(1)):raise HTTPException(409,'unit is immutable after stock activity')
    document(db,user,values.get('document_id'))
    return change(db,user,row,payload.expected_version,values,'asset.registry.update')

def retire_asset(db,user,key,payload):
    row=get_row(db,OperationalAsset,key,True);active(row);check_version(row,payload.expected_version)
    if db.scalar(select(AssetBalance.balance_id).where(AssetBalance.asset_id==key,AssetBalance.quantity>0).limit(1)) or db.scalar(select(AssetLoan.loan_id).where(AssetLoan.asset_id==key,AssetLoan.outstanding_quantity>0).limit(1)):raise HTTPException(409,'stock/outstanding loans must be zero before retirement')
    return change(db,user,row,payload.expected_version,{'active':False,'retired_reason':payload.reason,'retired_by':user.user_id},'asset.registry.retire')

def create_location(db,user,payload):
    data=payload.model_dump();location_links(db,user,data)
    row=AssetLocation(**data);db.add(row);flush(db);audit(db,user,'asset.location.create',row);return row

def patch_location(db,user,key,payload):
    row=get_row(db,AssetLocation,key,True);values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if any(values[k] is None for k in ['name','active'] if k in values):raise HTTPException(422,'name/active required')
    location_links(db,user,values)
    if values.get('active') is False and (db.scalar(select(AssetBalance.balance_id).where(AssetBalance.location_id==key,AssetBalance.quantity>0).limit(1)) or db.scalar(select(AssetLoan.loan_id).where(AssetLoan.location_id==key,AssetLoan.outstanding_quantity>0).limit(1))):raise HTTPException(409,'location has stock/outstanding loans')
    return change(db,user,row,payload.expected_version,values,'asset.location.update')

def create_lot(db,user,asset_id,payload):
    asset=get_row(db,OperationalAsset,asset_id,True);active(asset);check_version(asset,payload.expected_version);document(db,user,payload.document_id)
    data=payload.model_dump(exclude={'expected_version'})
    if payload.expires_on:need(db,user,'asset.approve');data.update(expiry_approved_by=user.user_id,expiry_approved_at=now_utc())
    row=AssetLot(asset_id=asset_id,**data);db.add(row);flush(db)
    change(db,user,asset,payload.expected_version,{},'asset.lot.add');audit(db,user,'asset.lot.create',row);return row

def balance(db,asset_id,lot_id,location_id):
    row=db.scalar(select(AssetBalance).where(AssetBalance.lot_id==lot_id,AssetBalance.location_id==location_id).with_for_update().execution_options(populate_existing=True))
    if row is None:
        row=AssetBalance(asset_id=asset_id,lot_id=lot_id,location_id=location_id,quantity=Decimal('0'));db.add(row);flush(db)
    return row

def movement_links(db,user,data):
    document(db,user,data.get('document_id'))
    if data.get('handler_employee_id'):need(db,user,'asset.borrower.manage','asset.borrower.read');employee(db,data['handler_employee_id'])
    if data.get('borrower_employee_id'):need(db,user,'asset.borrower.manage','asset.borrower.read');employee(db,data['borrower_employee_id'])
    if data.get('incident_id'):
        need(db,user,'incident.read');i=get_row(db,Incident,data['incident_id'],True);source_need(db,user,i)
        if i.status!='active':raise HTTPException(422,'incident cancelled')

def move(db,user,payload,writeoff=False):
    data=payload.model_dump() if hasattr(payload,'model_dump') else dict(payload)
    asset=get_row(db,OperationalAsset,data['asset_id'],True)
    movement_links(db,user,data)
    fingerprint=sha256(json.dumps({k:scalar(v) for k,v in data.items()},sort_keys=True).encode()).hexdigest()
    existing=db.scalar(select(AssetMovement).where(AssetMovement.idempotency_key==data['idempotency_key']))
    if existing:
        if existing.request_sha256!=fingerprint:raise HTTPException(409,'idempotency key already used for another request')
        if existing.kind in ['loan','return']:need(db,user,'asset.borrower.manage','asset.borrower.read')
        return existing
    active(asset);check_version(asset,data['expected_version']);kind=data['kind'];quantity=Decimal(data['quantity'])
    data['occurred_on']=data.get('occurred_on') or business_today()
    if data['occurred_on']>business_today():raise HTTPException(422,'stock occurrence date cannot be in the future')
    if kind=='loan' and data.get('due_on') and data['due_on']<data['occurred_on']:raise HTTPException(422,'loan due date precedes loan occurrence')
    lot=get_row(db,AssetLot,data['lot_id'],True)
    if lot.asset_id!=asset.asset_id:raise HTTPException(422,'lot belongs to another asset')
    active(lot)
    locations={key:get_row(db,AssetLocation,key,True) for key in sorted({data['location_id'],data.get('to_location_id')} - {None})}
    for location in locations.values():active(location)
    loc=locations[data['location_id']]
    if kind in ['issue','loan'] and lot.expires_on and lot.expires_on<business_today():raise HTTPException(422,'expired batch cannot be ordinarily issued/loaned; use approved expiry write-off')
    if kind in ['disposal','expiry_writeoff'] and not writeoff:raise HTTPException(403,'Human-approved service write-off required')
    if kind=='expiry_writeoff' and (not lot.expires_on or lot.expires_on>=business_today()):raise HTTPException(422,'batch has not expired')
    if kind=='transfer':active(get_row(db,AssetLocation,data['to_location_id'],True))
    loan=None
    if kind=='return':
        need(db,user,'asset.borrower.manage','asset.borrower.read');loan=get_row(db,AssetLoan,data['loan_id'],True)
        if (loan.asset_id,loan.lot_id)!=(asset.asset_id,lot.lot_id):raise HTTPException(422,'return asset/batch differs from loan')
        if quantity>loan.outstanding_quantity:raise HTTPException(422,'return exceeds outstanding loan quantity')
        if data['occurred_on']<loan.loaned_on:raise HTTPException(422,'return occurrence precedes loan')
    origin=balance(db,asset.asset_id,lot.lot_id,loc.location_id)
    incoming=kind in ['receive','return'];next_quantity=origin.quantity+(quantity if incoming else -quantity)
    if next_quantity<0:raise HTTPException(422,'insufficient batch stock at location')
    if next_quantity>=Decimal('100000000000'):raise HTTPException(422,'stock quantity exceeds supported range')
    change(db,user,origin,origin.version,{'quantity':next_quantity},'asset.balance.move')
    if kind=='transfer':
        destination=balance(db,asset.asset_id,lot.lot_id,data['to_location_id'])
        total=destination.quantity+quantity
        if total>=Decimal('100000000000'):raise HTTPException(422,'destination stock exceeds supported range')
        change(db,user,destination,destination.version,{'quantity':total},'asset.balance.transfer')
    if kind=='loan':
        loan=AssetLoan(asset_id=asset.asset_id,lot_id=lot.lot_id,location_id=loc.location_id,borrower_employee_id=data['borrower_employee_id'],quantity=quantity,outstanding_quantity=quantity,loaned_on=data['occurred_on'],due_on=data.get('due_on'),reason=data['reason'],created_by=user.user_id);db.add(loan);flush(db);audit(db,user,'asset.loan.create',loan)
    if kind=='return':change(db,user,loan,loan.version,{'outstanding_quantity':loan.outstanding_quantity-quantity},'asset.loan.return')
    fields={k:v for k,v in data.items() if k in AssetMovement.__table__.columns and k not in ['asset_version']}
    fields['loan_id']=loan.loan_id if loan else None
    row=AssetMovement(**fields,unit=asset.unit,request_sha256=fingerprint,asset_version=asset.version+1,created_by=user.user_id)
    db.add(row);flush(db);change(db,user,asset,data['expected_version'],{},'asset.stock.move');audit(db,user,'asset.movement.'+kind,row);return row

def create_service(db,user,asset_id,payload):
    asset=get_row(db,OperationalAsset,asset_id,True);active(asset);check_version(asset,payload.expected_version);document(db,user,payload.document_id)
    if payload.performed_on>business_today():raise HTTPException(422,'completed service date cannot be in the future')
    if payload.lot_id:
        lot=get_row(db,AssetLot,payload.lot_id,True)
        if lot.asset_id!=asset_id:raise HTTPException(422,'service batch belongs to another asset')
    if payload.location_id:get_row(db,AssetLocation,payload.location_id,True)
    row=AssetService(asset_id=asset_id,created_by=user.user_id,**payload.model_dump(exclude={'expected_version'}));db.add(row);flush(db)
    change(db,user,asset,payload.expected_version,{},'asset.service.add');audit(db,user,'asset.service.create',row);return row

def service_parent(db,key):
    initial=get_row(db,AssetService,key);asset=get_row(db,OperationalAsset,initial.asset_id,True);row=get_row(db,AssetService,key,True)
    return row,asset

def service_snapshot(db,user,row,asset):
    lot=get_row(db,AssetLot,row.lot_id,True) if row.lot_id else None
    keys={row.document_id,lot.document_id if lot else None}-{None}
    docs={}
    for key in sorted(keys):
        need(db,user,'document.read');docs[key]=get_row(db,Document,key,True)
    return {'asset_id':asset.asset_id,'asset_version':asset.version,'lot_id':row.lot_id,'lot_version':lot.version if lot else None,
            'document_sha256':docs[row.document_id].sha256 if row.document_id else None,
            'lot_document_id':lot.document_id if lot else None,
            'lot_document_sha256':docs[lot.document_id].sha256 if lot and lot.document_id else None,
            'fields':{k:scalar(getattr(row,k)) for k in ServiceInput.model_fields if k!='expected_version'}}

def service_action(db,user,key,payload,action):
    row,asset=service_parent(db,key);check_version(row,payload.expected_version)
    if action=='cancel':
        document(db,user,row.document_id)
        if row.status=='approved' or row.status=='cancelled':raise HTTPException(409,'approved/cancelled service history is immutable')
        return change(db,user,row,payload.expected_version,{'status':'cancelled','reason':payload.reason},'asset.service.cancel')
    active(asset);snapshot=service_snapshot(db,user,row,asset)
    if action=='review':
        if row.status!='draft':raise HTTPException(409,'service is not draft')
        return change(db,user,row,payload.expected_version,{'status':'reviewed','review_snapshot':snapshot,'reviewed_by':user.user_id,'reviewed_at':now_utc(),'reason':payload.reason},'asset.service.review')
    if row.status!='reviewed':raise HTTPException(409,'Human review is required before approval')
    check_version(asset,payload.expected_asset_version)
    if row.review_snapshot!=snapshot:raise HTTPException(409,'reviewed source changed; cancel and create a fresh service record')
    movement_id=None
    if row.kind in ['disposal','expiry_writeoff']:
        data={'asset_id':asset.asset_id,'expected_version':asset.version,'kind':row.kind,'lot_id':row.lot_id,'location_id':row.location_id,'quantity':row.quantity,'occurred_on':row.performed_on,'cost':row.cost,'document_id':row.document_id,'idempotency_key':'service:'+row.service_id,'reason':payload.reason}
        movement_id=move(db,user,data,writeoff=True).movement_id
    else:change(db,user,asset,asset.version,{k:getattr(row,k) for k in DATE_FIELDS if getattr(row,k) is not None},'asset.service.dates.approve')
    return change(db,user,row,payload.expected_version,{'status':'approved','approved_by':user.user_id,'approved_at':now_utc(),'reason':payload.reason,'movement_id':movement_id},'asset.service.approve')

def query_rows(db,user,dataset,q='',asset_id=None,limit=200,offset=0):
    model=MODELS[dataset];stmt=select(model)
    if asset_id:
        get_row(db,OperationalAsset,asset_id)
        if hasattr(model,'asset_id'):stmt=stmt.where(model.asset_id==asset_id)
    if q:
        names={'registry':['code','name','notes'],'locations':['code','name','notes'],'lots':['batch_code','provenance'],'services':['description','kind']}.get(dataset,[])
        if names:stmt=stmt.where(or_(*(getattr(model,k).contains(q) for k in names)))
    pk=list(model.__table__.primary_key.columns)[0]
    return db.scalars(stmt.order_by(model.created_at.desc(),pk).offset(offset).limit(limit)).all()

def history(db,user,key):
    get_row(db,OperationalAsset,key)
    out={kind:[visible(db,user,r) for r in query_rows(db,user,kind,asset_id=key,limit=10000)] for kind in ['movements','services']}
    if 'asset.borrower.read' in permission_codes(db,user.user_id):out['loans']=[visible(db,user,r) for r in query_rows(db,user,'loans',asset_id=key,limit=10000)]
    return out

def alerts(db,user,as_of,days):
    through=as_of+timedelta(days=days);out=[]
    for asset in db.scalars(select(OperationalAsset).where(OperationalAsset.active.is_(True))):
        for key in DATE_FIELDS:
            due=getattr(asset,key)
            if due and due<=through:out.append({'asset_id':asset.asset_id,'asset_name':asset.name,'kind':key[5:-3],'due_on':due.isoformat(),'overdue':due<as_of})
    for lot,asset in db.execute(select(AssetLot,OperationalAsset).join(OperationalAsset).where(OperationalAsset.active.is_(True),AssetLot.active.is_(True),AssetLot.expires_on<=through)):
        qty=db.scalar(select(func.sum(AssetBalance.quantity)).where(AssetBalance.lot_id==lot.lot_id)) or Decimal('0')
        outstanding=db.scalar(select(func.sum(AssetLoan.outstanding_quantity)).where(AssetLoan.lot_id==lot.lot_id)) or Decimal('0')
        if qty+outstanding>0:out.append({'asset_id':asset.asset_id,'asset_name':asset.name,'kind':'expiry','lot_id':lot.lot_id,'batch_code':lot.batch_code,'due_on':lot.expires_on.isoformat(),'quantity':str(qty),'unit':asset.unit,'overdue':lot.expires_on<as_of})
    if 'asset.borrower.read' in permission_codes(db,user.user_id):
        for loan in db.scalars(select(AssetLoan).where(AssetLoan.outstanding_quantity>0,AssetLoan.due_on<=through)):
            out.append({'asset_id':loan.asset_id,'kind':'loan_return','loan_id':loan.loan_id,'due_on':loan.due_on.isoformat(),'quantity':str(loan.outstanding_quantity),'overdue':loan.due_on<as_of})
    return sorted(out,key=lambda r:(r['due_on'],r['asset_id'],r['kind']))

def reorder(db):
    out=[]
    for asset in db.scalars(select(OperationalAsset).where(OperationalAsset.active.is_(True))):
        qty=db.scalar(select(func.sum(AssetBalance.quantity)).join(AssetLot,AssetLot.lot_id==AssetBalance.lot_id).where(AssetBalance.asset_id==asset.asset_id,AssetLot.active.is_(True),or_(AssetLot.expires_on.is_(None),AssetLot.expires_on>=business_today()))) or Decimal('0')
        if qty<asset.reorder_threshold:out.append({'asset_id':asset.asset_id,'code':asset.code,'name':asset.name,'available_quantity':str(qty.quantize(Decimal('.001'))),'threshold':str(asset.reorder_threshold),'candidate_quantity':str((asset.reorder_threshold-qty).quantize(Decimal('.001'))),'unit':asset.unit,'status':'candidate'})
    return out

# Explicit tabular schema. Export-only lineage columns are accepted but never applied as official state.
LINEAGE={'asset_id','location_id','lot_id','movement_id','service_id','created_at','source_document_id','file_sha256','import_preview_id','text_encoding','unit'}
def headers(dataset):
    if dataset not in INPUTS:raise HTTPException(422,'unknown import dataset')
    return ['schema_version']+(['asset_id'] if dataset in ['lots','services'] else [])+list(INPUTS[dataset].model_fields)

def parse_rows(dataset,content):
    if not content or len(content)>2*1024*1024:raise HTTPException(422,'empty file or exceeds 2 MiB')
    if content.startswith(b'PK'):
        try:
            from openpyxl import load_workbook
            import zipfile
            with zipfile.ZipFile(BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist())>10*1024*1024:raise ValueError('expanded workbook exceeds limit')
            w=load_workbook(BytesIO(content),read_only=True,data_only=False)
            if len(w.worksheets)!=1:raise ValueError('one sheet required')
            sheet=w.active;items=[]
            for cells in sheet.iter_rows():
                if any(c.data_type=='f' for c in cells):raise ValueError('formulas forbidden in import')
                items.append([c.value for c in cells])
                if len(items)>1001:raise ValueError('at most 1000 rows')
            w.close()
        except Exception as exc:raise HTTPException(422,'invalid XLSX content: '+str(exc)) from exc
    else:
        try:items=list(csv.reader(StringIO(content.decode('utf-8-sig'))))
        except (UnicodeDecodeError,csv.Error) as exc:raise HTTPException(422,'invalid UTF-8 CSV content') from exc
    if len(items)<2 or len(items)>1001:raise HTTPException(422,'1 to 1000 data rows required')
    cols=[str(x).strip() if x is not None else '' for x in items[0]]
    allowed=set(headers(dataset))|LINEAGE|set(MODELS[dataset].__table__.columns.keys())
    if len(cols)!=len(set(cols)) or '' in cols or not set(cols)<=allowed or 'schema_version' not in cols:raise HTTPException(422,'invalid explicit header/schema')
    required={k for k,f in INPUTS[dataset].model_fields.items() if f.is_required()}
    if dataset in ['lots','services']:required.add('asset_id')
    if not required<=set(cols):raise HTTPException(422,'required schema columns missing: '+','.join(sorted(required-set(cols))))
    rows=[]
    for index,values in enumerate(items[1:],2):
        if not any(v is not None and str(v).strip() for v in values):continue
        if len(values)>len(cols) and any(v not in [None,''] for v in values[len(cols):]):raise HTTPException(422,f'row {index}: extra cells')
        record=dict(zip(cols,values))
        encoding=record.pop('text_encoding',None)
        if encoding not in [None,'','formula-prefix-v1']:raise HTTPException(422,f'row {index}: unsupported text_encoding')
        if record.get('status') not in [None,'','draft']:raise HTTPException(422,f'row {index}: official/reviewed states cannot be imported')
        for key in ['approved_by','approved_at','reviewed_by','reviewed_at','expiry_approved_by','expiry_approved_at']:
            if record.get(key) not in [None,'']:raise HTTPException(422,f'row {index}: approval fields cannot be imported')
        if record.get('source_document_id') and not record.get('document_id'):record['document_id']=record['source_document_id']
        if record.pop('schema_version',None)!=SCHEMA_VERSION:raise HTTPException(422,f'row {index}: assets-v1 schema_version required')
        data={k:(v.isoformat() if isinstance(v,(date,datetime)) else v) for k,v in record.items() if v is not None and str(v).strip()!='' and (k in INPUTS[dataset].model_fields or (k=='asset_id' and dataset in ['lots','services']))}
        # CSV formula protection prefixes on export are data escaping, undo only the sentinel.
        for k,v in list(data.items()):
            if encoding=='formula-prefix-v1' and isinstance(v,str) and v.startswith("'") and (v[1:].startswith("'") or v[1:].lstrip().startswith(('=','+','-','@','\t','\r'))):data[k]=v[1:]
        try:
            parent=data.pop('asset_id') if dataset in ['lots','services'] else None
            parsed=INPUTS[dataset].model_validate(data).model_dump(mode='json')
            if parent:parsed['asset_id']=parent
        except (ValidationError,KeyError) as exc:raise HTTPException(422,f'row {index}: invalid fields: '+str(exc)) from exc
        rows.append(parsed)
    if not rows:raise HTTPException(422,'no data rows')
    return rows

def import_permissions(db,user,dataset):
    need(db,user,'asset.read','asset.import','asset.update' if dataset=='movements' else 'asset.create')

def apply_rows(db,user,dataset,rows,document_id):
    import_permissions(db,user,dataset)
    out=[]
    for data in rows:
        data=dict(data)
        if dataset!='locations' and document_id and not data.get('document_id'):data['document_id']=document_id
        parent=data.pop('asset_id') if dataset in ['lots','services'] else None
        payload=INPUTS[dataset].model_validate(data)
        fn={'registry':create_asset,'locations':create_location,'lots':create_lot,'movements':move,'services':create_service}[dataset]
        row=fn(db,user,parent,payload) if parent else fn(db,user,payload)
        out.append(getattr(row,list(row.__table__.primary_key.columns)[0].name))
    return out

def import_preview(db,user,dataset,content,filename,document_id):
    import_permissions(db,user,dataset);document(db,user,document_id)
    rows=parse_rows(dataset,content)
    # Same service validations inside a rolled-back savepoint; no dry-run stock/history persists.
    with db.begin_nested() as nested:
        apply_rows(db,user,dataset,rows,document_id);flush(db);nested.rollback()
    preview=AssetImportPreview(dataset=dataset,schema_version=SCHEMA_VERSION,filename=filename,file_sha256=sha256(content).hexdigest(),document_id=document_id,row_data=rows,created_by=user.user_id,expires_at=now_utc()+timedelta(minutes=30));db.add(preview);flush(db);audit(db,user,'asset.import.preview',preview);return preview

def confirm_import(db,user,key,payload):
    row=get_row(db,AssetImportPreview,key,True);import_permissions(db,user,row.dataset);check_version(row,payload.expected_version)
    if row.created_by!=user.user_id and 'asset.admin' not in permission_codes(db,user.user_id):raise HTTPException(403,'preview belongs to another user')
    expiry=row.expires_at.replace(tzinfo=timezone.utc) if row.expires_at.tzinfo is None else row.expires_at
    if row.status!='preview' or row.file_sha256!=payload.file_sha256 or expiry<=now_utc():raise HTTPException(409,'preview expired, applied, or content hash differs')
    document(db,user,row.document_id);ids=apply_rows(db,user,row.dataset,row.row_data,row.document_id)
    change(db,user,row,payload.expected_version,{'status':'applied','applied_ids':ids},'asset.import.confirm');return {'inserted':len(ids),'ids':ids,'schema_version':SCHEMA_VERSION,'file_sha256':row.file_sha256,'document_id':row.document_id}

def safe_cell(value):
    value=scalar(value)
    if value is None:return ''
    if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False,sort_keys=True)
    if isinstance(value,str) and (value.startswith("'") or value.lstrip().startswith(('=','+','-','@','\t','\r'))):return "'"+value
    return value

def tabular(columns,rows,format):
    cells=[[safe_cell(r.get(c)) for c in columns] for r in rows]
    if format=='csv':
        out=StringIO();w=csv.writer(out);w.writerow(columns);w.writerows(cells);return out.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    from openpyxl import Workbook
    w=Workbook();w.active.title='assets-v1';w.active.append(columns)
    for row in cells:w.active.append(row)
    out=BytesIO();w.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

def export_data(db,user,dataset,q,asset_id,format):
    if dataset not in MODELS:raise HTTPException(422,'unknown export dataset')
    if dataset=='loans':need(db,user,'asset.borrower.read')
    rows=query_rows(db,user,dataset,q,asset_id,limit=10001)
    if len(rows)>10000:raise HTTPException(422,'export exceeds 10000 rows; narrow query')
    lineage={}
    wanted={getattr(row,list(row.__table__.primary_key.columns)[0].name) for row in rows}
    perms=permission_codes(db,user.user_id)
    for preview in db.scalars(select(AssetImportPreview).where(AssetImportPreview.dataset==dataset,AssetImportPreview.status=='applied').order_by(AssetImportPreview.created_at)):
        if preview.document_id and 'document.read' not in perms:continue
        for key in set(preview.applied_ids)&wanted:
            lineage.setdefault(key,{'file_sha256':preview.file_sha256,'import_preview_id':preview.preview_id})
    values=[]
    for row in rows:
        value={'schema_version':SCHEMA_VERSION,'text_encoding':'formula-prefix-v1',**visible(db,user,row),**lineage.get(getattr(row,list(row.__table__.primary_key.columns)[0].name),{})}
        if hasattr(row,'asset_id') and dataset!='registry':value['unit']=get_row(db,OperationalAsset,row.asset_id).unit
        if dataset in ['lots','services']:value['expected_version']=get_row(db,OperationalAsset,row.asset_id).version
        # On replay movements require the current asset version, but never replay automatically.
        if dataset=='movements':value['expected_version']=get_row(db,OperationalAsset,row.asset_id).version
        if 'document_id' in value:value['source_document_id']=value['document_id']
        values.append(value)
    columns=['schema_version']+[c.name for c in MODELS[dataset].__table__.columns if any(c.name in v for v in values) or not values]
    for key in ['unit','expected_version','source_document_id','file_sha256','import_preview_id','text_encoding']:
        if any(key in v for v in values) and key not in columns:columns.append(key)
    return tabular(columns,values,format)

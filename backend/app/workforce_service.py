"""Workforce rules, roster, leave and attendance services. AI is never required."""
import csv
from datetime import date,datetime,timedelta,timezone
from hashlib import sha256
from io import BytesIO,StringIO
from zoneinfo import ZoneInfo,ZoneInfoNotFoundError
from openpyxl import Workbook,load_workbook
from fastapi import HTTPException
from sqlalchemy import and_,func,or_,select
from sqlalchemy.exc import IntegrityError
from .audit import write_audit
from .authz import permission_codes
from .models import Document,Employee,now_utc
from .personnel import EmployeeAssignment,OrganizationUnit
from .workforce_models import (
    WorkforceShiftType,WorkforceEmployeeQualification,WorkforceStaffingRule,
    WorkforceRosterEntry,WorkforceLeaveEntry,WorkforceAttendance,WorkforceTimeEntry,
    WorkforceImportPreview,
)
from .workforce_schemas import (
    ShiftTypeCreate,ShiftTypePatch,QualificationCreate,StaffingRuleCreate,RosterCreate,
    LeaveCreate,AttendanceCreate,AttendancePatch,TimeEntryCreate,
)

SCHEMA_VERSION='workforce-v1'
DEFAULT_TIMEZONE='Asia/Tokyo'

def scalar(value):
    if isinstance(value,(date,datetime)):return value.isoformat()
    return value

def row_dict(row):
    return {c.name:scalar(getattr(row,c.name)) for c in row.__table__.columns}

def need(db,user,*codes):
    if not set(codes).issubset(permission_codes(db,user.user_id)):
        raise HTTPException(403,'required workforce/source permission missing')

def save(db):
    try:db.commit()
    except IntegrityError as exc:
        db.rollback();raise HTTPException(409,'workforce data conflict') from exc

def get_row(db,model,key,lock=False):
    pk=list(model.__table__.primary_key.columns)[0]
    stmt=select(model).where(pk==key)
    if lock:stmt=stmt.with_for_update().execution_options(populate_existing=True)
    row=db.scalar(stmt)
    if row is None:raise HTTPException(404,'workforce record not found')
    return row

def check_version(row,expected):
    if row.version!=expected:raise HTTPException(409,'record version conflict; reload latest record')

def audit(db,user,action,row,before=None,after=None):
    pk=list(row.__table__.primary_key.columns)[0].name
    write_audit(db,user_id=user.user_id,action=action,entity_type=row.__tablename__,
                entity_id=getattr(row,pk),before=before,after=after if after is not None else row_dict(row))

def visible(db,user,row):
    data=row_dict(row)
    if isinstance(row,WorkforceLeaveEntry) and 'personnel.read' not in permission_codes(db,user.user_id):
        data.pop('private_reason',None)
    return data

def active_employee(db,employee_id,lock=False):
    stmt=select(Employee).where(Employee.employee_id==employee_id)
    if lock:stmt=stmt.with_for_update().execution_options(populate_existing=True)
    employee=db.scalar(stmt)
    if not employee or not employee.active:raise HTTPException(422,'employee is missing or inactive')
    return employee

def active_org(db,organization_id):
    row=db.get(OrganizationUnit,organization_id)
    if not row or not row.active:raise HTTPException(422,'organization is missing or inactive')
    return row

def require_document(db,user,document_id):
    if not document_id:return None
    need(db,user,'document.read')
    row=db.get(Document,document_id)
    if not row:raise HTTPException(422,'source document not found')
    return row

def effective_assignment(db,employee_id,on_date):
    return db.scalar(
        select(EmployeeAssignment)
        .where(
            EmployeeAssignment.employee_id==employee_id,
            EmployeeAssignment.kind=='primary',
            EmployeeAssignment.valid_from<=on_date,
            or_(EmployeeAssignment.valid_to.is_(None),EmployeeAssignment.valid_to>=on_date),
        )
        .order_by(EmployeeAssignment.valid_from.desc())
        .limit(1)
    )

def shift_window(shift,work_date):
    try:tz=ZoneInfo(shift.timezone_name)
    except ZoneInfoNotFoundError:raise HTTPException(422,'unknown shift timezone')
    start=datetime.combine(work_date,shift.start_time,tzinfo=tz)
    end_date=work_date+timedelta(days=1) if shift.cross_midnight or shift.end_time<=shift.start_time else work_date
    end=datetime.combine(end_date,shift.end_time,tzinfo=tz)
    if end<=start:raise HTTPException(422,'shift end must follow shift start')
    return start,end

def create_shift(db,user,payload:ShiftTypeCreate):
    try:ZoneInfo(payload.timezone_name)
    except ZoneInfoNotFoundError:raise HTTPException(422,'unknown shift timezone')
    row=WorkforceShiftType(**payload.model_dump())
    db.add(row);db.flush();audit(db,user,'workforce.shift.create',row);return row

def patch_shift(db,user,key,payload:ShiftTypePatch):
    row=get_row(db,WorkforceShiftType,key,True);check_version(row,payload.expected_version);before=row_dict(row)
    for k,v in payload.model_dump(exclude={'expected_version'},exclude_none=True).items():setattr(row,k,v)
    try:ZoneInfo(row.timezone_name)
    except ZoneInfoNotFoundError:raise HTTPException(422,'unknown shift timezone')
    row.version+=1;row.updated_at=now_utc();db.flush();audit(db,user,'workforce.shift.update',row,before);return row

def create_qualification(db,user,payload:QualificationCreate):
    active_employee(db,payload.employee_id);require_document(db,user,payload.document_id)
    row=WorkforceEmployeeQualification(**payload.model_dump(),created_by=user.user_id)
    db.add(row);db.flush();audit(db,user,'workforce.qualification.create',row);return row

def create_staffing_rule(db,user,payload:StaffingRuleCreate):
    active_org(db,payload.organization_id)
    shift=get_row(db,WorkforceShiftType,payload.shift_type_id)
    if not shift.active:raise HTTPException(422,'shift type inactive')
    row=WorkforceStaffingRule(**payload.model_dump(),created_by=user.user_id)
    db.add(row);db.flush();audit(db,user,'workforce.staffing_rule.create',row);return row

def roster_overlap(db,employee_id,start,end,exclude_id=None):
    stmt=select(WorkforceRosterEntry).where(
        WorkforceRosterEntry.employee_id==employee_id,
        WorkforceRosterEntry.status!='cancelled',
        WorkforceRosterEntry.starts_at<end,
        WorkforceRosterEntry.ends_at>start,
    )
    if exclude_id:stmt=stmt.where(WorkforceRosterEntry.roster_entry_id!=exclude_id)
    return db.scalar(stmt.limit(1))

def create_roster(db,user,payload:RosterCreate):
    active_employee(db,payload.employee_id);active_org(db,payload.organization_id);require_document(db,user,payload.document_id)
    shift=get_row(db,WorkforceShiftType,payload.shift_type_id)
    if not shift.active:raise HTTPException(422,'shift type inactive')
    assignment=effective_assignment(db,payload.employee_id,payload.work_date)
    if not assignment:raise HTTPException(409,'effective primary assignment is unavailable for work date')
    if not payload.support_placement and assignment.organization_id!=payload.organization_id:
        raise HTTPException(409,'normal placement must match effective primary assignment')
    start,end=shift_window(shift,payload.work_date)
    if roster_overlap(db,payload.employee_id,start,end):raise HTTPException(409,'employee already has an overlapping roster placement')
    row=WorkforceRosterEntry(
        **payload.model_dump(),assignment_id=assignment.assignment_id,assignment_version=assignment.version,
        starts_at=start,ends_at=end,payable_minutes=shift.payable_minutes,created_by=user.user_id,
    )
    db.add(row);db.flush();audit(db,user,'workforce.roster.create',row);return row

def leave_effect(row):
    return row.quantity_minutes if row.kind in ('grant','adjustment_add') else -row.quantity_minutes

def leave_balance(db,employee_id,leave_type,as_of):
    active_employee(db,employee_id)
    rows=db.scalars(select(WorkforceLeaveEntry).where(
        WorkforceLeaveEntry.employee_id==employee_id,
        WorkforceLeaveEntry.leave_type==leave_type,
        WorkforceLeaveEntry.status=='approved',
        WorkforceLeaveEntry.effective_on<=as_of,
    )).all()
    return sum(leave_effect(row) for row in rows)

def create_leave(db,user,payload:LeaveCreate):
    active_employee(db,payload.employee_id)
    row=WorkforceLeaveEntry(**payload.model_dump(),created_by=user.user_id)
    db.add(row);db.flush();audit(db,user,'workforce.leave.create',row);return row

def _aware(value):
    if value is None:return None
    if value.tzinfo is None or value.utcoffset() is None:
        # SQLite can return timezone-naive values for DateTime(timezone=True).
        # Inputs are required to be timezone-aware by schema; persisted naive values
        # are interpreted as UTC only for deterministic dev/test round-trips.
        return value.replace(tzinfo=timezone.utc)
    return value

def attendance_calculation(db,employee_id,roster_entry_id,check_in,check_out):
    _aware(check_in)
    if check_out:_aware(check_out)
    elapsed=None
    check_in=_aware(check_in)
    check_out=_aware(check_out)
    if check_out:elapsed=max(0,int((check_out.astimezone(timezone.utc)-check_in.astimezone(timezone.utc)).total_seconds()//60))
    planned=None;roster_version=None
    if roster_entry_id:
        roster=get_row(db,WorkforceRosterEntry,roster_entry_id)
        if roster.employee_id!=employee_id:raise HTTPException(422,'roster entry belongs to another employee')
        planned=roster.payable_minutes;roster_version=roster.version
    return {
        'elapsed_minutes':elapsed,
        'planned_payable_minutes':planned,
        'overtime_candidate_minutes':max(0,(elapsed or 0)-(planned or (elapsed or 0))) if elapsed is not None else None,
        'roster_entry_id':roster_entry_id,
        'roster_version':roster_version,
        'calculation_version':'workforce-time-v1',
    }

def create_attendance(db,user,payload:AttendanceCreate):
    active_employee(db,payload.employee_id)
    calculation=attendance_calculation(db,payload.employee_id,payload.roster_entry_id,payload.check_in_at,payload.check_out_at)
    row=WorkforceAttendance(**payload.model_dump(),worked_minutes=calculation['elapsed_minutes'],calculation=calculation,created_by=user.user_id)
    db.add(row);db.flush();audit(db,user,'workforce.attendance.create',row);return row

def patch_attendance(db,user,key,payload:AttendancePatch):
    row=get_row(db,WorkforceAttendance,key,True);check_version(row,payload.expected_version)
    if row.status!='draft':raise HTTPException(409,'only draft attendance can be edited')
    before=row_dict(row)
    check_in=payload.check_in_at or row.check_in_at
    check_out=payload.check_out_at if 'check_out_at' in payload.model_fields_set else row.check_out_at
    calculation=attendance_calculation(db,row.employee_id,row.roster_entry_id,check_in,check_out)
    row.check_in_at=check_in;row.check_out_at=check_out;row.worked_minutes=calculation['elapsed_minutes'];row.calculation=calculation
    row.version+=1;row.updated_at=now_utc();db.flush();audit(db,user,'workforce.attendance.update',row,before);return row

def create_time_entry(db,user,payload:TimeEntryCreate):
    active_employee(db,payload.employee_id)
    if payload.attendance_id:
        attendance=get_row(db,WorkforceAttendance,payload.attendance_id)
        if attendance.employee_id!=payload.employee_id:raise HTTPException(422,'attendance belongs to another employee')
    row=WorkforceTimeEntry(**payload.model_dump(),created_by=user.user_id)
    db.add(row);db.flush();audit(db,user,'workforce.time.create',row);return row

def time_balance(db,employee_id,as_of):
    rows=db.scalars(select(WorkforceTimeEntry).where(
        WorkforceTimeEntry.employee_id==employee_id,
        WorkforceTimeEntry.status=='approved',
        WorkforceTimeEntry.occurred_on<=as_of,
        WorkforceTimeEntry.kind.in_(['comp_grant','comp_use']),
    )).all()
    return sum(row.minutes if row.kind=='comp_grant' else -row.minutes for row in rows)

def transition(db,user,row,expected,action,note):
    check_version(row,expected);before=row_dict(row)
    allowed={'review':('draft','reviewed'),'approve':('reviewed','approved'),'cancel':(('draft','reviewed'),'cancelled')}
    if action not in allowed:raise HTTPException(422,'unknown Human action')
    required,target=allowed[action]
    valid=row.status in required if isinstance(required,tuple) else row.status==required
    if not valid:raise HTTPException(409,f'cannot {action} from current state')
    now=now_utc()
    if action=='review':row.reviewed_by=user.user_id;row.reviewed_at=now
    if action=='approve':row.approved_by=user.user_id;row.approved_at=now
    row.status=target;row.version+=1;row.updated_at=now
    if hasattr(row,'rule_note') and note:row.rule_note=((row.rule_note+'\n') if row.rule_note else '')+note
    db.flush();audit(db,user,f'workforce.{row.__tablename__}.{action}',row,before);return row

def staffing_action(db,user,key,expected,action,note):
    row=get_row(db,WorkforceStaffingRule,key,True)
    return transition(db,user,row,expected,action,note)

def roster_action(db,user,key,expected,action,note):
    row=get_row(db,WorkforceRosterEntry,key,True)
    if action in ('review','approve'):
        employee=active_employee(db,row.employee_id,True)
        assignment=db.get(EmployeeAssignment,row.assignment_id) if row.assignment_id else None
        if not assignment or assignment.version!=row.assignment_version:
            raise HTTPException(409,'assignment evidence changed; recreate roster placement')
        if roster_overlap(db,row.employee_id,row.starts_at,row.ends_at,row.roster_entry_id):
            raise HTTPException(409,'employee has overlapping roster placement')
        if action=='approve' and not row.support_placement and assignment.organization_id!=row.organization_id:
            raise HTTPException(409,'assignment no longer supports placement')
    return transition(db,user,row,expected,action,note)

def leave_action(db,user,key,expected,action,note):
    row=get_row(db,WorkforceLeaveEntry,key,True)
    if action=='approve' and row.kind in ('use','adjustment_subtract','expire'):
        balance=leave_balance(db,row.employee_id,row.leave_type,row.effective_on)
        if balance<row.quantity_minutes:raise HTTPException(409,'leave balance would become negative')
    return transition(db,user,row,expected,action,note)

def attendance_action(db,user,key,expected,action,note):
    row=get_row(db,WorkforceAttendance,key,True)
    if action in ('review','approve'):
        active_employee(db,row.employee_id,True)
        if row.check_out_at is None:raise HTTPException(409,'attendance requires check-out before Human review')
        calc=attendance_calculation(db,row.employee_id,row.roster_entry_id,row.check_in_at,row.check_out_at)
        if row.calculation.get('roster_version')!=calc.get('roster_version'):
            raise HTTPException(409,'roster evidence changed; refresh attendance')
        row.calculation=calc;row.worked_minutes=calc['elapsed_minutes']
    return transition(db,user,row,expected,action,note)

def time_action(db,user,key,expected,action,note):
    row=get_row(db,WorkforceTimeEntry,key,True)
    if action=='approve' and row.kind=='comp_use':
        if time_balance(db,row.employee_id,row.occurred_on)<row.minutes:raise HTTPException(409,'compensatory balance would become negative')
    return transition(db,user,row,expected,action,note)

def approved_leave_on(db,employee_id,on_date):
    return bool(db.scalar(select(WorkforceLeaveEntry.leave_entry_id).where(
        WorkforceLeaveEntry.employee_id==employee_id,
        WorkforceLeaveEntry.status=='approved',
        WorkforceLeaveEntry.kind=='use',
        WorkforceLeaveEntry.effective_on==on_date,
    ).limit(1)))

def valid_qualification(db,employee_id,code,on_date):
    return bool(db.scalar(select(WorkforceEmployeeQualification.qualification_id).where(
        WorkforceEmployeeQualification.employee_id==employee_id,
        WorkforceEmployeeQualification.code==code,
        WorkforceEmployeeQualification.active.is_(True),
        WorkforceEmployeeQualification.valid_from<=on_date,
        or_(WorkforceEmployeeQualification.valid_to.is_(None),WorkforceEmployeeQualification.valid_to>=on_date),
    ).limit(1)))

def staffing_warnings(db,on_date,organization_id=None):
    rules_stmt=select(WorkforceStaffingRule).where(
        WorkforceStaffingRule.status=='approved',
        WorkforceStaffingRule.effective_from<=on_date,
        or_(WorkforceStaffingRule.effective_to.is_(None),WorkforceStaffingRule.effective_to>=on_date),
    )
    if organization_id:rules_stmt=rules_stmt.where(WorkforceStaffingRule.organization_id==organization_id)
    rules=db.scalars(rules_stmt.order_by(WorkforceStaffingRule.organization_id)).all()
    out=[]
    for rule in rules:
        rosters=db.scalars(select(WorkforceRosterEntry).where(
            WorkforceRosterEntry.organization_id==rule.organization_id,
            WorkforceRosterEntry.shift_type_id==rule.shift_type_id,
            WorkforceRosterEntry.work_date==on_date,
            WorkforceRosterEntry.status=='approved',
        )).all()
        eligible=[]
        for roster in rosters:
            employee=db.get(Employee,roster.employee_id)
            if not employee or not employee.active or approved_leave_on(db,roster.employee_id,on_date):continue
            if rule.qualification_code and not valid_qualification(db,roster.employee_id,rule.qualification_code,on_date):continue
            eligible.append(roster)
        out.append({
            'staffing_rule_id':rule.staffing_rule_id,'organization_id':rule.organization_id,
            'shift_type_id':rule.shift_type_id,'qualification_code':rule.qualification_code,
            'required':rule.min_staff,'available':len(eligible),'shortage':max(0,rule.min_staff-len(eligible)),
            'status':'shortage' if len(eligible)<rule.min_staff else 'ok',
            'roster_entry_ids':[r.roster_entry_id for r in eligible],
            'rule_version':rule.version,
        })
    return out

def available_crew(db,on_date,organization_id=None,shift_type_id=None):
    stmt=select(WorkforceRosterEntry,Employee).join(Employee,Employee.employee_id==WorkforceRosterEntry.employee_id).where(
        WorkforceRosterEntry.work_date==on_date,
        WorkforceRosterEntry.status=='approved',
        Employee.active.is_(True),
    )
    if organization_id:stmt=stmt.where(WorkforceRosterEntry.organization_id==organization_id)
    if shift_type_id:stmt=stmt.where(WorkforceRosterEntry.shift_type_id==shift_type_id)
    out=[]
    for roster,employee in db.execute(stmt).all():
        if approved_leave_on(db,employee.employee_id,on_date):continue
        quals=db.scalars(select(WorkforceEmployeeQualification).where(
            WorkforceEmployeeQualification.employee_id==employee.employee_id,
            WorkforceEmployeeQualification.active.is_(True),
            WorkforceEmployeeQualification.valid_from<=on_date,
            or_(WorkforceEmployeeQualification.valid_to.is_(None),WorkforceEmployeeQualification.valid_to>=on_date),
        )).all()
        out.append({
            'employee_id':employee.employee_id,'employee_code':employee.employee_code,'display_name':employee.display_name,
            'organization_id':roster.organization_id,'shift_type_id':roster.shift_type_id,
            'roster_entry_id':roster.roster_entry_id,'support_placement':roster.support_placement,
            'assignment_id':roster.assignment_id,'assignment_version':roster.assignment_version,
            'qualifications':[q.code for q in quals],
        })
    return out

def statistics(db,year,month=None):
    start=date(year,month or 1,1)
    if month:
        end=(date(year+1,1,1) if month==12 else date(year,month+1,1))
    else:end=date(year+1,1,1)
    rosters=db.scalars(select(WorkforceRosterEntry).where(WorkforceRosterEntry.status=='approved',WorkforceRosterEntry.work_date>=start,WorkforceRosterEntry.work_date<end)).all()
    attendance=db.scalars(select(WorkforceAttendance).where(WorkforceAttendance.status=='approved',WorkforceAttendance.work_date>=start,WorkforceAttendance.work_date<end)).all()
    times=db.scalars(select(WorkforceTimeEntry).where(WorkforceTimeEntry.status=='approved',WorkforceTimeEntry.occurred_on>=start,WorkforceTimeEntry.occurred_on<end)).all()
    leaves=db.scalars(select(WorkforceLeaveEntry).where(WorkforceLeaveEntry.status=='approved',WorkforceLeaveEntry.effective_on>=start,WorkforceLeaveEntry.effective_on<end)).all()
    return {
        'year':year,'month':month,
        'approved_roster_entries':len(rosters),
        'worked_minutes':sum(x.worked_minutes or 0 for x in attendance),
        'overtime_minutes':sum(x.minutes for x in times if x.kind=='overtime'),
        'comp_grant_minutes':sum(x.minutes for x in times if x.kind=='comp_grant'),
        'comp_use_minutes':sum(x.minutes for x in times if x.kind=='comp_use'),
        'annual_leave_use_minutes':sum(x.quantity_minutes for x in leaves if x.leave_type=='annual' and x.kind=='use'),
        'special_leave_use_minutes':sum(x.quantity_minutes for x in leaves if x.leave_type=='special' and x.kind=='use'),
        'employee_count':len({x.employee_id for x in rosters}),
    }

def formula_safe(value):
    text='' if value is None else str(value)
    return "'"+text if text[:1] in ('=','+','-','@') else text

def _parse_table(blob):
    if len(blob)>2*1024*1024:raise HTTPException(413,'workforce import file too large')
    if blob[:2]==b'PK':
        try:
            wb=load_workbook(BytesIO(blob),read_only=True,data_only=False);ws=wb.active
            values=list(ws.iter_rows(values_only=True))
            if not values:return [],[]
            headers=[str(x or '').strip() for x in values[0]]
            return headers,[{headers[i]:('' if v is None else str(v)) for i,v in enumerate(row)} for row in values[1:]]
        except Exception as exc:raise HTTPException(422,'invalid XLSX workbook') from exc
    try:text=blob.decode('utf-8-sig')
    except UnicodeDecodeError as exc:raise HTTPException(422,'CSV must be UTF-8') from exc
    reader=csv.DictReader(StringIO(text));return reader.fieldnames or [],list(reader)

IMPORT_HEADERS=['schema_version','employee_code','organization_code','shift_code','work_date','support_placement','note']

def import_preview(db,user,blob,filename,document_id=None):
    require_document(db,user,document_id)
    headers,rows=_parse_table(blob)
    if headers!=IMPORT_HEADERS:raise HTTPException(422,{'message':'unexpected workforce roster headers','expected':IMPORT_HEADERS,'actual':headers})
    normalized=[]
    for index,row in enumerate(rows,2):
        if row.get('schema_version')!=SCHEMA_VERSION:raise HTTPException(422,f'row {index}: schema_version must be {SCHEMA_VERSION}')
        try:work_date=date.fromisoformat(row.get('work_date',''))
        except ValueError:raise HTTPException(422,f'row {index}: invalid work_date')
        flag=row.get('support_placement','').strip().lower()
        if flag not in ('true','false','1','0','yes','no'):raise HTTPException(422,f'row {index}: invalid support_placement')
        for key in ('employee_code','organization_code','shift_code'):
            if not row.get(key,'').strip():raise HTTPException(422,f'row {index}: {key} required')
        normalized.append({
            'employee_code':row['employee_code'].strip(),'organization_code':row['organization_code'].strip(),
            'shift_code':row['shift_code'].strip(),'work_date':work_date.isoformat(),
            'support_placement':flag in ('true','1','yes'),'note':row.get('note','').strip() or None,
        })
    digest=sha256(blob).hexdigest()
    preview=WorkforceImportPreview(dataset='rosters',schema_version=SCHEMA_VERSION,filename=filename,file_sha256=digest,
        document_id=document_id,row_data=normalized,created_by=user.user_id,expires_at=now_utc()+timedelta(hours=1))
    db.add(preview);db.flush();audit(db,user,'workforce.import.preview',preview,after={'file_sha256':digest,'rows':len(normalized),'document_id':document_id});return preview

def confirm_import(db,user,key,expected_version,file_sha256):
    preview=get_row(db,WorkforceImportPreview,key,True);check_version(preview,expected_version)
    if preview.status!='preview' or preview.expires_at<now_utc():raise HTTPException(409,'import preview expired or already applied')
    if preview.file_sha256!=file_sha256:raise HTTPException(409,'source file changed')
    created=[]
    for item in preview.row_data:
        employee=db.scalar(select(Employee).where(Employee.employee_code==item['employee_code']))
        org=db.scalar(select(OrganizationUnit).where(OrganizationUnit.code==item['organization_code']))
        shift=db.scalar(select(WorkforceShiftType).where(WorkforceShiftType.code==item['shift_code']))
        if not employee or not org or not shift:raise HTTPException(409,'import reference no longer resolves')
        row=create_roster(db,user,RosterCreate(employee_id=employee.employee_id,organization_id=org.organization_id,shift_type_id=shift.shift_type_id,
            work_date=date.fromisoformat(item['work_date']),support_placement=item['support_placement'],note=item['note'],document_id=preview.document_id))
        created.append(row.roster_entry_id)
    preview.status='applied';preview.version+=1;preview.updated_at=now_utc()
    audit(db,user,'workforce.import.apply',preview,after={'file_sha256':preview.file_sha256,'created_ids':created})
    return {'inserted':len(created),'roster_entry_ids':created,'file_sha256':preview.file_sha256}

def export_rosters(db,format,from_date=None,to_date=None):
    stmt=select(WorkforceRosterEntry,Employee,OrganizationUnit,WorkforceShiftType).join(Employee,Employee.employee_id==WorkforceRosterEntry.employee_id).join(OrganizationUnit,OrganizationUnit.organization_id==WorkforceRosterEntry.organization_id).join(WorkforceShiftType,WorkforceShiftType.shift_type_id==WorkforceRosterEntry.shift_type_id)
    if from_date:stmt=stmt.where(WorkforceRosterEntry.work_date>=from_date)
    if to_date:stmt=stmt.where(WorkforceRosterEntry.work_date<=to_date)
    rows=[]
    for roster,employee,org,shift in db.execute(stmt.order_by(WorkforceRosterEntry.work_date,org.code,employee.display_name)).all():
        rows.append([
            employee.employee_code or '',employee.display_name,org.code,org.name,shift.code,shift.name,
            roster.work_date.isoformat(),roster.starts_at.isoformat(),roster.ends_at.isoformat(),roster.payable_minutes,
            roster.support_placement,roster.status,roster.assignment_id or '',roster.assignment_version or '',roster.note or '',
        ])
    headers=['employee_code','employee_name','organization_code','organization_name','shift_code','shift_name','work_date','starts_at','ends_at','payable_minutes','support_placement','status','assignment_id','assignment_version','note']
    safe=[[formula_safe(v) for v in row] for row in rows]
    if format=='csv':
        stream=StringIO();writer=csv.writer(stream);writer.writerow(headers);writer.writerows(safe)
        return ('\ufeff'+stream.getvalue()).encode('utf-8'),'text/csv'
    wb=Workbook();ws=wb.active;ws.title='rosters';ws.append(headers)
    for row in safe:ws.append(row)
    out=BytesIO();wb.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

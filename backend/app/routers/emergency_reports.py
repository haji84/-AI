from datetime import date
from enum import Enum
import csv
from io import BytesIO, StringIO
from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy.orm import Session
from ..audit import write_audit
from ..authz import require_permission
from ..db import get_db
from ..emergency_reports import summary
from ..models import User

router = APIRouter(prefix='/emergency/reports', tags=['emergency reports'])


class Group(str, Enum):
    hospital = 'hospital'
    severity = 'severity'
    region = 'region'


class Format(str, Enum):
    csv = 'csv'
    xlsx = 'xlsx'


def report(db, start_date, end_date, group_by):
    if start_date and end_date and start_date > end_date:
        raise HTTPException(status_code=422, detail='開始日は終了日以前にしてください')
    return summary(db, start_date=start_date, end_date=end_date, group_by=group_by.value)


def audit_report(db, user, data, action, export_format=None):
    # Evidence contains only filters and aggregate totals, never patient records.
    evidence = {'period': data['period'], 'group_by': data['group_by'], 'totals': data['totals']}
    if export_format:
        evidence['format'] = export_format.value
    write_audit(db, user_id=user.user_id, action=action, entity_type='emergency_report', after=evidence)
    db.commit()


@router.get('/summary')
def get_summary(response: Response, start_date: date | None = None, end_date: date | None = None,
                group_by: Group = Group.hospital, db: Session = Depends(get_db),
                user: User = Depends(require_permission('emergency.report.read'))):
    data = report(db, start_date, end_date, group_by)
    response.headers['Cache-Control'] = 'no-store'
    audit_report(db, user, data, 'emergency.report.read')
    return data


@router.get('/export')
def export_report(format: Format = Format.xlsx, start_date: date | None = None,
                  end_date: date | None = None, group_by: Group = Group.hospital,
                  db: Session = Depends(get_db), user: User = Depends(require_permission('emergency.report.export'))):
    data = report(db, start_date, end_date, group_by)
    headers = ['コード', '救護者人数', '区分内事案件数', '全救護者に対する割合']
    rows = [[x['code'] if x['code'] is not None else '未入力', x['patients'], x['cases'], x['patient_share']] for x in data['groups']]
    if format == Format.csv:
        stream = StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(['開始日', '終了日', '分類', '事案件数', '救護者人数', '搬送先コード入力人数', '日付不明の除外事案'])
        writer.writerow([data['period']['start_date'], data['period']['end_date'], group_by.value,
                         *data['totals'].values(), data['undated_cases_excluded']])
        writer.writerow([])
        writer.writerow(headers)
        for row in rows:
            writer.writerow(["'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v for v in row])
        payload = stream.getvalue().encode('utf-8-sig')
        media_type = 'text/csv; charset=utf-8'
    else:
        from openpyxl import Workbook
        wb = Workbook()
        totals = wb.active
        totals.title = '合計'
        totals.append(['事案件数', '救護者人数', '搬送先コード入力人数', '日付不明の除外事案', '開始日', '終了日', '分類'])
        totals.append([*data['totals'].values(), data['undated_cases_excluded'], data['period']['start_date'], data['period']['end_date'], group_by.value])
        sheet = wb.create_sheet('内訳')
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
            # Codes are text even when a source cell starts with '='.
            sheet.cell(sheet.max_row, 1).data_type = 's'
        definitions = wb.create_sheet('集計定義')
        for key, value in data['definitions'].items():
            definitions.append([key, value])
        for s in wb:
            s.freeze_panes = 'A2'
        stream = BytesIO()
        wb.save(stream)
        payload = stream.getvalue()
        media_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    audit_report(db, user, data, 'emergency.report.export', format)
    return Response(payload, media_type=media_type,
                    headers={'Content-Disposition': 'attachment; filename="emergency-summary.' + format.value + '"', 'Cache-Control': 'no-store'})

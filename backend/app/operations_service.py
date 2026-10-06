"""Operations invariants, concurrency, explicit Human transitions and tabular exchange."""
import csv,json
from datetime import date,datetime,time,timezone,timedelta
from decimal import Decimal,ROUND_HALF_UP,ROUND_HALF_EVEN,ROUND_DOWN
from hashlib import sha256
from io import BytesIO,StringIO
from uuid import UUID
from fastapi import HTTPException
from pydantic import ValidationError
from sqlalchemy import select,update
from sqlalchemy.exc import IntegrityError
from .audit import write_audit
from .authz import permission_codes
from .models import Employee,Document,EmergencyCase,FireInvestigationCase,now_utc
from .operations_models import Incident,Vehicle,Dispatch,DispatchCrew,AllowanceRate,VehicleTrip,FuelEntry,VehicleService,OperationsImportPreview
from .operations_schemas import IncidentInput,DispatchFields,VehicleInput,DispatchInput,TripInput,FuelInput,ServiceInput,CrewInput


def scalar(value):
    if isinstance(value,(datetime,date,time)):return value.isoformat()
    if isinstance(value,Decimal):return str(value)
    return value

def row_dict(row):return {c.name:scalar(getattr(row,c.name)) for c in row.__table__.columns}

def need(db,user,*codes):
    if not set(codes).issubset(permission_codes(db,user.user_id)):raise HTTPException(403,'missing operations detail permission')

def get_row(db,model,key,lock=False):
    try:UUID(key)
    except (ValueError,TypeError):raise HTTPException(404,'record not found')
    pk=list(model.__table__.primary_key.columns)[0]
    stmt=select(model).where(pk==key)
    if lock:stmt=stmt.with_for_update().execution_options(populate_existing=True)
    row=db.scalar(stmt)
    if row is None:raise HTTPException(404,'record not found')
    return row

def audit(db,user,action,row,before=None,after=None):
    pk=list(row.__table__.primary_key.columns)[0].name
    write_audit(db,user_id=user.user_id,action=action,entity_type=row.__tablename__,entity_id=getattr(row,pk),before=before,after=after or {'version':getattr(row,'version',None)})

def flush(db):
    try:db.flush()
    except IntegrityError:
        db.rollback();raise HTTPException(409,'duplicate identity or invalid database relationship')

def save(db):
    try:db.commit()
    except IntegrityError:
        db.rollback();raise HTTPException(409,'duplicate identity or invalid database relationship')

def check_version(row,expected):
    if row.version!=expected:raise HTTPException(409,'record version conflict; reload latest record')

def change(db,user,row,expected,values,action):
    before=row_dict(row);pk=list(row.__table__.primary_key.columns)[0]
    result=db.execute(update(type(row)).where(pk==getattr(row,pk.name),type(row).version==expected).values(**values,version=expected+1,updated_at=now_utc()))
    if result.rowcount!=1:db.rollback();raise HTTPException(409,'record version conflict; reload latest record')
    db.refresh(row);audit(db,user,action,row,before);return row

def document(db,user,key):
    if key:need(db,user,'document.read');get_row(db,Document,key)

def active_employee(db,key):
    e=get_row(db,Employee,key)
    if not e.active:raise HTTPException(422,'employee is inactive')
    return e

def active_vehicle(db,key):
    v=get_row(db,Vehicle,key,True)
    if not v.active:raise HTTPException(422,'vehicle is inactive')
    return v

def source_permission(row):
    if row.emergency_case_id:return 'emergency.case.read'
    if row.fire_investigation_case_id:return 'fire_investigation.read'
    return None

def source_row(db,row,lock=False):
    if row.emergency_case_id:return get_row(db,EmergencyCase,row.emergency_case_id,lock)
    if row.fire_investigation_case_id:return get_row(db,FireInvestigationCase,row.fire_investigation_case_id,lock)
    return None

def incident_dict(db,row,perms):
    out=row_dict(row);permission=source_permission(row)
    if permission and permission not in perms:
        out.pop('emergency_case_id');out.pop('fire_investigation_case_id');out['source_restricted']=True
    elif permission:
        src=source_row(db,row)
        if isinstance(src,EmergencyCase):out['source']={'source_id':src.emergency_case_id,'version':src.version,'address':src.incident_address,'number':src.dispatch_number,'call_date':scalar(src.call_date),'call_time':scalar(src.call_time)}
        else:out['source']={'source_id':src.fire_investigation_case_id,'version':src.version,'address':src.location_text,'number':src.case_number,'occurred_at':scalar(src.occurred_at)}
    return out

def incident_date(db,row):
    src=source_row(db,row)
    if isinstance(src,EmergencyCase):return src.call_date
    if isinstance(src,FireInvestigationCase):return src.occurred_at.date() if src.occurred_at else None
    return row.occurred_at.date() if row.occurred_at else None

def source_need(db,user,row):
    permission=source_permission(row)
    if permission:
        need(db,user,permission);source_row(db,row,True)

def create_incident(db,user,payload):
    data=payload.model_dump()
    if payload.emergency_case_id:need(db,user,'emergency.case.read');get_row(db,EmergencyCase,payload.emergency_case_id,True)
    if payload.fire_investigation_case_id:need(db,user,'fire_investigation.read');get_row(db,FireInvestigationCase,payload.fire_investigation_case_id,True)
    row=Incident(**data);db.add(row);flush(db);audit(db,user,'incident.create',row);return row

def patch_incident(db,user,row,payload):
    source_need(db,user,row)
    if row.status!='active':raise HTTPException(409,'cancelled incident is immutable')
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    merged={**{k:(utc(getattr(row,k)) if isinstance(getattr(row,k),datetime) else getattr(row,k)) for k in IncidentInput.model_fields},**values}
    try:IncidentInput.model_validate(merged)
    except ValidationError:raise HTTPException(422,'invalid incident fields or linked source facts')
    return change(db,user,row,payload.expected_version,values,'incident.update')

def create_vehicle(db,user,payload):
    row=Vehicle(**payload.model_dump());db.add(row);flush(db);audit(db,user,'fleet.vehicle.create',row);return row

def dispatch_parent(db,user,key):
    initial=get_row(db,Dispatch,key)
    parent=get_row(db,Incident,initial.incident_id,True);source_need(db,user,parent)
    return get_row(db,Dispatch,key,True),parent

def draft(row):
    if row.status!='draft':raise HTTPException(409,'record is no longer a draft')

def create_dispatch(db,user,incident_id,payload):
    parent=get_row(db,Incident,incident_id,True);source_need(db,user,parent);check_version(parent,payload.expected_version)
    if parent.status!='active':raise HTTPException(409,'incident cancelled')
    if payload.vehicle_id:need(db,user,'fleet.read');active_vehicle(db,payload.vehicle_id)
    document(db,user,payload.document_id)
    row=Dispatch(incident_id=incident_id,**payload.model_dump(exclude={'expected_version'}));db.add(row);flush(db)
    change(db,user,parent,payload.expected_version,{},'incident.dispatch.add');audit(db,user,'incident.dispatch.create',row);return row

def patch_dispatch(db,user,key,payload):
    row,parent=dispatch_parent(db,user,key);draft(row)
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    try:DispatchFields.model_validate({**{k:(utc(getattr(row,k)) if isinstance(getattr(row,k),datetime) else getattr(row,k)) for k in DispatchFields.model_fields},**values})
    except ValidationError:raise HTTPException(422,'invalid dispatch fields or time order')
    if values.get('vehicle_id'):need(db,user,'fleet.read');active_vehicle(db,values['vehicle_id'])
    if 'vehicle_id' in values and values['vehicle_id']!=row.vehicle_id and db.scalar(select(VehicleTrip).where(VehicleTrip.dispatch_id==key)):raise HTTPException(409,'vehicle linked to recorded trips cannot change')
    document(db,user,values.get('document_id'))
    return change(db,user,row,payload.expected_version,{**values,'rate_id':None,'candidate_amount':None,'calculation':{}},'incident.dispatch.update')

def add_crew(db,user,key,payload):
    row,parent=dispatch_parent(db,user,key);draft(row);check_version(row,payload.expected_version);active_employee(db,payload.employee_id)
    crew=DispatchCrew(dispatch_id=key,employee_id=payload.employee_id,role=payload.role);db.add(crew);flush(db)
    change(db,user,row,payload.expected_version,{'rate_id':None,'candidate_amount':None,'calculation':{}},'incident.crew.add');audit(db,user,'incident.crew.create',crew);return crew

def dispatch_snapshot(db,row,parent):
    src=source_row(db,parent,True)
    crew=db.scalars(select(DispatchCrew).where(DispatchCrew.dispatch_id==row.dispatch_id).order_by(DispatchCrew.crew_id)).all()
    return {'incident_id':parent.incident_id,'incident_version':parent.version,'source_version':getattr(src,'version',None),'dispatch_fields':{k:scalar(utc(getattr(row,k)) if isinstance(getattr(row,k),datetime) else getattr(row,k)) for k in DispatchFields.model_fields},'crew':[row_dict(c) for c in crew]}

def calculate(db,user,key,payload):
    row,parent=dispatch_parent(db,user,key);draft(row);check_version(row,payload.expected_version)
    rate=get_row(db,AllowanceRate,payload.rate_id,True)
    if rate.status!='approved':raise HTTPException(422,'rate requires explicit Human approval')
    evidence=dispatch_snapshot(db,row,parent)
    if rate.basis=='per_dispatch':units=Decimal('1')
    elif rate.basis=='per_crew':
        units=Decimal(len(evidence['crew']))
        if units==0:raise HTTPException(422,'crew must be assigned before crew calculation')
    else:
        if not row.departed_at or not row.returned_at:raise HTTPException(422,'departure and return required for per-hour rate')
        units=Decimal(str((row.returned_at-row.departed_at).total_seconds()))/Decimal('3600')
    amount=(rate.amount*units).quantize(Decimal('.01'),rounding={'half_up':ROUND_HALF_UP,'half_even':ROUND_HALF_EVEN,'down':ROUND_DOWN}[rate.rounding])
    if amount>=Decimal('1000000000000'):raise HTTPException(422,'calculation exceeds amount range')
    evidence.update({'rate_id':rate.rate_id,'rate_version':rate.version,'rate_amount':str(rate.amount),'basis':rate.basis,'units':str(units),'rounding':rate.rounding})
    return change(db,user,row,payload.expected_version,{'rate_id':rate.rate_id,'candidate_amount':amount,'calculation':evidence},'incident.allowance.calculate')

def review_dispatch(db,user,key,payload,approve=False):
    row,parent=dispatch_parent(db,user,key);check_version(row,payload.expected_version)
    snapshot=dispatch_snapshot(db,row,parent)
    if approve:
        if row.status!='reviewed':raise HTTPException(409,'Human review required before approval')
        if row.review_snapshot!=snapshot:raise HTTPException(409,'review sources changed; cancel and create a new draft')
    else:
        draft(row)
        if row.rate_id:
            rate=get_row(db,AllowanceRate,row.rate_id)
            original={k:v for k,v in row.calculation.items() if k in snapshot}
            if original!=snapshot or rate.status!='approved' or rate.version!=row.calculation.get('rate_version'):raise HTTPException(409,'allowance sources changed; recalculate')
    values={'status':'approved' if approve else 'reviewed','review_note':payload.note}
    if approve:values.update(official_amount=row.candidate_amount,approved_by=user.user_id,approved_at=now_utc())
    else:values.update(reviewed_by=user.user_id,reviewed_at=now_utc(),review_snapshot=snapshot)
    return change(db,user,row,payload.expected_version,values,'incident.dispatch.approve' if approve else 'incident.dispatch.review')

def cancel_dispatch(db,user,key,payload):
    row,parent=dispatch_parent(db,user,key)
    if row.status=='cancelled':raise HTTPException(409,'already cancelled')
    # Preserve the typed approved amount and original approver; cancelled rows leave aggregates.
    return change(db,user,row,payload.expected_version,{'status':'cancelled','cancellation_note':payload.note,'cancelled_by':user.user_id,'cancelled_at':now_utc()},'incident.dispatch.cancel')

def trip(db,user,key,payload):
    # Always lock incident before vehicle when a dispatch link exists.
    if payload.dispatch_id:
        need(db,user,'incident.read');d,parent=dispatch_parent(db,user,payload.dispatch_id)
        if d.vehicle_id!=key or d.status=='cancelled' or parent.status=='cancelled':raise HTTPException(422,'dispatch vehicle mismatch or cancelled incident')
    v=active_vehicle(db,key);check_version(v,payload.expected_version)
    if payload.start_odometer!=v.odometer:raise HTTPException(422,'start odometer must equal current vehicle odometer')
    previous=db.scalar(select(VehicleTrip).where(VehicleTrip.vehicle_id==key).order_by(VehicleTrip.ended_at.desc()).limit(1))
    if previous and utc(payload.started_at)<utc(previous.ended_at):raise HTTPException(422,'trip overlaps earlier use history')
    if payload.driver_employee_id:need(db,user,'incident.crew.read');active_employee(db,payload.driver_employee_id)
    document(db,user,payload.document_id)
    row=VehicleTrip(vehicle_id=key,created_by=user.user_id,**payload.model_dump(exclude={'expected_version'}));db.add(row);flush(db)
    change(db,user,v,payload.expected_version,{'odometer':payload.end_odometer},'fleet.odometer.update');audit(db,user,'fleet.trip.create',row);return row

def utc(value):return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)

def fuel(db,user,key,payload):
    v=active_vehicle(db,key);check_version(v,payload.expected_version);document(db,user,payload.document_id)
    stock=v.fuel_stock
    if payload.kind=='receipt':stock+=payload.liters
    if payload.kind=='issue':stock-=payload.liters
    if stock<0:raise HTTPException(422,'fuel stock underflow')
    if stock>=Decimal('1000000000000'):raise HTTPException(422,'fuel stock range exceeded')
    row=FuelEntry(vehicle_id=key,created_by=user.user_id,**payload.model_dump(exclude={'expected_version'}));db.add(row);flush(db)
    change(db,user,v,payload.expected_version,{'fuel_stock':stock},'fleet.stock.update');audit(db,user,'fleet.fuel.create',row);return row

def service(db,user,key,payload):
    v=active_vehicle(db,key);check_version(v,payload.expected_version);document(db,user,payload.document_id)
    if payload.next_service_odometer is not None and payload.next_service_odometer<v.odometer:raise HTTPException(422,'next service mileage precedes current mileage')
    if payload.resolves_fault_id:
        fault=get_row(db,VehicleService,payload.resolves_fault_id,True)
        if payload.kind!='repair' or fault.kind!='fault' or fault.vehicle_id!=key or fault.status=='cancelled' or fault.resolved_by_service_id or payload.performed_on<fault.performed_on:raise HTTPException(422,'repair must explicitly reference an unresolved fault of this vehicle')
    row=VehicleService(vehicle_id=key,created_by=user.user_id,**payload.model_dump(exclude={'expected_version'}));db.add(row);flush(db)
    change(db,user,v,payload.expected_version,{},'fleet.service.add');audit(db,user,'fleet.service.create',row);return row

def review_service(db,user,key,payload,approve=False):
    initial=get_row(db,VehicleService,key);v=get_row(db,Vehicle,initial.vehicle_id,True);row=get_row(db,VehicleService,key,True)
    check_version(row,payload.expected_version)
    if row.status!=('reviewed' if approve else 'draft'):raise HTTPException(409,'service review state conflict')
    values={'status':'approved' if approve else 'reviewed','review_note':payload.note}
    if approve:
        check_version(v,payload.expected_vehicle_version)
        if row.resolves_fault_id:
            if payload.expected_fault_version is None:raise HTTPException(422,'expected_fault_version required for fault resolution')
            fault=get_row(db,VehicleService,row.resolves_fault_id,True);check_version(fault,payload.expected_fault_version)
            if fault.resolved_by_service_id or fault.status=='cancelled' or fault.vehicle_id!=v.vehicle_id or fault.kind!='fault':raise HTTPException(409,'fault is already resolved or no longer eligible')
            change(db,user,fault,payload.expected_fault_version,{'resolved_by_service_id':row.service_id,'resolved_at':now_utc()},'fleet.fault.resolve')
        latest=db.scalar(select(VehicleService).where(VehicleService.vehicle_id==v.vehicle_id,VehicleService.status=='approved').order_by(VehicleService.performed_on.desc()).limit(1))
        if latest and latest.performed_on>row.performed_on:raise HTTPException(409,'a newer approved service exists')
        if row.next_service_odometer is not None and row.next_service_odometer<v.odometer:raise HTTPException(409,'vehicle mileage changed beyond proposed next service')
        values.update(approved_by=user.user_id,approved_at=now_utc())
        dates={k:getattr(row,k) for k in ['next_inspection_on','next_service_on','next_service_odometer'] if getattr(row,k) is not None}
        change(db,user,v,v.version,dates,'fleet.schedule.update')
    else:values['reviewed_by']=user.user_id
    return change(db,user,row,payload.expected_version,values,'fleet.service.approve' if approve else 'fleet.service.review')

def statistics(db,user,year,month=None):
    perms=permission_codes(db,user.user_id);rows=[i for i in db.scalars(select(Incident)) if (d:=incident_date(db,i)) and d.year==year and (month is None or d.month==month)]
    ids=[i.incident_id for i in rows];counts={}
    for i in rows:
        if i.status=='active':counts[i.kind]=counts.get(i.kind,0)+1
    dispatches=list(db.scalars(select(Dispatch).where(Dispatch.incident_id.in_(ids)))) if ids else []
    out={'year':year,'month':month,'incident_counts':counts,'incidents':sum(counts.values()),'dispatches':sum(d.status!='cancelled' for d in dispatches),'approved_dispatches':sum(d.status=='approved' for d in dispatches),'official_allowance':str(sum((d.official_amount or Decimal('0') for d in dispatches if d.status=='approved'),Decimal('0')).quantize(Decimal('.01')))}
    if 'incident.read' in perms:out['source_incident_ids']=[i.incident_id for i in rows if not source_permission(i) or source_permission(i) in perms]
    return out

def alerts(db,as_of):
    out=[]
    for v in db.scalars(select(Vehicle).where(Vehicle.active.is_(True))):
        for kind,key in [('inspection','next_inspection_on'),('service','next_service_on')]:
            due=getattr(v,key)
            if due and (due-as_of).days<=30:out.append({'vehicle_id':v.vehicle_id,'code':v.code,'kind':kind,'due_on':due.isoformat(),'days_remaining':(due-as_of).days,'overdue':due<as_of})
        if v.next_service_odometer is not None and v.odometer>=v.next_service_odometer:out.append({'vehicle_id':v.vehicle_id,'code':v.code,'kind':'service_mileage','odometer':str(v.odometer),'due_odometer':str(v.next_service_odometer),'overdue':True})
        for s in db.scalars(select(VehicleService).where(VehicleService.vehicle_id==v.vehicle_id,VehicleService.kind=='fault',VehicleService.status!='cancelled')):
            if not s.resolved_by_service_id:out.append({'vehicle_id':v.vehicle_id,'code':v.code,'kind':'unresolved_fault','service_id':s.service_id,'description':s.description})
    return out

IMPORT_SCHEMA={'incidents':IncidentInput,'vehicles':VehicleInput,'dispatches':DispatchInput,'trips':TripInput,'fuel':FuelInput,'services':ServiceInput,'crew':CrewInput}
IMPORT_PARENT={'crew':'dispatch_id','dispatches':'incident_id','trips':'vehicle_id','fuel':'vehicle_id','services':'vehicle_id'}

def read_table(raw,filename):
    if len(raw)>5*1024*1024:raise HTTPException(413,'table exceeds 5 MB')
    try:
        if filename.lower().endswith('.csv'):
            reader=csv.reader(StringIO(raw.decode('utf-8-sig')));table=list(reader)
        elif filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            from zipfile import ZipFile
            with ZipFile(BytesIO(raw)) as z:
                if sum(x.file_size for x in z.infolist())>30*1024*1024:raise ValueError('expanded workbook too large')
            w=load_workbook(BytesIO(raw),read_only=True,data_only=False);table=list(w.active.values);w.close()
        else:raise ValueError('csv or xlsx required')
        if not table:raise ValueError('empty table')
        headers=[str(x).strip() if x is not None else '' for x in table[0]]
        if any(not x for x in headers) or len(headers)!=len(set(headers)):raise ValueError('invalid headers')
        if len(table)>1001:raise ValueError('maximum 1000 rows')
        rows=[]
        for values in table[1:]:
            if not any(x is not None and str(x)!='' for x in values):continue
            if len(values)>len(headers):raise ValueError('extra cells')
            if any(isinstance(x,str) and x.startswith('=') for x in values):raise ValueError('formula cells forbidden')
            rows.append({k:v for k,v in zip(headers,values) if v is not None and v!=''})
        return rows,headers
    except HTTPException:raise
    except Exception:raise HTTPException(422,'invalid UTF-8 CSV/XLSX table; formulas and oversized workbooks are forbidden')

def apply_rows(db,user,dataset,rows):
    domain='incident' if dataset in ['incidents','dispatches','crew'] else 'fleet'
    need(db,user,f'{domain}.import',f'{domain}.create',f'{domain}.read')
    schema=IMPORT_SCHEMA[dataset];parent_key=IMPORT_PARENT.get(dataset);ids=[]
    if dataset=='crew':need(db,user,'incident.crew.read','incident.crew.manage')
    for n,source in enumerate(rows,2):
        data=dict(source);parent=data.pop(parent_key,None) if parent_key else None
        payload=schema.model_validate(data)
        if parent_key and not parent:raise HTTPException(422,f'row {n}: parent ID required')
        if dataset=='incidents':row=create_incident(db,user,payload)
        elif dataset=='vehicles':row=create_vehicle(db,user,payload)
        elif dataset=='dispatches':row=create_dispatch(db,user,parent,payload)
        elif dataset=='crew':row=add_crew(db,user,parent,payload)
        elif dataset=='trips':row=trip(db,user,parent,payload)
        elif dataset=='fuel':row=fuel(db,user,parent,payload)
        else:row=service(db,user,parent,payload)
        ids.append(getattr(row,list(row.__table__.primary_key.columns)[0].name))
    return ids


def import_preview(db,user,dataset,raw,filename):
    if dataset not in IMPORT_SCHEMA:raise HTTPException(422,'unknown import dataset')
    domain='incident' if dataset in ['incidents','dispatches','crew'] else 'fleet';need(db,user,f'{domain}.import',f'{domain}.create',f'{domain}.read')
    rows,headers=read_table(raw,filename);schema=IMPORT_SCHEMA[dataset];parent_key=IMPORT_PARENT.get(dataset)
    if set(headers)-set(schema.model_fields)-({parent_key} if parent_key else set()):raise HTTPException(422,'unknown columns; tenant, user and official decision fields are forbidden')
    if not rows:raise HTTPException(422,'table has no data rows')
    # Run the real invariant checks transactionally, then discard every domain/audit write.
    try:apply_rows(db,user,dataset,rows)
    except ValidationError:db.rollback();raise HTTPException(422,'table validation failed; no domain records changed')
    except HTTPException:db.rollback();raise
    db.rollback()
    serial_rows=[{k:scalar(v) for k,v in row.items()} for row in rows]
    preview=OperationsImportPreview(dataset=dataset,filename=filename,file_sha256=sha256(raw).hexdigest(),row_data=serial_rows,created_by=user.user_id,expires_at=now_utc()+timedelta(hours=1))
    db.add(preview);flush(db);audit(db,user,f'{domain}.import.preview',preview,after={'dataset':dataset,'rows':len(rows),'file_sha256':preview.file_sha256});save(db)
    return {'status':'preview','preview_id':preview.preview_id,'version':preview.version,'dataset':dataset,'rows':len(rows),'file_sha256':preview.file_sha256,'expires_at':scalar(preview.expires_at),'columns':headers,'sample':serial_rows[:20]}


def confirm_import(db,user,key,payload):
    preview=get_row(db,OperationsImportPreview,key,True)
    if preview.created_by!=user.user_id:raise HTTPException(403,'import preview belongs to another user')
    check_version(preview,payload.expected_version)
    if preview.status!='preview' or utc(preview.expires_at)<now_utc() or preview.file_sha256!=payload.file_sha256:raise HTTPException(409,'preview expired, changed or already applied; upload again')
    domain='incident' if preview.dataset in ['incidents','dispatches','crew'] else 'fleet'
    try:
        ids=apply_rows(db,user,preview.dataset,preview.row_data)
        change(db,user,preview,payload.expected_version,{'status':'applied','applied_ids':ids},f'{domain}.import.confirm')
        write_audit(db,user_id=user.user_id,action=f'{domain}.import',entity_type=f'operations_{preview.dataset}',entity_id=key,after={'inserted':len(ids),'file_sha256':preview.file_sha256,'source_ids':ids})
        save(db)
    except ValidationError:db.rollback();raise HTTPException(422,'table validation failed; entire import rolled back')
    except HTTPException:db.rollback();raise
    return {'status':'applied','inserted':len(ids),'source_ids':ids,'file_sha256':preview.file_sha256}


def safe_cell(value):
    if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False)
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
    return value

def render_table(rows,format,headers=None):
    # Union only permission-filtered row keys so later source projections are retained.
    headers=headers or list(dict.fromkeys(key for row in rows for key in row)) or ['source_id']
    values=[[safe_cell(r.get(k)) for k in headers] for r in rows]
    if format=='csv':
        out=StringIO();writer=csv.writer(out);writer.writerow(headers);writer.writerows(values);return out.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    from openpyxl import Workbook
    w=Workbook();s=w.active;s.title='Operations';s.append(headers)
    for v in values:s.append(v)
    out=BytesIO();w.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

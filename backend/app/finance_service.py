"""Financial proposals and posting transactions. AI never posts authority."""
from datetime import date, datetime
from decimal import Decimal
from uuid import UUID
from fastapi import HTTPException
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from .audit import write_audit
from .authz import permission_codes
from .models import now_utc
from .finance_models import FinanceYear, BudgetAccount

SCHEMA_VERSION = 'finance-v1'
def scalar(value):
    if isinstance(value,Decimal): return format(value,'f')
    if isinstance(value,(date,datetime)): return value.isoformat()
    return value

def row_dict(row): return {c.name:scalar(getattr(row,c.name)) for c in row.__table__.columns}
def need(db,user,*codes):
    if not set(codes).issubset(permission_codes(db,user.user_id)): raise HTTPException(403,'required finance/source permission missing')
def get_row(db,model,key,lock=False):
    try: UUID(key)
    except (ValueError,TypeError): raise HTTPException(404,'record not found')
    pk=list(model.__table__.primary_key.columns)[0]
    stmt=select(model).where(pk==key)
    if lock: stmt=stmt.with_for_update().execution_options(populate_existing=True)
    row=db.scalar(stmt)
    if row is None: raise HTTPException(404,'record not found')
    return row

def check_version(row,expected):
    if row.version!=expected: raise HTTPException(409,'version conflict; reload current record')
def bump(db,row,expected,values=None):
    pk=list(row.__table__.primary_key.columns)[0]
    result=db.execute(update(type(row)).where(pk==getattr(row,pk.name),type(row).version==expected).values(**(values or {}),version=expected+1,updated_at=now_utc()).execution_options(synchronize_session=False))
    if result.rowcount!=1: raise HTTPException(409,'version conflict; reload current record')
    db.refresh(row)
def audit(db,user,action,row,before=None):
    pk=list(row.__table__.primary_key.columns)[0].name
    write_audit(db,user_id=user.user_id,action='finance.'+action,entity_type=row.__tablename__,entity_id=getattr(row,pk),before=before,after=row_dict(row))
def save(db):
    try: db.commit()
    except IntegrityError:
        db.rollback();raise HTTPException(409,'duplicate or conflicting financial record') from None

def create_year(db,user,payload):
    row=FinanceYear(**payload.model_dump());db.add(row);db.flush();audit(db,user,'year.create',row);return row

def patch_year(db,user,key,payload):
    row=get_row(db,FinanceYear,key,True);check_version(row,payload.expected_version)
    if row.status!='draft': raise HTTPException(409,'approved fiscal policy is immutable')
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if any(v is None for v in values.values()):raise HTTPException(422,'policy fields cannot be cleared')
    before=row_dict(row);bump(db,row,row.version,values);audit(db,user,'year.update',row,before);return row

def approve_year(db,user,key,payload):
    row=get_row(db,FinanceYear,key,True);check_version(row,payload.expected_version)
    if row.status!='draft': raise HTTPException(409,'policy already approved')
    bump(db,row,row.version,{'status':'approved','approved_by':user.user_id,'approved_at':now_utc(),'reason':payload.reason});audit(db,user,'year.approve',row);return row

def create_account(db,user,payload):
    year=get_row(db,FinanceYear,payload.year_id,True)
    if year.status!='approved' or payload.level>len(year.account_levels):raise HTTPException(422,'account depth must follow Human-approved hierarchy policy')
    if payload.level==1:
        if payload.parent_id: raise HTTPException(422,'root account has no parent')
    else:
        if not payload.parent_id: raise HTTPException(422,'account parent required')
        parent=get_row(db,BudgetAccount,payload.parent_id,True)
        if parent.year_id!=payload.year_id or parent.level!=payload.level-1: raise HTTPException(422,'parent must be preceding level in same fiscal year')
    row=BudgetAccount(**payload.model_dump());db.add(row);db.flush();audit(db,user,'account.create',row);return row

def patch_account(db,user,key,payload):
    row=get_row(db,BudgetAccount,key,True);check_version(row,payload.expected_version);before=row_dict(row);bump(db,row,row.version,{'name':payload.name});audit(db,user,'account.update',row,before);return row

from .models import Document, ContractCase
from .finance_models import FinanceProposal, FinanceJournal
from .finance_schemas import ProposalInput

def money(value):return format(Decimal(value).quantize(Decimal('.01')),'f')
def source_document(db,user,key,lock=False):
    need(db,user,'document.read');doc=get_row(db,Document,key,lock)
    return {'document_id':doc.document_id,'sha256':doc.sha256,'filename':doc.original_filename}
def lock_sources(db,user,contract_ids=(),document_ids=(),counterparty_ids=()):
    """Refresh and hold all common evidence in contract/vendor/Document UUID order."""
    contracts=[];parties=set(counterparty_ids);docs=set(document_ids)
    for key in sorted(set(contract_ids)):
        need(db,user,'contract.read');row=get_row(db,ContractCase,key,True);contracts.append(row)
        if row.counterparty_id:parties.add(row.counterparty_id)
        docs.update(db.scalars(select(ContractDocument.document_id).where(ContractDocument.contract_case_id==key)))
    for key in sorted(parties):
        need(db,user,'contract.read');get_row(db,ContractCounterparty,key,True)
    for key in sorted(docs):source_document(db,user,key,True)
    return contracts

def account_balance(db,key):
    amounts=[Decimal('0'),Decimal('0'),Decimal('0')]
    for r in db.scalars(select(FinanceJournal).where(FinanceJournal.account_id==key)):
        amounts=[a+v for a,v in zip(amounts,[r.allocated,r.reserved,r.spent])]
    allocated,reserved,spent=amounts
    return {'account_id':key,'allocated':money(allocated),'reserved':money(reserved),'spent':money(spent),'available':money(allocated-reserved-spent)}
def visible(db,user,row):
    perms=permission_codes(db,user.user_id)
    hidden=set()
    if 'document.read' not in perms:hidden.update({'document','target_document','documents','document_id','source_document_id','file_sha256','filename','sha256'})
    if 'contract.read' not in perms:hidden.update({'contract','target_contract','contract_case_id','counterparty_id','counterparty_name','counterparty'})
    def redact(value):
        if isinstance(value,dict):return {k:redact(v) for k,v in value.items() if k not in hidden}
        if isinstance(value,list):return [redact(v) for v in value]
        return value
    return redact(row_dict(row))

def proposal_accounts(db,row):
    ids=sorted({row.account_id,*([row.to_account_id] if row.to_account_id else [])})
    return {key:get_row(db,BudgetAccount,key,True) for key in ids}
def validate_proposal(db,user,data):
    account=get_row(db,BudgetAccount,data['account_id'])
    year=get_row(db,FinanceYear,account.year_id)
    if account.level!=len(year.account_levels):raise HTTPException(422,'financial postings require configured leaf account')
    if year.status!='approved':raise HTTPException(422,'fiscal policy requires Human approval')
    amount=Decimal(data['amount'])
    if not amount.is_finite() or amount!=amount.quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'amount exceeds configured exact precision; implicit rounding forbidden')
    if data['currency']!=year.currency:raise HTTPException(422,'currency mismatch')
    kind=data['kind']
    if amount<0 and kind!='amendment':raise HTTPException(422,'only amendment permits signed adjustment')
    source_document(db,user,data['document_id'])
    if kind=='transfer':
        if not data.get('to_account_id') or data['to_account_id']==account.account_id:raise HTTPException(422,'different destination account required')
        other=get_row(db,BudgetAccount,data['to_account_id'])
        if other.year_id!=account.year_id or other.level!=len(year.account_levels):raise HTTPException(422,'destination must be leaf in same fiscal year')
    elif data.get('to_account_id'):raise HTTPException(422,'destination allowed only for transfer')
    if kind=='reversal':
        if not data.get('reverses_id'):raise HTTPException(422,'approved reversal target required')
        target=get_row(db,FinanceProposal,data['reverses_id'],True)
        if target.contract_case_id:need(db,user,'contract.read');source_document(db,user,target.document_id)
        if target.status!='approved' or target.kind=='reversal' or target.account_id!=account.account_id or abs(target.amount)!=amount:raise HTTPException(422,'reversal must match approved target account and exact amount')
        if db.scalar(select(FinanceProposal).where(FinanceProposal.reverses_id==target.proposal_id,FinanceProposal.status!='cancelled')):raise HTTPException(409,'target already has a reversal')
    elif data.get('reverses_id'):raise HTTPException(422,'reversal reference allowed only for reversal')
    if kind not in ('commitment','payment') and (data.get('contract_case_id') or data.get('commitment_id') or data.get('invoice_id')):raise HTTPException(422,'contract/commitment references allowed only for procurement')
    procurement_validate(db,user,data,account,year)
    return account,year

def create_proposal(db,user,payload):
    if db.scalar(select(FinanceProposal).where(FinanceProposal.idempotency_key==payload.idempotency_key)):raise HTTPException(409,'idempotency key already exists')
    data=payload.model_dump();validate_proposal(db,user,data)
    row=FinanceProposal(**data,created_by=user.user_id);db.add(row);db.flush();audit(db,user,'proposal.create',row);return row

def patch_proposal(db,user,key,payload):
    row=get_row(db,FinanceProposal,key,True);check_version(row,payload.expected_version)
    if row.status!='draft':raise HTTPException(409,'reviewed/approved proposal cannot be rewritten')
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if any(v is None for v in values.values()):raise HTTPException(422,'financial fields cannot be cleared')
    data={**row_dict(row),**values};validate_proposal(db,user,data);before=row_dict(row);bump(db,row,row.version,values);audit(db,user,'proposal.update',row,before);return row

def proposal_snapshot(db,user,row,accounts):
    snapshot = {'accounts':{k:{'version':r.version,'year_id':r.year_id,'code':r.code} for k,r in accounts.items()},'year':row_dict(get_row(db,FinanceYear,accounts[row.account_id].year_id)),'document':source_document(db,user,row.document_id,True),'amount':money(row.amount),'currency':row.currency,'kind':row.kind}
    if row.contract_case_id:
        need(db,user,'contract.read');snapshot['contract']=contract_dict(db,user,get_row(db,ContractCase,row.contract_case_id,True),True)
    if row.invoice_id:
        snapshot['invoice']=visible(db,user,get_row(db,ProcurementEvent,row.invoice_id,True))
    if row.commitment_id:
        snapshot['commitment']=visible(db,user,get_row(db,FinanceProposal,row.commitment_id,True))
    if row.reverses_id:
        target=get_row(db,FinanceProposal,row.reverses_id,True);snapshot['target']=visible(db,user,target)
        snapshot['target_document']=source_document(db,user,target.document_id,True)
        if target.contract_case_id:
            need(db,user,'contract.read');snapshot['target_contract']=contract_dict(db,user,get_row(db,ContractCase,target.contract_case_id,True),True)
    return snapshot

def journal_changes(db,row):
    amount=row.amount
    if row.kind in ('initial','amendment'):return {row.account_id:(amount,Decimal('0'),Decimal('0'))}
    if row.kind=='transfer':return {row.account_id:(-amount,Decimal('0'),Decimal('0')),row.to_account_id:(amount,Decimal('0'),Decimal('0'))}
    if row.kind=='commitment':return {row.account_id:(Decimal('0'),amount,Decimal('0'))}
    if row.kind=='payment':return {row.account_id:(Decimal('0'),-amount,amount)}
    target=get_row(db,FinanceProposal,row.reverses_id,True)
    return {r.account_id:(-r.allocated,-r.reserved,-r.spent) for r in db.scalars(select(FinanceJournal).where(FinanceJournal.proposal_id==target.proposal_id))}

def proposal_action(db,user,key,payload,action):
    # Lock accounts first and proposals second, consistently across approval/reversal.
    initial=get_row(db,FinanceProposal,key)
    target=get_row(db,FinanceProposal,initial.reverses_id) if initial.reverses_id else None
    ids={initial.account_id,*([initial.to_account_id] if initial.to_account_id else [])}
    if target:ids.update([target.account_id,*([target.to_account_id] if target.to_account_id else [])])
    accounts={k:get_row(db,BudgetAccount,k,True) for k in sorted(ids)}
    source_contract=initial.contract_case_id or (target.contract_case_id if target else None)
    lock_sources(db,user,[source_contract] if source_contract else [],[initial.document_id,*([target.document_id] if target else [])])
    row=get_row(db,FinanceProposal,key,True);check_version(row,payload.expected_version)
    if action=='cancel':
        if row.status not in ('draft','reviewed'):raise HTTPException(409,'approved journal requires compensating reversal')
        bump(db,row,row.version,{'status':'cancelled','reason':payload.reason});audit(db,user,'proposal.cancel',row);return row
    if row.status not in ('draft','reviewed') or (action=='approve' and row.status!='reviewed'):raise HTTPException(409,'Human review required before approval')
    # Reversal existence check excludes this proposal.
    if row.kind=='reversal':
        target=get_row(db,FinanceProposal,row.reverses_id,True)
        if db.scalar(select(FinanceProposal).where(FinanceProposal.reverses_id==row.reverses_id,FinanceProposal.proposal_id!=row.proposal_id,FinanceProposal.status=='approved')):raise HTTPException(409,'target already reversed')
    else:validate_proposal(db,user,row_dict(row))
    snapshot=proposal_snapshot(db,user,row,accounts)
    if action=='review':
        bump(db,row,row.version,{'status':'reviewed','review_snapshot':snapshot,'reviewed_by':user.user_id,'reviewed_at':now_utc(),'reason':payload.reason});audit(db,user,'proposal.review',row);return row
    if row.review_snapshot!=snapshot:raise HTTPException(409,'financial source changed since review; review again')
    account=accounts[row.account_id];year=get_row(db,FinanceYear,account.year_id)
    procurement_validate(db,user,row_dict(row),account,year,posting=True)
    if target and target.kind=='commitment' and commitment_remaining(db,target)!=target.amount:raise HTTPException(422,'reverse executed payments before commitment reversal')
    changes=journal_changes(db,row)
    for account_id,(allocated,reserved,spent) in changes.items():
        b=account_balance(db,account_id)
        new=[Decimal(b[k])+v for k,v in zip(['allocated','reserved','spent'],[allocated,reserved,spent])]
        if min(new)<0 or new[0]-new[1]-new[2]<0:raise HTTPException(422,'insufficient budget or invalid reservation/spending reversal')
    for account_id,(allocated,reserved,spent) in changes.items():
        db.add(FinanceJournal(proposal_id=row.proposal_id,account_id=account_id,allocated=allocated,reserved=reserved,spent=spent))
        bump(db,accounts[account_id],accounts[account_id].version)
    bump(db,row,row.version,{'status':'approved','approved_by':user.user_id,'approved_at':now_utc(),'reason':payload.reason});audit(db,user,'proposal.approve',row);return row

from .models import ContractCounterparty, ContractDocument, ContractChange
from .finance_models import ProcurementProfile, FinanceCandidate, FinanceContractAmendment

def contract_dict(db,user,row,lock=False):
    if lock:lock_sources(db,user,[row.contract_case_id])
    need(db,user,'contract.read');data=row_dict(row)
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==row.contract_case_id))
    data['profile']=visible(db,user,profile) if profile else None
    party=get_row(db,ContractCounterparty,row.counterparty_id,lock) if row.counterparty_id else None
    data['counterparty_name']=party.name if party else None
    data['counterparty']=row_dict(party) if party else None
    if 'document.read' in permission_codes(db,user.user_id):
        data['documents']=[{'document_id':link.document_id,'document_role':link.document_role,'sha256':get_row(db,Document,link.document_id,lock).sha256} for link in db.scalars(select(ContractDocument).where(ContractDocument.contract_case_id==row.contract_case_id))]
    return data

def record_contract_change(db,user,row,before,reason):
    from sqlalchemy import func
    sequence=(db.scalar(select(func.max(ContractChange.sequence_no)).where(ContractChange.contract_case_id==row.contract_case_id)) or 0)+1
    db.add(ContractChange(contract_case_id=row.contract_case_id,sequence_no=sequence,reason=reason,before_data=before,after_data=row_dict(row),created_by=user.user_id))
def verify_contract_money(db,row,expected):
    # Existing SQLite NUMERIC-affinity tables cannot become TEXT by ORM declaration.
    # Verify the persisted value before commit; never silently manufacture lost cents.
    db.refresh(row)
    if row.amount != expected:
        raise HTTPException(422,'existing database cannot preserve exact contract amount; migrate verified original evidence to PostgreSQL')

def create_contract(db,user,payload):
    need(db,user,'contract.read','contract.create');year=get_row(db,FinanceYear,payload.year_id,True)
    if year.status!='approved' or year.currency!=payload.currency or payload.amount!=payload.amount.quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'approved matching fiscal currency/precision policy required')
    source_document(db,user,payload.document_id)
    if payload.counterparty_id:
        party=get_row(db,ContractCounterparty,payload.counterparty_id)
        if not party.active:raise HTTPException(422,'inactive contractor')
    if payload.start_date and payload.end_date and payload.start_date>payload.end_date:raise HTTPException(422,'contract dates out of order')
    values=payload.model_dump(exclude={'year_id','document_id','renewal_on'})
    row=ContractCase(**values,created_by=user.user_id);db.add(row);db.flush();verify_contract_money(db,row,payload.amount)
    db.add(ProcurementProfile(contract_case_id=row.contract_case_id,year_id=year.year_id,renewal_on=payload.renewal_on))
    db.add(ContractDocument(contract_case_id=row.contract_case_id,document_id=payload.document_id,document_role='contract_source'))
    audit(db,user,'contract.create',row);record_contract_change(db,user,row,{},'Created from exact evidence-linked finance adapter');return row

def validate_contract_policy(db,row,amount):
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==row.contract_case_id))
    if profile:
        year=get_row(db,FinanceYear,profile.year_id)
        if year.status!='approved' or year.currency!=row.currency or amount is None or not amount.is_finite() or amount<0 or abs(amount)>=Decimal('10000000000000000') or amount!=amount.quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'contract must match approved fiscal currency/precision; amount cannot be cleared')
    return profile

def approve_contract(db,user,key,payload):
    need(db,user,'contract.read','contract.approve');row=get_row(db,ContractCase,key,True);check_version(row,payload.expected_version)
    if row.status=='approved':raise HTTPException(409,'approved source is immutable; use reviewed amendment')
    if not validate_contract_policy(db,row,row.amount):raise HTTPException(422,'financial contract fiscal configuration required')
    if row.start_date and row.end_date and row.start_date>row.end_date:raise HTTPException(422,'contract dates out of order')
    docs=list(db.scalars(select(ContractDocument).where(ContractDocument.contract_case_id==key)))
    if not docs:raise HTTPException(422,'financial contract requires linked source evidence')
    snapshot=contract_dict(db,user,row,True)
    for doc in docs:source_document(db,user,doc.document_id,True)
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==key))
    if profile:bump(db,profile,profile.version,{'approved_evidence':{**snapshot,'profile':{k:v for k,v in snapshot['profile'].items() if k!='approved_evidence'}}})
    before=row_dict(row);bump(db,row,row.version,{'status':'approved','approved_by':user.user_id,'approved_at':now_utc()});audit(db,user,'contract.approve',row,before);record_contract_change(db,user,row,before,payload.reason);return row

def reversed_target(db,key):
    return db.scalar(select(FinanceProposal).where(FinanceProposal.reverses_id==key,FinanceProposal.status=='approved')) is not None

def commitment_remaining(db,commit):
    paid=sum((r.amount for r in db.scalars(select(FinanceProposal).where(FinanceProposal.commitment_id==commit.proposal_id,FinanceProposal.kind=='payment',FinanceProposal.status=='approved')) if not reversed_target(db,r.proposal_id)),Decimal('0'))
    return commit.amount-paid

def procurement_validate(db,user,data,account,year,posting=False):
    if data['kind'] not in ('commitment','payment'):return
    need(db,user,'contract.read')
    if not data.get('contract_case_id'):raise HTTPException(422,'common approved contract reference required')
    contract=get_row(db,ContractCase,data['contract_case_id'],True)
    if contract.status!='approved' or contract.amount is None or not contract.amount.is_finite() or contract.amount<0 or contract.amount>=Decimal('10000000000000000') or contract.currency!=year.currency:raise HTTPException(422,'approved matching contract amount/currency required')
    contract_docs=list(db.scalars(select(ContractDocument).where(ContractDocument.contract_case_id==contract.contract_case_id)))
    if not contract_docs:raise HTTPException(422,'approved common contract source evidence required')
    for link in contract_docs:source_document(db,user,link.document_id)
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==contract.contract_case_id))
    if not profile or profile.year_id!=year.year_id:raise HTTPException(422,'contract fiscal year mismatch; configure common contract profile')
    validate_contract_policy(db,contract,contract.amount)
    if data['kind']=='commitment':
        if data.get('commitment_id') or data.get('invoice_id'):raise HTTPException(422,'commitment cannot reference another commitment/invoice')
        if posting:
            committed=sum((r.amount for r in db.scalars(select(FinanceProposal).where(FinanceProposal.contract_case_id==contract.contract_case_id,FinanceProposal.kind=='commitment',FinanceProposal.status=='approved')) if not reversed_target(db,r.proposal_id)),Decimal('0'))
            if committed+Decimal(data['amount'])>contract.amount:raise HTTPException(422,'commitments exceed approved contract amount')
    else:
        if not data.get('commitment_id'):raise HTTPException(422,'approved commitment required')
        commit=get_row(db,FinanceProposal,data['commitment_id'],True)
        if commit.kind!='commitment' or commit.status!='approved' or reversed_target(db,commit.proposal_id) or commit.account_id!=account.account_id or commit.contract_case_id!=contract.contract_case_id:raise HTTPException(422,'payment must match active approved commitment/account/contract')
        if posting and Decimal(data['amount'])>commitment_remaining(db,commit):raise HTTPException(422,'payment exceeds remaining approved commitment')
        if year.require_invoice_on_payment and not data.get('invoice_id'):raise HTTPException(422,'approved fiscal policy requires linked approved invoice')
        if data.get('invoice_id'):
            invoice=get_row(db,ProcurementEvent,data['invoice_id'],True)
            if invoice.kind!='invoice' or invoice.status!='approved' or invoice.contract_case_id!=contract.contract_case_id or invoice.currency!=year.currency:raise HTTPException(422,'payment invoice must be approved and match common contract/currency')
            if posting:
                paid=sum((r.amount for r in db.scalars(select(FinanceProposal).where(FinanceProposal.invoice_id==invoice.event_id,FinanceProposal.kind=='payment',FinanceProposal.status=='approved')) if not reversed_target(db,r.proposal_id)),Decimal('0'))
                if paid+Decimal(data['amount'])>invoice.amount:raise HTTPException(422,'payment exceeds approved invoice balance')

def create_candidate(db,user,payload):
    account=get_row(db,BudgetAccount,payload.account_id);year=get_row(db,FinanceYear,payload.year_id)
    if account.year_id!=year.year_id or payload.currency!=year.currency or payload.amount!=payload.amount.quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'candidate account/year/currency/precision mismatch')
    source_document(db,user,payload.document_id)
    if payload.contract_case_id:need(db,user,'contract.read');get_row(db,ContractCase,payload.contract_case_id)
    if payload.counterparty_id:need(db,user,'contract.read');get_row(db,ContractCounterparty,payload.counterparty_id)
    row=FinanceCandidate(**payload.model_dump());db.add(row);db.flush();audit(db,user,'candidate.create',row);return row

def candidate_action(db,user,key,payload,action):
    initial=get_row(db,FinanceCandidate,key)
    lock_sources(db,user,[initial.contract_case_id] if initial.contract_case_id else [],[initial.document_id],[initial.counterparty_id] if initial.counterparty_id else [])
    row=get_row(db,FinanceCandidate,key,True);check_version(row,payload.expected_version)
    if row.status=='cancelled' or (action=='review' and row.status=='reviewed'):raise HTTPException(409,'reviewed/cancelled candidate immutable')
    doc=source_document(db,user,row.document_id,True)
    snapshot={'document':doc,'amount':money(row.amount),'currency':row.currency}
    if row.contract_case_id:snapshot['contract']=contract_dict(db,user,get_row(db,ContractCase,row.contract_case_id,True),True)
    if row.counterparty_id:
        need(db,user,'contract.read');snapshot['counterparty']=row_dict(get_row(db,ContractCounterparty,row.counterparty_id,True))
    bump(db,row,row.version,{'status':'reviewed' if action=='review' else 'cancelled','review_snapshot':row.review_snapshot if action=='cancel' else snapshot,'reviewed_by':user.user_id,'reason':payload.reason});audit(db,user,'candidate.'+action,row);return row

def create_amendment(db,user,key,payload):
    need(db,user,'contract.read','contract.update');contract=get_row(db,ContractCase,key,True);check_version(contract,payload.expected_contract_version)
    if contract.status!='approved':raise HTTPException(422,'amendment requires approved contract')
    source_document(db,user,payload.document_id)
    row=FinanceContractAmendment(contract_case_id=key,**payload.model_dump());db.add(row);db.flush();audit(db,user,'amendment.create',row);return row

def amendment_action(db,user,key,payload,action):
    need(db,user,'contract.read')
    initial=get_row(db,FinanceContractAmendment,key);lock_sources(db,user,[initial.contract_case_id],[initial.document_id]);contract=get_row(db,ContractCase,initial.contract_case_id,True)
    row=get_row(db,FinanceContractAmendment,key,True);check_version(row,payload.expected_version)
    if row.status=='approved' or (action=='approve' and row.status!='reviewed'):raise HTTPException(409,'reviewed amendment required')
    check_version(contract,row.expected_contract_version)
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==contract.contract_case_id))
    if profile:
        year=get_row(db,FinanceYear,profile.year_id)
        if row.amount!=row.amount.quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'amendment exceeds configured currency precision')
    snapshot={'contract':contract_dict(db,user,contract,True),'document':source_document(db,user,row.document_id,True),'amount':money(row.amount),'end_date':scalar(row.end_date)}
    if action=='review':bump(db,row,row.version,{'status':'reviewed','review_snapshot':snapshot,'reviewed_by':user.user_id,'reason':payload.reason})
    else:
        need(db,user,'contract.approve')
        if row.review_snapshot!=snapshot:raise HTTPException(409,'contract amendment evidence changed')
        committed=sum((r.amount for r in db.scalars(select(FinanceProposal).where(FinanceProposal.contract_case_id==contract.contract_case_id,FinanceProposal.kind=='commitment',FinanceProposal.status=='approved')) if not reversed_target(db,r.proposal_id)),Decimal('0'))
        if committed>row.amount:raise HTTPException(422,'amendment below existing approved commitments; compensate commitments first')
        before=row_dict(contract);bump(db,contract,contract.version,{'amount':row.amount,'end_date':row.end_date,'approved_by':user.user_id,'approved_at':now_utc()});verify_contract_money(db,contract,row.amount);record_contract_change(db,user,contract,before,payload.reason)
        db.add(ContractDocument(contract_case_id=contract.contract_case_id,document_id=row.document_id,document_role='approved_amendment')) if not db.scalar(select(ContractDocument).where(ContractDocument.contract_case_id==contract.contract_case_id,ContractDocument.document_id==row.document_id)) else None
        bump(db,row,row.version,{'status':'approved','approved_by':user.user_id,'approved_at':now_utc(),'reason':payload.reason})
    audit(db,user,'amendment.'+action,row);return row

import csv, json
from datetime import timedelta, timezone
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4
from pydantic import ValidationError
from .finance_models import FinanceImportPreview, FinanceRenderedForm
from .finance_schemas import YearInput, AccountInput, CandidateInput, ContractInput
from .settings import settings
from .storage import store_upload
from .models import FormTemplate
from .official_form_renderer import render_template, TemplateRenderError
from fastapi import UploadFile

IMPORTS={'years':(YearInput,create_year),'accounts':(AccountInput,create_account),'contracts':(ContractInput,create_contract),'candidates':(CandidateInput,create_candidate),'proposals':(ProposalInput,create_proposal)}
EXPORTS={'years':FinanceYear,'accounts':BudgetAccount,'contracts':ContractCase,'counterparties':ContractCounterparty,'candidates':FinanceCandidate,'proposals':FinanceProposal,'journal':FinanceJournal,'amendments':FinanceContractAmendment}

def import_need(db,user,dataset):
    if dataset not in IMPORTS:raise HTTPException(422,'unknown finance import dataset')
    need(db,user,'finance.read','finance.import','finance.admin' if dataset=='years' else 'finance.create','document.read','document.create')
    if dataset=='contracts':need(db,user,'contract.read','contract.create')
def headers(dataset):
    if dataset not in IMPORTS:raise HTTPException(422,'unknown finance import dataset')
    return ['schema_version',*IMPORTS[dataset][0].model_fields]

def safe_cell(value):
    text='' if value is None else str(value)
    return "'"+text if text.startswith("'") or text.lstrip().startswith(('=','+','-','@','\t','\r')) else text

def tabular(columns,rows,format):
    if format=='csv':
        out=StringIO();writer=csv.DictWriter(out,fieldnames=columns,extrasaction='ignore');writer.writeheader()
        writer.writerows({k:safe_cell(r.get(k)) for k in columns} for r in rows)
        return out.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if format=='xlsx':
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.append(columns)
        for row in rows:sheet.append([safe_cell(row.get(k)) for k in columns])
        out=BytesIO();book.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    raise HTTPException(422,'format must be csv or xlsx')

def read_import(raw):
    if len(raw)>8*1024*1024:raise HTTPException(422,'finance import exceeds 8 MiB')
    try:
        if raw.startswith(b'PK'):
            from openpyxl import load_workbook
            import zipfile
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                if sum(f.file_size for f in archive.infolist())>32*1024*1024:raise ValueError('expanded workbook too large')
            sheet=load_workbook(BytesIO(raw),read_only=True,data_only=False).active
            rows=list(sheet.iter_rows(values_only=True))
            if not rows:raise ValueError('empty workbook')
            names=[str(c or '') for c in rows[0]]
            if 'amount' in names and any(isinstance(row[names.index('amount')],float) for row in rows[1:] if len(row)>names.index('amount')):raise ValueError('Excel money must be text decimal or exact integer; floating cells are not financial authority')
            for column in ('code','contract_no'):
                if column in names and any(row[names.index(column)] is not None and not isinstance(row[names.index(column)],str) for row in rows[1:] if len(row)>names.index(column)):raise ValueError('Excel codes/numbers must be text; displayed leading-zero formats cannot establish source codes')
            values=[dict(zip(names,['' if v is None else v.date().isoformat() if isinstance(v,datetime) and v.time()==datetime.min.time() else str(v) for v in row])) for row in rows[1:] if any(v is not None for v in row)]
        else:
            reader=csv.DictReader(StringIO(raw.decode('utf-8-sig')));names=reader.fieldnames or [];values=list(reader)
        if len(names)!=len(set(names)) or 'schema_version' not in names or not values or len(values)>1000:raise ValueError('unique headers and 1–1000 rows required')
        if any(None in r or any(v is None for v in r.values()) for r in values):raise ValueError('row column count mismatch')
        return values
    except Exception as exc:raise HTTPException(422,'invalid finance CSV/Excel: '+str(exc)) from None

READONLY={'reason','level_label','hierarchy_code','schema_version','text_encoding','source_document_id','file_sha256','import_preview_id','proposal_id','candidate_id','account_id','year_id','contract_case_id','counterparty_id','journal_id','amendment_id','event_id','approved_evidence','version','created_at','updated_at','status','review_snapshot','reviewed_by','reviewed_at','approved_by','approved_at','created_by','profile','counterparty_name','counterparty','documents'}

def import_payload(dataset,raw):
    schema=IMPORTS[dataset][0]
    if raw.get('schema_version')!=SCHEMA_VERSION:raise HTTPException(422,'schema_version finance-v1 required')
    unknown=set(raw)-set(schema.model_fields)-READONLY
    if unknown:raise HTTPException(422,'unknown finance columns: '+','.join(sorted(unknown)))
    values={k:v for k,v in raw.items() if k in schema.model_fields and v!=''}
    # Exports carry safe-text marker; remove exactly the escape added by this exporter.
    if raw.get('text_encoding')=='apostrophe-v2':
        values={k:(v[1:] if isinstance(v,str) and v.startswith("'") else v) for k,v in values.items()}
    elif raw.get('text_encoding')=='apostrophe-v1':
        values={k:(v[1:] if isinstance(v,str) and v.startswith("'") and v[1:].lstrip().startswith(('=','+','-','@','\t','\r')) else v) for k,v in values.items()}
    if 'account_levels' in values and isinstance(values['account_levels'],str):
        try:values['account_levels']=json.loads(values['account_levels'])
        except ValueError:raise HTTPException(422,'account_levels requires JSON array of code/label objects') from None
    try:return schema.model_validate(values)
    except ValidationError as exc:raise HTTPException(422,str(exc)) from None

class PreviewRollback(Exception):pass

def preview_import(db,user,dataset,upload,document_id=None):
    import_need(db,user,dataset);raw=upload.file.read(8*1024*1024+1);digest=sha256(raw).hexdigest();rows=read_import(raw)
    if db.scalar(select(FinanceImportPreview).where(FinanceImportPreview.dataset==dataset,FinanceImportPreview.file_sha256==digest)):raise HTTPException(409,'file already previewed/imported; duplicate import blocked')
    # Validate the real writes within a savepoint, then discard every row/audit mutation.
    samples=[]
    try:
        with db.begin_nested():
            for number,data in enumerate(rows,2):
                payload=import_payload(dataset,data);row=IMPORTS[dataset][1](db,user,payload);samples.append(row_dict(row))
            raise PreviewRollback()
    except PreviewRollback:pass
    if document_id:
        doc=get_row(db,Document,document_id);source_document(db,user,document_id)
        if doc.sha256!=digest:raise HTTPException(422,'original Document hash must equal uploaded import bytes')
    else:
        path,stored_hash,size=store_upload(UploadFile(filename=upload.filename,file=BytesIO(raw)))
        doc=Document(storage_path=path,original_filename=upload.filename or 'finance-import',sha256=stored_hash,size_bytes=size,document_type='finance_import_original',created_by=user.user_id);db.add(doc);db.flush()
    row=FinanceImportPreview(dataset=dataset,file_sha256=digest,source_document_id=doc.document_id,row_data=rows,created_by=user.user_id,expires_at=now_utc()+timedelta(hours=1));db.add(row);db.flush();audit(db,user,'import.preview',row)
    return row, samples[:10]

def confirm_import(db,user,key,payload):
    row=get_row(db,FinanceImportPreview,key,True);import_need(db,user,row.dataset);check_version(row,payload.expected_version)
    if row.created_by!=user.user_id:raise HTTPException(403,'import preview must be confirmed by its creating Human')
    expiry=row.expires_at.replace(tzinfo=timezone.utc) if row.expires_at.tzinfo is None else row.expires_at
    if row.status!='preview' or row.file_sha256!=payload.file_sha256 or expiry<now_utc():raise HTTPException(409,'preview consumed, expired or file hash changed')
    payloads=[import_payload(row.dataset,data) for data in row.row_data]
    contract_ids={p.contract_case_id for p in payloads if getattr(p,'contract_case_id',None)}
    party_ids={p.counterparty_id for p in payloads if getattr(p,'counterparty_id',None)}
    doc_ids={row.source_document_id,*[p.document_id for p in payloads if getattr(p,'document_id',None)]}
    lock_sources(db,user,contract_ids,doc_ids,party_ids)
    original=source_document(db,user,row.source_document_id,True)
    if original['sha256']!=row.file_sha256:raise HTTPException(409,'source original hash changed')
    # All rows and all audits commit together; a later failure closes/rolls back session.
    ids=[]
    for validated in payloads:
        entity=IMPORTS[row.dataset][1](db,user,validated);pk=list(entity.__table__.primary_key.columns)[0].name;ids.append(getattr(entity,pk))
    bump(db,row,row.version,{'status':'applied','applied_ids':ids});audit(db,user,'import.confirm',row)
    return {'inserted':len(ids),'applied_ids':ids,'source_document_id':row.source_document_id,'file_sha256':row.file_sha256}

def export_rows(db,user,dataset,year_id=None):
    if dataset not in EXPORTS:raise HTTPException(422,'unknown finance export dataset')
    need(db,user,'finance.read','finance.export')
    if dataset in ('contracts','counterparties','amendments','procurement-events'):need(db,user,'contract.read')
    model=EXPORTS[dataset];values=[]
    previews=list(db.scalars(select(FinanceImportPreview).where(FinanceImportPreview.status=='applied')))
    for row in db.scalars(select(model).order_by(list(model.__table__.primary_key.columns)[0])):
        if year_id:
            related_year=getattr(row,'year_id',None)
            if isinstance(row,(FinanceProposal,FinanceJournal)):related_year=db.get(BudgetAccount,row.account_id).year_id
            if isinstance(row,(ContractCase,ProcurementEvent)):
                profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==row.contract_case_id));related_year=profile.year_id if profile else None
            if related_year!=year_id:continue
        data=visible(db,user,row)
        if isinstance(row,ContractCase):
            data=contract_dict(db,user,row);profile=data.pop('profile')
            if profile:data.update(year_id=profile['year_id'],renewal_on=profile['renewal_on'])
            docs=data.pop('documents',[])
            if docs:data['document_id']=docs[0]['document_id']
        pk=list(row.__table__.primary_key.columns)[0].name
        if 'document.read' in permission_codes(db,user.user_id):
            preview=next((p for p in previews if p.dataset==dataset and getattr(row,pk) in p.applied_ids),None)
            if preview:data.update(source_document_id=preview.source_document_id,file_sha256=preview.file_sha256,import_preview_id=preview.preview_id)
        values.append({'schema_version':SCHEMA_VERSION,'text_encoding':'apostrophe-v2',**{k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in data.items()}})
    write_audit(db,user_id=user.user_id,action='finance.export',entity_type='finance_export',after={'dataset':dataset,'year_id':year_id,'count':len(values)})
    return values

def finance_summary(db,year_id):
    totals={k:Decimal('0') for k in ['allocated','reserved','spent','available']};accounts=[]
    for account in db.scalars(select(BudgetAccount).where(BudgetAccount.year_id==year_id,BudgetAccount.level==len(get_row(db,FinanceYear,year_id).account_levels)).order_by(BudgetAccount.code)):
        b=account_balance(db,account.account_id);accounts.append({**row_dict(account),**b})
        for k in totals:totals[k]+=Decimal(b[k])
    return {'year_id':year_id,**{k:money(v) for k,v in totals.items()},'accounts':accounts}

def render_proposal(db,user,key,payload):
    need(db,user,'finance.read','finance.export','document.read','document.create','template.read')
    row=get_row(db,FinanceProposal,key)
    if row.status!='approved':raise HTTPException(409,'formal output requires approved journal')
    source=row if not row.reverses_id else get_row(db,FinanceProposal,row.reverses_id)
    if row.contract_case_id or source.contract_case_id:need(db,user,'contract.read')
    source_document(db,user,source.document_id)
    template=get_row(db,FormTemplate,payload.form_template_id)
    if template.module_code not in ('budget','procurement','contracts') or template.status!='active' or template.modification_policy!='fill_only':raise HTTPException(422,'active original finance fill-only template required')
    doc=get_row(db,Document,template.document_id);root=Path(settings.storage_root).resolve();source=(root/doc.storage_path).resolve()
    if root not in source.parents or not source.is_file() or sha256(source.read_bytes()).hexdigest()!=doc.sha256:raise HTTPException(409,'original template hash/path cannot be verified')
    folder=root/'derived'/'finance';folder.mkdir(parents=True,exist_ok=True);destination=folder/(str(uuid4())+source.suffix.lower())
    values={**visible(db,user,row),'amount':money(row.amount),'account_code':get_row(db,BudgetAccount,row.account_id).code,'fiscal_year':get_row(db,FinanceYear,get_row(db,BudgetAccount,row.account_id).year_id).fiscal_year}
    try:manifest=render_template(source,destination,template.field_mapping,values)
    except TemplateRenderError as exc:raise HTTPException(422,str(exc)) from None
    raw=destination.read_bytes();output=Document(storage_path=str(destination.relative_to(root)),original_filename='finance-'+destination.name,sha256=sha256(raw).hexdigest(),size_bytes=len(raw),document_type='finance_rendered_original',created_by=user.user_id);db.add(output);db.flush()
    manifest.update(template_sha256=doc.sha256,template_document_id=doc.document_id,proposal_id=row.proposal_id,proposal_version=row.version,source_snapshot=visible(db,user,row)['review_snapshot'],output_sha256=output.sha256)
    rendered=FinanceRenderedForm(proposal_id=row.proposal_id,form_template_id=template.form_template_id,document_id=output.document_id,manifest=manifest,created_by=user.user_id);db.add(rendered);db.flush();audit(db,user,'template.render',rendered);return rendered

def patch_contract(db,user,key,payload):
    need(db,user,'contract.read','contract.update');row=get_row(db,ContractCase,key,True);check_version(row,payload.expected_version)
    if row.status=='approved':raise HTTPException(409,'approved contract source requires reviewed amendment')
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if 'amount' in values and values['amount'] is None:raise HTTPException(422,'contract amount cannot be cleared')
    start=values.get('start_date',row.start_date);end=values.get('end_date',row.end_date)
    if start and end and start>end:raise HTTPException(422,'contract dates out of order')
    validate_contract_policy(db,row,values.get('amount',row.amount))
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==key))
    if profile and values.get('amount') is not None:
        year=get_row(db,FinanceYear,profile.year_id)
        if values['amount']!=values['amount'].quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'amount exceeds configured precision')
    before=row_dict(row);bump(db,row,row.version,values)
    if 'amount' in values:verify_contract_money(db,row,values['amount'])
    record_contract_change(db,user,row,before,'Exact finance draft amendment');audit(db,user,'contract.update',row,before);return row

def account_dict(db,row):
    year=get_row(db,FinanceYear,row.year_id);policy=year.account_levels[row.level-1]
    return {**row_dict(row),'hierarchy_code':policy['code'],'level_label':policy['label']}

def patch_candidate(db,user,key,payload):
    row=get_row(db,FinanceCandidate,key,True);check_version(row,payload.expected_version)
    if row.status!='candidate':raise HTTPException(409,'reviewed candidate values immutable; create new evidence-linked candidate')
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if any(v is None for v in values.values()):raise HTTPException(422,'candidate fields cannot be cleared')
    year=get_row(db,FinanceYear,row.year_id)
    if values.get('amount') is not None and values['amount']!=values['amount'].quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'candidate exceeds configured precision')
    if values.get('document_id'):source_document(db,user,values['document_id'])
    before=row_dict(row);bump(db,row,row.version,values);audit(db,user,'candidate.update',row,before);return row

from .finance_models import ProcurementEvent
from .finance_schemas import EventInput
IMPORTS['procurement-events']=(EventInput,lambda db,user,payload:create_event(db,user,payload))
EXPORTS['procurement-events']=ProcurementEvent

def validate_event(db,user,data,contract):
    need(db,user,'contract.read');source_document(db,user,data['document_id'])
    if contract.status!='approved' or contract.currency!=data['currency']:raise HTTPException(422,'approved matching common contract required')
    profile=db.scalar(select(ProcurementProfile).where(ProcurementProfile.contract_case_id==contract.contract_case_id))
    if not profile:raise HTTPException(422,'contract fiscal configuration required')
    year=get_row(db,FinanceYear,profile.year_id)
    if Decimal(data['amount'])!=Decimal(data['amount']).quantize(Decimal(10)**-year.decimal_places):raise HTTPException(422,'procurement event exceeds fiscal precision')
    if data.get('related_event_id'):
        related=get_row(db,ProcurementEvent,data['related_event_id'],True)
        allowed={'inspection':'delivery','invoice':'inspection'}
        if related.status!='approved' or related.contract_case_id!=contract.contract_case_id or related.kind!=allowed.get(data['kind']):raise HTTPException(422,'related stage must be approved preceding delivery/inspection for same contract')

def create_event(db,user,payload):
    contract=get_row(db,ContractCase,payload.contract_case_id,True);validate_event(db,user,payload.model_dump(),contract)
    row=ProcurementEvent(**payload.model_dump());db.add(row);db.flush();audit(db,user,'procurement.create',row);return row

def patch_event(db,user,key,payload):
    initial=get_row(db,ProcurementEvent,key);contract=get_row(db,ContractCase,initial.contract_case_id,True);row=get_row(db,ProcurementEvent,key,True);check_version(row,payload.expected_version)
    if row.status!='draft':raise HTTPException(409,'reviewed/approved procurement values immutable')
    values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
    if any(v is None for v in values.values()):raise HTTPException(422,'event fields cannot be cleared')
    validate_event(db,user,{**row_dict(row),**values},contract);before=row_dict(row);bump(db,row,row.version,values);audit(db,user,'procurement.update',row,before);return row

def event_action(db,user,key,payload,action):
    initial=get_row(db,ProcurementEvent,key);lock_sources(db,user,[initial.contract_case_id],[initial.document_id]);contract=get_row(db,ContractCase,initial.contract_case_id,True);row=get_row(db,ProcurementEvent,key,True);check_version(row,payload.expected_version)
    if row.status in ('approved','cancelled') or (action=='approve' and row.status!='reviewed'):raise HTTPException(409,'reviewed procurement event required; approved history immutable')
    if action=='cancel':
        bump(db,row,row.version,{'status':'cancelled','reason':payload.reason});audit(db,user,'procurement.cancel',row);return row
    validate_event(db,user,row_dict(row),contract)
    snapshot={'contract':contract_dict(db,user,contract,True),'document':source_document(db,user,row.document_id,True),'amount':money(row.amount),'occurred_on':scalar(row.occurred_on),'kind':row.kind,'description':row.description}
    if row.related_event_id:snapshot['related_event']=visible(db,user,get_row(db,ProcurementEvent,row.related_event_id,True))
    if action=='review':bump(db,row,row.version,{'status':'reviewed','review_snapshot':snapshot,'reviewed_by':user.user_id,'reason':payload.reason})
    else:
        if row.review_snapshot!=snapshot:raise HTTPException(409,'procurement/vendor evidence changed since review; review again')
        before=row_dict(contract)
        if not db.scalar(select(ContractDocument).where(ContractDocument.contract_case_id==contract.contract_case_id,ContractDocument.document_id==row.document_id)):db.add(ContractDocument(contract_case_id=contract.contract_case_id,document_id=row.document_id,document_role=row.kind))
        bump(db,contract,contract.version);record_contract_change(db,user,contract,before,payload.reason)
        bump(db,row,row.version,{'status':'approved','approved_by':user.user_id,'approved_at':now_utc(),'reason':payload.reason})
    audit(db,user,'procurement.'+action,row);return row

def extract_finance_document(db,user,key):
    need(db,user,'finance.read','document.read');doc=get_row(db,Document,key);source=source_document(db,user,key)
    from .document_intake import extract_document
    try:text,method,pages,evidence=extract_document(doc)
    except (ValueError,FileNotFoundError,RuntimeError):raise HTTPException(422,'original cannot be extracted with configured common document tools') from None
    # Extraction tools can report local paths; expose only source identity and bounded text.
    return {'authority':'candidate','source':source,'text':text[:200000],'method':method,'page_count':pages,'truncated':len(text)>200000}

def contract_support(db,user,key):
    need(db,user,'finance.read','contract.read','document.read');row=get_row(db,ContractCase,key);source=contract_dict(db,user,row)
    roles={d['document_role'] for d in source['documents']}
    required=['specification','decision','contract_source']
    similar=select(ContractCase).where(ContractCase.contract_case_id!=key)
    if row.counterparty_id:similar=similar.where(ContractCase.counterparty_id==row.counterparty_id)
    elif row.contract_method:similar=similar.where(ContractCase.contract_method==row.contract_method)
    else:similar=similar.where(ContractCase.title==row.title)
    candidates=[contract_dict(db,user,r) for r in db.scalars(similar.order_by(ContractCase.updated_at.desc()).limit(20))]
    return {'authority':'candidate','source':source,'missing_document_candidates':[r for r in required if r not in roles],'similar_cases':candidates,'draft':f'契約件名: {row.title}\n契約先: {source.get("counterparty_name", "未設定")}\n契約額: {money(row.amount) if row.amount is not None else "未設定"} {row.currency}\n契約方法: {row.contract_method or "未設定"}\nHuman確認事項: 仕様・決定根拠・契約条件を原本と照合してください。'}

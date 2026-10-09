"""Live typed sources; snapshots prove values and authorization before derived text leaves the server."""
from datetime import date,datetime
from decimal import Decimal,localcontext,InvalidOperation
from hashlib import sha256
from pathlib import Path
from uuid import uuid4
from io import BytesIO,StringIO
import csv,json,re,unicodedata,zipfile
from fastapi import HTTPException,UploadFile
from sqlalchemy import select,update
from sqlalchemy.exc import IntegrityError
from .authz import permission_codes
from .audit import write_audit
from .db import Base
from .models import Document,Employee,EmergencyCase,Facility,ContractCase,LegalSourceDocument,FireInvestigationCase,FormTemplate,now_utc
from .finance_models import FinanceProposal,ProcurementEvent
from .operations_models import Incident,Vehicle
from .assets_models import OperationalAsset
from .workforce_models import WorkforceRosterEntry
from .inquiries_models import Inquiry,InquiryEvidence,InquiryCandidate,InquiryRenderedForm
from .inquiries_schemas import InquiryInput,InquiryPatch,Claim
from .finance_service import get_row,check_version,bump,row_dict,tabular
from .settings import settings
from .unified_search import lexical_score
SCHEMA_VERSION='inquiries-v1'
# Whitelist typed records and explicit fields, never accept client-authored numeric sources.
SOURCES={
 'document':(Document,'document.read',('original_filename',)),
 'personnel':(Employee,'personnel.read',('display_name','employee_code','organization_unit','title','active')),
 'emergency':(EmergencyCase,'emergency.case.read',('source_case_key','dispatch_number','call_date','incident_type','incident_address','activity_type')),
 'facility':(Facility,'facility.read',('name','address','status')),
 'contract':(ContractCase,'contract.read',('title','contract_no','amount','currency','status')),
 'finance':(FinanceProposal,'finance.read',('kind','amount','currency','reason','status')),
 'incident':(Incident,'incident.read',('title','kind','number','occurred_at','notes','status')),
 'vehicle':(Vehicle,'fleet.read',('code','name','odometer','fuel_stock','notes','active')),
 'asset':(OperationalAsset,'asset.read',('code','name','reorder_threshold','unit','notes','active')),
 'workforce':(WorkforceRosterEntry,'workforce.read',('work_date','starts_at','ends_at','payable_minutes','support_placement','note','status')),
 'legal':(LegalSourceDocument,'legal_source.read',('title',)),
 'fire':(FireInvestigationCase,'fire_investigation.read',('title','location_text','status')),
}
WRITTEN_NUMBERS=re.compile(r'\b(?:zero|one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|thirteen|fourteen|fifteen|sixteen|seventeen|eighteen|nineteen|twenty|thirty|forty|fifty|sixty|seventy|eighty|ninety|hundred|thousand|million|billion|trillion|half|quarter)\b|[〇零一二三四五六七八九十百千万億兆壱弐参拾佰阡]+(?=[人名件台円回個本棟戸歳年月日時分秒]|[。、,. ]|$)',re.I)
NUMBER=re.compile(r'(?<![\d.])[-+]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?(?![\d]|\.\d)')
def normalized(s):return unicodedata.normalize('NFKC',s).replace('−','-')
def numeric_tokens(text):
 # UUID identifiers in exported originals are not numerical quantities. Mask
 # their complete structural literals so hex fragments are not scientific values.
 value=normalized(text)
 value=re.sub(r'\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}\b',lambda m:' '*len(m.group()),value)
 def mask_digest(match):
  token=match.group();quantity=NUMBER.match(token)
  # Even hex-looking text may be a scientific quantity plus a unit (1e...f).
  # Preserve that complete supported lexical form for the bounded Decimal gate.
  if quantity and re.fullmatch(r'[A-Za-z%]+(?:[23])?',token[quantity.end():]):return token
  return ' '*len(token)
 value=re.sub(r'\b(?=[0-9a-fA-F]{0,63}[a-dfA-DF])[0-9a-fA-F]{64}\b',mask_digest,value)
 # An e-only 64-character literal may be a quantity and is never masked.
 out=[];consumed_until=0
 # Unsupported fractions/localized separators/signs must never be split into
 # independently proven integers. The same lexical rules apply to live sources.
 if '⁄' in value or re.search(r'\d\s*[/\\]\s*[-+]?(?:\d|\.\d)|[‐‑‒–—―]\s*(?:\d|\.\d)|\d\s*[٫٬·]\s*\d|\d[.,]{2,}\d|(?<!\d)\.{2,}\d|\d[ \t]+\d{3}(?!\d)',value):raise HTTPException(422,'unsupported numerical notation; use exact decimal literals and declared formulas')
 for span in re.finditer(r'(?<![\w.])\d+(?:[.,]\d+)+',value):
  if NUMBER.fullmatch(span.group()) is None:raise HTTPException(422,'ambiguous decimal/grouping notation; use exact decimal literals')
 for m in NUMBER.finditer(value):
  if m.start()<consumed_until:continue
  token=m.group();tail=value[m.end():];unit=re.match(r'[ \t]*([A-Za-z%]+(?:[23])?|[人名件台円回個本棟戸歳年月日時分秒]+)',tail)
  consumed_until=m.end()+(unit.end() if unit else 0)
  # Scientific notation is accepted only as an exact Decimal, without float conversion.
  try:decimal=Decimal(token.replace(',',''))
  except InvalidOperation:raise HTTPException(422,'source number is not an exact finite decimal') from None
  if not decimal.is_finite() or len(decimal.as_tuple().digits)>60 or abs(decimal.adjusted())>60:raise HTTPException(422,'source number exceeds bounded exact decimal limits')
  out.append({'text':token,'value':format(decimal,'f'),'unit':unit.group(1) if unit else ''})
 return out

def need(db,user,*codes):
 if not set(codes)<=permission_codes(db,user.user_id):raise HTTPException(403,'required inquiry/source permission missing')
def audit_metadata(data):
 # audit.read is deliberately independent of inquiry/source rights. Never copy
 # prose, exact values, query context, adapter labels or reasons into its payload.
 digest=lambda value:sha256(json.dumps(value,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()).hexdigest()
 metadata={'schema_version':'inquiry-audit-v1','change_sha256':digest(data)}
 for key in ('inquiry_id','evidence_id','candidate_id','rendered_id','document_id','form_template_id','revision_of','source_id','source_type','status','version','deleted','created_by','reviewed_by','approved_by','reviewed_at','approved_at','retrieved_at','generated_at','created_at','updated_at'):
  if key in data:metadata[key]=data[key]
 if 'claims' in data:metadata['claim_count']=len(data['claims'])
 sources={key:data[key] for key in ('snapshot','input_provenance','review_snapshot','manifest','provenance') if key in data}
 if sources:metadata['source_sha256']=digest(sources)
 return metadata
def audit(db,user,action,row,before=None):
 key=list(row.__table__.primary_key.columns)[0].name
 write_audit(db,user_id=user.user_id,action='inquiry.'+action,entity_type=row.__tablename__,entity_id=getattr(row,key),before=audit_metadata(before) if before is not None else None,after=audit_metadata(row_dict(row)),ai_used=isinstance(row,InquiryCandidate) and row.model!='deterministic-evidence-extract',ai_model_version='sha256:'+sha256(row.model_version.encode()).hexdigest() if isinstance(row,InquiryCandidate) else None)
def save(db):
 try:db.commit()
 except IntegrityError:db.rollback();raise HTTPException(409,'conflicting inquiry record') from None

def _document_local_permissions(db,doc):
 """Only local edges; graph closure below follows derived originals iteratively."""
 required={'document.read'}
 # This exact upload type owns its boundary even before it is linked to a record.
 if doc.document_type=='hazardous_evidence':required.add('hazardous.read')
 if doc.building_id:required.add('facility.read')
 prefixes={'emergency_patients':'emergency.patient.read','emergency_crews':'emergency.crew.read','emergency_':'emergency.case.read','employee_':'personnel.read','personnel_':'personnel.read','human_role_rules':'personnel.read','temporary_role_grants':'personnel.read','workforce_leave_entries':'personnel.read','workforce_':'workforce.read','finance_':'finance.read','contract_':'contract.read','fire_':'fire_investigation.read','operation_':'incident.read','asset_':'asset.read','operational_assets':'asset.read','facility_':'facility.read','legal_':'legal_source.read','drawing_':'drawing.read'}
 for table in Base.metadata.tables.values():
  if table.name.startswith('inquiry_'):continue
  columns=[c for c in table.columns if any(f.target_fullname=='documents.document_id' for f in c.foreign_keys)]
  if not columns:continue
  for c in columns:
   if db.scalar(select(c).where(c==doc.document_id).limit(1)) is not None:
    for prefix,permission in prefixes.items():
     if table.name.startswith(prefix):required.add(permission)
 return required

def permission_closure(db,nodes):
 """Visit every permission-bearing edge once, including cyclic import ownership.

 A revisit avoids recursion only: local rights and every outward edge of the
 first visit are processed before marking it visited. No cycle discards rights.
 """
 required=set();pending=list(nodes);visited=set()
 while pending:
  node=pending.pop()
  if node in visited:continue
  tag,*identity=node
  if tag=='document':
   doc=db.get(Document,identity[0])
   if doc is not None:
    required.update(_document_local_permissions(db,doc))
    for inquiry in db.scalars(select(Inquiry)):
     if inquiry.provenance.get('source_document_id')==doc.document_id:pending.append(('inquiry',inquiry.inquiry_id))
    for rendered in db.scalars(select(InquiryRenderedForm).where(InquiryRenderedForm.document_id==doc.document_id)):
     pending.append(('inquiry',rendered.inquiry_id))
     required.update(rendered.manifest.get('template_required_permissions',[]))
     if rendered.manifest.get('template_document_id'):pending.append(('document',rendered.manifest['template_document_id']))
     template=db.get(FormTemplate,rendered.form_template_id)
     if template:pending.append(('document',template.document_id))
  elif tag=='inquiry':
   required.add('inquiry.read');row=db.get(Inquiry,identity[0])
   if row is not None:
    required.update(row.provenance.get('required_permissions',[]))
    if row.provenance.get('source_document_id'):pending.append(('document',row.provenance['source_document_id']))
    links=[*row.provenance.get('security_sources',[]),*[{'source_type':e.source_type,'source_id':e.source_id,**e.snapshot} for e in evidence_rows(db,row)]]
    for link in links:
     required.update(link.get('required_permissions',[]));pending.append(('source',link['source_type'],link['source_id']))
     pending.extend(('document',d['document_id']) for d in link.get('documents',[]))
  elif tag=='finance_event':
   required.update({'finance.read','contract.read'});row=db.get(ProcurementEvent,identity[0])
   if row is not None:
    pending.append(('document',row.document_id));pending.append(('source','contract',row.contract_case_id))
    if row.related_event_id:pending.append(('finance_event',row.related_event_id))
  elif tag=='source':
   kind,key=identity
   if kind not in SOURCES:raise HTTPException(403,'source authorization unavailable')
   model,permission,_=SOURCES[kind];required.add(permission);row=db.get(model,key)
   if row is not None:
    if kind=='document':pending.append(('document',key))
    else:
     for c in row.__table__.columns:
      if any(f.target_fullname=='documents.document_id' for f in c.foreign_keys) and getattr(row,c.name):pending.append(('document',getattr(row,c.name)))
     if kind=='finance':
      if row.contract_case_id:pending.append(('source','contract',row.contract_case_id))
      for key in (row.commitment_id,row.reverses_id):
       if key:pending.append(('source','finance',key))
      if row.invoice_id:pending.append(('finance_event',row.invoice_id))
     if kind=='contract':
      from .models import ContractDocument
      pending.extend(('document',did) for did in db.scalars(select(ContractDocument.document_id).where(ContractDocument.contract_case_id==key)))
     if kind=='incident':
      if row.emergency_case_id:pending.append(('source','emergency',row.emergency_case_id))
      if row.fire_investigation_case_id:pending.append(('source','fire',row.fire_investigation_case_id))
     if kind=='legal':
      from .models import LegalSourceDocumentVersion
      revision=db.scalar(select(LegalSourceDocumentVersion).where(LegalSourceDocumentVersion.legal_source_document_id==key).order_by(LegalSourceDocumentVersion.retrieved_at.desc(),LegalSourceDocumentVersion.legal_source_document_version_id).limit(1))
      if revision and revision.raw_document_id:pending.append(('document',revision.raw_document_id))
  else:raise HTTPException(403,'source authorization unavailable')
  visited.add(node)
 return required

def document_permissions(db,doc):return {'document.read',*permission_closure(db,[('document',doc.document_id)])}

def guard_document(db,user,doc):
 """Narrow common hook for protected typed originals; generic originals stay shared."""
 if doc.document_type in ('inquiry_import_original','inquiry_rendered_original','hazardous_evidence'):need(db,user,*document_permissions(db,doc))
 return doc

def document_path(doc):
 root=Path(settings.storage_root).resolve();path=(root/doc.storage_path).resolve()
 if root not in path.parents or not path.is_file() or sha256(path.read_bytes()).hexdigest()!=doc.sha256:raise HTTPException(409,'source original missing or SHA changed')
 return path

def source(db,user,kind,key,query=None,lock=False):
 if kind not in SOURCES:raise HTTPException(422,'unknown typed source')
 model,permission,fields=SOURCES[kind];need(db,user,permission)
 row=get_row(db,model,key,lock);required={permission};docs=set()
 if getattr(row,'deleted',False) or getattr(row,'status','') in ('cancelled','deleted') or getattr(row,'active',True) is False:raise HTTPException(409,'source is inactive or deleted')
 if kind=='document':
  guard_document(db,user,row)
  required|=document_permissions(db,row);need(db,user,*required);path=document_path(row)
  from .document_intake import extract_document
  try:text,method,pages,extra=extract_document(row)
  except (ValueError,RuntimeError,OSError):raise HTTPException(422,'common original extraction unavailable') from None
  if len(text)>200000:raise HTTPException(422,'source text exceeds evidence retrieval limit')
  docs.add(row.document_id)
 else:
  values={field:str(getattr(row,field)) for field in fields if getattr(row,field,None) is not None}
  text='\n'.join(f'{k}: {v}' for k,v in values.items())
  if kind=='incident':
   for attr,child in [('emergency_case_id','emergency'),('fire_investigation_case_id','fire')]:
    if getattr(row,attr,None):
     linked=source(db,user,child,getattr(row,attr),lock=lock);required.update(linked['required_permissions']);text+='\n'+linked['text'];docs.update(d['document_id'] for d in linked['documents'])
  for column in row.__table__.columns:
   if any(f.target_fullname=='documents.document_id' for f in column.foreign_keys) and getattr(row,column.name):docs.add(getattr(row,column.name))
  if kind=='contract':
   from .models import ContractDocument
   docs.update(db.scalars(select(ContractDocument.document_id).where(ContractDocument.contract_case_id==key)))
 source_revision=None
 if kind=='legal':
  from .models import LegalSourceDocumentVersion
  statement=select(LegalSourceDocumentVersion).where(LegalSourceDocumentVersion.legal_source_document_id==key).order_by(LegalSourceDocumentVersion.retrieved_at.desc(),LegalSourceDocumentVersion.legal_source_document_version_id).limit(1)
  if lock:statement=statement.with_for_update().execution_options(populate_existing=True)
  revision=db.scalar(statement)
  if revision:
   source_revision={'version_id':revision.legal_source_document_version_id,'version_label':revision.version_label,'sha256':revision.sha256,'source_date':str(revision.source_current_date or revision.effective_from or revision.retrieved_at)}
   text+='\n'+revision.normalized_text
   if revision.raw_document_id:docs.add(revision.raw_document_id)
 documents=[]
 for document_id in sorted(docs):
  doc=get_row(db,Document,document_id,lock);required|=document_permissions(db,doc);need(db,user,*required);document_path(doc)
  documents.append({'document_id':document_id,'sha256':doc.sha256})
 # Exact value strings are derived from server-selected fields/text, not supplied by the caller.
 values=numeric_tokens(text)
 if kind in ('contract','finance'):
  amount=getattr(row,'amount',None)
  if amount is not None:values.append({'text':str(amount),'value':format(Decimal(amount),'f'),'unit':row.currency})
 if kind=='vehicle':
  for field,unit in [('odometer','km'),('fuel_stock','L')]:values.append({'text':str(getattr(row,field)),'value':format(getattr(row,field),'f'),'unit':unit})
 if kind=='workforce':values.append({'text':str(row.payable_minutes),'value':str(row.payable_minutes),'unit':'minutes'})
 if kind=='asset':values.append({'text':str(row.reorder_threshold),'value':format(row.reorder_threshold,'f'),'unit':row.unit})
 return {'source_type':kind,'source_id':key,'record_version':getattr(row,'version',None),'source_revision':source_revision,'source_date':str(getattr(row,'updated_at',None) or getattr(row,'created_at',None) or getattr(row,'call_date',None) or ''),'documents':documents,'text':text,'content_sha256':sha256(text.encode()).hexdigest(),'values':values,'required_permissions':sorted(required),'query_parameters':query or {},'navigation':{'surface':kind,'source_id':key,'href':f'/documents/{key}/download' if kind=='document' else None}}

def evidence_rows(db,row):return list(db.scalars(select(InquiryEvidence).where(InquiryEvidence.inquiry_id==row.inquiry_id).order_by(InquiryEvidence.evidence_id)))
def source_access(db,user,kind,key):need(db,user,*permission_closure(db,[('source',kind,key)]))

def authorized(db,user,row):
 need(db,user,*permission_closure(db,[('inquiry',row.inquiry_id)]))
 return row

def visible(db,user,row):
 authorized(db,user,row)
 return {**row_dict(row),'evidence':[row_dict(e) for e in evidence_rows(db,row)]}
def get_inquiry(db,user,key,lock=False):
 row=get_row(db,Inquiry,key,lock);authorized(db,user,row)
 if row.deleted:raise HTTPException(404,'inquiry not found')
 return row

def list_rows(db,user,q='',year=None):
 stmt=select(Inquiry).where(Inquiry.deleted.is_(False))
 if year is not None:stmt=stmt.where(Inquiry.year==year)
 out=[]
 for row in db.scalars(stmt.order_by(Inquiry.updated_at.desc(),Inquiry.inquiry_id)):
  try:authorized(db,user,row)
  except HTTPException as exc:
   if exc.status_code==403:continue
   raise
  if not q or normalized(q).lower() in normalized(row.question+' '+row.draft).lower():out.append(row)
 return out

def create(db,user,payload,provenance=None):
 need(db,user,'inquiry.read','inquiry.create');row=Inquiry(**payload.model_dump(),created_by=user.user_id,provenance=provenance or {'method':'manual'});db.add(row);db.flush();audit(db,user,'create',row);return row

def validate_evidence(db,user,row,lock=False):
 snapshot={}
 for e in evidence_rows(db,row):
  try:live=source(db,user,e.source_type,e.source_id,e.query_parameters,lock)
  except HTTPException as exc:
   if exc.status_code==404:raise HTTPException(409,'evidence source missing; cannot confirm answer') from None
   raise
  if live!=e.snapshot or e.excerpt not in live['text']:raise HTTPException(409,'evidence source changed; replace evidence and review again')
  snapshot[e.evidence_id]=live
 return snapshot

def validate_claims(draft,claims,snapshots):
 if WRITTEN_NUMBERS.search(normalized(draft)) or '⁄' in normalized(draft):raise HTTPException(422,'numerical assertions must use exact digit decimal notation with evidence claims')
 def proven(op):
  snap=snapshots.get(op['evidence_id'])
  if not snap:raise HTTPException(422,'claim needs inquiry-bound evidence')
  if not any(Decimal(v['value'])==Decimal(op['value']) and v['unit']==op.get('unit','') for v in snap['values']):raise HTTPException(422,'claim value/unit does not equal retrieved source')
  return Decimal(op['value'])
 supported=[]
 with localcontext() as ctx:
  ctx.prec=200
  for c in claims:
   text_tokens=numeric_tokens(c['text'])
   if len(text_tokens)!=1 or normalized(c['text'])!=text_tokens[0]['text'] or Decimal(text_tokens[0]['value'])!=Decimal(c['value']):raise HTTPException(422,'claim text must be one exact numerical token equal to claim value')
   if c.get('formula','identity')=='identity':
    if c.get('operands'):raise HTTPException(422,'identity claim has no formula operands')
    proven(c)
   else:
    ops=c.get('operands',[])
    if not ops:raise HTTPException(422,'declared formula requires source-bound operands')
    if c['evidence_id'] not in snapshots:raise HTTPException(422,'formula anchor evidence missing')
    numbers=[proven(op) for op in ops];formula=c['formula']
    if formula in ('sum','difference') and any(op['unit']!=c['unit'] for op in ops):raise HTTPException(422,'formula units must agree')
    if formula=='sum':result=sum(numbers,Decimal(0))
    elif formula=='difference' and len(numbers)==2:result=numbers[0]-numbers[1]
    elif formula=='product' and len(numbers)==2 and sum(bool(op['unit']) for op in ops)<=1 and c['unit']==next((op['unit'] for op in ops if op['unit']),''):result=numbers[0]*numbers[1]
    elif formula=='quotient' and len(numbers)==2 and numbers[1]!=0 and ops[0]['unit']==ops[1]['unit'] and not c['unit']:result=numbers[0]/numbers[1]
    else:raise HTTPException(422,'unsupported formula operands/units')
    if result!=Decimal(c['value']):raise HTTPException(422,'exact declared formula result mismatch')
    if formula=='quotient' and result*numbers[1]!=numbers[0]:raise HTTPException(422,'non-terminating exact division requires another declared representation')
   supported.append(c)
 for token in numeric_tokens(draft):
  if not any(normalized(c['text'])==token['text'] and Decimal(c['value'])==Decimal(token['value']) and c['unit']==token['unit'] for c in supported):raise HTTPException(422,'unsupported numerical claim in answer prose')
 for c in supported:
  if not any(normalized(c['text'])==t['text'] and c['unit']==t['unit'] for t in numeric_tokens(draft)):raise HTTPException(422,'claim must occur in answer prose with its unit')

def mutable(row):
 if row.status=='approved':raise HTTPException(409,'approved answer immutable; create a revision')
def patch(db,user,key,payload):
 need(db,user,'inquiry.update');row=get_inquiry(db,user,key,True);check_version(row,payload.expected_version);mutable(row)
 values=payload.model_dump(exclude={'expected_version'},exclude_unset=True)
 if any(v is None for v in values.values()):raise HTTPException(422,'inquiry fields cannot be null')
 claims=values.get('claims',row.claims);draft=values.get('draft',row.draft)
 snapshots=validate_evidence(db,user,row,True);validate_claims(draft,claims,snapshots)
 before=row_dict(row);values.update(status='draft',reviewed_by=None,reviewed_at=None,review_snapshot={});bump(db,row,row.version,values);audit(db,user,'update',row,before);return row

def add_evidence(db,user,key,payload):
 need(db,user,'inquiry.update');row=get_inquiry(db,user,key,True);check_version(row,payload.expected_version);mutable(row)
 if set(payload.query_parameters)-{'query','field'}:raise HTTPException(422,'only recorded query/field retrieval parameters supported')
 if any(not isinstance(v,str) or len(v)>500 for v in payload.query_parameters.values()):raise HTTPException(422,'bounded string query parameters required')
 snapshot=source(db,user,payload.source_type,payload.source_id,payload.query_parameters,True)
 if payload.excerpt not in snapshot['text']:raise HTTPException(422,'quote must be verbatim from retrieved source')
 if payload.query_parameters.get('query') and normalized(payload.query_parameters['query']).lower() not in normalized(snapshot['text']).lower():raise HTTPException(422,'query does not match selected source')
 if payload.query_parameters.get('field') and (payload.source_type=='document' or payload.query_parameters['field'] not in SOURCES[payload.source_type][2]):raise HTTPException(422,'field not exposed by typed adapter')
 docs=snapshot['documents'];e=InquiryEvidence(inquiry_id=key,source_type=payload.source_type,source_id=payload.source_id,document_id=docs[0]['document_id'] if docs else None,excerpt=payload.excerpt,snapshot=snapshot,query_parameters=payload.query_parameters,created_by=user.user_id);db.add(e);db.flush();bump(db,row,row.version,{'status':'draft','reviewed_by':None,'reviewed_at':None,'review_snapshot':{},'provenance':{**row.provenance,'security_sources':[*row.provenance.get('security_sources',[]),{'source_type':e.source_type,'source_id':e.source_id,'required_permissions':snapshot['required_permissions'],'documents':snapshot['documents']}]}});audit(db,user,'evidence.add',e);return e

def remove_evidence(db,user,key,evidence_id,payload):
 need(db,user,'inquiry.update');row=get_inquiry(db,user,key,True);check_version(row,payload.expected_version);mutable(row);e=get_row(db,InquiryEvidence,evidence_id,True)
 if e.inquiry_id!=key:raise HTTPException(404,'evidence not found')
 if any(c['evidence_id']==evidence_id or any(o['evidence_id']==evidence_id for o in c.get('operands',[])) for c in row.claims):raise HTTPException(409,'remove dependent claims before removing evidence')
 audit(db,user,'evidence.remove',e);db.delete(e);bump(db,row,row.version,{'status':'draft','reviewed_by':None,'reviewed_at':None,'review_snapshot':{}})

def action(db,user,key,payload,name):
 need(db,user,'inquiry.'+name);row=get_inquiry(db,user,key,True);check_version(row,payload.expected_version);mutable(row)
 if not row.draft or not evidence_rows(db,row):raise HTTPException(422,'answer and live evidence required for Human review')
 snapshots=validate_evidence(db,user,row,True);validate_claims(row.draft,row.claims,snapshots)
 snap={'draft':row.draft,'claims':row.claims,'evidence':snapshots,'year':row.year,'question':row.question}
 if name=='review':values={'status':'reviewed','review_snapshot':snap,'reviewed_by':user.user_id,'reviewed_at':now_utc()}
 else:
  if row.status!='reviewed' or not row.reviewed_by or row.review_snapshot!=snap:raise HTTPException(409,'explicit Human review of unchanged evidence required')
  values={'status':'approved','approved_by':user.user_id,'approved_at':now_utc()}
 bump(db,row,row.version,values);audit(db,user,name,row,{'reason':payload.reason});return row

def revision(db,user,key,payload):
 need(db,user,'inquiry.create');old=get_inquiry(db,user,key,True);check_version(old,payload.expected_version)
 if old.status!='approved':raise HTTPException(409,'revision requires approved parent')
 row=Inquiry(year=old.year,question=old.question,draft=old.draft,claims=[],revision_of=key,created_by=user.user_id,provenance={**old.provenance,'method':'revision','parent_version':old.version,'reason':payload.reason});db.add(row);db.flush();mapping={}
 for e in evidence_rows(db,old):
  copy=InquiryEvidence(inquiry_id=row.inquiry_id,source_type=e.source_type,source_id=e.source_id,document_id=e.document_id,excerpt=e.excerpt,snapshot=e.snapshot,query_parameters=e.query_parameters,created_by=user.user_id);db.add(copy);db.flush();mapping[e.evidence_id]=copy.evidence_id
 claims=json.loads(json.dumps(old.claims))
 for c in claims:
  c['evidence_id']=mapping[c['evidence_id']]
  for op in c.get('operands',[]):op['evidence_id']=mapping[op['evidence_id']]
 row.claims=claims;audit(db,user,'revision.create',row);return row

# Optional in-process local-model adapter. Deployment supplies the callable; unavailable AI never blocks CRUD.
local_model_adapter=None

def generate(db,user,key,payload):
 # Snapshot reads and optional model work do not hold the account/inquiry/source write locks.
 need(db,user,'inquiry.create','inquiry.update')
 row=get_inquiry(db,user,key);check_version(row,payload.expected_version);mutable(row)
 snapshots=validate_evidence(db,user,row)
 if not snapshots:raise HTTPException(422,'select evidence before AI Draft')
 inputs={'inquiry_version':row.version,'question':row.question,'year':row.year,'sources':snapshots}
 excerpts=[row_dict(e) for e in evidence_rows(db,row)]
 db.commit()
 if local_model_adapter is None:
  result={'model':'deterministic-evidence-extract','model_version':'1','draft':'\n'.join(e['excerpt'] for e in excerpts),'confidence':None,'claims':[]}
  for token in numeric_tokens(result['draft']):
   e=next((e for e in excerpts if any(Decimal(v['value'])==Decimal(token['value']) and v['unit']==token['unit'] for v in e['snapshot']['values'])),None)
   if e:result['claims'].append({**token,'evidence_id':e['evidence_id'],'formula':'identity','operands':[]})
 else:
  try:result=local_model_adapter(json.loads(json.dumps(inputs)))
  except Exception:raise HTTPException(503,'local model unavailable; manual evidence-backed drafting remains usable') from None
 try:
  if not isinstance(result,dict) or set(result)-{'draft','claims','model','model_version','confidence'} or not isinstance(result.get('draft'),str) or len(result['draft'])>50000 or not all(isinstance(result.get(k),str) and 0<len(result[k])<=200 for k in ('model','model_version')):raise ValueError()
  claims=[Claim.model_validate(c).model_dump() for c in result.get('claims',[])]
  if len(claims)>500:raise ValueError()
  confidence=result.get('confidence')
  if confidence is not None and (not isinstance(confidence,str) or len(confidence)>50 or not Decimal(confidence).is_finite() or not Decimal(0)<=Decimal(confidence)<=Decimal(1)):raise ValueError()
 except (ValueError,InvalidOperation,TypeError):raise HTTPException(422,'invalid local candidate/provenance') from None
 # Final short Human/session gate protects the candidate write after potentially slow model work.
 from .authz import require_mutation_permission
 user=require_mutation_permission('inquiry.create')(user=user,db=db)
 need(db,user,'inquiry.update')
 current=get_inquiry(db,user,key,True);check_version(current,inputs['inquiry_version']);mutable(current)
 if validate_evidence(db,user,current,True)!=snapshots:raise HTTPException(409,'candidate source inputs changed during generation')
 rowc=InquiryCandidate(inquiry_id=key,draft=result['draft'],claims=claims,model=result['model'],model_version=result['model_version'],confidence=confidence,input_provenance=inputs,created_by=user.user_id)
 db.add(rowc);db.flush();audit(db,user,'candidate.generate',rowc);return rowc

def adopt(db,user,key,payload):
 need(db,user,'inquiry.update');row=get_inquiry(db,user,key,True);check_version(row,payload.expected_version);mutable(row);c=get_row(db,InquiryCandidate,payload.candidate_id)
 if c.inquiry_id!=key or c.input_provenance.get('inquiry_version')!=row.version:raise HTTPException(409,'candidate input version changed')
 snapshots=validate_evidence(db,user,row,True)
 if c.input_provenance.get('sources')!=snapshots:raise HTTPException(409,'candidate source inputs changed')
 validate_claims(c.draft,c.claims,snapshots);before=row_dict(row);bump(db,row,row.version,{'draft':c.draft,'claims':c.claims,'status':'draft','reviewed_by':None,'reviewed_at':None,'review_snapshot':{},'provenance':{**row.provenance,'candidate_id':c.candidate_id,'model':c.model,'model_version':c.model_version,'input_provenance':c.input_provenance,'confidence':c.confidence,'generated_at':c.generated_at.isoformat(),'adopted_by':user.user_id,'adopted_at':now_utc().isoformat(),'reason':payload.reason}});audit(db,user,'candidate.adopt',row,before);return row

HEADERS=['schema_version','year','question','draft','claims','evidence','history']
HISTORY_FIELDS={'inquiry_id','version','status','revision_of','provenance','created_at','updated_at','created_by','reviewed_by','reviewed_at','approved_by','approved_at'}
def export_rows(db,user,year=None):
 need(db,user,'inquiry.read','inquiry.export');rows=[]
 for row in list_rows(db,user,year=year):
  evidence=[{'source_type':e.source_type,'source_id':e.source_id,'query_parameters':e.query_parameters,'excerpt':e.excerpt,'snapshot':e.snapshot,'export_evidence_id':e.evidence_id} for e in evidence_rows(db,row)]
  rows.append({'schema_version':SCHEMA_VERSION,'year':row.year,'question':row.question,'draft':row.draft,'claims':json.dumps(row.claims,ensure_ascii=False),'evidence':json.dumps(evidence,ensure_ascii=False),'history':json.dumps({k:v for k,v in row_dict(row).items() if k in HISTORY_FIELDS},ensure_ascii=False)})
 return rows

def json_cell(value):
 def reject_constant(value):raise ValueError('non-finite JSON number forbidden')
 try:parsed=json.loads(value,parse_constant=reject_constant)
 except RecursionError:raise ValueError('JSON nesting exceeds limit') from None
 pending=[(parsed,0)];count=0
 while pending:
  item,depth=pending.pop();count+=1
  if depth>50 or count>100000:raise ValueError('JSON nesting/content exceeds limits')
  if isinstance(item,dict):pending.extend((v,depth+1) for v in item.values())
  elif isinstance(item,list):pending.extend((v,depth+1) for v in item)
 return parsed

def read_tabular(raw,filename):
 if len(raw)>8*1024*1024:raise HTTPException(422,'import exceeds 8 MiB')
 try:
  if filename.lower().endswith('.xlsx'):
   from openpyxl import load_workbook
   with zipfile.ZipFile(BytesIO(raw)) as z:
    if sum(f.file_size for f in z.infolist())>32*1024*1024:raise ValueError('expanded workbook too large')
   book=load_workbook(BytesIO(raw),read_only=True,data_only=False)
   if len(book.worksheets)!=1:raise ValueError('one sheet required')
   matrix=[]
   for cells in book.active.iter_rows():
    if any(c.data_type=='f' for c in cells):raise ValueError('formula cells forbidden')
    matrix.append(['' if c.value is None else str(c.value) for c in cells])
    if len(matrix)>1001:raise ValueError('at most 1000 rows')
   names=matrix[0];rows=[dict(zip(names,values)) for values in matrix[1:] if any(values)]
   if any(len(values)!=len(names) for values in matrix[1:]):raise ValueError('row width mismatch')
  elif filename.lower().endswith('.csv'):
   reader=csv.DictReader(StringIO(raw.decode('utf-8-sig')));names=reader.fieldnames or [];rows=list(reader)
  else:raise ValueError('CSV or XLSX required')
  if len(names)!=len(set(names)) or set(names)!=set(HEADERS) or not 1<=len(rows)<=1000 or any(None in r or any(v is None for v in r.values()) for r in rows):raise ValueError('exact schema and 1–1000 complete rows required')
  # Reversible export neutralization; content is never executed on import.
  decoded=[{k:v[1:] if v.startswith("'") else v for k,v in r.items()} for r in rows]
  for row in decoded:
   for field,empty in [('claims','[]'),('evidence','[]'),('history','{}')]:row[field+'_json']=json_cell(row[field] or empty)
  return decoded
 except Exception as exc:raise HTTPException(422,'invalid inquiries CSV/Excel: '+str(exc)) from None

def import_file(db,user,upload):
 need(db,user,'inquiry.read','inquiry.import','inquiry.create','inquiry.update','document.read','document.create')
 raw=upload.file.read(8*1024*1024+1);db.commit()
 rows=read_tabular(raw,upload.filename or '');digest=sha256(raw).hexdigest();created=[];restrictions=set()
 from .authz import require_mutation_permission
 user=require_mutation_permission('inquiry.import')(user=user,db=db)
 need(db,user,'inquiry.create','inquiry.update','document.read','document.create')
 # DB transaction is owned by request. Validate all rows before original registration or commit.
 for number,data in enumerate(rows,2):
  try:
   if data['schema_version']!=SCHEMA_VERSION:raise ValueError('schema_version inquiries-v1 required')
   p=InquiryInput(year=data['year'],question=data['question']);history=data['history_json']
   if not isinstance(history,dict) or set(history)-HISTORY_FIELDS or history.get('status','draft') not in ('draft','reviewed','approved') or not isinstance(history.get('provenance',{}),dict):raise ValueError('invalid non-authoritative exported history')
   from uuid import UUID
   for k in ('inquiry_id','revision_of','created_by','reviewed_by','approved_by'):
    if history.get(k):UUID(history[k])
   for k in ('created_at','updated_at','reviewed_at','approved_at'):
    if history.get(k):datetime.fromisoformat(history[k])
   if history.get('version') is not None and (not isinstance(history['version'],int) or history['version']<1):raise ValueError('invalid exported version')
   history_lineage=[]
   for item in history.get('provenance',{}).get('security_sources',[]):
    if not isinstance(item,dict) or set(item)!={'source_type','source_id','required_permissions','documents'}:raise ValueError('invalid historical source lineage')
    historical_source=source(db,user,item['source_type'],item['source_id'],lock=True)
    documents={d['document_id']:d for d in historical_source['documents']}
    if not isinstance(item['documents'],list) or len(item['documents'])>100:raise ValueError('invalid historical originals')
    for d in item['documents']:
     if not isinstance(d,dict) or set(d)!={'document_id','sha256'}:raise ValueError('invalid historical original identity')
     doc=get_row(db,Document,d['document_id']);permissions=document_permissions(db,doc);need(db,user,*permissions);restrictions.update(permissions);documents[doc.document_id]={'document_id':doc.document_id,'sha256':doc.sha256}
    history_lineage.append({'source_type':item['source_type'],'source_id':item['source_id'],'required_permissions':historical_source['required_permissions'],'documents':list(documents.values())});restrictions.update(historical_source['required_permissions'])
   claims=[Claim.model_validate(c).model_dump() for c in data['claims_json']];inputs=data['evidence_json']
   if not isinstance(inputs,list) or len(inputs)>100 or len(claims)>500 or len(data['draft'])>50000:raise ValueError('bounded evidence, claims and draft required')
   row=create(db,user,p,{'method':'import','file_sha256':digest,'row_number':number,'imported_history':history,'imported_history_authority':'unverified historical metadata; current answer remains draft','security_sources':history_lineage});mapping={};snapshots={}
   for item in inputs:
    if set(item)!={'source_type','source_id','query_parameters','excerpt','snapshot','export_evidence_id'}:raise ValueError('evidence schema mismatch')
    from .inquiries_schemas import EvidenceInput
    ep=EvidenceInput(expected_version=row.version,**{k:item[k] for k in ('source_type','source_id','query_parameters','excerpt')})
    e=add_evidence(db,user,row.inquiry_id,ep)
    if e.snapshot!=item['snapshot']:raise HTTPException(409,'imported source snapshot is stale')
    if item['export_evidence_id'] in mapping:raise ValueError('duplicate evidence ID')
    mapping[item['export_evidence_id']]=e.evidence_id;snapshots[e.evidence_id]=e.snapshot;restrictions.update(e.snapshot['required_permissions'])
   for claim in claims:
    claim['evidence_id']=mapping[claim['evidence_id']]
    for op in claim['operands']:op['evidence_id']=mapping[op['evidence_id']]
   validate_claims(data['draft'],claims,snapshots);row.draft=data['draft'];row.claims=claims;created.append(row)
  except HTTPException:raise
  except (ValueError,TypeError,KeyError):raise HTTPException(422,f'invalid inquiry import row {number}') from None
 from .storage import store_upload
 path,h,size=store_upload(UploadFile(filename=upload.filename,file=BytesIO(raw)));doc=Document(storage_path=path,original_filename=upload.filename,sha256=h,size_bytes=size,document_type='inquiry_import_original',created_by=user.user_id);db.add(doc);db.flush()
 for row in created:row.provenance={**row.provenance,'source_document_id':doc.document_id,'required_permissions':sorted(restrictions)};audit(db,user,'import',row)
 return {'inserted':len(created),'inquiry_ids':[r.inquiry_id for r in created],'source_document_id':doc.document_id,'file_sha256':digest}

def render(db,user,key,payload):
 need(db,user,'inquiry.export','document.read','document.create','template.read');row=get_inquiry(db,user,key,True)
 if row.status!='approved':raise HTTPException(409,'formal output requires Human approved answer')
 snapshots=validate_evidence(db,user,row,True);validate_claims(row.draft,row.claims,snapshots)
 if row.review_snapshot.get('evidence')!=snapshots:raise HTTPException(409,'approved answer evidence changed')
 template=get_row(db,FormTemplate,payload.form_template_id,True)
 if template.module_code!='inquiries' or template.status!='active' or template.modification_policy!='fill_only':raise HTTPException(422,'active original inquiries fill-only template required')
 doc=get_row(db,Document,template.document_id,True);template_permissions=sorted(document_permissions(db,doc));need(db,user,*template_permissions);path=document_path(doc);root=Path(settings.storage_root).resolve();folder=root/'derived'/'inquiries';folder.mkdir(parents=True,exist_ok=True);dest=folder/(str(uuid4())+path.suffix.lower())
 from .official_form_renderer import render_template,TemplateRenderError
 values={**row_dict(row),'answer':row.draft,'evidence_quotes':'\n'.join(e.excerpt for e in evidence_rows(db,row))}
 try:manifest=render_template(path,dest,template.field_mapping,values)
 except TemplateRenderError as exc:raise HTTPException(422,str(exc)) from None
 if sha256(path.read_bytes()).hexdigest()!=doc.sha256:
  dest.unlink(missing_ok=True);raise HTTPException(409,'original template changed while rendering')
 raw=dest.read_bytes();output=Document(storage_path=str(dest.relative_to(root)),original_filename='inquiry-'+dest.name,sha256=sha256(raw).hexdigest(),size_bytes=len(raw),document_type='inquiry_rendered_original',created_by=user.user_id);db.add(output);db.flush()
 manifest.update(template_document_id=doc.document_id,template_sha256=doc.sha256,template_required_permissions=template_permissions,inquiry_id=key,inquiry_version=row.version,source_snapshot=snapshots,output_sha256=output.sha256)
 record=InquiryRenderedForm(inquiry_id=key,form_template_id=template.form_template_id,document_id=output.document_id,manifest=manifest,created_by=user.user_id);db.add(record);db.flush();audit(db,user,'template.render',record);return record

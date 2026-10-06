from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter


RENDERER_VERSION = "official-form-renderer-v1"


class TemplateRenderError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _required_value(values: dict[str, Any], key: str, mapping: dict) -> Any:
    if key in values and values[key] is not None:
        return values[key]
    if mapping.get("required", True):
        raise TemplateRenderError(f"required template value is missing: {key}")
    return ""


def _docx_paragraphs(doc: DocxDocument):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested in cell.tables:
                    for nrow in nested.rows:
                        for ncell in nrow.cells:
                            for paragraph in ncell.paragraphs:
                                yield paragraph
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                yield paragraph
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            yield paragraph


def _replace_docx_placeholder(doc: DocxDocument, placeholder: str, value: str) -> int:
    replaced = 0
    for paragraph in _docx_paragraphs(doc):
        if placeholder not in paragraph.text:
            continue
        run_hits = [run for run in paragraph.runs if placeholder in run.text]
        if not run_hits:
            raise TemplateRenderError(
                f"DOCX placeholder spans multiple runs and cannot be safely replaced: {placeholder}"
            )
        for run in run_hits:
            count = run.text.count(placeholder)
            run.text = run.text.replace(placeholder, value)
            replaced += count
    return replaced


def render_docx(
    source: Path,
    destination: Path,
    field_mapping: dict,
    values: dict[str, Any],
) -> dict:
    doc = DocxDocument(source)
    applied: list[dict] = []
    for key, mapping in field_mapping.items():
        if not isinstance(mapping, dict):
            raise TemplateRenderError(f"invalid DOCX mapping for {key}")
        placeholder = mapping.get("placeholder")
        if not placeholder:
            raise TemplateRenderError(f"DOCX mapping requires placeholder: {key}")
        value = str(_required_value(values, key, mapping))
        count = _replace_docx_placeholder(doc, placeholder, value)
        if count == 0 and mapping.get("required", True):
            raise TemplateRenderError(f"DOCX placeholder not found: {placeholder}")
        applied.append({"key": key, "placeholder": placeholder, "occurrences": count})
    doc.save(destination)
    return {"format": "docx", "applied": applied}


def render_excel(
    source: Path,
    destination: Path,
    field_mapping: dict,
    values: dict[str, Any],
) -> dict:
    keep_vba = source.suffix.lower() == ".xlsm"
    wb = load_workbook(source, keep_vba=keep_vba, data_only=False)
    applied: list[dict] = []
    for key, mapping in field_mapping.items():
        if not isinstance(mapping, dict):
            raise TemplateRenderError(f"invalid spreadsheet mapping for {key}")
        sheet = mapping.get("sheet")
        cell = mapping.get("cell")
        if not sheet or not cell:
            raise TemplateRenderError(f"spreadsheet mapping requires sheet and cell: {key}")
        if sheet not in wb.sheetnames:
            raise TemplateRenderError(f"worksheet not found: {sheet}")
        value = _required_value(values, key, mapping)
        ws = wb[sheet]
        ws[cell] = value
        if isinstance(value, str):
            ws[cell].data_type = "s"
        applied.append({"key": key, "sheet": sheet, "cell": cell})
    wb.save(destination)
    return {
        "format": "xlsm" if keep_vba else "xlsx",
        "applied": applied,
        "keep_vba": keep_vba,
    }


def render_pdf_form(
    source: Path,
    destination: Path,
    field_mapping: dict,
    values: dict[str, Any],
) -> dict:
    reader = PdfReader(str(source))
    fields = reader.get_fields() or {}
    if not fields:
        raise TemplateRenderError("PDF has no AcroForm fields; coordinate guessing is not allowed")
    field_values: dict[str, str] = {}
    applied: list[dict] = []
    for key, mapping in field_mapping.items():
        if not isinstance(mapping, dict):
            raise TemplateRenderError(f"invalid PDF mapping for {key}")
        field = mapping.get("field")
        if not field:
            raise TemplateRenderError(f"PDF mapping requires field: {key}")
        if field not in fields:
            if mapping.get("required", True):
                raise TemplateRenderError(f"PDF form field not found: {field}")
            continue
        value = str(_required_value(values, key, mapping))
        field_values[field] = value
        applied.append({"key": key, "field": field})

    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    for page in writer.pages:
        writer.update_page_form_field_values(
            page,
            field_values,
            auto_regenerate=False,
        )
    with destination.open("wb") as f:
        writer.write(f)
    return {"format": "pdf", "applied": applied, "acroform": True}


def render_template(
    source: Path,
    destination: Path,
    field_mapping: dict,
    values: dict[str, Any],
) -> dict:
    if not field_mapping:
        raise TemplateRenderError("form template field_mapping is empty")
    suffix = source.suffix.lower()
    if suffix == ".docx":
        manifest = render_docx(source, destination, field_mapping, values)
    elif suffix in {".xlsx", ".xlsm"}:
        manifest = render_excel(source, destination, field_mapping, values)
    elif suffix == ".pdf":
        manifest = render_pdf_form(source, destination, field_mapping, values)
    else:
        raise TemplateRenderError(f"unsupported official template format: {suffix or '(none)'}")
    manifest["renderer_version"] = RENDERER_VERSION
    manifest["template_suffix"] = suffix
    return manifest


def flatten_values(value: Any, prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    if isinstance(value, dict):
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else str(key)
            out.update(flatten_values(child, child_prefix))
    elif isinstance(value, list):
        out[prefix] = "\n".join(str(x) for x in value)
    else:
        out[prefix] = value
    return out


"""Synthetic finance authority regressions; no operational data."""
import os
os.environ['FIRE_AI_DATABASE_URL']='sqlite+pysqlite:///:memory:'
from io import BytesIO
from uuid import uuid4
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from app.db import Base,engine,SessionLocal
from app.main import app
from app.models import Employee,User,Role,Permission,RolePermission,UserRole,Document,AuditLog
from app.security import hash_password
CODES=[f'finance.{p}' for p in ['read','create','update','review','approve','admin','import','export']]+['search.use','contract.read','contract.create','contract.update','contract.approve','document.read','document.create','template.read']
@pytest.fixture
def client():
 Base.metadata.drop_all(engine);Base.metadata.create_all(engine)
 with SessionLocal() as db:
  e=Employee(display_name='Synthetic finance');db.add(e);db.flush();u=User(username='finance',employee_id=e.employee_id,password_hash=hash_password('synthetic-password'));r=Role(code='finance-test',name='Synthetic');db.add_all([u,r]);db.flush();db.add(UserRole(user_id=u.user_id,role_id=r.role_id))
  for code in CODES:
   p=Permission(code=code);db.add(p);db.flush();db.add(RolePermission(role_id=r.role_id,permission_id=p.permission_id))
  db.commit()
 with TestClient(app) as c:
  assert c.post('/auth/login',json={'username':'finance','password':'synthetic-password'}).status_code==200
  yield c

def post(c,path,data,status=201):
 r=c.post('/finance'+path,json=data);assert r.status_code==status,r.text;return r.json()
def remove(code):
 with SessionLocal() as db:
  p=db.scalar(select(Permission).where(Permission.code==code))
  for r in db.scalars(select(RolePermission).where(RolePermission.permission_id==p.permission_id)):db.delete(r)
  db.commit()
def document(c):
 return c.post('/documents/upload',files={'file':('synthetic.txt',b'Synthetic financial evidence','text/plain')}).json()['document_id']
def year(c,code=2026):
 y=post(c,'/years',{'fiscal_year':code,'currency':'JPY','decimal_places':2,'reason':'Synthetic explicit policy'})
 return post(c,'/years/'+y['year_id']+'/approve',{'expected_version':1,'reason':'Human policy confirmed'},200)
def accounts(c,y,label='SYN'):
 parent=None
 for level in range(1,6):
  parent=post(c,'/accounts',{'year_id':y['year_id'],'parent_id':parent['account_id'] if parent else None,'level':level,'code':f'0{level}{label}','name':f'Synthetic {label} level {level}'})
 return parent

def test_policy_accounts_decimal_codes_version_permissions(client):
 y=year(client);a=accounts(client,y)
 assert a['code']=='05SYN' and a['level']==5
 assert client.get('/finance/accounts').headers['cache-control']=='no-store'
 assert client.patch('/finance/accounts/'+a['account_id'],json={'expected_version':1,'name':'Corrected'}).json()['version']==2
 assert client.patch('/finance/accounts/'+a['account_id'],json={'expected_version':1,'name':'Stale'}).status_code==409
 assert client.patch('/finance/accounts/'+a['account_id'],json={'expected_version':2,'parent_id':a['account_id']}).status_code==422
 assert client.post('/finance/accounts',json={'year_id':y['year_id'],'level':5,'code':'000','name':'Missing parent'}).status_code==422
 y2=year(client,2027)
 assert client.post('/finance/accounts',json={'year_id':y2['year_id'],'parent_id':a['account_id'],'level':1,'code':'000','name':'Cross year'}).status_code==422
 assert client.post('/finance/years',json={'fiscal_year':2028,'currency':'UNKNOWN','decimal_places':2,'reason':'Invalid'}).status_code==422
 remove('finance.update');assert client.patch('/finance/accounts/'+a['account_id'],json={'expected_version':2,'name':'Denied'}).status_code==403
 with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='finance.account.update'))

def proposal(c,a,d,kind='initial',amount='1000.00',**extra):
 return post(c,'/proposals',{'kind':kind,'account_id':a['account_id'],'amount':amount,'currency':'JPY','document_id':d,'reason':'Synthetic evidence-linked decision','idempotency_key':str(uuid4()),**extra})
def approve(c,p):
 r=post(c,'/proposals/'+p['proposal_id']+'/review',{'expected_version':p['version'],'reason':'Human checked exact values'},200)
 return post(c,'/proposals/'+p['proposal_id']+'/approve',{'expected_version':r['version'],'reason':'Human formal decision'},200)
def balance(c,a):return c.get('/finance/accounts/'+a['account_id']+'/balance').json()

def test_journal_human_gate_transfer_conservation_and_compensation(client):
 y=year(client);a=accounts(client,y);b=accounts(client,y,'OTHER');d=document(client)
 p=proposal(client,a,d)
 assert balance(client,a)['allocated']=='0.00'
 assert client.post('/finance/proposals/'+p['proposal_id']+'/approve',json={'expected_version':1,'reason':'Bypass'}).status_code==409
 p=approve(client,p);assert balance(client,a)['available']=='1000.00'
 assert client.patch('/finance/proposals/'+p['proposal_id'],json={'expected_version':p['version'],'amount':'1'}).status_code==409
 t=approve(client,proposal(client,a,d,'transfer','300.00',to_account_id=b['account_id']))
 assert balance(client,a)['available']=='700.00' and balance(client,b)['available']=='300.00'
 bad=proposal(client,a,d,'transfer','701.00',to_account_id=b['account_id'])
 r=post(client,'/proposals/'+bad['proposal_id']+'/review',{'expected_version':1,'reason':'Checked'},200)
 assert client.post('/finance/proposals/'+bad['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Overspend'}).status_code==422
 assert balance(client,b)['allocated']=='300.00'
 rev=approve(client,proposal(client,a,d,'reversal','300.00',reverses_id=t['proposal_id']))
 assert balance(client,a)['available']=='1000.00' and balance(client,b)['available']=='0.00'
 assert client.get('/finance/proposals/'+t['proposal_id']).json()['status']=='approved'
 assert client.post('/finance/proposals',json={'kind':'reversal','account_id':a['account_id'],'amount':'300.00','currency':'JPY','document_id':d,'reason':'Duplicate reversal','idempotency_key':'duplicate','reverses_id':t['proposal_id']}).status_code==409
 assert len(client.get('/finance/journal').json())==5
 with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='finance.proposal.approve'))

def test_financial_invalid_decimal_policy_and_stale_review(client):
 y=year(client);a=accounts(client,y);d=document(client)
 for value in ['NaN','Infinity','-1','1.001',0.1]:
  assert client.post('/finance/proposals',json={'kind':'initial','account_id':a['account_id'],'amount':value,'currency':'JPY','document_id':d,'reason':'Invalid','idempotency_key':str(uuid4())}).status_code==422
 p=proposal(client,a,d,'initial','0.00');assert approve(client,p)['amount']=='0.00'
 assert client.post('/finance/proposals',json={'kind':'initial','account_id':a['account_id'],'amount':'1','currency':'USD','document_id':d,'reason':'Currency','idempotency_key':'currency'}).status_code==422
 p=proposal(client,a,d,'amendment','-1.00');r=post(client,'/proposals/'+p['proposal_id']+'/review',{'expected_version':1,'reason':'Checked'},200)
 assert client.patch('/finance/accounts/'+a['account_id'],json={'expected_version':a['version']+1,'name':'Changed after review'}).status_code==200
 assert client.post('/finance/proposals/'+p['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Stale'}).status_code==409
 remove('document.read');assert client.post('/finance/proposals',json={'kind':'initial','account_id':a['account_id'],'amount':'1','currency':'JPY','document_id':d,'reason':'Denied','idempotency_key':'denied'}).status_code==403

def contract(c,y,d,amount='500.00'):
 cp=post(c,'/counterparties',{'name':'Synthetic supplier'})
 case=post(c,'/contracts',{'title':'Synthetic equipment contract','counterparty_id':cp['counterparty_id'],'amount':amount,'currency':'JPY','year_id':y['year_id'],'document_id':d,'start_date':'2026-04-01','end_date':'2027-03-31','renewal_on':'2027-02-01'})
 return post(c,'/contracts/'+case['contract_case_id']+'/approve',{'expected_version':case['version'],'reason':'Human source approval'},200)

def test_legacy_contract_status_and_approved_source_bypasses_closed(client):
 cp=client.post('/contracts/counterparties',json={'name':'Synthetic legacy supplier'}).json()
 case=client.post('/contracts',json={'title':'Synthetic legacy contract','counterparty_id':cp['counterparty_id'],'amount':100}).json();cid=case['contract_case_id']
 assert client.patch('/contracts/'+cid,json={'expected_version':1,'status':'approved'}).status_code==422
 approved=client.post('/contracts/'+cid+'/approve',json={'expected_version':1});assert approved.status_code==200
 assert client.patch('/contracts/'+cid,json={'expected_version':2,'amount':200}).status_code==409
 assert client.post('/contracts/'+cid+'/approve',json={'expected_version':2}).status_code==409
 assert client.get('/contracts/'+cid).json()['amount']==100

def test_common_contract_commitment_payment_caps_exact_once_and_reversal(client):
 y=year(client);a=accounts(client,y);d=document(client);approve(client,proposal(client,a,d));c=contract(client,y,d)
 assert client.get('/contracts/'+c['contract_case_id']).status_code==200
 assert client.get('/finance/contracts?q=Synthetic').json()[0]['amount']=='500.00'
 commit=proposal(client,a,d,'commitment','400.00',contract_case_id=c['contract_case_id']);assert balance(client,a)['reserved']=='0.00'
 commit=approve(client,commit);assert balance(client,a)['reserved']=='400.00'
 pay=approve(client,proposal(client,a,d,'payment','125.25',contract_case_id=c['contract_case_id'],commitment_id=commit['proposal_id']))
 assert balance(client,a)=={'account_id':a['account_id'],'allocated':'1000.00','reserved':'274.75','spent':'125.25','available':'600.00'}
 assert client.post('/finance/proposals/'+pay['proposal_id']+'/approve',json={'expected_version':pay['version'],'reason':'Replay'}).status_code==409
 assert client.post('/finance/proposals',json={'kind':'payment','account_id':a['account_id'],'amount':'1.00','currency':'JPY','document_id':d,'reason':'Replay','idempotency_key':pay['idempotency_key'],'commitment_id':commit['proposal_id'],'contract_case_id':c['contract_case_id']}).status_code==409
 too=proposal(client,a,d,'payment','300.00',contract_case_id=c['contract_case_id'],commitment_id=commit['proposal_id']);r=post(client,'/proposals/'+too['proposal_id']+'/review',{'expected_version':1,'reason':'Check'},200)
 assert client.post('/finance/proposals/'+r['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Overpay'}).status_code==422
 early=proposal(client,a,d,'reversal','400.00',reverses_id=commit['proposal_id']);r=post(client,'/proposals/'+early['proposal_id']+'/review',{'expected_version':1,'reason':'Check'},200)
 assert client.post('/finance/proposals/'+r['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Unpaid remains'}).status_code==422
 approve(client,proposal(client,a,d,'reversal','125.25',reverses_id=pay['proposal_id']));assert balance(client,a)['spent']=='0.00'
 # Reviewed reversal may be refreshed after the compensating payment transaction changes account version.
 early=client.get('/finance/proposals/'+early['proposal_id']).json();approve(client,early);assert balance(client,a)['reserved']=='0.00'
 assert client.get('/finance/proposals/'+pay['proposal_id']).json()['amount']=='125.25'

def test_quotes_requests_amendments_source_version_and_permissions(client):
 y=year(client);a=accounts(client,y);d=document(client);c=contract(client,y,d)
 for kind in ['quote','request','estimate']:
  candidate=post(client,'/candidates',{'kind':kind,'year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic next-year evidence','amount':'12.34','currency':'JPY','document_id':d,'contract_case_id':c['contract_case_id']})
  r=post(client,'/candidates/'+candidate['candidate_id']+'/review',{'expected_version':1,'reason':'Human verified actual evidence'},200)
  assert r['status']=='reviewed' and balance(client,a)['allocated']=='0.00'
 approve(client,proposal(client,a,d));commit=approve(client,proposal(client,a,d,'commitment','400.00',contract_case_id=c['contract_case_id']))
 pay=proposal(client,a,d,'payment','10.00',contract_case_id=c['contract_case_id'],commitment_id=commit['proposal_id']);r=post(client,'/proposals/'+pay['proposal_id']+'/review',{'expected_version':1,'reason':'Check'},200)
 amendment=post(client,'/contracts/'+c['contract_case_id']+'/amendments',{'expected_contract_version':c['version'],'amount':'600.00','end_date':'2027-04-01','document_id':d,'reason':'Synthetic authorized price amendment'})
 r2=post(client,'/amendments/'+amendment['amendment_id']+'/review',{'expected_version':1,'reason':'Human verified revision'},200)
 post(client,'/amendments/'+amendment['amendment_id']+'/approve',{'expected_version':r2['version'],'reason':'Human amended source'},200)
 assert client.post('/finance/proposals/'+pay['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Stale contract'}).status_code==409
 assert client.get('/finance/contracts/'+c['contract_case_id']).json()['amount']=='600.00'
 approve(client,client.get('/finance/proposals/'+pay['proposal_id']).json())
 assert len(client.get('/finance/contracts/'+c['contract_case_id']+'/history').json()['changes'])>=2
 remove('contract.read');assert client.get('/finance/contracts').status_code==403
 detail=client.get('/finance/proposals/'+commit['proposal_id']).json();assert 'contract_case_id' not in detail and 'contract' not in detail['review_snapshot']
 remove('document.read');assert 'document' not in client.get('/finance/proposals/'+commit['proposal_id']).json()['review_snapshot']

def test_largest_exact_money_repeated_fraction_transfer_and_roundtrip(client):
 y=year(client);a=accounts(client,y);b=accounts(client,y,'FRACTION');d=document(client)
 p=approve(client,proposal(client,a,d,'initial','9999999999999999.99'))
 assert p['amount']=='9999999999999999.99' and balance(client,a)['allocated']=='9999999999999999.99'
 for _ in range(3):approve(client,proposal(client,a,d,'transfer','0.01',to_account_id=b['account_id']))
 assert balance(client,a)['available']=='9999999999999999.96' and balance(client,b)['available']=='0.03'
 assert client.post('/finance/proposals',json={'kind':'initial','account_id':a['account_id'],'amount':'10000000000000000.00','currency':'JPY','document_id':d,'reason':'Overflow','idempotency_key':'overflow'}).status_code==422

def test_legacy_profile_approval_requires_source_read_and_reversal_privacy(client):
 y=year(client);a=accounts(client,y);d=document(client);approve(client,proposal(client,a,d));c=contract(client,y,d);commit=approve(client,proposal(client,a,d,'commitment','10.00',contract_case_id=c['contract_case_id']))
 draft=post(client,'/contracts',{'title':'Synthetic draft source','amount':'1.00','currency':'JPY','year_id':y['year_id'],'document_id':d})
 remove('document.read')
 assert client.post('/contracts/'+draft['contract_case_id']+'/approve',json={'expected_version':draft['version']}).status_code==403
 assert d not in str(client.get('/finance/proposals/'+commit['proposal_id']).json())

def test_csv_excel_import_export_real_roundtrip_provenance_atomic_permissions(client):
 y=year(client);a=accounts(client,y);d=document(client)
 data=f'schema_version,kind,account_id,amount,currency,document_id,reason,idempotency_key\nfinance-v1,initial,{a["account_id"]},12.34,JPY,{d},=Synthetic(),csv-1\n'.encode()
 p=client.post('/finance/import/proposals',files={'file':('synthetic.csv',data,'text/csv')});assert p.status_code==200,p.text;p=p.json()
 assert balance(client,a)['allocated']=='0.00' and client.get('/finance/proposals').json()==[]
 assert p['source_document_id'] and p['file_sha256']
 assert client.post('/finance/import-previews/'+p['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':'0'*64}).status_code==409
 r=post(client,'/import-previews/'+p['preview_id']+'/confirm',{'expected_version':1,'file_sha256':p['file_sha256']},200);assert r['inserted']==1
 assert client.post('/finance/import-previews/'+p['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':p['file_sha256']}).status_code==409
 out=client.get('/finance/export/proposals');assert "'=Synthetic()" in out.text and p['file_sha256'] in out.text and p['source_document_id'] in out.text
 from openpyxl import load_workbook,Workbook
 sheet=load_workbook(BytesIO(client.get('/finance/export/proposals?format=xlsx').content)).active;assert sheet.max_row==2
 # Export roundtrip makes a fresh draft, preserving exact zeros and formula-like text.
 import csv,io
 rows=list(csv.DictReader(io.StringIO(out.content.decode('utf-8-sig'))));rows[0]['idempotency_key']='roundtrip'
 raw=io.StringIO();w=csv.DictWriter(raw,fieldnames=rows[0]);w.writeheader();w.writerows(rows)
 p=client.post('/finance/import/proposals',files={'file':('roundtrip.csv',raw.getvalue().encode(),'text/csv')}).json();post(client,'/import-previews/'+p['preview_id']+'/confirm',{'expected_version':1,'file_sha256':p['file_sha256']},200)
 assert client.get('/finance/proposals').json()[0]['reason']=='=Synthetic()'
 # Confirmation revalidates all rows; a duplicate in row two rolls back row one.
 raw=data.replace(b'csv-1',b'first-new')+data.splitlines()[1].replace(b'csv-1',b'conflict')+b'\n'
 p=client.post('/finance/import/proposals',files={'file':('rollback.csv',raw,'text/csv')}).json()
 proposal(client,a,d,idempotency_key='conflict')
 assert client.post('/finance/import-previews/'+p['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':p['file_sha256']}).status_code==409
 assert 'first-new' not in [r['idempotency_key'] for r in client.get('/finance/proposals').json()]
 p=client.post('/finance/import/proposals',files={'file':('deny.csv',data.replace(b'csv-1',b'denied'),'text/csv')}).json()
 remove('finance.create');assert client.post('/finance/import-previews/'+p['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':p['file_sha256']}).status_code==403

def test_search_deadlines_fiscal_totals_and_registered_original_template(client):
 y=year(client);a=accounts(client,y);d=document(client);c=contract(client,y,d);p=approve(client,proposal(client,a,d))
 assert client.get('/search?q=Synthetic&modules=budget').status_code==200
 assert client.get('/finance/summary?year_id='+y['year_id']).json()['allocated']=='1000.00'
 assert client.get('/finance/alerts?as_of=2027-03-01&days=60').json()['contracts']
 from openpyxl import Workbook,load_workbook
 w=Workbook();w.active['A1']='Official synthetic original';raw=BytesIO();w.save(raw)
 doc=client.post('/documents/upload',files={'file':('synthetic-original.xlsx',raw.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json()
 t=client.post('/templates',json={'template_code':'SYN-FINANCE','name':'Synthetic official original','module_code':'budget','document_id':doc['document_id'],'version_label':'1','field_mapping':{'amount':{'sheet':'Sheet','cell':'A2'},'currency':{'sheet':'Sheet','cell':'A3'}},'print_settings':{},'modification_policy':'fill_only'})
 # Test fixture explicitly grants template registration only for this synthetic original.
 if t.status_code==403:
  with SessionLocal() as db:
   role=db.scalar(select(Role));perm=Permission(code='template.manage');db.add(perm);db.flush();db.add(RolePermission(role_id=role.role_id,permission_id=perm.permission_id));db.commit()
  t=client.post('/templates',json={'template_code':'SYN-FINANCE','name':'Synthetic official original','module_code':'budget','document_id':doc['document_id'],'version_label':'1','field_mapping':{'amount':{'sheet':'Sheet','cell':'A2'},'currency':{'sheet':'Sheet','cell':'A3'}},'print_settings':{},'modification_policy':'fill_only'})
 assert t.status_code==201,t.text
 r=post(client,'/proposals/'+p['proposal_id']+'/render',{'form_template_id':t.json()['form_template_id']},201)
 rendered=client.get('/documents/'+r['document_id']+'/download');sheet=load_workbook(BytesIO(rendered.content)).active
 assert sheet['A1'].value=='Official synthetic original' and sheet['A2'].value=='1000.00'
 assert client.get('/documents/'+doc['document_id']+'/download').content==raw.getvalue()
 remove('finance.read');assert client.get('/search?q=Synthetic&modules=budget').status_code==403

def test_bootstrap_registers_finance_in_fresh_process_and_shared_shell(tmp_path):
 import subprocess,sys,sqlite3
 from pathlib import Path
 root=Path(__file__).resolve().parents[2];env={**os.environ,'PYTHONPATH':str(root/'backend'),'FIRE_AI_DATABASE_URL':'sqlite+pysqlite:///'+str(tmp_path/'fresh.db'),'FIRE_AI_STORAGE_ROOT':str(tmp_path/'storage'),'FIRE_AI_PRODUCTION_MODE':'false'};env.pop('FIRE_AI_TENANT_ID',None)
 r=subprocess.run([sys.executable,'-m','app.bootstrap','--username','synthetic-finance','--display-name','Synthetic finance','--password','synthetic-password'],cwd=tmp_path,env=env,capture_output=True,text=True);assert r.returncode==0,r.stderr
 with sqlite3.connect(tmp_path/'fresh.db') as db:
  assert db.execute("SELECT count(*) FROM sqlite_master WHERE name='finance_journal'").fetchone()[0]==1
  assert db.execute("SELECT count(*) FROM permissions WHERE code='finance.approve'").fetchone()[0]==1
 html=(root/'frontend/index.html').read_text();assert 'id="financeBtn"' in html and '/ui/finance.js' in html and 'await initFinance()' in html

def test_common_contract_exact_maximum_and_draft_patch(client):
 y=year(client);d=document(client)
 c=post(client,'/contracts',{'title':'Synthetic maximum common contract','amount':'9999999999999999.99','currency':'JPY','year_id':y['year_id'],'document_id':d})
 assert c['amount']=='9999999999999999.99'
 assert client.get('/finance/contracts/'+c['contract_case_id']).json()['amount']=='9999999999999999.99'
 r=client.patch('/finance/contracts/'+c['contract_case_id'],json={'expected_version':1,'amount':'9999999999999999.98','title':'Synthetic corrected','end_date':'2027-03-30'});assert r.status_code==200,r.text;assert r.json()['amount']=='9999999999999999.98'
 assert client.patch('/finance/contracts/'+c['contract_case_id'],json={'expected_version':1,'amount':'1'}).status_code==409

def test_signed_amendment_reversal_cross_year_and_denied_gates(client):
 y=year(client);a=accounts(client,y);d=document(client);approve(client,proposal(client,a,d))
 negative=approve(client,proposal(client,a,d,'amendment','-0.01'));assert balance(client,a)['available']=='999.99'
 approve(client,proposal(client,a,d,'reversal','0.01',reverses_id=negative['proposal_id']));assert balance(client,a)['available']=='1000.00'
 y2=year(client,2027);b=accounts(client,y2)
 assert client.post('/finance/proposals',json={'kind':'transfer','account_id':a['account_id'],'to_account_id':b['account_id'],'amount':'1','currency':'JPY','document_id':d,'reason':'Invalid','idempotency_key':'cross-year'}).status_code==422
 c=contract(client,y,d,amount='50.00');commit=proposal(client,a,d,'commitment','60.00',contract_case_id=c['contract_case_id']);r=post(client,'/proposals/'+commit['proposal_id']+'/review',{'expected_version':1,'reason':'Check'},200)
 assert client.post('/finance/proposals/'+commit['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Over contract'}).status_code==422
 remove('finance.approve');assert client.post('/finance/proposals/'+commit['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Denied'}).status_code==403

def test_legacy_exact_strings_and_snapshot_change_detection(client):
 cp=client.post('/contracts/counterparties',json={'name':'Synthetic legacy exact'}).json()
 c=client.post('/contracts',json={'title':'Synthetic legacy exact maximum','amount':'9999999999999999.99','counterparty_id':cp['counterparty_id']}).json();cid=c['contract_case_id']
 assert client.get('/finance/contracts/'+cid).json()['amount']=='9999999999999999.99'
 assert client.patch('/contracts/'+cid,json={'expected_version':1,'amount':'9999999999999999.98'}).status_code==200
 assert client.get('/finance/contracts/'+cid).json()['amount']=='9999999999999999.98'
 for value in ['NaN','Infinity','-1','1.001','10000000000000000.00']:
  assert client.post('/contracts',json={'title':'Synthetic invalid','amount':value}).status_code==422
 # Source document hash mutation invalidates Human review rather than committing stale evidence.
 y=year(client);a=accounts(client,y);d=document(client);p=proposal(client,a,d);r=post(client,'/proposals/'+p['proposal_id']+'/review',{'expected_version':1,'reason':'Human checked'},200)
 with SessionLocal() as db:db.get(Document,d).sha256='a'*64;db.commit()
 assert client.post('/finance/proposals/'+p['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Stale evidence'}).status_code==409

def test_finance_migration_and_exact_live_ui_form_behavior(tmp_path):
 from pathlib import Path
 import subprocess
 root=Path(__file__).resolve().parents[2];migration=root/'db/migrations/048_run_b_finance.sql';assert migration.exists()
 from app.migrations import split_sql
 assert len(split_sql(migration.read_text()))>=15
 # Execute the real form builder and submission: 18 digits, 0 and leading zeros stay strings.
 script=tmp_path/'finance-form.cjs';script.write_text(r'''
const fs=require('fs'),vm=require('vm'),assert=require('assert');
const nodes=new Map(),get=id=>{if(!nodes.has(id))nodes.set(id,{innerHTML:'',textContent:'',value:'',classList:{toggle(){}}});return nodes.get(id)};
const ctx={console,$:get,esc:x=>String(x??''),api:async()=>({permissions:['finance.read','document.read']})};vm.createContext(ctx);vm.runInContext(fs.readFileSync(process.argv[2],'utf8'),ctx);
(async()=>{let submitted;ctx.financeForm('Synthetic',[['amount','Amount','money'],['code','Code'],['reason','Reason']],{amount:'0',code:'0001'},async data=>{submitted=data},()=>{});assert(get('financeContent').innerHTML.includes('value="0"'));get('financeField_amount').value='9999999999999999.99';get('financeField_code').value='0001';get('financeField_reason').value='Synthetic';get('financeForm').onsubmit({preventDefault(){}});await new Promise(r=>setTimeout(r,0));assert.strictEqual(submitted.amount,'9999999999999999.99');assert.strictEqual(submitted.code,'0001');get('financeField_amount').value='0';get('financeForm').onsubmit({preventDefault(){}});await new Promise(r=>setTimeout(r,0));assert.strictEqual(submitted.amount,'0');})().catch(e=>{console.error(e);process.exit(1)});
''')
 r=subprocess.run(['node',str(script),str(root/'frontend/finance.js')],capture_output=True,text=True);assert r.returncode==0,r.stderr

def test_finance_two_department_databases_same_ids_sessions_and_sources(tmp_path):
 from fastapi import FastAPI
 from sqlalchemy import create_engine
 from sqlalchemy.orm import sessionmaker
 from sqlalchemy.pool import StaticPool
 from app.settings import Settings
 from app.tenant import initialize_tenant,TenantBoundaryMiddleware
 from app.rbac_seed import seed_rbac
 from app.routers import auth,finance,search
 from app.finance_models import FinanceYear,BudgetAccount
 from app.models import now_utc
 clients=[];factories=[];ids=[str(uuid4()) for _ in range(7)]
 try:
  for slug in ['alpha','beta']:
   cfg=Settings(_env_file=None,database_url='sqlite+pysqlite:///:memory:',tenant_id=str(uuid4()),storage_root=str(tmp_path/slug),trusted_hosts=[slug+'.test']);eng=create_engine(cfg.database_url,connect_args={'check_same_thread':False},poolclass=StaticPool);Base.metadata.create_all(eng);initialize_tenant(eng,cfg,'Synthetic '+slug);factory=sessionmaker(bind=eng,expire_on_commit=False,autoflush=False);factories.append((factory,eng,cfg))
   with factory() as db:
    roles=seed_rbac(db);u=User(username='same-finance',password_hash=hash_password(slug+'-synthetic-password'));db.add(u);db.flush();db.add(UserRole(user_id=u.user_id,role_id=roles['system_admin'].role_id));db.add(Document(document_id=ids[0],original_filename=slug+'-private-evidence',storage_path='synthetic',sha256='a'*64));db.add(FinanceYear(year_id=ids[1],fiscal_year=2026,currency='JPY',decimal_places=2,status='approved',approved_by=u.user_id,approved_at=now_utc(),reason=slug+'-policy'));db.flush()
    for level in range(1,6):db.add(BudgetAccount(account_id=ids[level+1],year_id=ids[1],parent_id=ids[level] if level>1 else None,level=level,code=f'00{level}',name=slug+'-private-account'));db.flush()
    db.commit()
   application=FastAPI();application.add_middleware(TenantBoundaryMiddleware,engine=eng,config=cfg)
   for router in [auth.router,finance.router,search.router]:application.include_router(router)
   def dependency_for(selected_factory):
    def dependency():
     with selected_factory() as db:yield db
    return dependency
   from app.db import get_db
   application.dependency_overrides[get_db]=dependency_for(factory)
   c=TestClient(application,base_url='http://'+slug+'.test');clients.append(c)
  alpha,beta=clients
  assert alpha.post('/auth/login',json={'username':'same-finance','password':'alpha-synthetic-password'}).status_code==200
  assert beta.get('/finance/years',headers={'Cookie':'fire_ai_session='+alpha.cookies.get('fire_ai_session')}).status_code==401
  assert beta.post('/auth/login',json={'username':'same-finance','password':'alpha-synthetic-password'}).status_code==401
  assert beta.post('/auth/login',json={'username':'same-finance','password':'beta-synthetic-password'}).status_code==200
  a={'account_id':ids[6]}
  for c,slug,amount in [(alpha,'alpha','111.11'),(beta,'beta','222.22')]:
   p=approve(c,proposal(c,a,ids[0],amount=amount));assert p['review_snapshot']['document']['filename']==slug+'-private-evidence';assert balance(c,a)['available']==amount
   other='beta' if slug=='alpha' else 'alpha';assert c.get('/search',params={'q':other+'-private-account','modules':'budget'}).json()['hits']==[]
   assert other+'-private-evidence' not in c.get('/finance/export/proposals').text
  assert alpha.post('/finance/years',json={'fiscal_year':2027,'currency':'JPY','decimal_places':0,'reason':'Spoof','tenant_id':factories[1][2].tenant_id}).status_code==422
  assert alpha.get('/finance/years',headers={'Host':'beta.test'}).status_code==421
  b=alpha.get('/finance/accounts/'+ids[6]+'/balance',headers={'X-Tenant-ID':factories[1][2].tenant_id},params={'tenant_id':factories[1][2].tenant_id}).json();assert b['available']=='111.11'
  # Canonical runtime session rejects even a manually supplied wrong database bind.
  import app.db as db_module
  from app.tenant import TenantBoundaryError
  old=db_module.settings;db_module.settings=factories[0][2]
  try:
   with pytest.raises(TenantBoundaryError):db_module.BoundSession(bind=factories[1][1])
  finally:db_module.settings=old
 finally:
  for c in clients:c.close()
  for _,eng,_ in factories:eng.dispose()

def postgresql_finance_original(kind,sha256='c'*64):
 # Registered synthetic originals for the native fixture, never operational files.
 return Document(original_filename='Synthetic PG '+kind,storage_path='synthetic/finance-'+kind+'-'+uuid4().hex+'.txt',sha256=sha256)

def test_actual_postgresql_finance_migration_locks_concurrent_payment_and_maximum():
 url=os.environ.get('FIRE_AI_TEST_POSTGRES_URL')
 if not url:pytest.skip('actual PostgreSQL financial locks/concurrency execute in CI')
 from pathlib import Path
 import json
 from decimal import Decimal
 from concurrent.futures import ThreadPoolExecutor
 from threading import Barrier
 from sqlalchemy import create_engine,event,text,update
 from sqlalchemy.orm import sessionmaker
 from sqlalchemy.exc import DBAPIError,StatementError,IntegrityError
 from fastapi import HTTPException
 from app.finance_models import FinanceYear,BudgetAccount,FinanceProposal,FinanceJournal
 from app.models import ContractCase,ContractCounterparty
 from app.finance_schemas import YearInput,AccountInput,Action,ProposalInput,ContractInput
 from app.rbac_seed import seed_rbac
 from app import finance_service as svc
 from app.migrations import split_sql
 schema='synthetic_finance_'+uuid4().hex+'-q';shadow=schema+'_shadow';eng=create_engine(url)
 @event.listens_for(eng,'connect')
 def scope(dbapi_connection,record):
  with dbapi_connection.cursor() as cursor:cursor.execute('SET search_path TO "'+schema+'",public')
 scoped=eng.execution_options(schema_translate_map={None:schema})
 try:
  with eng.begin() as db:db.execute(text('CREATE SCHEMA "'+schema+'"'))
  Base.metadata.create_all(scoped,tables=[t for t in Base.metadata.sorted_tables if not t.name.startswith('finance_')])
  with eng.begin() as db:
   for statement in split_sql((Path(__file__).resolve().parents[2]/'db/migrations/048_run_b_finance.sql').read_text()):db.exec_driver_sql(statement,execution_options={'no_parameters':True})
  factory=sessionmaker(bind=scoped,expire_on_commit=False,autoflush=False)
  with factory() as db:
   roles=seed_rbac(db);u=User(username='synthetic-finance-pg',password_hash=hash_password('synthetic-password'));db.add(u);db.flush();db.add(UserRole(user_id=u.user_id,role_id=roles['system_admin'].role_id));d=postgresql_finance_original('proof',sha256='a'*64);db.add(d);db.flush()
   y=svc.create_year(db,u,YearInput(fiscal_year=2026,currency='JPY',decimal_places=2,account_levels=[{'code':'project','label':'事業'},{'code':'section','label':'節'}],reason='Explicit policy'));svc.approve_year(db,u,y.year_id,Action(expected_version=1,reason='Human policy'))
   parent=None
   for level in range(1,3):parent=svc.create_account(db,u,AccountInput(year_id=y.year_id,parent_id=parent.account_id if parent else None,level=level,code=f'00{level}',name='Synthetic PG account'))
   aid=parent.account_id;uid=u.user_id;did=d.document_id;yid=y.year_id
   other=svc.create_account(db,u,AccountInput(year_id=y.year_id,parent_id=parent.parent_id,level=2,code='OTHER',name='Synthetic exact transfer account'));bid=other.account_id
   vendor=ContractCounterparty(name='Synthetic PG vendor',registration_no='001');db.add(vendor);db.flush();vendorid=vendor.counterparty_id;db.commit()
  # A qualified ORM write must use its own schema's policy/parent, irrespective
  # of caller search_path. Shadow rows deliberately disagree in both directions.
  with eng.begin() as db:
   db.execute(text('CREATE SCHEMA "'+shadow+'"'))
   for table in ['finance_years','finance_accounts']:
    db.execute(text(f'CREATE TABLE "{shadow}".{table} AS TABLE "{schema}".{table}'))
  with factory() as db:
   db.execute(text('SET LOCAL search_path TO pg_catalog'))
   db.execute(update(BudgetAccount).where(BudgetAccount.account_id==aid).values(name='Valid without caller finance schema'))
   db.execute(text('SET LOCAL search_path TO "'+shadow+'",pg_catalog'))
   db.execute(text(f"UPDATE \"{shadow}\".finance_years SET status='draft'"))
   db.execute(update(BudgetAccount).where(BudgetAccount.account_id==aid).values(name='Real approved policy wins'))
   db.execute(text(f"UPDATE \"{shadow}\".finance_years SET status='approved'"))
   db.execute(text(f'UPDATE "{shadow}".finance_accounts SET level=32'))
   db.execute(update(BudgetAccount).where(BudgetAccount.account_id==aid).values(name='Real preceding parent wins'))
   # Direct writes bypass service validation, so require the native trigger's
   # check violation and message, not merely a later foreign-key rejection.
   def rejected_account(message,**values):
    with pytest.raises(DBAPIError) as rejected:
     with db.begin_nested():
      db.add(BudgetAccount(code='REJECT-'+uuid4().hex,name='Synthetic shadow rejection',**values));db.flush()
    assert rejected.value.orig.sqlstate=='23514'
    assert message in str(rejected.value.orig)
   policy_message='account depth follows approved configured policy'
   parent_message='parent must be preceding level in same year'
   db.execute(text(f'UPDATE "{shadow}".finance_years SET account_levels=CAST(:levels AS json)'),{'levels':json.dumps([{'code':str(i),'label':'Synthetic'} for i in range(32)])})
   rejected_account(policy_message,year_id=yid,parent_id=aid,level=3)
   draft=svc.create_year(db,db.get(User,uid),YearInput(fiscal_year=2027,currency='JPY',decimal_places=2,reason='Synthetic unapproved policy'))
   db.execute(text(f'UPDATE "{shadow}".finance_years SET year_id=:id'),{'id':draft.year_id})
   rejected_account(policy_message,year_id=draft.year_id,level=1)
   missing_year=str(uuid4());db.execute(text(f'UPDATE "{shadow}".finance_years SET year_id=:id'),{'id':missing_year})
   rejected_account(policy_message,year_id=missing_year,level=1)
   db.execute(text(f'UPDATE "{shadow}".finance_years SET year_id=:id'),{'id':yid})
   db.execute(text(f'DELETE FROM "{shadow}".finance_accounts WHERE account_id<>:id'),{'id':aid})
   db.execute(text(f'UPDATE "{shadow}".finance_accounts SET level=1'))
   rejected_account(parent_message,year_id=yid,parent_id=aid,level=2)
   missing_parent=str(uuid4());db.execute(text(f'UPDATE "{shadow}".finance_accounts SET account_id=:id'),{'id':missing_parent})
   rejected_account(parent_message,year_id=yid,parent_id=missing_parent,level=2)
   db.rollback()
  def new(kind,amount,**extra):
   with factory() as db:
    u=db.get(User,uid);p=svc.create_proposal(db,u,ProposalInput(kind=kind,account_id=aid,amount=amount,currency='JPY',document_id=did,reason='Synthetic PG decision',idempotency_key=str(uuid4()),**extra));db.commit();p=svc.proposal_action(db,u,p.proposal_id,Action(expected_version=p.version,reason='Human reviewed'),'review');db.commit();return p.proposal_id,p.version
  def approve_one(key,version):
   with factory() as db:
    db.execute(text('SET LOCAL search_path TO pg_catalog'))
    u=db.get(User,uid);svc.proposal_action(db,u,key,Action(expected_version=version,reason='Human approved'),'approve');db.commit()
  key,v=new('initial','9999999999999999.99');approve_one(key,v)
  for _ in range(3):
   transfer,tv=new('transfer','0.01',to_account_id=bid);approve_one(transfer,tv)
  with factory() as db:
   u=db.get(User,uid);c=svc.create_contract(db,u,ContractInput(title='Synthetic PG common contract',counterparty_id=vendorid,amount='400.03',currency='JPY',year_id=yid,document_id=did));db.commit();svc.approve_contract(db,u,c.contract_case_id,Action(expected_version=1,reason='Human source'));db.commit();cid=c.contract_case_id
  commit,cv=new('commitment','400.03',contract_case_id=cid);approve_one(commit,cv)
  # Repeated smallest fraction payments preserve max-magnitude available balance exactly.
  for _ in range(3):
   p,pv=new('payment','0.01',contract_case_id=cid,commitment_id=commit);approve_one(p,pv)
  with factory() as db:
   assert svc.account_balance(db,aid)=={'account_id':aid,'allocated':'9999999999999999.96','reserved':'400.00','spent':'0.03','available':'9999999999999599.93'}
   column=db.execute(text("SELECT numeric_precision,numeric_scale FROM information_schema.columns WHERE table_schema=:s AND table_name='finance_proposals' AND column_name='amount'"),{'s':schema}).one();assert column==(18,2)
   rows=svc.export_rows(db,db.get(User,uid),'journal');assert any(r['allocated']=='9999999999999999.99' for r in rows)
  with factory() as db:
   from app.routers.contracts import create_contract as legacy_create,patch_contract as legacy_patch
   from app.schemas import ContractCreate,ContractPatch
   u=db.get(User,uid);created=legacy_create(ContractCreate(title='Synthetic PG legacy max',amount='9999999999999999.99'),db,u);legacy_id=created.contract_case_id
   assert svc.money(db.get(ContractCase,legacy_id).amount)=='9999999999999999.99'
   legacy_patch(legacy_id,ContractPatch(expected_version=1,amount='9999999999999999.98'),db,u);db.expire_all();assert svc.money(db.get(ContractCase,legacy_id).amount)=='9999999999999999.98'
   from app.finance_models import FinanceCandidate
   from app.finance_schemas import CandidateInput
   candidate=svc.create_candidate(db,u,CandidateInput(kind='quote',year_id=yid,account_id=aid,title='Synthetic PG exact validation',amount='0',currency='JPY',document_id=did));candidateid=candidate.candidate_id;db.commit()
   for value in ['NaN','Infinity','0.001','10000000000000000.00']:
    db.get(FinanceCandidate,candidateid).amount=Decimal(value)
    with pytest.raises(StatementError):db.flush()
    db.rollback()
   db.get(FinanceCandidate,candidateid).amount=Decimal('-0.01')
   with pytest.raises(IntegrityError):db.flush()
   db.rollback()
   with pytest.raises(DBAPIError):db.execute(text("UPDATE finance_candidates SET amount='NaN'::numeric WHERE candidate_id=:id"),{'id':candidateid})
   db.rollback()
   exact=svc.create_contract(db,u,ContractInput(title='Synthetic PG adapter max',amount='9999999999999999.99',currency='JPY',year_id=yid,document_id=did));db.commit();db.expire_all();assert svc.money(db.get(ContractCase,exact.contract_case_id).amount)=='9999999999999999.99'
   from starlette.datastructures import UploadFile
   for format in ['csv','xlsx']:
    values=next(r for r in svc.export_rows(db,u,'proposals') if r['kind']=='initial');values['idempotency_key']='pg-roundtrip-'+format;raw,mime=svc.tabular(list(values),[values],format)
    preview,_=svc.preview_import(db,u,'proposals',UploadFile(filename='SyntheticPG.'+format,file=BytesIO(raw)));db.commit()
    from app.finance_schemas import ImportConfirm
    imported=svc.confirm_import(db,u,preview.preview_id,ImportConfirm(expected_version=preview.version,file_sha256=preview.file_sha256));db.commit();assert svc.money(db.get(FinanceProposal,imported['applied_ids'][0]).amount)=='9999999999999999.99'
  p1,v1=new('payment','300.00',contract_case_id=cid,commitment_id=commit);p2,v2=new('payment','300.00',contract_case_id=cid,commitment_id=commit)
  barrier=Barrier(2)
  def race(item):
   barrier.wait()
   try:approve_one(*item);return 200
   except HTTPException as exc:return exc.status_code
  with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(race,[(p1,v1),(p2,v2)]))
  assert results.count(200)==1 and any(r in (409,422) for r in results)
  with factory() as db:assert svc.account_balance(db,aid)['spent']=='300.03'
  # Replay same posted payment also cannot create a second posting.
  posted=p1 if results[0]==200 else p2
  with pytest.raises(HTTPException):approve_one(posted,2)
  with factory() as first:
   svc.proposal_action(first,first.get(User,uid),p2 if posted==p1 else p1,Action(expected_version=2,reason='Hold real account and source locks'),'review')
   with factory() as second:
    second.execute(text("SET LOCAL lock_timeout='150ms'"))
    with pytest.raises(DBAPIError):second.execute(update(ContractCase).where(ContractCase.contract_case_id==cid).values(title='Blocked concurrent mutation'))
    second.rollback()
   with factory() as second:
    second.execute(text("SET LOCAL lock_timeout='150ms'"))
    with pytest.raises(DBAPIError):second.execute(update(ContractCounterparty).where(ContractCounterparty.counterparty_id==vendorid).values(name='Blocked vendor mutation'))
    second.rollback()
   with factory() as second:
    second.execute(text("SET LOCAL lock_timeout='150ms'"))
    with pytest.raises(DBAPIError):second.execute(update(Document).where(Document.document_id==did).values(sha256='b'*64))
    second.rollback()
   first.rollback()
  # Distinct standalone originals: hold fresh evidence locks through review and approval.
  from app.finance_schemas import AmendmentInput,EventInput,ImportConfirm
  from app.finance_models import FinanceContractAmendment,ProcurementEvent
  from starlette.datastructures import UploadFile
  def blocked_document(first,document_id):
   with factory() as second:
    second.execute(text("SET LOCAL lock_timeout='150ms'"))
    with pytest.raises(DBAPIError):second.execute(update(Document).where(Document.document_id==document_id).values(sha256='b'*64))
    second.rollback()
  for kind in ['initial','payment','amendment','event','import','candidate']:
   with factory() as db:
    u=db.get(User,uid);own=postgresql_finance_original(kind);db.add(own);db.flush();ownid=own.document_id
    if kind in ['initial','payment']:
     refs={'contract_case_id':cid,'commitment_id':commit} if kind=='payment' else {}
     row=svc.create_proposal(db,u,ProposalInput(kind=kind,account_id=aid,amount='0',currency='JPY',document_id=ownid,reason='Distinct source',idempotency_key=str(uuid4()),**refs));identity=row.proposal_id;action=svc.proposal_action
    elif kind=='amendment':
     contract=db.get(ContractCase,cid);row=svc.create_amendment(db,u,cid,AmendmentInput(expected_contract_version=contract.version,amount='400.03',document_id=ownid,reason='Distinct amendment'));identity=row.amendment_id;action=svc.amendment_action
    elif kind=='candidate':
     row=svc.create_candidate(db,u,CandidateInput(kind='quote',year_id=yid,account_id=aid,title='Distinct quote original',amount='0',currency='JPY',document_id=ownid,contract_case_id=cid,counterparty_id=vendorid));identity=row.candidate_id;action=svc.candidate_action
    elif kind=='event':
     row=svc.create_event(db,u,EventInput(contract_case_id=cid,kind='delivery',occurred_on='2026-10-06',description='Distinct delivery',amount='0',currency='JPY',document_id=ownid));identity=row.event_id;action=svc.event_action
    else:
     raw=f'schema_version,kind,account_id,amount,currency,document_id,reason,idempotency_key\nfinance-v1,initial,{aid},0,JPY,{ownid},Distinct import,{uuid4()}\n'.encode();row,_=svc.preview_import(db,u,'proposals',UploadFile(filename='SyntheticDistinct.csv',file=BytesIO(raw)));identity=row.preview_id;originalid=row.source_document_id;digest=row.file_sha256
    db.commit()
   if kind=='import':
    with factory() as first:
     svc.confirm_import(first,first.get(User,uid),identity,ImportConfirm(expected_version=1,file_sha256=digest));blocked_document(first,ownid);blocked_document(first,originalid);first.rollback()
    continue
   for mode,version in ([('review',1)] if kind=='candidate' else [('review',1),('approve',2)]):
    with factory() as first:
     action(first,first.get(User,uid),identity,Action(expected_version=version,reason='Held distinct source'),mode);blocked_document(first,ownid);first.rollback()
    if mode=='review' and kind!='candidate':
     with factory() as db:action(db,db.get(User,uid),identity,Action(expected_version=1,reason='Persist Human source review'),'review');db.commit()
   # Previously loaded identity-map evidence must refresh and reject a post-review SHA change.
   with factory() as first:
    assert first.get(Document,ownid).sha256=='c'*64
    with factory() as second:second.execute(update(Document).where(Document.document_id==ownid).values(sha256='d'*64));second.commit()
    if kind=='candidate':
     refreshed=action(first,first.get(User,uid),identity,Action(expected_version=1,reason='Fresh changed quote original'),'review');assert refreshed.review_snapshot['document']['sha256']=='d'*64;blocked_document(first,ownid);first.rollback();continue
    with pytest.raises(HTTPException) as stale:action(first,first.get(User,uid),identity,Action(expected_version=2,reason='Reject changed original'),'approve')
    assert stale.value.status_code==409;first.rollback()
  with factory() as db:
   with pytest.raises(DBAPIError):db.execute(update(FinanceJournal).where(FinanceJournal.proposal_id==key).values(allocated=Decimal('1.00')));db.flush()
   db.rollback()
   with pytest.raises(DBAPIError):db.execute(update(FinanceProposal).where(FinanceProposal.proposal_id==key).values(amount=Decimal('1.00')));db.flush()
   db.rollback()
   # Native trigger behavior after the complete migration was actually executed.
   for model,identity,field,value in [(FinanceYear,yid,'currency','USD'),(BudgetAccount,aid,'code','INVALID-IDENTITY')]:
    pk=list(model.__table__.primary_key.columns)[0]
    with pytest.raises(DBAPIError):db.execute(update(model).where(pk==identity).values(**{field:value}));db.flush()
    db.rollback()
   u=db.get(User,uid);event=svc.create_event(db,u,EventInput(contract_case_id=cid,kind='delivery',occurred_on='2026-10-06',description='Synthetic guard acceptance',amount='0',currency='JPY',document_id=did));eventid=event.event_id;db.commit()
   svc.event_action(db,u,eventid,Action(expected_version=1,reason='Human guard evidence review'),'review');db.commit();svc.event_action(db,u,eventid,Action(expected_version=2,reason='Human guard approval'),'approve');db.commit()
   case=db.get(ContractCase,cid);amendment=svc.create_amendment(db,u,cid,AmendmentInput(expected_contract_version=case.version,amount='400.03',document_id=did,reason='Synthetic guard acceptance'));amendmentid=amendment.amendment_id;db.commit()
   svc.amendment_action(db,u,amendmentid,Action(expected_version=1,reason='Human guard evidence review'),'review');db.commit();svc.amendment_action(db,u,amendmentid,Action(expected_version=2,reason='Human guard approval'),'approve');db.commit()
   from sqlalchemy import delete
   for model,identity in [(ProcurementEvent,eventid),(FinanceContractAmendment,amendmentid)]:
    pk=list(model.__table__.primary_key.columns)[0]
    with pytest.raises(DBAPIError):db.execute(update(model).where(pk==identity).values(reason='Blocked approved rewrite'));db.flush()
    db.rollback()
    with pytest.raises(DBAPIError):db.execute(delete(model).where(pk==identity));db.flush()
    db.rollback()
   with pytest.raises(ValueError):ProposalInput(kind='initial',account_id=aid,amount='10000000000000000.00',currency='JPY',document_id=did,reason='Overflow',idempotency_key='overflow')
 finally:
  with eng.begin() as db:
   db.execute(text('DROP SCHEMA IF EXISTS "'+shadow+'" CASCADE'))
   db.execute(text('DROP SCHEMA IF EXISTS "'+schema+'" CASCADE'))
  eng.dispose()

def test_existing_sqlite_numeric_contract_affinity_cannot_silently_lose_cents(tmp_path):
 import subprocess,sys
 from pathlib import Path
 root=Path(__file__).resolve().parents[2];env={**os.environ,'PYTHONPATH':str(root/'backend'),'FIRE_AI_DATABASE_URL':'sqlite+pysqlite:///'+str(tmp_path/'legacy.db'),'FIRE_AI_STORAGE_ROOT':str(tmp_path/'storage'),'FIRE_AI_PRODUCTION_MODE':'false'};env.pop('FIRE_AI_TENANT_ID',None)
 code=r'''
from sqlalchemy.schema import CreateTable
from sqlalchemy import text
from fastapi.testclient import TestClient
from app.main import app
from app.db import Base,engine,SessionLocal
from app.models import User,UserRole,ContractCase
from app.rbac_seed import seed_rbac
from app.security import hash_password
sql=str(CreateTable(ContractCase.__table__).compile(engine)).replace('amount VARCHAR(24)','amount NUMERIC(18, 2)')
assert 'amount NUMERIC' in sql
with engine.begin() as db:db.execute(text(sql))
Base.metadata.create_all(engine)
with SessionLocal() as db:
 roles=seed_rbac(db);u=User(username='synthetic-legacy',password_hash=hash_password('synthetic-password'));db.add(u);db.flush();db.add(UserRole(user_id=u.user_id,role_id=roles['system_admin'].role_id));db.commit()
with TestClient(app) as c:
 assert c.post('/auth/login',json={'username':'synthetic-legacy','password':'synthetic-password'}).status_code==200
 r=c.post('/contracts',json={'title':'Synthetic impossible numeric affinity','amount':'9999999999999999.99'})
 assert r.status_code==422,r.text
 assert c.get('/finance/contracts').json()==[]
 r=c.post('/contracts',json={'title':'Synthetic compatible legacy','amount':'12.34'});assert r.status_code==201,r.text;key=r.json()['contract_case_id']
 assert c.get('/finance/contracts/'+key).json()['amount']=='12.34'
 r=c.patch('/contracts/'+key,json={'expected_version':1,'amount':'9999999999999999.99'});assert r.status_code==422,r.text
 assert c.get('/finance/contracts/'+key).json()['amount']=='12.34'
'''
 r=subprocess.run([sys.executable,'-c',code],cwd=tmp_path,env=env,capture_output=True,text=True);assert r.returncode==0,r.stderr

def test_maximum_full_journal_csv_xlsx_roundtrip_and_leading_zero_dates(client):
 import csv,io
 from openpyxl import Workbook,load_workbook
 y=year(client);a=accounts(client,y);b=accounts(client,y,'OTHER');d=document(client);approve(client,proposal(client,a,d,'initial','9999999999999999.99'))
 for _ in range(3):approve(client,proposal(client,a,d,'transfer','0.01',to_account_id=b['account_id']))
 c=contract(client,y,d,'500.01');commit=approve(client,proposal(client,a,d,'commitment','500.01',contract_case_id=c['contract_case_id']))
 for _ in range(3):approve(client,proposal(client,a,d,'payment','0.01',contract_case_id=c['contract_case_id'],commitment_id=commit['proposal_id']))
 assert balance(client,a)['available']=='9999999999999499.95' and balance(client,a)['spent']=='0.03'
 for format in ['csv','xlsx']:
  raw=client.get('/finance/export/proposals?format='+format).content
  if format=='csv':rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
  else:
   cells=list(load_workbook(BytesIO(raw)).active.values);rows=[dict(zip(cells[0],r)) for r in cells[1:]]
  initial=next(r for r in rows if r['kind']=='initial');assert initial['amount']=='9999999999999999.99';initial['idempotency_key']='roundtrip-'+format
  if format=='csv':out=io.StringIO();w=csv.DictWriter(out,fieldnames=initial);w.writeheader();w.writerow(initial);raw=out.getvalue().encode()
  else:book=Workbook();book.active.append(list(initial));book.active.append(list(initial.values()));out=BytesIO();book.save(out);raw=out.getvalue()
  preview=client.post('/finance/import/proposals',files={'file':('roundtrip.'+format,raw,'application/octet-stream')});assert preview.status_code==200,preview.text;preview=preview.json();result=post(client,'/import-previews/'+preview['preview_id']+'/confirm',{'expected_version':1,'file_sha256':preview['file_sha256']},200)
  draft=client.get('/finance/proposals/'+result['applied_ids'][0]).json();assert draft['amount']=='9999999999999999.99' and draft['status']=='draft'
 assert balance(client,a)['available']=='9999999999999499.95'
 # A real Excel account exchange retains text codes, including leading zeros.
 book=Workbook();book.active.append(['schema_version','year_id','level','code','name']);book.active.append(['finance-v1',y['year_id'],1,'000099','Synthetic text code']);out=BytesIO();book.save(out);preview=client.post('/finance/import/accounts',files={'file':('codes.xlsx',out.getvalue(),'application/octet-stream')}).json();post(client,'/import-previews/'+preview['preview_id']+'/confirm',{'expected_version':1,'file_sha256':preview['file_sha256']},200);assert any(r['code']=='000099' for r in client.get('/finance/accounts').json())
 exported=list(csv.DictReader(io.StringIO(client.get('/finance/export/contracts').content.decode('utf-8-sig'))))[0];assert exported['start_date']=='2026-04-01' and exported['end_date']=='2027-03-31'
 exported['contract_no']='ROUNDTRIP';exported['title']='Synthetic dates roundtrip';out=io.StringIO();w=csv.DictWriter(out,fieldnames=exported);w.writeheader();w.writerow(exported);preview=client.post('/finance/import/contracts',files={'file':('contracts.csv',out.getvalue().encode(),'text/csv')});assert preview.status_code==200,preview.text;preview=preview.json();post(client,'/import-previews/'+preview['preview_id']+'/confirm',{'expected_version':1,'file_sha256':preview['file_sha256']},200)
 created=next(r for r in client.get('/finance/contracts').json() if r['title']=='Synthetic dates roundtrip');assert created['start_date']=='2026-04-01' and created['end_date']=='2027-03-31' and created['profile']['renewal_on']=='2027-02-01'

def test_candidate_draft_update_cancel_and_legacy_missing_contract_evidence(client):
 y=year(client);a=accounts(client,y);d=document(client);candidate=post(client,'/candidates',{'kind':'request','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic request','amount':'0','currency':'JPY','document_id':d})
 r=client.patch('/finance/candidates/'+candidate['candidate_id'],json={'expected_version':1,'amount':'12.34','title':'Synthetic corrected request'});assert r.status_code==200,r.text
 reviewed=post(client,'/candidates/'+candidate['candidate_id']+'/review',{'expected_version':2,'reason':'Human checked'},200)
 assert client.patch('/finance/candidates/'+candidate['candidate_id'],json={'expected_version':3,'amount':'0'}).status_code==409
 cancelled=post(client,'/candidates/'+candidate['candidate_id']+'/cancel',{'expected_version':reviewed['version'],'reason':'Human candidate withdrawn'},200);assert cancelled['status']=='cancelled' and cancelled['review_snapshot']==reviewed['review_snapshot']
 legacy=client.post('/contracts',json={'title':'Synthetic missing original','amount':'10'}).json();client.post('/contracts/'+legacy['contract_case_id']+'/approve',json={'expected_version':1})
 r=client.put('/finance/contracts/'+legacy['contract_case_id']+'/profile',json={'expected_version':2,'year_id':y['year_id']});assert r.status_code==200,r.text
 assert client.post('/finance/proposals',json={'kind':'commitment','account_id':a['account_id'],'amount':'1','currency':'JPY','document_id':d,'contract_case_id':legacy['contract_case_id'],'reason':'Missing approved contract evidence','idempotency_key':'missing-source'}).status_code==422

def test_invalid_import_rows_duplicate_account_rollback_and_xlsx_float_rejected(client):
 from openpyxl import Workbook
 y=year(client)
 raw=f'schema_version,year_id,level,code,name\nfinance-v1,{y["year_id"]},1,NEW,Synthetic\nfinance-v1,{y["year_id"]},1,NEW,Duplicate\n'.encode()
 r=client.post('/finance/import/accounts',files={'file':('invalid.csv',raw,'text/csv')});assert r.status_code==409,r.text
 assert client.get('/finance/accounts').json()==[]
 a=accounts(client,y);d=document(client);w=Workbook();w.active.append(['schema_version','kind','account_id','amount','currency','document_id','reason','idempotency_key']);w.active.append(['finance-v1','initial',a['account_id'],0.1,'JPY',d,'Numeric floating source','float']);out=BytesIO();w.save(out)
 r=client.post('/finance/import/proposals',files={'file':('numeric.xlsx',out.getvalue(),'application/octet-stream')});assert r.status_code==422,r.text

def test_configurable_account_hierarchy_alternate_depth_and_policy_provenance(client):
 levels=[{'code':'project','label':'事業'},{'code':'section','label':'節'}]
 y=post(client,'/years',{'fiscal_year':2029,'currency':'JPY','decimal_places':2,'account_levels':levels,'reason':'Headquarters configured hierarchy'})
 y=post(client,'/years/'+y['year_id']+'/approve',{'expected_version':1,'reason':'Human policy accepted'},200)
 root=post(client,'/accounts',{'year_id':y['year_id'],'level':1,'code':'0001','name':'Synthetic project'})
 leaf=post(client,'/accounts',{'year_id':y['year_id'],'parent_id':root['account_id'],'level':2,'code':'0002','name':'Synthetic section'})
 assert client.get('/finance/accounts?leaf=true&year_id='+y['year_id']).json()[0]['level_label']=='節'
 assert client.post('/finance/accounts',json={'year_id':y['year_id'],'parent_id':leaf['account_id'],'level':3,'code':'0003','name':'Outside policy'}).status_code==422
 d=document(client);p=approve(client,proposal(client,leaf,d,'initial','12.34'));assert p['review_snapshot']['year']['account_levels']==levels
 assert client.get('/finance/summary?year_id='+y['year_id']).json()['allocated']=='12.34'
 assert client.patch('/finance/years/'+y['year_id'],json={'expected_version':2,'account_levels':levels[:1]}).status_code==409
 extended=[{'code':f'L{i}','label':label} for i,label in enumerate(['款','項','目','事業','節','細節','細々節'])]
 y=post(client,'/years',{'fiscal_year':2030,'currency':'JPY','decimal_places':0,'account_levels':extended,'reason':'Synthetic extended configured hierarchy'});y=post(client,'/years/'+y['year_id']+'/approve',{'expected_version':1,'reason':'Human extended policy'},200);parent=None
 for index in range(7):parent=post(client,'/accounts',{'year_id':y['year_id'],'parent_id':parent['account_id'] if parent else None,'level':index+1,'code':f'000{index}','name':f'Synthetic configured {index}'})
 approve(client,proposal(client,parent,d,'initial','100'));assert balance(client,parent)['allocated']=='100.00'
 assert '細々節' in client.get('/finance/export/years').text and '0006' in client.get('/finance/export/accounts').text

def event(c,contract,d,kind,amount='0.00',**extra):
 return post(c,'/procurement-events',{'contract_case_id':contract['contract_case_id'],'kind':kind,'occurred_on':'2026-10-06','description':'Synthetic '+kind,'amount':amount,'currency':'JPY','document_id':d,**extra})
def approve_event(c,row):
 r=post(c,'/procurement-events/'+row['event_id']+'/review',{'expected_version':row['version'],'reason':'Human inspected source'},200)
 return post(c,'/procurement-events/'+row['event_id']+'/approve',{'expected_version':r['version'],'reason':'Human official procurement event'},200)

def test_procurement_delivery_inspection_invoice_payment_and_vendor_freshness(client):
 y=post(client,'/years',{'fiscal_year':2026,'currency':'JPY','decimal_places':2,'require_invoice_on_payment':True,'reason':'Human invoice policy'});y=post(client,'/years/'+y['year_id']+'/approve',{'expected_version':1,'reason':'Human approved invoice policy'},200);a=accounts(client,y);d=document(client);c=contract(client,y,d);approve(client,proposal(client,a,d))
 delivery=event(client,c,d,'delivery');r=post(client,'/procurement-events/'+delivery['event_id']+'/review',{'expected_version':1,'reason':'Human delivery check'},200)
 vendor=c['counterparty_id'];assert client.patch('/finance/counterparties/'+vendor,json={'expected_version':1,'registration_no':'SYN-NEW'}).status_code==200
 assert client.post('/finance/procurement-events/'+delivery['event_id']+'/approve',json={'expected_version':r['version'],'reason':'Stale identity'}).status_code==409
 delivery=approve_event(client,r);assert delivery['review_snapshot']['contract']['counterparty']['registration_no']=='SYN-NEW'
 inspection=approve_event(client,event(client,c,d,'inspection',related_event_id=delivery['event_id']))
 invoice=approve_event(client,event(client,c,d,'invoice','100.00',related_event_id=inspection['event_id']))
 commit=approve(client,proposal(client,a,d,'commitment','200.00',contract_case_id=c['contract_case_id']))
 assert client.post('/finance/proposals',json={'kind':'payment','account_id':a['account_id'],'amount':'1.00','currency':'JPY','document_id':d,'contract_case_id':c['contract_case_id'],'commitment_id':commit['proposal_id'],'reason':'Missing required invoice','idempotency_key':'missing-invoice'}).status_code==422
 payment=approve(client,proposal(client,a,d,'payment','90.00',contract_case_id=c['contract_case_id'],commitment_id=commit['proposal_id'],invoice_id=invoice['event_id']))
 over=proposal(client,a,d,'payment','11.00',contract_case_id=c['contract_case_id'],commitment_id=commit['proposal_id'],invoice_id=invoice['event_id']);r=post(client,'/proposals/'+over['proposal_id']+'/review',{'expected_version':1,'reason':'Human checked'},200)
 assert client.post('/finance/proposals/'+over['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Exceeds invoice'}).status_code==422
 assert client.get('/finance/procurement-events/'+invoice['event_id']).json()['payments'][0]['proposal_id']==payment['proposal_id']
 assert any(r['kind']=='invoice' for r in client.get('/finance/procurement-events').json())
 remove('contract.read');assert client.get('/finance/procurement-events/'+invoice['event_id']).status_code==403


def test_vendor_change_invalidates_financial_review_and_quote_snapshot(client):
 y=year(client);a=accounts(client,y);d=document(client);c=contract(client,y,d);approve(client,proposal(client,a,d));p=proposal(client,a,d,'commitment','100',contract_case_id=c['contract_case_id']);r=post(client,'/proposals/'+p['proposal_id']+'/review',{'expected_version':1,'reason':'Human reviewed vendor'},200)
 assert client.patch('/finance/counterparties/'+c['counterparty_id'],json={'expected_version':1,'name':'Synthetic revised identity'}).status_code==200
 assert client.post('/finance/proposals/'+p['proposal_id']+'/approve',json={'expected_version':r['version'],'reason':'Stale vendor'}).status_code==409
 candidate=post(client,'/candidates',{'kind':'quote','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic supplier quote','amount':'1','currency':'JPY','document_id':d,'counterparty_id':c['counterparty_id']});reviewed=post(client,'/candidates/'+candidate['candidate_id']+'/review',{'expected_version':1,'reason':'Human quote check'},200)
 assert reviewed['review_snapshot']['counterparty']['name']=='Synthetic revised identity'

def test_finance_source_assistance_is_evidence_bound_and_never_posts(client):
 y=year(client);a=accounts(client,y);d=document(client);c=contract(client,y,d)
 support=client.get('/finance/contracts/'+c['contract_case_id']+'/support');assert support.status_code==200,support.text
 assert 'specification' in support.json()['missing_document_candidates'] and 'decision' in support.json()['missing_document_candidates']
 assert support.json()['authority']=='candidate' and support.json()['source']['counterparty']['version']==1
 first=client.post('/documents/upload',files={'file':('quote-a.txt',b'Synthetic quote amount 10.01','text/plain')}).json()['document_id'];second=client.post('/documents/upload',files={'file':('quote-b.txt',b'Synthetic quote amount 10.02','text/plain')}).json()['document_id']
 extracted=client.get('/finance/documents/'+first+'/extract');assert extracted.status_code==200,extracted.text;assert '10.01' in extracted.json()['text'] and extracted.json()['source']['sha256']
 diff=post(client,'/document-diff',{'before_document_id':first,'after_document_id':second},200);assert '-Synthetic quote amount 10.01' in diff['diff'] and '+Synthetic quote amount 10.02' in diff['diff']
 result=post(client,'/calculate',{'amounts':['9999999999999999.96','0.01','0.01','0.01']},200);assert result['total']=='9999999999999999.99' and result['authority']=='candidate'
 assert client.post('/finance/calculate',json={'amounts':['9999999999999999.99','0.01']}).status_code==422
 assert balance(client,a)['allocated']=='0.00'
 remove('document.read');assert client.get('/finance/documents/'+first+'/extract').status_code==403;assert client.get('/finance/contracts/'+c['contract_case_id']+'/support').status_code==403

def test_finance_v2_ui_has_usable_stage_and_invoice_reference_controls():
 from pathlib import Path
 js=(Path(__file__).resolve().parents[2]/'frontend/finance.js').read_text()
 for function in ['financeEvents','financeEventForm','financeEventDetail','financeSupport']:assert 'function '+function+'(' in js
 assert "['invoice_id'" in js and '/finance/procurement-events?kind=invoice&status=approved' in js
 assert "['specification','仕様書']" in js and "['decision','決裁文書']" in js

def test_financial_contract_legacy_approval_binds_evidence_and_reversal_requires_source_read(client):
 y=year(client);a=accounts(client,y);d=document(client)
 c=post(client,'/contracts',{'title':'Synthetic approval binding','amount':'500','currency':'JPY','year_id':y['year_id'],'document_id':d})
 assert client.post('/contracts/'+c['contract_case_id']+'/approve',json={'expected_version':c['version']}).status_code==200
 c=client.get('/finance/contracts/'+c['contract_case_id']).json();assert c['profile']['approved_evidence']['documents'][0]['sha256']
 approve(client,proposal(client,a,d));commit=approve(client,proposal(client,a,d,'commitment','100',contract_case_id=c['contract_case_id']))
 remove('contract.read')
 assert client.post('/finance/proposals',json={'kind':'reversal','account_id':a['account_id'],'amount':'100','currency':'JPY','document_id':d,'reason':'No source authority','reverses_id':commit['proposal_id'],'idempotency_key':'unauthorized-reversal'}).status_code==403

def test_finance_import_preview_original_hash_owner_and_event_roundtrip(client):
 import csv,io
 from app.finance_models import FinanceImportPreview
 y=year(client);d=document(client)
 raw=('schema_version,year_id,level,code,name\nfinance-v1,'+y['year_id']+',1,001,Synthetic original\n').encode()
 assert client.post('/finance/import/accounts',data={'document_id':d},files={'file':('import.csv',raw,'text/csv')}).status_code==422
 preview=client.post('/finance/import/accounts',files={'file':('import.csv',raw,'text/csv')}).json()
 with SessionLocal() as db:
  db.get(FinanceImportPreview,preview['preview_id']).created_by=str(uuid4());db.commit()
 assert client.post('/finance/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==403
 c=contract(client,y,d);e=event(client,c,d,'delivery');rows=list(csv.DictReader(io.StringIO(client.get('/finance/export/procurement-events').content.decode('utf-8-sig'))));rows[0]['description']='Synthetic import event'
 out=io.StringIO();w=csv.DictWriter(out,fieldnames=rows[0]);w.writeheader();w.writerow(rows[0]);preview=client.post('/finance/import/procurement-events',files={'file':('events.csv',out.getvalue().encode(),'text/csv')});assert preview.status_code==200,preview.text

def test_procurement_event_unified_search_links_and_source_permission(client):
 y=year(client);d=document(client);c=contract(client,y,d);e=event(client,c,d,'invoice','10.00')
 r=client.get('/search',params={'q':'Synthetic invoice','modules':'procurement'});assert r.status_code==200,r.text
 hit=next((h for h in r.json()['hits'] if h['source_id']==e['event_id']),None);assert hit and hit['navigation']['dataset']=='procurement-events'
 remove('contract.read');r=client.get('/search',params={'q':'Synthetic invoice','modules':'procurement'});assert not any(h['source_id']==e['event_id'] for h in r.json()['hits'])

def test_sqlite_canonical_money_constraints_never_use_numeric_text_comparison(client):
 from sqlalchemy import text
 from sqlalchemy.exc import IntegrityError
 y=year(client);a=accounts(client,y);d=document(client);row=post(client,'/candidates',{'kind':'quote','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic sign constraint','amount':'9999999999999999.99','currency':'JPY','document_id':d})
 with SessionLocal() as db:
  key=db.execute(text('SELECT candidate_id FROM finance_candidates')).scalar()
  with pytest.raises(IntegrityError):db.execute(text("UPDATE finance_candidates SET amount='-0.01' WHERE candidate_id=:key"),{'key':key})
  db.rollback()
  assert db.execute(text('SELECT amount FROM finance_candidates')).scalar()=='9999999999999999.99'
  ddl=db.execute(text("SELECT sql FROM sqlite_master WHERE name='finance_candidates'")).scalar();assert 'amount >= 0' not in ddl

def test_sqlite_raw_money_format_orm_api_scale_and_signed_zero(client):
 from sqlalchemy import text
 from sqlalchemy.exc import IntegrityError,StatementError
 from decimal import Decimal
 from app.finance_models import FinanceCandidate
 y=year(client);a=accounts(client,y);d=document(client);row=post(client,'/candidates',{'kind':'quote','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic strict raw text','amount':'0','currency':'JPY','document_id':d})
 with SessionLocal() as db:
  key=db.execute(text('SELECT candidate_id FROM finance_candidates')).scalar()
  for value in ['NaN','Infinity','invalid','1e2','0.001','10000000000000000.00','01.00','1.0','1..00','--1.00','-.01']:
   with pytest.raises(IntegrityError):db.execute(text('UPDATE finance_candidates SET amount=:v WHERE candidate_id=:key'),{'v':value,'key':key})
   db.rollback()
  candidate=db.get(FinanceCandidate,row['candidate_id']);candidate.amount=Decimal('0.001')
  with pytest.raises(StatementError):db.flush()
  db.rollback()
 assert client.post('/finance/candidates',json={'kind':'quote','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic invalid API scale','amount':'0.001','currency':'JPY','document_id':d}).status_code==422
 signed=post(client,'/candidates',{'kind':'quote','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic mathematical zero','amount':'-0.00','currency':'JPY','document_id':d});assert Decimal(signed['amount'])==0

def test_legacy_financial_draft_money_cannot_bypass_approved_fiscal_precision(client):
 y=post(client,'/years',{'fiscal_year':2026,'currency':'JPY','decimal_places':0,'reason':'Synthetic no fraction policy'});y=post(client,'/years/'+y['year_id']+'/approve',{'expected_version':1,'reason':'Human fiscal policy'},200);d=document(client)
 c=post(client,'/contracts',{'title':'Synthetic configured legacy finance','amount':'10','currency':'JPY','year_id':y['year_id'],'document_id':d})
 for value in ['10.01',None]:
  r=client.patch('/contracts/'+c['contract_case_id'],json={'expected_version':1,'amount':value});assert r.status_code==422,r.text
 assert client.get('/finance/contracts/'+c['contract_case_id']).json()['amount']=='10.00'
 # Existing incompatible source cannot be silently certified through either approval entry point.
 with SessionLocal() as db:
  from app.models import ContractCase
  from decimal import Decimal
  db.get(ContractCase,c['contract_case_id']).amount=Decimal('10.01');db.commit()
 assert client.post('/contracts/'+c['contract_case_id']+'/approve',json={'expected_version':1}).status_code==422
 assert client.post('/finance/contracts/'+c['contract_case_id']+'/approve',json={'expected_version':1,'reason':'Invalid precision cannot approve'}).status_code==422

def test_legacy_unsafe_json_float_rejects_before_decimal_and_exact_strings_survive(client):
 # This literal already loses .99 during JSON float decoding; reject before Decimal conversion.
 r=client.post('/contracts',content='{"title":"Synthetic unsafe float","amount":9999999999999998.99}',headers={'Content-Type':'application/json'});assert r.status_code==422,r.text
 r=client.post('/contracts',json={'title':'Synthetic exact decimal source','amount':'9999999999999998.99'});assert r.status_code==201,r.text;key=r.json()['contract_case_id'];assert client.get('/finance/contracts/'+key).json()['amount']=='9999999999999998.99'
 r=client.patch('/contracts/'+key,content='{"expected_version":1,"amount":9999999999999998.99}',headers={'Content-Type':'application/json'});assert r.status_code==422,r.text
 assert client.get('/finance/contracts/'+key).json()['amount']=='9999999999999998.99'
 assert client.patch('/contracts/'+key,json={'expected_version':1,'amount':'9999999999999999.99'}).status_code==200
 assert client.get('/finance/contracts/'+key).json()['amount']=='9999999999999999.99'
 for value in [12.34,0.0,123]:assert client.post('/contracts',json={'title':'Synthetic compatible ordinary numeric','amount':value}).status_code==201
 # Integers remain lossless JSON numbers even above the conservative legacy float range.
 r=client.post('/contracts',json={'title':'Synthetic exact integer','amount':9999999999999999});assert r.status_code==201,r.text;assert client.get('/finance/contracts/'+r.json()['contract_case_id']).json()['amount']=='9999999999999999.00'
 y=year(client);d=document(client);assert client.post('/finance/contracts',json={'title':'Synthetic strict finance float','amount':12.34,'currency':'JPY','year_id':y['year_id'],'document_id':d}).status_code==422

def test_excel_numeric_display_codes_require_lossless_text_source(client):
 from openpyxl import Workbook
 y=year(client);book=Workbook();book.active.append(['schema_version','year_id','level','code','name']);book.active.append(['finance-v1',y['year_id'],1,1,'Synthetic displayed code']);book.active['D2'].number_format='0000';out=BytesIO();book.save(out)
 r=client.post('/finance/import/accounts',files={'file':('numeric-code.xlsx',out.getvalue(),'application/octet-stream')});assert r.status_code==422,r.text
 assert client.get('/finance/accounts').json()==[]

def test_review_picker_uses_selected_entity_identity(tmp_path):
 import subprocess
 from pathlib import Path
 root=Path(__file__).resolve().parents[2];script=tmp_path/'picker.cjs'
 script.write_text(r'''
const fs=require('fs'),vm=require('vm'),assert=require('assert');const nodes=new Map();const get=id=>{if(!nodes.has(id))nodes.set(id,{value:'',innerHTML:'',remove(){},closest(){return {append(){}}}});return nodes.get(id)};
let rows=[];const ctx={URLSearchParams,$:get,esc:String,document:{createElement(){return {}},contains:node=>[...nodes.values()].includes(node)},api:async()=>rows};vm.createContext(ctx);vm.runInContext(fs.readFileSync(process.argv[2],'utf8'),ctx);
(async()=>{for(const [key,row,want] of [['commitment_id',{proposal_id:'COMMITMENT',commitment_id:null,account_id:'ACCOUNT',kind:'commitment',amount:'1',currency:'JPY'},'COMMITMENT'],['invoice_id',{event_id:'INVOICE',contract_case_id:'CONTRACT',kind:'invoice',amount:'1',currency:'JPY'},'INVOICE'],['related_event_id',{event_id:'INSPECTION',related_event_id:'DELIVERY',kind:'inspection',amount:'1',currency:'JPY'},'INSPECTION']]){rows=[row];await ctx.financePicker(key,'/synthetic','Synthetic');assert(get('financeField_'+key).innerHTML.includes('value="'+want+'"'),get('financeField_'+key).innerHTML)}})().catch(e=>{console.error(e);process.exit(1)});
''')
 r=subprocess.run(['node',str(script),str(root/'frontend/finance.js')],capture_output=True,text=True);assert r.returncode==0,r.stderr

@pytest.mark.parametrize('surface',['read','render'])
def test_review_reversal_transitive_source_privacy_and_render(client,surface):
 from app.models import FormTemplate
 from openpyxl import Workbook
 y=year(client);a=accounts(client,y);d=document(client);c=contract(client,y,d);approve(client,proposal(client,a,d))
 commit=approve(client,proposal(client,a,d,'commitment','400',contract_case_id=c['contract_case_id']))
 reversal=approve(client,proposal(client,a,d,'reversal','400',reverses_id=commit['proposal_id']))
 book=Workbook();book.active['A1']='Synthetic official original';raw=BytesIO();book.save(raw)
 original=client.post('/documents/upload',files={'file':('SyntheticOriginal.xlsx',raw.getvalue())}).json()['document_id']
 with SessionLocal() as db:
  template=FormTemplate(template_code='SyntheticReview',name='Synthetic original',module_code='budget',document_id=original,version_label='1',field_mapping={'amount':{'sheet':'Sheet','cell':'A2'}},status='active');db.add(template);db.commit();tid=template.form_template_id
 remove('contract.read')
 if surface=='render':
  response=client.post('/finance/proposals/'+reversal['proposal_id']+'/render',json={'form_template_id':tid});assert response.status_code==403,response.text
 else:
  remove('document.read')
  for response in [client.get('/finance/proposals/'+reversal['proposal_id']),client.get('/finance/proposals'),client.get('/finance/export/proposals')]:
   assert response.status_code==200,response.text
   assert 'Synthetic equipment contract' not in response.text and '2027-03-31' not in response.text and 'Synthetic supplier' not in response.text and 'target_contract' not in response.text,response.text
   assert c['contract_case_id'] not in response.text and d not in response.text

@pytest.mark.parametrize('kind',['proposal','amendment','event','import','candidate'])
def test_review_standalone_evidence_uses_fresh_lock(client,monkeypatch,kind):
 from app import finance_service as svc
 y=year(client);a=accounts(client,y);linked=document(client);own=document(client);seen=[];original=svc.get_row
 def tracked(db,model,key,lock=False):
  if model is Document:seen.append((key,lock))
  return original(db,model,key,lock)
 if kind=='proposal':
  row=proposal(client,a,own);path='/proposals/'+row['proposal_id']+'/review';payload={'expected_version':1,'reason':'Synthetic source review'}
 elif kind=='import':
  raw=f'schema_version,year_id,level,code,name\nfinance-v1,{y["year_id"]},1,009,Synthetic import\n'.encode();row=client.post('/finance/import/accounts',files={'file':('Synthetic.csv',raw)}).json();own=row['source_document_id'];path='/import-previews/'+row['preview_id']+'/confirm';payload={'expected_version':1,'file_sha256':row['file_sha256']}
 else:
  c=contract(client,y,linked)
  if kind=='candidate':
   row=post(client,'/candidates',{'kind':'quote','year_id':y['year_id'],'account_id':a['account_id'],'title':'Synthetic distinct candidate','amount':'0','currency':'JPY','document_id':own,'contract_case_id':c['contract_case_id'],'counterparty_id':c['counterparty_id']});path='/candidates/'+row['candidate_id']+'/review'
  elif kind=='amendment':
   row=post(client,'/contracts/'+c['contract_case_id']+'/amendments',{'expected_contract_version':c['version'],'amount':'500','document_id':own,'reason':'Synthetic amended original'});path='/amendments/'+row['amendment_id']+'/review'
  else:
   row=post(client,'/procurement-events',{'contract_case_id':c['contract_case_id'],'kind':'delivery','occurred_on':'2026-10-06','description':'Synthetic standalone delivery','amount':'0','currency':'JPY','document_id':own});path='/procurement-events/'+row['event_id']+'/review'
  payload={'expected_version':1,'reason':'Synthetic source review'}
 monkeypatch.setattr(svc,'get_row',tracked);post(client,path,payload,200)
 assert (own,True) in seen,seen
 locked=list(dict.fromkeys(key for key,lock in seen if lock));assert locked==sorted(locked),locked

@pytest.mark.parametrize('format',['csv','xlsx'])
def test_review_literal_apostrophe_formula_roundtrip(client,format):
 import csv,io
 from openpyxl import load_workbook
 from app import finance_service as svc
 y=year(client);a=accounts(client,y);d=document(client)
 originals=["'=Literal original apostrophe","''=Two apostrophes","'plain",'=Formula', '+Formula', '-Formula', '@Formula','  =Formula']
 originals=[proposal(client,a,d,reason=text)['reason'] for text in originals]
 response=client.get('/finance/export/proposals?format='+format);assert response.status_code==200
 rows=svc.read_import(response.content)
 assert sorted(svc.import_payload('proposals',row).reason for row in rows)==sorted(originals)
 for row in rows:row['idempotency_key']=str(uuid4())
 raw,mime=svc.tabular(list(rows[0]),[{**r,'reason':svc.import_payload('proposals',r).reason} for r in rows],format)
 preview=client.post('/finance/import/proposals',files={'file':('SyntheticRoundtrip.'+format,raw,mime)});assert preview.status_code==200,preview.text
 row=preview.json();post(client,'/import-previews/'+row['preview_id']+'/confirm',{'expected_version':1,'file_sha256':row['file_sha256']},200)
 assert sorted(r['reason'] for r in client.get('/finance/proposals').json())==sorted(originals*2)

def test_finance_guard_functions_are_complete_canonical_migration_statements():
 from pathlib import Path
 from app.migrations import split_sql
 sql=(Path(__file__).resolve().parents[2]/'db/migrations/048_run_b_finance.sql').read_text();statements=split_sql(sql)
 expected={'finance_immutable_guard','finance_account_identity_guard'}
 functions=[s for s in statements if 'CREATE FUNCTION ' in s]
 assert len(functions)==2
 for name in expected:
  statement=next(s for s in functions if 'CREATE FUNCTION '+name+'()' in s)
  assert 'RETURN NEW;\nEND;' in statement,statement
  assert statement.rstrip().endswith("'"),statement
  assert 'RAISE EXCEPTION' in statement and '23514' in statement
 triggers={name:function for name,function in [('finance_journal_immutable','finance_immutable_guard'),('finance_proposal_immutable','finance_immutable_guard'),('finance_policy_immutable','finance_immutable_guard'),('finance_amendment_immutable','finance_immutable_guard'),('finance_event_immutable','finance_immutable_guard'),('finance_account_identity','finance_account_identity_guard')]}
 for name,function in triggers.items():
  statement=next(s for s in statements if 'CREATE TRIGGER '+name+' ' in s)
  assert 'EXECUTE FUNCTION '+function+'()' in statement,statement
 assert not any(s.strip().startswith(('RETURN NEW','END IF','$$','DECLARE p ')) for s in statements)
 assert 'uq_finance_approved_reversal' in sql


def test_finance_sql_passes_actual_psycopg_driver_conversion_without_parameters():
 from pathlib import Path
 from psycopg._queries import PostgresQuery
 from psycopg.adapt import Transformer
 from app.migrations import split_sql
 statements=split_sql((Path(__file__).resolve().parents[2]/'db/migrations/048_run_b_finance.sql').read_text())
 for statement in statements:
  # SQLAlchemy exec_driver_sql invokes the same psycopg converter with empty params.
  for parameters in [None,(),{}]:
   query=PostgresQuery(Transformer());query.convert(statement,parameters)
   assert query.query==statement.encode('utf-8') and query.params==(None if parameters is None else []),statement


def test_finance_account_guard_binds_lookups_to_trigger_table_schema():
 from pathlib import Path
 from app.migrations import split_sql
 statements=split_sql((Path(__file__).resolve().parents[2]/'db/migrations/048_run_b_finance.sql').read_text())
 statement=next(s for s in statements if 'CREATE FUNCTION finance_account_identity_guard()' in s)
 body=statement.split("AS '",1)[1].rstrip().removesuffix("'").replace("''", "'")
 # Both authoritative reads must ignore caller search_path and bind UUID values.
 for table,key,target,source in [('finance_years','year_id','y','NEW.year_id'),('finance_accounts','account_id','p','NEW.parent_id')]:
  assert ("EXECUTE 'SELECT * FROM ' || pg_catalog.quote_ident(TG_TABLE_SCHEMA) || "
          +f"'.{table} WHERE {key} = $1' INTO {target} USING {source};") in body
  assert f'{target}.{key} IS NULL' in body
 # Dynamic EXECUTE does not set FOUND: a previous statement cannot decide absence.
 assert 'NOT FOUND' not in body


def test_finance_postgresql_original_fixtures_preserve_storage_identity(client):
 originals=[postgresql_finance_original(kind) for kind in ['proof','initial','payment','amendment','event','import','candidate']]
 with SessionLocal() as db:
  db.add_all(originals);db.flush()
  assert len({r.storage_path for r in originals})==7
  assert len({r.document_id for r in originals})==7
  db.rollback()

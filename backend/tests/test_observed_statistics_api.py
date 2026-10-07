"""Saved statistics use current rights, frozen values and append-only Human decisions."""
from datetime import date, datetime, timezone
from io import BytesIO
import importlib.util
import json
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select, text
from sqlalchemy.orm import Session
from app.main import app as registered_models
from app.db import Base, get_db
from app.models import EmergencyCase, EmergencyPatient, User, UserRole, Role, RolePermission, Permission, AuditLog, UserSession
from app.security import hash_password
from app.rbac_seed import seed_rbac
from app.routers import auth

QUERY = {'start_date':'2026-01-01','end_date':'2026-01-31','metric_keys':['emergency.cases','emergency.patient_records']}
AGGREGATE = {'statistics.read','statistics.record','statistics.export','emergency.report.read','emergency.report.export','incident.aggregate','incident.export','fleet.aggregate','fleet.export'}


@pytest.fixture
def statistics_api(tmp_path):
    if importlib.util.find_spec('app.statistics_models'):
        import app.statistics_models
    engine=create_engine('sqlite+pysqlite:///'+str(tmp_path/'api.db'),connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        seed_rbac(db)
        role=Role(code='statistics-test-aggregate',name='Synthetic aggregate only');db.add(role);db.flush()
        for code in AGGREGATE:
            permission=db.scalar(select(Permission).where(Permission.code==code))
            if permission is None:
                permission=Permission(code=code,description='Synthetic');db.add(permission);db.flush()
            db.add(RolePermission(role_id=role.role_id,permission_id=permission.permission_id))
        user=User(username='statistics-user',password_hash=hash_password('synthetic-statistics-password'));db.add(user);db.flush()
        db.add(UserRole(user_id=user.user_id,role_id=role.role_id))
        case=EmergencyCase(source_case_key='synthetic',call_date=date(2026,1,2));db.add(case);db.flush()
        patient=EmergencyPatient(emergency_case_id=case.emergency_case_id,patient_number=1,diagnosis_text='PRIVATE_MEDICAL');db.add(patient);db.flush()
        refs={'user_id':user.user_id,'role_id':role.role_id,'case_id':case.emergency_case_id,'patient_id':patient.emergency_patient_id};db.commit()
    application=FastAPI();application.include_router(auth.router)
    if importlib.util.find_spec('app.routers.statistics'):
        from app.routers import statistics
        application.include_router(statistics.router)
    def dependency():
        with Session(engine,expire_on_commit=False) as db:yield db
    application.dependency_overrides[get_db]=dependency
    with TestClient(application) as client:
        response=client.post('/auth/login',json={'username':'statistics-user','password':'synthetic-statistics-password'})
        assert response.status_code==200,response.text
        yield client,engine,refs
    engine.dispose()


def save(client, query=None):
    response=client.post('/statistics/reports',json=query or QUERY)
    assert response.status_code==201,response.text
    return response.json()


def revoke(engine, refs, code):
    with Session(engine) as db:
        permission=db.scalar(select(Permission).where(Permission.code==code))
        db.delete(db.get(RolePermission,(refs['role_id'],permission.permission_id)));db.commit()


def test_strict_query_and_aggregate_output_never_contains_private_identifiers(statistics_api):
    client,engine,refs=statistics_api
    response=client.post('/statistics/query',json=QUERY)
    assert response.status_code==200,response.text
    assert [metric['value'] for metric in response.json()['metrics']]==['1','1']
    for value in refs.values():assert value not in response.text
    assert 'PRIVATE' not in response.text and 'content_hash' not in response.text
    assert client.post('/statistics/query',json={**QUERY,'evidence':{}}).status_code==422
    revoke(engine,refs,'emergency.report.read')
    assert client.post('/statistics/query',json=QUERY).status_code==403


def test_immutable_snapshot_confirmation_and_replacement_preserve_old_facts(statistics_api):
    client,engine,refs=statistics_api
    row=save(client);key=row['report_id'];snapshot=row['snapshot']
    assert row['state']=='saved' and row['version']==1
    confirm={'expected_version':1,'acknowledged':True,'review_note':'Human checked synthetic sources'}
    response=client.post('/statistics/reports/'+key+'/confirm',json=confirm)
    assert response.status_code==200,response.text
    confirmed=response.json()
    assert confirmed['version']==2 and confirmed['state']=='confirmed'
    assert confirmed['snapshot']==snapshot and confirmed['snapshot']['coverage_status']=='unknown'
    assert client.post('/statistics/reports/'+key+'/confirm',json=confirm).status_code==409
    with Session(engine) as db:
        db.add(EmergencyCase(source_case_key='later',call_date=date(2026,1,4)));db.commit()
    assert client.get('/statistics/reports/'+key).json()['snapshot']==snapshot
    replacement=client.post('/statistics/reports/'+key+'/replacements',json={'expected_version':2,'reason':'Synthetic source correction'})
    assert replacement.status_code==201,replacement.text
    new=replacement.json()
    assert new['report_id']!=key and new['predecessor_id']==key and new['state']=='saved'
    assert new['snapshot']['metrics'][0]['value']=='2'
    old=client.get('/statistics/reports/'+key).json()
    assert old['version']==3 and old['state']=='confirmed' and old['snapshot']==snapshot
    assert old['successor_id']==new['report_id']
    assert client.post('/statistics/reports/'+key+'/replacements',json={'expected_version':3,'reason':'Second successor'}).status_code==409
    history=client.get('/statistics/reports/'+key+'/history').json()
    assert [item['action'] for item in history['items']]==['saved','confirmed','replaced']


@pytest.mark.parametrize('mutation',['new_match','new_patient','date_change','delete','new_unknown'])
def test_confirmation_recomputes_full_membership_and_exclusions(statistics_api,mutation):
    client,engine,refs=statistics_api
    row=save(client)
    with Session(engine) as db:
        case=db.get(EmergencyCase,refs['case_id'])
        if mutation=='new_match':db.add(EmergencyCase(source_case_key='inserted',call_date=date(2026,1,3)))
        elif mutation=='new_patient':db.add(EmergencyPatient(emergency_case_id=case.emergency_case_id,patient_number=2))
        elif mutation=='date_change':case.call_date=date(2026,2,1)
        elif mutation=='delete':db.delete(db.get(EmergencyPatient,refs['patient_id']))
        elif mutation=='new_unknown':db.add(EmergencyCase(source_case_key='unknown',call_date=None))
        db.commit()
    response=client.post('/statistics/reports/'+row['report_id']+'/confirm',json={'expected_version':1,'acknowledged':True,'review_note':'Synthetic'})
    assert response.status_code==409,response.text
    assert 'content_hash' not in response.text and refs['case_id'] not in response.text


def test_complete_source_authorization_filters_reports_before_pagination(statistics_api):
    client,engine,refs=statistics_api
    first=save(client)
    save(client,{**QUERY,'metric_keys':['emergency.cases','fleet.trips']})
    third=save(client)
    revoke(engine,refs,'fleet.aggregate')
    page=client.get('/statistics/reports?limit=1&offset=1')
    assert page.status_code==200,page.text
    assert page.json()['total']==2 and len(page.json()['items'])==1
    assert {page.json()['items'][0]['report_id']} <= {first['report_id'],third['report_id']}
    revoke(engine,refs,'emergency.report.read')
    assert client.get('/statistics/reports/'+first['report_id']).status_code==403
    assert client.get('/statistics/reports/'+first['report_id']+'/history').status_code==403
    assert client.get('/statistics/reports/'+first['report_id']+'/export').status_code==403
    assert client.get('/statistics/reports').json()['total']==0


def test_saved_exports_are_exact_generic_safe_and_source_authorized(statistics_api):
    from openpyxl import load_workbook
    client,engine,refs=statistics_api
    row=save(client);key=row['report_id']
    client.post('/statistics/reports/'+key+'/confirm',json={'expected_version':1,'acknowledged':True,'review_note':'=PRIVATE_NOTE()'})
    with Session(engine) as db:
        db.add(EmergencyCase(source_case_key='after-export',call_date=date(2026,1,3)));db.commit()
    csv=client.get('/statistics/reports/'+key+'/export?format=csv')
    assert csv.status_code==200,csv.text
    assert '汎用統計出力（正式様式ではありません）' in csv.content.decode('utf-8-sig')
    assert 'PRIVATE_NOTE' not in csv.text and refs['case_id'] not in csv.text
    xlsx=client.get('/statistics/reports/'+key+'/export?format=xlsx')
    assert xlsx.status_code==200,xlsx.text
    workbook=load_workbook(BytesIO(xlsx.content),data_only=False)
    flat=[cell.value for sheet in workbook for cells in sheet for cell in cells]
    assert '汎用統計出力（正式様式ではありません）' in flat
    assert not any(cell.data_type=='f' for sheet in workbook for cells in sheet for cell in cells)
    assert '1' in flat and 'PRIVATE_NOTE' not in str(flat)
    revoke(engine,refs,'emergency.report.export')
    assert client.get('/statistics/reports/'+key+'/export?format=csv').status_code==403
    with Session(engine) as db:
        logs=list(db.scalars(select(AuditLog).where(AuditLog.action.like('statistics.%'))))
        assert logs
        for log in logs:
            assert refs['patient_id'] not in json.dumps(log.after_data) and refs['case_id'] not in json.dumps(log.after_data)
            assert 'fingerprint' not in json.dumps(log.after_data) and 'PRIVATE' not in json.dumps(log.after_data)


def test_aggregate_drilldown_denied_then_original_source_permissions_allow(statistics_api):
    client,engine,refs=statistics_api
    row=save(client);url='/statistics/reports/'+row['report_id']+'/drilldown?metric_key=emergency.patient_records'
    assert client.get(url).status_code==403
    with Session(engine) as db:
        for code in ['emergency.case.read','emergency.patient.read']:
            permission=db.scalar(select(Permission).where(Permission.code==code))
            db.add(RolePermission(role_id=refs['role_id'],permission_id=permission.permission_id))
        db.commit()
    response=client.get(url)
    assert response.status_code==200,response.text
    assert response.json()['items'][0]['record_id']==refs['patient_id']
    assert response.json()['items'][0]['parent_id']==refs['case_id']
    assert 'PRIVATE_MEDICAL' not in response.text


def test_sql_updates_cannot_rewrite_snapshot_evidence_or_history(statistics_api):
    from sqlalchemy.exc import DBAPIError
    client,engine,refs=statistics_api
    row=save(client)
    for statement in ["UPDATE statistics_reports SET snapshot='{}'", "UPDATE statistics_evidence SET evidence='{}'", "DELETE FROM statistics_reports", "DELETE FROM statistics_history", "UPDATE statistics_history SET action='tampered'", "DELETE FROM statistics_evidence"]:
        with pytest.raises(DBAPIError):
            with engine.begin() as db:db.execute(text(statement))
    assert client.get('/statistics/reports/'+row['report_id']).json()['snapshot']==row['snapshot']


@pytest.mark.parametrize('loss',['permission','session','account','password'])
def test_post_capture_authority_loss_blocks_release(statistics_api,monkeypatch,loss):
    from app import statistics_service as service
    client,engine,refs=statistics_api
    real=service.capture_in_snapshot
    def revoke_after(*args,**kwargs):
        result=real(*args,**kwargs)
        if loss=='permission':revoke(engine,refs,'emergency.report.read')
        else:
            with Session(engine) as db:
                if loss=='session':db.scalar(select(UserSession)).revoked_at=datetime.now(timezone.utc)
                elif loss=='account':db.get(User,refs['user_id']).active=False
                elif loss=='password':db.get(User,refs['user_id']).password_expires_at=datetime(2020,1,1,tzinfo=timezone.utc)
                db.commit()
        return result
    monkeypatch.setattr(service,'capture_in_snapshot',revoke_after)
    response=client.post('/statistics/query',json=QUERY)
    assert response.status_code in (401,403),response.text


def test_post_export_preparation_revocation_blocks_bytes(statistics_api,monkeypatch):
    from app import statistics_exports
    client,engine,refs=statistics_api
    row=save(client)
    real=statistics_exports.render_export
    def revoke_after(*args,**kwargs):
        result=real(*args,**kwargs);revoke(engine,refs,'emergency.report.export');return result
    monkeypatch.setattr(statistics_exports,'render_export',revoke_after)
    assert client.get('/statistics/reports/'+row['report_id']+'/export?format=csv').status_code==403


def test_department_bound_capture_uses_verified_engine_not_connection_as_factory(statistics_api,tmp_path,monkeypatch):
    from uuid import uuid4
    from app.settings import Settings,settings
    from app.tenant import initialize_tenant
    client,engine,refs=statistics_api
    config=Settings(_env_file=None,tenant_id=str(uuid4()),trusted_hosts=['testserver'],storage_root=str(tmp_path/'department'))
    initialize_tenant(engine,config,'Synthetic department',adopt_existing=True)
    monkeypatch.setattr(settings,'tenant_id',config.tenant_id)
    monkeypatch.setattr(settings,'trusted_hosts',config.trusted_hosts)
    monkeypatch.setattr(settings,'storage_root',config.storage_root)
    result=client.post('/statistics/query',json=QUERY)
    assert result.status_code==200,result.text


def test_confirmation_command_gate_precedes_report_existence_lookup(statistics_api):
    client,engine,refs=statistics_api
    row=save(client)
    revoke(engine,refs,'statistics.record')
    from uuid import uuid4
    payload={'expected_version':1,'acknowledged':True,'review_note':'Synthetic'}
    for key in [row['report_id'],str(uuid4())]:
        assert client.post('/statistics/reports/'+key+'/confirm',json=payload).status_code==403


def test_confirmation_keeps_pinned_timezone_after_server_setting_change(statistics_api,monkeypatch):
    from app.settings import settings
    client,engine,refs=statistics_api
    row=save(client)
    monkeypatch.setattr(settings,'statistics_business_timezone','America/New_York')
    response=client.post('/statistics/reports/'+row['report_id']+'/confirm',json={'expected_version':1,'acknowledged':True,'review_note':'Synthetic'})
    assert response.status_code==200,response.text
    assert response.json()['snapshot']['period']==row['snapshot']['period']


def test_two_request_factories_do_not_use_global_engine_or_accept_tenant_selector(statistics_api,tmp_path):
    from app.routers import statistics
    from app.statistics_models import StatisticsReport
    client,first,refs=statistics_api
    first_saved=save(client)
    second=create_engine('sqlite+pysqlite:///'+str(tmp_path/'second-department.db'),connect_args={'check_same_thread':False})
    Base.metadata.create_all(second)
    with Session(first) as original, Session(second) as db:
        roles=seed_rbac(db)
        # Matching user ID does not make the first department's session valid here.
        user=User(user_id=refs['user_id'],username='second-user',password_hash=hash_password('synthetic-statistics-password'));db.add(user);db.flush()
        role=Role(code='second-statistics',name='Synthetic second role');db.add(role);db.flush()
        for code in AGGREGATE:
            permission=db.scalar(select(Permission).where(Permission.code==code))
            if permission is None:permission=Permission(code=code);db.add(permission);db.flush()
            db.add(RolePermission(role_id=role.role_id,permission_id=permission.permission_id))
        db.add(UserRole(user_id=user.user_id,role_id=role.role_id))
        db.add_all([EmergencyCase(source_case_key='second-'+str(n),call_date=date(2026,1,3)) for n in range(2)])
        snapshot=json.loads(json.dumps(first_saved['snapshot']));snapshot['metrics'][0]['value']='2'
        db.add(StatisticsReport(report_id=first_saved['report_id'],snapshot=snapshot,metric_keys=QUERY['metric_keys'],created_by=user.user_id))
        db.commit()
    application=FastAPI();application.include_router(auth.router);application.include_router(statistics.router)
    def second_db():
        with Session(second,expire_on_commit=False) as db:yield db
    application.dependency_overrides[get_db]=second_db
    with TestClient(application) as other:
        other.cookies.update(client.cookies)
        assert other.get('/statistics/reports/'+first_saved['report_id']).status_code==401
        assert other.post('/auth/login',json={'username':'second-user','password':'synthetic-statistics-password'}).status_code==200
        assert other.post('/statistics/query',json=QUERY).json()['metrics'][0]['value']=='2'
        assert other.get('/statistics/reports/'+first_saved['report_id']).json()['snapshot']['metrics'][0]['value']=='2'
        assert client.get('/statistics/reports/'+first_saved['report_id']).json()['snapshot']['metrics'][0]['value']=='1'
        assert client.post('/statistics/query',json={**QUERY,'tenant_id':'second'}).status_code==422
    second.dispose()


def test_export_formula_safety_applies_to_every_metadata_and_value_cell():
    from app.statistics_exports import render_export
    from openpyxl import load_workbook
    from app.statistics_sources import DEFINITIONS
    metric={**DEFINITIONS['emergency.cases'],'value':'=1+1','status':'observed','coverage_status':'unknown','exclusions':[],'limitations':['@unsafe']}
    report={'report_id':'=REPORT','version':1,'state':'saved','confirmed_at':None,'snapshot':{'metrics':[metric],'coverage_status':'unknown','captured_at':'@time','query_version':'+query','public_checksum':'-checksum','period':{'start_date':'=date'}}}
    csv,_=render_export(report,'csv')
    assert "'=REPORT" in csv.decode('utf-8-sig') and "'=1+1" in csv.decode('utf-8-sig')
    raw,_=render_export(report,'xlsx')
    workbook=load_workbook(BytesIO(raw),data_only=False)
    assert not any(cell.data_type=='f' for sheet in workbook for row in sheet for cell in row)


def test_server_generated_lifecycle_times_keep_explicit_utc_across_save_get_confirm_history(statistics_api):
    from datetime import timedelta
    client,engine,refs=statistics_api
    saved=save(client)
    path='/statistics/reports/'+saved['report_id']
    loaded=client.get(path).json()
    response=client.post(path+'/confirm',json={'expected_version':1,'acknowledged':True,'review_note':'Synthetic lifecycle timestamp check'})
    assert response.status_code==200,response.text
    confirmed=response.json()
    reloaded=client.get(path).json()
    history=client.get(path+'/history').json()['items']
    times={**{stage+'.created_at':row['created_at'] for stage,row in [('save',saved),('get',loaded),('confirm',confirmed),('confirmed_get',reloaded)]},
           'confirm.confirmed_at':confirmed['confirmed_at'],'confirmed_get.confirmed_at':reloaded['confirmed_at'],
           **{'history.'+item['action']:item['occurred_at'] for item in history}}
    ambiguous={key:value for key,value in times.items() if datetime.fromisoformat(value).tzinfo is None}
    assert not ambiguous,ambiguous
    assert all(datetime.fromisoformat(value).utcoffset()==timedelta(0) for value in times.values())
    assert saved['created_at']==loaded['created_at']==confirmed['created_at']==reloaded['created_at']
    assert confirmed['confirmed_at']==reloaded['confirmed_at']
    assert saved['confirmed_at'] is None and loaded['confirmed_at'] is None
    assert all(row['snapshot']==saved['snapshot'] for row in [loaded,confirmed,reloaded])


def test_lifecycle_public_serialization_converts_aware_server_times_to_utc():
    from types import SimpleNamespace
    from app.statistics_reports import public_report
    row=SimpleNamespace(report_id='synthetic',version=2,state='confirmed',predecessor_id=None,successor_id=None,
        created_at=datetime.fromisoformat('2026-10-07T18:00:00+09:00'),
        confirmed_at=datetime.fromisoformat('2026-10-07T19:00:00+09:00'),snapshot={'untouched':'synthetic'})
    public=public_report(row)
    assert public['created_at']=='2026-10-07T09:00:00+00:00'
    assert public['confirmed_at']=='2026-10-07T10:00:00+00:00'
    assert row.created_at.isoformat()=='2026-10-07T18:00:00+09:00'

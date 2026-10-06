import os
os.environ['FIRE_AI_DATABASE_URL'] = 'sqlite+pysqlite:///:memory:'

from datetime import date
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from app.db import Base, engine, SessionLocal
from app.main import app
from app.models import Employee, User, Role, Permission, RolePermission, UserRole, EmergencyCase, EmergencyPatient, AuditLog
from app.security import hash_password

@pytest.fixture
def client():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        e=Employee(display_name='synthetic'); db.add(e); db.flush()
        u=User(employee_id=e.employee_id, username='run-b', password_hash=hash_password('synthetic-password'))
        r=Role(code='run-b',name='Run B'); db.add_all([u,r]); db.flush()
        db.add(UserRole(user_id=u.user_id,role_id=r.role_id))
        for code in ['emergency.case.read','emergency.case.create','emergency.case.update','emergency.patient.read','emergency.patient.create','emergency.patient.update','emergency.clinical.review','emergency.clinical.generate','emergency.report.read','emergency.crew.read','emergency.crew.manage','emergency.report.create','emergency.report.review','emergency.import','search.use']:
            p=Permission(code=code); db.add(p); db.flush(); db.add(RolePermission(role_id=r.role_id,permission_id=p.permission_id))
        db.commit()
    with TestClient(app) as c:
        assert c.post('/auth/login',json={'username':'run-b','password':'synthetic-password'}).status_code==200
        yield c

def make_patient(client):
    c=client.post('/emergency/cases',json={'call_date':'2026-10-01','station_code':'S','dispatch_number':'1','call_time':'23:50','dispatch_time':'23:55','scene_arrival_time':'00:10'} )
    assert c.status_code==201, c.text
    cid=c.json()['emergency_case_id']
    p=client.post(f'/emergency/cases/{cid}/patients',json={'patient_number':1,'hospital_code':'H','severity_code':'unknown','diagnosis_text':'アレルギー疑い'})
    assert p.status_code==201, p.text
    return cid,p.json()['emergency_patient_id']

def test_candidate_requires_human_and_cannot_overwrite_source(client):
    cid,pid=make_patient(client)
    t=client.post(f'/emergency/patients/{pid}/treatments',json={'code':'cpr','performed_at':'2026-10-02T00:15:00+09:00','notes':'synthetic'})
    assert t.status_code==201, t.text
    result=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1})
    assert result.status_code==200, result.text
    flags=result.json()
    assert {f['flag_type'] for f in flags}=={'cpa','allergy'}
    summary=client.get('/emergency/clinical-statistics?year=2026&month=10').json()
    assert summary['verified_clinical_counts']=={}
    f=next(f for f in flags if f['flag_type']=='cpa')
    assert client.post(f"/emergency/clinical-flags/{f['flag_id']}/review",json={'expected_version':1,'decision':'confirmed'}).status_code==200
    assert client.post(f"/emergency/clinical-flags/{f['flag_id']}/review",json={'expected_version':1,'decision':'rejected'}).status_code==409
    summary=client.get('/emergency/clinical-statistics?year=2026&month=10').json()
    assert summary['verified_clinical_counts']=={'cpa':1}
    with SessionLocal() as db:
        assert db.get(EmergencyPatient,pid).diagnosis_text=='アレルギー疑い'
        assert db.scalar(select(AuditLog).where(AuditLog.action=='emergency.clinical.review')) is not None

def test_stale_candidate_and_patient_version_conflict(client):
    cid,pid=make_patient(client)
    flags=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1}).json()
    change=client.patch(f'/emergency/patients/{pid}',json={'expected_version':1,'diagnosis_text':'確認後訂正'})
    assert change.status_code==200 and change.json()['version']==2
    assert client.patch(f'/emergency/patients/{pid}',json={'expected_version':1,'age':45}).status_code==409
    assert client.post(f"/emergency/clinical-flags/{flags[0]['flag_id']}/review",json={'expected_version':1,'decision':'confirmed'}).status_code==409

def test_checks_ambiguous_midnight_and_exports(client):
    cid,pid=make_patient(client)
    r=client.get(f'/emergency/cases/{cid}/checks')
    assert r.status_code==200, r.text
    assert any(x['code']=='time_day_offset_unresolved' for x in r.json()['issues'])
    assert client.get('/emergency/clinical-statistics?year=2026&month=13').status_code==422

def test_permissions_prevent_patient_detail_and_candidate_generation(client):
    cid,pid=make_patient(client)
    with SessionLocal() as db:
        p=db.scalar(select(Permission).where(Permission.code=='emergency.patient.read'))
        for link in db.scalars(select(RolePermission).where(RolePermission.permission_id==p.permission_id)):
            db.delete(link)
        db.commit()
    assert client.get(f'/emergency/cases/{cid}/patients').status_code==403
    assert client.get(f'/emergency/cases/{cid}/checks').status_code==403
    assert client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1}).status_code==403
    assert client.get('/emergency/clinical-statistics?year=2026').status_code==200

def test_report_review_snapshot_is_immutable(client):
    cid,pid=make_patient(client)
    draft=client.post(f'/emergency/cases/{cid}/reports',json={'kind':'post_review','content':{'assessment':'synthetic draft'}})
    assert draft.status_code==201, draft.text
    rid=draft.json()['report_id']
    review=client.post(f'/emergency/reports/{rid}/review',json={'expected_version':1,'decision':'confirmed'})
    assert review.status_code==200, review.text
    assert review.json()['source_snapshot']['case']['version']==1
    assert client.post(f'/emergency/reports/{rid}/review',json={'expected_version':2,'decision':'rejected'}).status_code==409

def test_duplicate_candidate_idempotency_and_negation(client):
    cid,pid=make_patient(client)
    first=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1}).json()
    second=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1}).json()
    assert [x['flag_id'] for x in first]==[x['flag_id'] for x in second]
    assert client.patch(f'/emergency/patients/{pid}',json={'expected_version':1,'diagnosis_text':'アレルギーなし。CPA否定'}).status_code==200
    candidates=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':2}).json()
    assert candidates==[]

def test_summary_never_discloses_source_ids_without_detail_permissions(client):
    make_patient(client)
    with SessionLocal() as db:
        for code in ['emergency.case.read','emergency.patient.read']:
            p=db.scalar(select(Permission).where(Permission.code==code))
            for link in db.scalars(select(RolePermission).where(RolePermission.permission_id==p.permission_id)): db.delete(link)
        db.commit()
    result=client.get('/emergency/clinical-statistics?year=2026').json()
    assert 'source_case_ids' not in result and 'source_patient_ids' not in result
    assert client.get('/search?q=S&modules=emergency').status_code==403


def test_emergency_search_uses_existing_case_and_never_patient_text(client):
    cid,pid=make_patient(client)
    assert client.patch(f'/emergency/cases/{cid}',json={'expected_version':1,'incident_address':'synthetic-address'}).status_code==200
    result=client.get('/search?q=synthetic-address&modules=emergency')
    assert result.status_code==200, result.text
    assert result.json()['hits'][0]['source_id']==cid
    assert 'アレルギー' not in result.text


def test_treatment_change_stales_previously_confirmed_candidate(client):
    cid,pid=make_patient(client)
    flag=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1}).json()[0]
    assert client.post(f"/emergency/clinical-flags/{flag['flag_id']}/review",json={'expected_version':1,'decision':'confirmed'}).status_code==200
    assert client.post(f'/emergency/patients/{pid}/treatments',json={'code':'oxygen'}).status_code==201
    stats=client.get('/emergency/clinical-statistics?year=2026').json()
    assert stats['verified_clinical_counts']=={} and stats['stale_confirmed_flags']==1


def test_report_captures_treatments_and_detects_stale_snapshot(client):
    cid,pid=make_patient(client)
    draft=client.post(f'/emergency/cases/{cid}/reports',json={'kind':'lifesaving_record'}).json()
    assert 'treatments' in draft['source_snapshot']
    assert client.post(f'/emergency/patients/{pid}/treatments',json={'code':'cpr'}).status_code==201
    assert client.post(f"/emergency/reports/{draft['report_id']}/review",json={'expected_version':1,'decision':'confirmed'}).status_code==409


def test_ui_emergency_surface_shares_existing_shell():
    from pathlib import Path
    html=(Path(__file__).resolve().parents[2]/'frontend/index.html').read_text()
    assert 'emergencyBtn' in html and 'openEmergency' in html
    js=(Path(__file__).resolve().parents[2]/'frontend/emergency.js').read_text()
    assert '/clinical-candidates' in js and '/review' in js and '/checks' in js

def test_import_audited_workbook_into_existing_shared_models(client):
    from openpyxl import Workbook
    from io import BytesIO
    w=Workbook(); w.remove(w.active)
    case=w.create_sheet('CSV_事案台帳');case.append(['覚知年月','署所ｺｰﾄﾞ','出場番号','覚知年月日']);case.append(['202610','S','50','2026-10-01'])
    patient=w.create_sheet('CSV_救護者台帳');patient.append(['覚知年月','署所ｺｰﾄﾞ','出場番号','救護者番号','病院ｺｰﾄﾞ']);patient.append(['202610','S','50','1','H'])
    crew=w.create_sheet('CSV_出動隊員');crew.append(['覚知年月','署所ｺｰﾄﾞ','出場番号','隊員種別','隊員ｺｰﾄﾞ'])
    out=BytesIO();w.save(out)
    result=client.post('/emergency/import/workbook',files={'file':('synthetic.xlsx',out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert result.status_code==200, result.text
    assert result.json()['cases_inserted']==1 and result.json()['patients_inserted']==1
    again=client.post('/emergency/import/workbook',files={'file':('synthetic.xlsx',out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert again.status_code==200 and again.json()['status']=='already_imported'
    assert client.get('/emergency/reports/summary').json()['totals']['patients']==1


def test_common_employee_reference_and_rbac_no_cache(client):
    cid,pid=make_patient(client)
    with SessionLocal() as db:
        employee=db.scalar(select(Employee))
        eid=employee.employee_id
    assigned=client.post(f'/emergency/cases/{cid}/crew',json={'employee_id':eid,'crew_role':'leader'})
    assert assigned.status_code==201, assigned.text
    assert assigned.json()['employee_id']==eid
    assert client.post(f'/emergency/cases/{cid}/crew',json={'employee_id':eid,'crew_role':'leader'}).status_code==409
    assert client.get(f'/emergency/cases/{cid}').headers['Cache-Control']=='no-store'

def test_candidate_fingerprint_has_database_uniqueness(client):
    from app.models import EmergencyClinicalFlag
    from sqlalchemy.exc import IntegrityError
    cid,pid=make_patient(client)
    flags=client.post(f'/emergency/patients/{pid}/clinical-candidates',json={'expected_version':1}).json()
    assert len(flags[0]['candidate_fingerprint'])==64
    with SessionLocal() as db:
        existing=db.get(EmergencyClinicalFlag,flags[0]['flag_id'])
        duplicate=EmergencyClinicalFlag(emergency_patient_id=pid,flag_type=existing.flag_type,flag_value='candidate',derivation_method='rule',candidate_fingerprint=existing.candidate_fingerprint)
        db.add(duplicate)
        with pytest.raises(IntegrityError):db.flush()
        db.rollback()

def test_existing_module_identity_is_extended_not_duplicated(client):
    from app.module_seed import seed_modules
    with SessionLocal() as db:
        modules=seed_modules(db); original=modules['emergency_reporting'].module_id;db.commit()
        modules=seed_modules(db);db.commit()
        assert modules['emergency_reporting'].module_id==original
        assert modules['emergency_reporting'].version=='1.1.0'
        assert modules['emergency_reporting'].manifest['operational_api']=='/emergency'

def test_infant_age_zero_survives_real_frontend_field_render():
    from pathlib import Path
    import subprocess
    js=Path(__file__).resolve().parents[2]/'frontend/emergency.js'
    result=subprocess.run(['node','-e',"const vm=require('vm'),fs=require('fs');const ctx=vm.createContext({});vm.runInContext(fs.readFileSync(process.argv[2],'utf8').match(/^const esc=.*$/m)[0],ctx);vm.runInContext(fs.readFileSync(process.argv[1],'utf8'),ctx);process.stdout.write(vm.runInContext(`emergencyFields([['age','年齢','number']],{age:0})`,ctx));",str(js),str(js.parent/'index.html')],capture_output=True,text=True,check=True)
    assert 'value="0"' in result.stdout

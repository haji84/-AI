"""Synthetic stock integrity and operational asset lifecycle regressions."""
import os
os.environ['FIRE_AI_DATABASE_URL']='sqlite+pysqlite:///:memory:'
from datetime import date
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from app.db import Base,engine,SessionLocal
from app.main import app
from app.models import Employee,User,Role,Permission,RolePermission,UserRole,Document,Facility,AuditLog
from app.security import hash_password
CODES=[f'asset.{p}' for p in ['read','create','update','review','approve','admin','import','export','borrower.read','borrower.manage']]+['search.use','document.read','facility.read','fleet.read','incident.read']
@pytest.fixture
def client():
    Base.metadata.drop_all(engine);Base.metadata.create_all(engine)
    with SessionLocal() as db:
        e=Employee(display_name='Synthetic PrivateBorrowerSecret',employee_code='SYN');db.add(e);db.flush()
        u=User(username='assets',employee_id=e.employee_id,password_hash=hash_password('synthetic-password'));r=Role(code='assets-test',name='Synthetic');db.add_all([u,r]);db.flush();db.add(UserRole(user_id=u.user_id,role_id=r.role_id))
        for code in CODES:
            p=Permission(code=code);db.add(p);db.flush();db.add(RolePermission(role_id=r.role_id,permission_id=p.permission_id))
        db.commit()
    with TestClient(app) as c:
        assert c.post('/auth/login',json={'username':'assets','password':'synthetic-password'}).status_code==200
        yield c

def post(c,path,data,status=201):
    r=c.post('/assets'+path,json=data);assert r.status_code==status,r.text;return r.json()
def remove(code):
    with SessionLocal() as db:
        p=db.scalar(select(Permission).where(Permission.code==code))
        for link in db.scalars(select(RolePermission).where(RolePermission.permission_id==p.permission_id)):db.delete(link)
        db.commit()
def setup_stock(c,category='drug',expiry='2099-01-01',quantity='10.125'):
    a=post(c,'/registry',{'code':'SYN','name':'Synthetic stock','category':category,'unit':'vial','reorder_threshold':'5'})
    l=post(c,'/locations',{'code':'A','name':'Synthetic store'})
    lot=post(c,'/registry/'+a['asset_id']+'/lots',{'expected_version':1,'batch_code':'B1','expires_on':expiry,'provenance':'Synthetic source'})
    m=post(c,'/movements',{'asset_id':a['asset_id'],'expected_version':2,'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'receive','quantity':quantity,'idempotency_key':'receive-1','reason':'Synthetic receipt'})
    return c.get('/assets/registry/'+a['asset_id']).json(),l,lot,m

def test_registry_crud_stale_rbac_audit_and_decimal_validation(client):
    a=post(client,'/registry',{'code':'S','name':'Synthetic','category':'durable','unit':'piece','reorder_threshold':'0'})
    r=client.patch('/assets/registry/'+a['asset_id'],json={'expected_version':1,'name':'Updated'});assert r.status_code==200 and r.json()['version']==2 and r.json()['asset_id']==a['asset_id']
    assert client.patch('/assets/registry/'+a['asset_id'],json={'expected_version':1,'name':'Stale'}).status_code==409
    for value in ['-1','NaN','Infinity','0.00001']:
        assert client.post('/assets/registry',json={'code':value,'name':'Invalid','category':'drug','unit':'vial','reorder_threshold':value}).status_code==422
    assert client.post('/assets/registry',json={'code':'T','name':'Spoof','category':'drug','unit':'vial','tenant_id':'chosen'}).status_code==422
    assert client.get('/assets/registry').headers['cache-control']=='no-store'
    with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='asset.registry.update'))
    remove('asset.update');assert client.patch('/assets/registry/'+a['asset_id'],json={'expected_version':2,'name':'Denied'}).status_code==403

def test_lot_transfer_conservation_underflow_and_idempotency(client):
    a,l,lot,m=setup_stock(client);other=post(client,'/locations',{'code':'B','name':'Synthetic vehicle bay'})
    data={'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':lot['lot_id'],'location_id':l['location_id'],'to_location_id':other['location_id'],'kind':'transfer','quantity':'3.125','idempotency_key':'transfer-1','reason':'Synthetic transfer'}
    first=post(client,'/movements',data);same=post(client,'/movements',data);assert same['movement_id']==first['movement_id']
    assert client.post('/assets/movements',json={**data,'quantity':'1'}).status_code==409
    rows=client.get('/assets/balances?asset_id='+a['asset_id']).json();assert {r['quantity'] for r in rows}=={'7.000','3.125'} and {r['lot_id'] for r in rows}=={lot['lot_id']}
    bad={**data,'expected_version':a['version']+1,'kind':'issue','quantity':'8','idempotency_key':'underflow'};bad.pop('to_location_id')
    assert client.post('/assets/movements',json=bad).status_code==422
    assert len(client.get('/assets/registry/'+a['asset_id']+'/history').json()['movements'])==2
    assert client.post('/assets/registry/'+a['asset_id']+'/retire',json={'expected_version':a['version']+1,'reason':'Stock present'}).status_code==409

def test_loan_return_outstanding_employee_and_privacy(client):
    a,l,lot,m=setup_stock(client)
    with SessionLocal() as db:eid=db.scalar(select(Employee.employee_id))
    data={'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'loan','quantity':'2.125','borrower_employee_id':eid,'idempotency_key':'loan-1','reason':'Synthetic loan','due_on':'2099-01-01'}
    loan=post(client,'/movements',data);key=loan['loan_id'];r=client.get('/assets/loans/'+key).json();assert r['outstanding_quantity']=='2.125'
    back={'asset_id':a['asset_id'],'expected_version':a['version']+1,'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'return','quantity':'3','loan_id':key,'idempotency_key':'return-over','reason':'Return'}
    assert client.post('/assets/movements',json=back).status_code==422
    returned=post(client,'/movements',{**back,'quantity':'1.125','idempotency_key':'return-part'});assert returned['loan_id']==key
    assert client.get('/assets/loans/'+key).json()['outstanding_quantity']=='1.000'
    remove('asset.borrower.read');assert client.get('/assets/loans').status_code==403
    history=client.get('/assets/registry/'+a['asset_id']+'/history').json();assert eid not in str(history) and 'borrower_employee_id' not in str(history)
    assert client.get('/search?q=PrivateBorrowerSecret&modules=operational_assets').json()['hits']==[]
    assert eid not in client.get('/assets/export/movements').text

def test_inactive_missing_employee_links_and_document_permissions(client):
    from app.operations_models import Vehicle,Incident
    with SessionLocal() as db:
        e=Employee(display_name='Synthetic inactive borrower',active=False);db.add(e);db.flush();eid=e.employee_id
        f=Facility(name='Synthetic facility');v=Vehicle(code='V',name='Synthetic vehicle');i=Incident(kind='other',title='Synthetic incident');d=Document(original_filename='Synthetic.pdf',mime_type='application/pdf',storage_path='synthetic',sha256='a'*64,size_bytes=1);db.add_all([f,v,i,d]);db.commit();fid,vid,iid,did=f.building_id,v.vehicle_id,i.incident_id,d.document_id
    location=post(client,'/locations',{'code':'LINK','name':'Linked','building_id':fid,'vehicle_id':vid});assert location['building_id']==fid
    a,l,lot,m=setup_stock(client)
    for employee in [eid,str(uuid4())]:
        r=client.post('/assets/movements',json={'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'loan','quantity':'1','borrower_employee_id':employee,'idempotency_key':employee,'reason':'Loan'});assert r.status_code in [404,422]
    x=post(client,'/movements',{'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'issue','quantity':'1','handler_employee_id':None,'incident_id':iid,'document_id':did,'idempotency_key':'linked','reason':'Synthetic issue'});assert x['incident_id']==iid and x['document_id']==did
    remove('document.read');assert did not in str(client.get('/assets/registry/'+a['asset_id']+'/history').json())
    assert client.post('/assets/registry',json={'code':'DOC','name':'Restricted','category':'durable','unit':'piece','document_id':did}).status_code==403

def test_expiry_is_lot_specific_and_ordinary_issue_rejected(client):
    a,l,lot,m=setup_stock(client,expiry='2000-01-01')
    assert client.post('/assets/movements',json={'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'issue','quantity':'1','idempotency_key':'expired','reason':'Issue'}).status_code==422
    good=post(client,'/registry/'+a['asset_id']+'/lots',{'expected_version':a['version'],'batch_code':'B2','expires_on':'2099-01-01','provenance':'Different batch'})
    a=client.get('/assets/registry/'+a['asset_id']).json();post(client,'/movements',{'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':good['lot_id'],'location_id':l['location_id'],'kind':'receive','quantity':'2','idempotency_key':'good','reason':'Receipt'})
    assert len(client.get('/assets/registry').json())==1 and len(client.get('/assets/registry/'+a['asset_id']+'/lots').json())==2
    alerts=client.get('/assets/alerts?as_of=2026-10-06').json();assert any(x['kind']=='expiry' and x['lot_id']==lot['lot_id'] for x in alerts)

def test_service_human_gate_stale_source_disposal_and_immutable_history(client):
    a,l,lot,m=setup_stock(client,category='durable');key=a['asset_id']
    s=post(client,'/registry/'+key+'/services',{'expected_version':a['version'],'kind':'pressure_test','performed_on':'2026-10-01','description':'Synthetic pressure test','cost':'12.34','next_pressure_test_on':'2027-10-01'})
    assert client.get('/assets/registry/'+key).json()['next_pressure_test_on'] is None
    post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'reason':'Checked'},200)
    v=client.get('/assets/registry/'+key).json()['version']
    assert client.post('/assets/services/'+s['service_id']+'/approve',json={'expected_version':2,'expected_asset_version':v-1,'reason':'Stale'}).status_code==409
    approved=post(client,'/services/'+s['service_id']+'/approve',{'expected_version':2,'expected_asset_version':v,'reason':'Human approved'},200);assert approved['approved_by'] and approved['status']=='approved'
    assert client.patch('/assets/services/'+s['service_id'],json={'expected_version':3,'description':'Overwrite'}).status_code in [404,405]
    current=client.get('/assets/registry/'+key).json();assert current['next_pressure_test_on']=='2027-10-01'
    d=post(client,'/registry/'+key+'/services',{'expected_version':current['version'],'kind':'disposal','performed_on':'2026-10-06','description':'Synthetic write-off','cost':'0','lot_id':lot['lot_id'],'location_id':l['location_id'],'quantity':'10.125'})
    post(client,'/services/'+d['service_id']+'/review',{'expected_version':1,'reason':'Checked stock'},200)
    v=client.get('/assets/registry/'+key).json()['version'];post(client,'/services/'+d['service_id']+'/approve',{'expected_version':2,'expected_asset_version':v,'reason':'Human write-off'},200)
    assert client.get('/assets/balances?asset_id='+key).json()[0]['quantity']=='0.000'
    v=client.get('/assets/registry/'+key).json()['version'];retired=post(client,'/registry/'+key+'/retire',{'expected_version':v,'reason':'Human retirement'},200);assert not retired['active']
    with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='asset.service.approve'))

def test_deadline_reorder_candidates_and_review_source_conflict(client):
    a,l,lot,m=setup_stock(client,quantity='1')
    candidates=client.get('/assets/reorder').json();assert candidates[0]['asset_id']==a['asset_id'] and candidates[0]['candidate_quantity']=='4.000' and candidates[0]['status']=='candidate'
    s=post(client,'/registry/'+a['asset_id']+'/services',{'expected_version':a['version'],'kind':'calibration','performed_on':'2026-10-01','description':'Synthetic','cost':'0','next_calibration_on':'2026-10-05','next_use_on':'2026-10-03','next_service_on':'2026-10-04'})
    post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'reason':'Checked'},200)
    v=client.get('/assets/registry/'+a['asset_id']).json()['version'];assert client.patch('/assets/registry/'+a['asset_id'],json={'expected_version':v,'name':'Changed source'}).status_code==200
    assert client.post('/assets/services/'+s['service_id']+'/approve',json={'expected_version':2,'expected_asset_version':v+1,'reason':'Approve stale review'}).status_code==409

def test_real_csv_excel_roundtrip_dry_run_confirm_provenance_rollback(client):
    with SessionLocal() as db:
        doc=Document(original_filename='Synthetic.csv',storage_path='synthetic-import',sha256='b'*64,size_bytes=1);db.add(doc);db.commit();did=doc.document_id
    csv=b'schema_version,code,name,category,unit,reorder_threshold\nassets-v1,CSV,=Synthetic(),drug,vial,0\n'
    p=client.post('/assets/import/registry',data={'document_id':did},files={'file':('synthetic.csv',csv,'text/csv')});assert p.status_code==200,p.text;preview=p.json()
    assert client.get('/assets/registry').json()==[]
    path='/assets/import-previews/'+preview['preview_id']+'/confirm'
    assert client.post(path,json={'expected_version':1,'file_sha256':'0'*64}).status_code==409
    applied=client.post(path,json={'expected_version':1,'file_sha256':preview['file_sha256']});assert applied.status_code==200 and applied.json()['inserted']==1
    assert client.post(path,json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==409
    out=client.get('/assets/export/registry?format=csv');assert "'=Synthetic()" in out.text and 'assets-v1' in out.text and did in out.text and preview['file_sha256'] in out.text
    from openpyxl import Workbook,load_workbook
    w=Workbook();w.active.append(['schema_version','code','name','category','unit']);w.active.append(['assets-v1','XLSX','Synthetic workbook','consumable','box']);raw=BytesIO();w.save(raw)
    p=client.post('/assets/import/registry',files={'file':('misleading.bin',raw.getvalue(),'application/octet-stream')});assert p.status_code==200,p.text
    preview=p.json();assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    sheet=load_workbook(BytesIO(client.get('/assets/export/registry?format=xlsx').content)).active;assert sheet.max_row==3
    bad=b'schema_version,code,name,category,unit,reorder_threshold\nassets-v1,OK,Synthetic,durable,piece,0\nassets-v1,BAD,Synthetic,drug,vial,-1\n';assert client.post('/assets/import/registry',files={'file':('synthetic.csv',bad,'text/csv')}).status_code==422
    assert len(client.get('/assets/registry').json())==2
    assert client.get('/search?q=Synthetic workbook&modules=operational_assets').json()['hits'][0]['source_type']=='operational_asset'

def test_import_confirmation_rechecks_rolls_back_and_lot_exchange(client):
    a,l,lot,m=setup_stock(client)
    raw=f'schema_version,asset_id,expected_version,batch_code,expires_on,provenance\nassets-v1,{a["asset_id"]},{a["version"]},IMP,2099-01-01,Synthetic batch\n'.encode()
    p=client.post('/assets/import/lots',files={'file':('synthetic.csv',raw,'text/csv')});assert p.status_code==200,p.text
    preview=p.json();assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    assert 'IMP' in client.get('/assets/export/lots').text and '2099-01-01' in client.get('/assets/export/lots').text
    raw=b'schema_version,code,name,category,unit\nassets-v1,FIRST,Synthetic,durable,piece\nassets-v1,CONFLICT,Synthetic,durable,piece\n';preview=client.post('/assets/import/registry',files={'file':('synthetic.csv',raw,'text/csv')}).json()
    post(client,'/registry',{'code':'CONFLICT','name':'Concurrent','category':'durable','unit':'piece'})
    assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==409
    assert 'FIRST' not in [x['code'] for x in client.get('/assets/registry').json()]

def test_approval_permission_and_expired_writeoff(client):
    a,l,lot,m=setup_stock(client,expiry='2000-01-01')
    s=post(client,'/registry/'+a['asset_id']+'/services',{'expected_version':a['version'],'kind':'expiry_writeoff','performed_on':'2026-10-06','description':'Synthetic expired batch','lot_id':lot['lot_id'],'location_id':l['location_id'],'quantity':'1'})
    post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'reason':'Verified expiry'},200)
    v=client.get('/assets/registry/'+a['asset_id']).json()['version'];remove('asset.approve')
    assert client.post('/assets/services/'+s['service_id']+'/approve',json={'expected_version':2,'expected_asset_version':v,'reason':'Denied'}).status_code==403
    assert client.get('/assets/balances?asset_id='+a['asset_id']).json()[0]['quantity']=='10.125'

def test_nested_document_redaction_and_import_stock_update_gate(client):
    a,l,lot,m=setup_stock(client)
    with SessionLocal() as db:
        doc=Document(original_filename='Synthetic.pdf',storage_path='nested-proof',sha256='d'*64,size_bytes=1);db.add(doc);db.commit();did=doc.document_id
    s=post(client,'/registry/'+a['asset_id']+'/services',{'expected_version':a['version'],'kind':'repair','performed_on':'2026-10-01','description':'Synthetic repair','document_id':did})
    post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'reason':'Checked'},200)
    remove('document.read');assert did not in str(client.get('/assets/services/'+s['service_id']).json()) and did not in client.get('/assets/export/services').text
    v=client.get('/assets/registry/'+a['asset_id']).json()['version']
    raw=f'schema_version,asset_id,expected_version,lot_id,location_id,kind,quantity,idempotency_key,reason\nassets-v1,{a["asset_id"]},{v},{lot["lot_id"]},{l["location_id"]},issue,1,forbidden-import,Synthetic\n'.encode()
    remove('asset.update');assert client.post('/assets/import/movements',files={'file':('synthetic.csv',raw,'text/csv')}).status_code==403

def test_csv_export_roundtrip_keeps_zero_and_source_lineage(client):
    a=post(client,'/registry',{'code':'ZERO','name':'=Synthetic formula label','category':'durable','unit':'piece','reorder_threshold':'0'})
    raw=client.get('/assets/export/registry').content
    import csv
    rows=list(csv.DictReader(raw.decode('utf-8-sig').splitlines()));assert rows[0]['reorder_threshold']=='0.000'
    rows[0]['code']='ROUNDTRIP';out=__import__('io').StringIO();w=csv.DictWriter(out,fieldnames=rows[0].keys());w.writeheader();w.writerows(rows)
    p=client.post('/assets/import/registry',files={'file':('synthetic.csv',out.getvalue().encode(),'text/csv')});assert p.status_code==200,p.text
    preview=p.json();assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    assert next(x for x in client.get('/assets/registry').json() if x['code']=='ROUNDTRIP')['name']=='=Synthetic formula label'

def test_confirm_import_rechecks_update_permission_and_preserves_stock(client):
    a,l,lot,m=setup_stock(client)
    raw=f'schema_version,asset_id,expected_version,lot_id,location_id,kind,quantity,idempotency_key,reason\nassets-v1,{a["asset_id"]},{a["version"]},{lot["lot_id"]},{l["location_id"]},issue,1,preview-permission,Synthetic\n'.encode()
    p=client.post('/assets/import/movements',files={'file':('synthetic.csv',raw,'text/csv')});assert p.status_code==200,p.text
    preview=p.json();remove('asset.update');assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==403
    assert client.get('/assets/balances?asset_id='+a['asset_id']).json()[0]['quantity']=='10.125'

def test_shared_shell_assets_ui_escapes_zero_and_routes_human_review(tmp_path):
    import subprocess
    root=Path(__file__).resolve().parents[2]
    script=tmp_path/'synthetic-ui.cjs'
    script.write_text('''const fs=require('fs'),vm=require('vm'),assert=require('assert');
const nodes=new Map();const get=id=>{if(!nodes.has(id))nodes.set(id,{innerHTML:'',textContent:'',value:'',classList:{toggle(){},add(){},remove(){}}});return nodes.get(id)};
let requests=[];
const ctx={console,URLSearchParams,crypto:require('crypto').webcrypto,$:get,esc:x=>String(x??'').replaceAll('&','&amp;').replaceAll('<','&lt;').replaceAll('>','&gt;').replaceAll('"','&quot;'),document:{querySelectorAll:()=>[]},api:async(path)=>{requests.push(path);if(path==='/auth/permissions')return {permissions:['asset.read','asset.review','asset.approve']};if(path.includes('/history'))return {movements:[],services:[]};if(path.includes('/lots')||path.includes('/balances')||path.includes('/locations')||path.includes('/loans'))return [];return {asset_id:'synthetic',code:'S',name:'<script>unsafe</script>',category:'drug',unit:'vial',active:true,version:1,reorder_threshold:'0.000'}}};
vm.createContext(ctx);vm.runInContext(fs.readFileSync(process.argv[2],'utf8'),ctx);
(async()=>{await ctx.initAssets();await ctx.assetsDetail('synthetic');assert(get('assetsContent').innerHTML.includes('&lt;script&gt;unsafe'));assert(get('assetsContent').innerHTML.includes('0.000'));await ctx.assetsServiceAction({service_id:'service',version:2},'approve');assert(get('assetsContent').innerHTML.includes('Human'));assert(requests.some(p=>p==='/assets/registry/synthetic'));await ctx.assetsDetail('synthetic','UniqueBatch');assert(requests.some(p=>p.includes('/lots?q=UniqueBatch&')));const html=fs.readFileSync(process.argv[3],'utf8');vm.runInContext(html.slice(html.indexOf('function unifiedSearchAction(hit)'),html.indexOf('async function runUnifiedSearch')),ctx);ctx.closeUnifiedSearch=()=>{};ctx.openAssets=async()=>{};const action=ctx.unifiedSearchAction({module:'operational_assets',source_type:'asset_lot',source_id:'lot-child',navigation:{surface:'operational_assets',asset_id:'asset-parent'}});assert(action.includes('asset-parent'));await ctx.openUnifiedSearchHit('operational_assets','asset_lot','lot-child','','asset-parent');assert(requests.some(p=>p==='/assets/registry/asset-parent'));})().catch(e=>{console.error(e);process.exit(1)});
''')
    r=subprocess.run(['node',str(script),str(root/'frontend/assets.js'),str(root/'frontend/index.html')],capture_output=True,text=True);assert r.returncode==0,r.stderr

def test_expiry_enforcement_uses_japan_business_date_at_utc_boundary(client,monkeypatch):
    from datetime import datetime,timezone
    from app import assets_service as svc
    monkeypatch.setattr(svc,'now_utc',lambda:datetime(2026,10,6,16,0,tzinfo=timezone.utc))
    a,l,lot,m=setup_stock(client,expiry='2026-10-06')
    assert client.post('/assets/movements',json={'asset_id':a['asset_id'],'expected_version':a['version'],'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'issue','quantity':'1','idempotency_key':'midnight','reason':'Synthetic'}).status_code==422
    assert client.get('/assets/reorder').json()[0]['available_quantity']=='0.000'
    assert any(x['overdue'] for x in client.get('/assets/alerts').json() if x['lot_id']==lot['lot_id'])

def test_asset_stock_and_documents_isolate_server_bound_departments(tmp_path):
    from test_run_b_operations_tenant_integration import departments,login,DOCUMENT_ID
    from app.routers import assets
    with departments(tmp_path) as [(alpha,_,_,_), (beta,_,beta_cfg,_)]:
        for c in [alpha,beta]:c.app.include_router(assets.router)
        login(alpha,'alpha');token=alpha.cookies.get('fire_ai_session');assert beta.get('/assets/registry',headers={'Cookie':'fire_ai_session='+token}).status_code==401
        login(beta,'beta');a,l,lot,m=setup_stock(alpha)
        assert beta.get('/assets/registry/'+a['asset_id']).status_code==404
        assert beta.get('/assets/export/balances').status_code==200 and a['asset_id'] not in beta.get('/assets/export/balances').text
        assert beta.get('/search?q=SYN&modules=operational_assets').json()['hits']==[]
        r=alpha.get('/assets/registry/'+a['asset_id'],headers={'X-Tenant-ID':beta_cfg.tenant_id},params={'tenant_id':beta_cfg.tenant_id});assert r.status_code==200
        assert alpha.get('/assets/registry',headers={'Host':'beta.test'}).status_code==421
        assert alpha.post('/assets/registry',json={'code':'SPOOF','name':'Synthetic','category':'drug','unit':'vial','tenant_id':beta_cfg.tenant_id}).status_code==422
        raw=b'schema_version,code,name,category,unit,tenant_id\nassets-v1,SPOOF,Synthetic,drug,vial,beta\n';assert alpha.post('/assets/import/registry',files={'file':('synthetic.csv',raw,'text/csv')}).status_code==422

def test_postgresql_migration_locks_atomic_issue_duplicate_and_stale_version():
    """CI disposable PostgreSQL schema: applies real append-only 042 DDL and exercises locks."""
    url=os.environ.get('FIRE_AI_TEST_POSTGRES_URL')
    if not url:pytest.skip('actual PostgreSQL inventory lock/migration acceptance runs in CI')
    from sqlalchemy import create_engine,text,update,inspect
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.exc import DBAPIError
    from fastapi import HTTPException
    from app.assets_models import OperationalAsset,AssetLot,AssetLocation,AssetBalance,AssetMovement
    from app.assets_schemas import AssetInput,LocationInput,LotInput,MovementInput
    from app.assets_service import create_asset,create_location,create_lot,move
    from app.migrations import split_sql
    from app.rbac_seed import seed_rbac
    schema='synthetic_assets_'+uuid4().hex;pg=create_engine(url);scoped=pg.execution_options(schema_translate_map={None:schema})
    asset_names={'operational_assets','asset_locations','asset_lots','asset_balances','asset_loans','asset_movements','asset_services','asset_import_previews'}
    try:
        with pg.begin() as connection:connection.execute(text(f'CREATE SCHEMA "{schema}"'))
        Base.metadata.create_all(scoped,tables=[t for t in Base.metadata.tables.values() if t.name not in asset_names])
        with pg.begin() as connection:
            connection.execute(text(f'SET LOCAL search_path TO "{schema}"'))
            sql=(Path(__file__).resolve().parents[2]/'db/migrations/042_run_b_operational_assets.sql').read_text()
            for statement in split_sql(sql):connection.exec_driver_sql(statement)
        reflected=inspect(pg)
        assert asset_names<=set(reflected.get_table_names(schema=schema))
        assert {f['referred_table'] for f in reflected.get_foreign_keys('asset_movements',schema=schema)}>={'asset_lots','asset_locations','asset_loans','documents','employees','operation_incidents'}
        sessions=sessionmaker(bind=scoped,expire_on_commit=False,autoflush=False)
        with sessions() as db:
            roles=seed_rbac(db);operator=Employee(display_name='Synthetic PG operator',active=True);db.add(operator);db.flush();u=User(employee_id=operator.employee_id,username='synthetic-pg-assets',password_hash=hash_password('synthetic-password'));db.add(u);db.flush();db.add(UserRole(user_id=u.user_id,role_id=roles['system_admin'].role_id));db.commit();uid=u.user_id
            a=create_asset(db,u,AssetInput(code='PG',name='Synthetic PG inventory',category='drug',unit='vial'));l=create_location(db,u,LocationInput(code='A',name='Synthetic PG store'));lot=create_lot(db,u,a.asset_id,LotInput(expected_version=1,batch_code='PG-B1',provenance='Synthetic PG source',expires_on=date(2099,1,1)));db.commit();aid,lid,lotid=a.asset_id,l.location_id,lot.lot_id
            receipt=MovementInput(asset_id=aid,expected_version=2,lot_id=lotid,location_id=lid,kind='receive',quantity='2.125',idempotency_key='pg-receive',reason='Synthetic');move(db,u,receipt);db.commit()
        issue=MovementInput(asset_id=aid,expected_version=3,lot_id=lotid,location_id=lid,kind='issue',quantity='1.125',idempotency_key='pg-issue',reason='Synthetic')
        with sessions() as issuer:
            first=move(issuer,issuer.get(User,uid),issue)
            with scoped.connect() as contender:
                with pytest.raises(DBAPIError) as locked:
                    with contender.begin():
                        contender.execute(text("SET LOCAL lock_timeout = '250ms'"));contender.execute(update(AssetBalance).where(AssetBalance.asset_id==aid).values(quantity=0))
                assert getattr(locked.value.orig,'sqlstate',None)=='55P03'
            issuer.commit();movement_id=first.movement_id
        with sessions() as retry:
            same=move(retry,retry.get(User,uid),issue);assert same.movement_id==movement_id;retry.commit()
            stale=issue.model_copy(update={'idempotency_key':'pg-stale'})
            with pytest.raises(HTTPException) as conflict:move(retry,retry.get(User,uid),stale)
            assert conflict.value.status_code==409;retry.rollback()
            assert retry.scalar(select(AssetBalance).where(AssetBalance.asset_id==aid)).quantity==__import__('decimal').Decimal('1.000')
            assert len(retry.scalars(select(AssetMovement).where(AssetMovement.asset_id==aid)).all())==2
            assert retry.scalar(select(AuditLog).where(AuditLog.action=='asset.movement.issue')).entity_id==movement_id
    finally:
        with pg.begin() as connection:connection.execute(text(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE'))
        pg.dispose()

def test_dated_stock_import_and_return_chronology_preserve_occurrence(client):
    a,l,lot,m=setup_stock(client)
    raw=f'schema_version,asset_id,expected_version,lot_id,location_id,kind,quantity,occurred_on,idempotency_key,reason\nassets-v1,{a["asset_id"]},{a["version"]},{lot["lot_id"]},{l["location_id"]},receive,1.125,2025-04-01,dated-receipt,Synthetic historical receipt\n'.encode()
    p=client.post('/assets/import/movements',files={'file':('synthetic.csv',raw,'text/csv')});assert p.status_code==200,p.text;preview=p.json();assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    history=client.get('/assets/registry/'+a['asset_id']+'/history').json()['movements'];dated=next(x for x in history if x['idempotency_key']=='dated-receipt');assert dated['occurred_on']=='2025-04-01' and dated['created_at'][:10]!='2025-04-01'
    assert '2025-04-01' in client.get('/assets/export/movements').text
    with SessionLocal() as db:eid=db.scalar(select(Employee.employee_id))
    v=client.get('/assets/registry/'+a['asset_id']).json()['version']
    loan=post(client,'/movements',{'asset_id':a['asset_id'],'expected_version':v,'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'loan','quantity':'1','borrower_employee_id':eid,'occurred_on':'2026-10-01','due_on':'2026-10-06','idempotency_key':'dated-loan','reason':'Synthetic'})
    assert client.get('/assets/loans/'+loan['loan_id']).json()['loaned_on']=='2026-10-01'
    back={'asset_id':a['asset_id'],'expected_version':v+1,'lot_id':lot['lot_id'],'location_id':l['location_id'],'kind':'return','quantity':'1','loan_id':loan['loan_id'],'occurred_on':'2026-09-30','idempotency_key':'dated-return','reason':'Synthetic'}
    assert client.post('/assets/movements',json=back).status_code==422
    assert client.post('/assets/movements',json={**back,'occurred_on':'2026-10-02'}).status_code==201
    assert client.post('/assets/movements',json={**back,'expected_version':v+2,'kind':'receive','loan_id':None,'occurred_on':'2099-01-01','idempotency_key':'future'}).status_code==422

def test_human_service_review_binds_document_hash_and_rejects_changed_proof(client):
    a,l,lot,m=setup_stock(client,category='durable')
    with SessionLocal() as db:
        doc=Document(original_filename='Synthetic.pdf',storage_path='hash-proof',sha256='e'*64,size_bytes=1);db.add(doc);db.commit();did=doc.document_id
    s=post(client,'/registry/'+a['asset_id']+'/services',{'expected_version':a['version'],'kind':'inspection','performed_on':'2026-10-01','description':'Synthetic proof','document_id':did,'next_service_on':'2026-10-07'})
    reviewed=post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'reason':'Checked original'},200);assert reviewed['review_snapshot']['document_sha256']=='e'*64
    v=client.get('/assets/registry/'+a['asset_id']).json()['version']
    with SessionLocal() as db:db.get(Document,did).sha256='f'*64;db.commit()
    assert client.post('/assets/services/'+s['service_id']+'/approve',json={'expected_version':2,'expected_asset_version':v,'reason':'Changed proof'}).status_code==409
    assert client.get('/assets/registry/'+a['asset_id']).json()['next_service_on'] is None
    with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='asset.service.approve')) is None
    remove('document.read');assert 'e'*64 not in str(client.get('/assets/services/'+s['service_id']).json())

def test_lot_source_document_hash_is_bound_in_human_service_snapshot(client):
    a=post(client,'/registry',{'code':'PROOF','name':'Synthetic lot proof','category':'drug','unit':'vial'})
    with SessionLocal() as db:
        doc=Document(original_filename='Synthetic.pdf',storage_path='lot-hash-proof',sha256='1'*64,size_bytes=1);db.add(doc);db.commit();did=doc.document_id
    lot=post(client,'/registry/'+a['asset_id']+'/lots',{'expected_version':1,'batch_code':'PROOF-B','provenance':'Synthetic lot original','document_id':did,'expires_on':'2099-01-01'})
    s=post(client,'/registry/'+a['asset_id']+'/services',{'expected_version':2,'kind':'inspection','performed_on':'2026-10-01','description':'Synthetic batch inspection','lot_id':lot['lot_id']})
    reviewed=post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'reason':'Checked batch proof'},200);assert reviewed['review_snapshot']['lot_document_sha256']=='1'*64
    with SessionLocal() as db:db.get(Document,did).sha256='2'*64;db.commit()
    assert client.post('/assets/services/'+s['service_id']+'/approve',json={'expected_version':2,'expected_asset_version':3,'reason':'Changed batch proof'}).status_code==409
    remove('document.read');out=client.get('/assets/services/'+s['service_id']).json();assert did not in str(out) and '1'*64 not in str(out)

def test_batch_search_finds_lot_beyond_first_page_without_borrower_or_proof_leak(client):
    from app.assets_models import AssetLot
    a=post(client,'/registry',{'code':'SEARCH','name':'Synthetic batch search','category':'drug','unit':'vial'})
    with SessionLocal() as db:
        db.add_all([AssetLot(asset_id=a['asset_id'],batch_code=f'SYN-{i:03}',provenance='Synthetic source') for i in range(205)])
        db.add(AssetLot(asset_id=a['asset_id'],batch_code='UniqueBatchBeyondPage',provenance='UniqueBatchProvenance'));db.commit()
    assert len(client.get('/assets/registry/'+a['asset_id']+'/lots?limit=200').json())==200
    found=client.get('/assets/registry/'+a['asset_id']+'/lots?q=UniqueBatchBeyondPage&limit=50').json();assert len(found)==1
    search=client.get('/search',params={'q':'UniqueBatchBeyondPage','modules':'operational_assets'});assert search.status_code==200,search.text
    hit=search.json()['hits'][0];assert hit['source_type']=='asset_lot' and hit['source_id']==found[0]['lot_id'] and hit['navigation']['asset_id']==a['asset_id']
    assert 'document_id' not in str(hit) and client.get('/search?q=PrivateBorrowerSecret&modules=operational_assets').json()['hits']==[]
    remove('asset.read');assert client.get('/search?q=UniqueBatchBeyondPage&modules=operational_assets').status_code==403

def test_formula_export_encoding_roundtrip_preserves_literal_apostrophe(client):
    import csv
    from io import StringIO
    raw=b"schema_version,code,name,category,unit\nassets-v1,QUOTE,'=Synthetic literal,durable,piece\n"
    p=client.post('/assets/import/registry',files={'file':('synthetic.csv',raw,'text/csv')});assert p.status_code==200,p.text;preview=p.json();assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    row=client.get('/assets/registry').json()[0];assert row['name']=="'=Synthetic literal"
    exported=client.get('/assets/export/registry').content.decode('utf-8-sig');rows=list(csv.DictReader(StringIO(exported)));assert rows[0]['text_encoding']=='formula-prefix-v1'
    rows[0]['code']='QUOTE-COPY';out=StringIO();w=csv.DictWriter(out,fieldnames=rows[0]);w.writeheader();w.writerows(rows)
    p=client.post('/assets/import/registry',files={'file':('synthetic.csv',out.getvalue().encode(),'text/csv')});assert p.status_code==200,p.text;preview=p.json();assert client.post('/assets/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    assert next(x for x in client.get('/assets/registry').json() if x['code']=='QUOTE-COPY')['name']=="'=Synthetic literal"

def test_fresh_cli_bootstrap_registers_operational_domain_tables(tmp_path):
    """A fresh process, so app.main's already-loaded metadata cannot conceal missing registration."""
    import subprocess,sqlite3
    root=Path(__file__).resolve().parents[2];database=tmp_path/'synthetic-fresh-bootstrap.db'
    env={**os.environ,'PYTHONPATH':str(root/'backend'),'FIRE_AI_DATABASE_URL':'sqlite+pysqlite:///'+str(database),'FIRE_AI_STORAGE_ROOT':str(tmp_path/'storage'),'FIRE_AI_PRODUCTION_MODE':'false'}
    env.pop('FIRE_AI_TENANT_ID',None)
    r=subprocess.run([__import__('sys').executable,'-m','app.bootstrap','--username','synthetic-fresh','--display-name','Synthetic bootstrap','--password','synthetic-bootstrap-password'],cwd=tmp_path,env=env,capture_output=True,text=True);assert r.returncode==0,r.stderr
    with sqlite3.connect(database) as db:
        tables={name for name, in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert {'operational_assets','asset_lots','asset_balances','asset_movements','asset_services','operation_incidents','operation_vehicles','emergency_treatments','emergency_report_drafts','organization_units','employee_assignments','employee_assignment_roles','account_password_history'}<=tables
        assert db.execute("SELECT count(*) FROM app_users WHERE username='synthetic-fresh'").fetchone()[0]==1
        assert db.execute("SELECT count(*) FROM permissions WHERE code='asset.update'").fetchone()[0]==1
        assert db.execute("SELECT count(*) FROM operational_assets").fetchone()[0]==0

def test_fresh_app_startup_uses_declared_timezone_data_without_os_database(tmp_path):
    import subprocess,sys
    root=Path(__file__).resolve().parents[2]
    env={**os.environ,'PYTHONPATH':str(root/'backend'),'PYTHONTZPATH':'','FIRE_AI_DATABASE_URL':'sqlite+pysqlite:///:memory:','FIRE_AI_PRODUCTION_MODE':'false','FIRE_AI_ASSET_BUSINESS_TIMEZONE':'Asia/Tokyo'}
    env.pop('FIRE_AI_TENANT_ID',None)
    code="from datetime import datetime,timezone; from zoneinfo import TZPATH; import app.main; from app.assets_service import BUSINESS_TIMEZONE; assert TZPATH==(); assert datetime(2026,10,6,16,tzinfo=timezone.utc).astimezone(BUSINESS_TIMEZONE).date().isoformat()=='2026-10-07'; from fastapi.testclient import TestClient; client=TestClient(app.main.app); assert client.get('/assets/policy').status_code==401; client.close()"
    r=subprocess.run([sys.executable,'-c',code],cwd=tmp_path,env=env,capture_output=True,text=True);assert r.returncode==0,r.stderr
    invalid=subprocess.run([sys.executable,'-c','import app.main'],cwd=tmp_path,env={**env,'FIRE_AI_ASSET_BUSINESS_TIMEZONE':'Synthetic/Unknown'},capture_output=True,text=True);assert invalid.returncode!=0 and 'ZoneInfoNotFoundError' in invalid.stderr

def test_real_service_draft_ui_shows_proposed_dates_and_permission_gated_proof(tmp_path):
    import subprocess
    root=Path(__file__).resolve().parents[2];script=tmp_path/'synthetic-draft.cjs'
    script.write_text('''const fs=require('fs'),vm=require('vm'),assert=require('assert');
const nodes=new Map(),get=id=>{if(!nodes.has(id))nodes.set(id,{innerHTML:'',textContent:'',classList:{toggle(){}}});return nodes.get(id)};
let permitted=true;const draft={service_id:'synthetic',asset_id:'asset',kind:'calibration',status:'draft',performed_on:'2031-09-01',description:'<unsafe>',cost:'0.00',lot_id:'lot',location_id:null,quantity:null,next_pressure_test_on:'2031-09-17',next_use_on:'2031-09-18',next_calibration_on:'2031-09-19',next_service_on:'2031-09-20',document_id:'synthetic-proof',review_snapshot:{}};
const ctx={console,$:get,esc:x=>String(x??'').replaceAll('<','&lt;').replaceAll('>','&gt;').replaceAll('"','&quot;'),api:async path=>path==='/auth/permissions'?{permissions:['asset.read','asset.review',...(permitted?['document.read']:[])]}:draft};vm.createContext(ctx);vm.runInContext(fs.readFileSync(process.argv[2],'utf8'),ctx);
(async()=>{await ctx.initAssets();await ctx.assetsServiceDetail('synthetic');let h=get('assetsContent').innerHTML;for(const date of ['2031-09-17','2031-09-18','2031-09-19','2031-09-20'])assert(h.includes(date));assert(h.includes('synthetic-proof'));assert(h.includes('0.00'));assert(h.includes('assetsReview'));assert(h.includes('&lt;unsafe&gt;'));
permitted=false;await ctx.initAssets();await ctx.assetsServiceDetail('synthetic');h=get('assetsContent').innerHTML;assert(!h.includes('synthetic-proof'));assert(h.includes('2031-09-20'));draft.kind='disposal';draft.quantity='12.125';draft.location_id='synthetic-location';await ctx.assetsServiceDetail('synthetic');assert(get('assetsContent').innerHTML.includes('12.125'));assert(get('assetsContent').innerHTML.includes('synthetic-location'));})().catch(e=>{console.error(e);process.exit(1)});
''')
    r=subprocess.run(['node',str(script),str(root/'frontend/assets.js')],capture_output=True,text=True);assert r.returncode==0,r.stderr

def test_borrower_picker_api_search_and_paging_reaches_employee_after_200(client):
    with SessionLocal() as db:
        db.add_all([Employee(display_name=f'AA Synthetic {i:03}',employee_code=f'PAGE-{i:03}') for i in range(205)])
        target=Employee(display_name='ZZ UniqueBorrowerBeyondPage',employee_code='LAST');db.add(target);db.commit();eid=target.employee_id
    first=client.get('/assets/employees?limit=200&offset=0').json();assert len(first)==200 and eid not in [x['employee_id'] for x in first]
    later=client.get('/assets/employees?limit=200&offset=200').json();assert eid in [x['employee_id'] for x in later]
    search=client.get('/assets/employees?q=UniqueBorrowerBeyondPage&limit=50&offset=0').json();assert len(search)==1 and search[0]['employee_id']==eid
    remove('asset.borrower.read');assert client.get('/assets/employees?q=UniqueBorrowerBeyondPage').status_code==403

def test_real_loan_picker_pages_searches_and_submits_later_employee(tmp_path):
    import subprocess
    root=Path(__file__).resolve().parents[2];script=tmp_path/'synthetic-borrower.cjs'
    script.write_text('''const fs=require('fs'),vm=require('vm'),assert=require('assert');
const nodes=new Map(),get=id=>{if(!nodes.has(id))nodes.set(id,{innerHTML:'',textContent:'',value:'',disabled:false,classList:{toggle(){}},insertAdjacentHTML(_position,value){this.innerHTML+=value;if(id==='assetsForm')get('assetsContent').innerHTML+=value}});return nodes.get(id)};
let permitted=true,posted=null,requests=[];const ctx={console,URLSearchParams,crypto:require('crypto').webcrypto,$:get,esc:x=>String(x??'').replaceAll('<','&lt;').replaceAll('>','&gt;').replaceAll('"','&quot;'),api:async(path,options)=>{requests.push(path);if(path==='/auth/permissions')return {permissions:['asset.read','asset.update',...(permitted?['asset.borrower.read','asset.borrower.manage']:[])]};if(path.includes('/employees')){const q=new URL('http://synthetic'+path).searchParams;if(q.get('q')==='UniqueBorrowerBeyondPage'||Number(q.get('offset'))>=200)return [{employee_id:'later-employee',employee_code:'LAST',display_name:'ZZ UniqueBorrowerBeyondPage'}];return Array.from({length:200},(_,i)=>({employee_id:'first-'+i,employee_code:'S'+i,display_name:'AA '+i}))}if(path==='/assets/policy')return {business_date:'2026-10-06'};if(path.includes('/lots'))return [{lot_id:'lot',batch_code:'B'}];if(path.includes('/locations'))return [{location_id:'store',code:'A',name:'Synthetic',active:true}];if(path==='/assets/movements'){posted=JSON.parse(options.body);return {movement_id:'done'}};return []}};
vm.createContext(ctx);vm.runInContext(fs.readFileSync(process.argv[2],'utf8'),ctx);
(async()=>{await ctx.initAssets();await ctx.assetsMovementForm({asset_id:'asset',unit:'vial',version:3},'loan');assert(get('assetsContent').innerHTML.includes('assetsBorrowerQuery'));await get('assetsBorrowerNext').onclick();assert(get('assetsField_borrower_employee_id').innerHTML.includes('later-employee'));get('assetsBorrowerQuery').value='UniqueBorrowerBeyondPage';await get('assetsBorrowerSearch').onclick();assert(requests.some(p=>p.includes('/employees?')&&p.includes('q=UniqueBorrowerBeyondPage')));assert(get('assetsField_borrower_employee_id').innerHTML.includes('later-employee'));for(const [key,value] of Object.entries({kind:'loan',lot_id:'lot',location_id:'store',quantity:'1.125',occurred_on:'2026-10-06',borrower_employee_id:'later-employee',reason:'Synthetic loan'}))get('assetsField_'+key).value=value;ctx.assetsDetail=async()=>{};await get('assetsForm').onsubmit({preventDefault(){}});assert(posted.borrower_employee_id==='later-employee'&&posted.quantity==='1.125');
permitted=false;await ctx.initAssets();let count=requests.filter(p=>p.includes('/employees')).length;await ctx.assetsMovementForm({asset_id:'asset',unit:'vial',version:3},'receive');assert(requests.filter(p=>p.includes('/employees')).length===count);assert(!get('assetsContent').innerHTML.includes('assetsBorrowerQuery'));})().catch(e=>{console.error(e);process.exit(1)});
''')
    r=subprocess.run(['node',str(script),str(root/'frontend/assets.js')],capture_output=True,text=True);assert r.returncode==0,r.stderr

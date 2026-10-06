import os
os.environ['FIRE_AI_DATABASE_URL']='sqlite+pysqlite:///:memory:'
from io import BytesIO
from datetime import date
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from app.db import Base,engine,SessionLocal
from app.main import app
from app.models import Employee,User,Role,Permission,RolePermission,UserRole,EmergencyCase,FireInvestigationCase,AuditLog
from app.security import hash_password
CODES=[f'{domain}.{p}' for domain in ['incident','fleet'] for p in ['read','create','update','review','approve','admin','aggregate','export','import']]+['incident.crew.read','incident.crew.manage','emergency.case.read','fire_investigation.read','document.read','search.use']
@pytest.fixture
def client():
    Base.metadata.drop_all(engine);Base.metadata.create_all(engine)
    with SessionLocal() as db:
        e=Employee(display_name='Synthetic crew',employee_code='SYN');db.add(e);db.flush()
        u=User(username='operations',employee_id=e.employee_id,password_hash=hash_password('synthetic-password'));r=Role(code='operations-test',name='Synthetic');db.add_all([u,r]);db.flush();db.add(UserRole(user_id=u.user_id,role_id=r.role_id))
        for code in CODES:
            p=Permission(code=code);db.add(p);db.flush();db.add(RolePermission(role_id=r.role_id,permission_id=p.permission_id))
        db.commit()
    with TestClient(app) as c:
        assert c.post('/auth/login',json={'username':'operations','password':'synthetic-password'}).status_code==200
        yield c

def post(c,path,data,status=201):
    r=c.post('/operations'+path,json=data);assert r.status_code==status,r.text;return r.json()
def incident(c):return post(c,'/incidents',{'kind':'rescue','title':'Synthetic rescue','occurred_at':'2026-10-01T10:00:00Z','address':'Synthetic address','number':'S1'})
def vehicle(c):return post(c,'/vehicles',{'code':'V1','name':'Synthetic vehicle','odometer':'0'})
def remove(code):
    with SessionLocal() as db:
        p=db.scalar(select(Permission).where(Permission.code==code))
        for l in db.scalars(select(RolePermission).where(RolePermission.permission_id==p.permission_id)):db.delete(l)
        db.commit()

def test_incident_crud_stale_identity_rbac(client):
    i=incident(client);p='/operations/incidents/'+i['incident_id']
    r=client.patch(p,json={'expected_version':1,'title':'Changed'});assert r.status_code==200 and r.json()['version']==2
    assert client.patch(p,json={'expected_version':1,'title':'Stale'}).status_code==409
    assert client.post('/operations/incidents',json={'kind':'watch','title':'X','tenant_id':'chosen'}).status_code==422
    assert client.get(p).headers['cache-control']=='no-store'
    remove('incident.update');assert client.patch(p,json={'expected_version':2,'title':'Denied'}).status_code==403

def test_source_preservation_uniqueness_privacy(client):
    with SessionLocal() as db:
        s=EmergencyCase(source_case_key='synthetic',incident_address='Protected address',call_date=date(2026,10,1),dispatch_number='E1');f=FireInvestigationCase(title='Synthetic fire',case_number='F1',location_text='Fire address');db.add_all([s,f]);db.commit();eid=s.emergency_case_id;fid=f.fire_investigation_case_id
    i=post(client,'/incidents',{'kind':'emergency_support','title':'Support','emergency_case_id':eid})
    assert i['address'] is None and i['source']['address']=='Protected address'
    assert client.post('/operations/incidents',json={'kind':'emergency_support','title':'Duplicate','emergency_case_id':eid}).status_code==409
    assert client.post('/operations/incidents',json={'kind':'emergency_support','title':'Copied','emergency_case_id':eid,'address':'Copy'}).status_code==422
    fi=post(client,'/incidents',{'kind':'fire','title':'Fire response','fire_investigation_case_id':fid});assert fi['source']['number']=='F1'
    remove('emergency.case.read');detail=client.get('/operations/incidents/'+i['incident_id']).json();assert 'source' not in detail and 'emergency_case_id' not in detail
    r=client.get('/search?q=Protected&modules=operations');assert r.status_code==200 and r.json()['hits']==[]
    with SessionLocal() as db:assert db.get(EmergencyCase,eid).incident_address=='Protected address'

def test_dispatch_human_gate_crew_audit(client):
    i=incident(client);v=vehicle(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'Synthetic','vehicle_id':v['vehicle_id'],'departed_at':'2026-10-01T10:00:00Z','returned_at':'2026-10-01T11:00:00Z','activity':'rescue','report':'Synthetic report'})
    with SessionLocal() as db:eid=db.scalar(select(Employee.employee_id))
    post(client,'/dispatches/'+d['dispatch_id']+'/crew',{'expected_version':1,'employee_id':eid,'role':'leader'})
    rate=post(client,'/allowance-rates',{'rounding':'half_up','code':'synthetic','label':'Synthetic fixed authorized amount','amount':'12.34','basis':'per_dispatch','approval_reference':'Synthetic authorization'})
    assert rate['status']=='draft'
    assert client.post('/operations/dispatches/'+d['dispatch_id']+'/calculate',json={'expected_version':2,'rate_id':rate['rate_id']}).status_code==422
    post(client,'/allowance-rates/'+rate['rate_id']+'/approve',{'expected_version':1,'note':'Synthetic approval'},200)
    calc=post(client,'/dispatches/'+d['dispatch_id']+'/calculate',{'expected_version':2,'rate_id':rate['rate_id']},200)
    assert calc['candidate_amount']=='12.34' and calc['official_amount'] is None
    reviewed=post(client,'/dispatches/'+d['dispatch_id']+'/review',{'expected_version':3,'note':'Checked'},200)
    approved=post(client,'/dispatches/'+d['dispatch_id']+'/approve',{'expected_version':reviewed['version'],'note':'Human approval'},200)
    assert approved['official_amount']=='12.34' and approved['approved_by']
    assert client.patch('/operations/dispatches/'+d['dispatch_id'],json={'expected_version':approved['version'],'report':'Overwrite'}).status_code==409
    assert client.get('/operations/statistics?year=2026&month=10').json()['official_allowance']=='12.34'
    remove('incident.crew.read');assert client.get('/operations/dispatches/'+d['dispatch_id']+'/crew').status_code==403
    with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='incident.dispatch.approve'))

def test_inactive_crew_and_mismatched_vehicle(client):
    i=incident(client);v=vehicle(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U','vehicle_id':v['vehicle_id']})
    # Validate an inactive crew target independently from the active operator session.
    with SessionLocal() as db:
        e=Employee(display_name='Synthetic inactive crew',employee_code='SYN-INACTIVE',active=False)
        db.add(e);db.flush();eid=e.employee_id;db.commit()
    assert client.post('/operations/dispatches/'+d['dispatch_id']+'/crew',json={'expected_version':1,'employee_id':eid,'role':'crew'}).status_code==422
    v2=post(client,'/vehicles',{'code':'V2','name':'Other'})
    assert client.post('/operations/vehicles/'+v2['vehicle_id']+'/trips',json={'expected_version':1,'dispatch_id':d['dispatch_id'],'started_at':'2026-10-01T10:00:00Z','ended_at':'2026-10-01T11:00:00Z','start_odometer':'0','end_odometer':'10','purpose':'Synthetic'}).status_code==422

def test_fleet_history_odometer_stock_deadlines(client):
    v=vehicle(client);p='/vehicles/'+v['vehicle_id']
    tripdata={'expected_version':1,'started_at':'2026-10-01T10:00:00Z','ended_at':'2026-10-01T11:00:00Z','start_odometer':'0','end_odometer':'10.5','purpose':'Synthetic'}
    trip=post(client,p+'/trips',tripdata);assert trip['end_odometer']=='10.5'
    assert client.post('/operations'+p+'/trips',json=tripdata).status_code==409
    assert client.post('/operations'+p+'/trips',json={**tripdata,'expected_version':2,'start_odometer':'9'}).status_code==422
    post(client,p+'/fuel',{'expected_version':2,'kind':'receipt','liters':'20.25','amount':'30.01','occurred_at':'2026-10-01T12:00:00Z'})
    post(client,p+'/fuel',{'expected_version':3,'kind':'issue','liters':'5.25','amount':'0','occurred_at':'2026-10-01T13:00:00Z'})
    assert client.post('/operations'+p+'/fuel',json={'expected_version':4,'kind':'issue','liters':'16','amount':'0','occurred_at':'2026-10-01T14:00:00Z'}).status_code==422
    s=post(client,p+'/services',{'expected_version':4,'kind':'repair','performed_on':'2026-10-01','description':'Synthetic repair','cost':'123.45','next_service_on':'2026-10-03','next_inspection_on':'2026-10-05','next_service_odometer':'20'})
    post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    post(client,'/services/'+s['service_id']+'/approve',{'expected_version':2,'expected_vehicle_version':5,'note':'Approved'},200)
    assert {x['kind'] for x in client.get('/operations/alerts?as_of=2026-10-06').json()} >= {'service','inspection'}
    h=client.get('/operations'+p+'/history').json();assert len(h['trips'])==1 and len(h['fuel'])==2 and len(h['services'])==1
    assert client.get('/operations'+p).json()['fuel_stock']=='15.00'
    assert client.patch('/operations/trips/'+trip['trip_id'],json={'end_odometer':'1'}).status_code in [404,405]

def test_real_csv_xlsx_import_export_search(client):
    csv=b'code,name,odometer\nCSV1,Synthetic imported,0\n'
    r=client.post('/operations/import/vehicles',files={'file':('synthetic.csv',csv,'text/csv')});assert r.status_code==200,r.text;preview=r.json();applied=client.post('/operations/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']});assert applied.status_code==200;assert applied.json()['inserted']==1
    assert client.post('/operations/import/vehicles',files={'file':('synthetic.csv',csv,'text/csv')}).status_code==409
    assert client.get('/search?q=Synthetic imported&modules=fleet').json()['hits'][0]['source_type']=='vehicle'
    from openpyxl import Workbook,load_workbook
    w=Workbook();w.active.append(['kind','title','occurred_at','address','number']);w.active.append(['watch','Synthetic workbook','2026-10-01T10:00:00Z','A','W1']);out=BytesIO();w.save(out)
    r=client.post('/operations/import/incidents',files={'file':('synthetic.xlsx',out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')});assert r.status_code==200,r.text
    preview=r.json();assert client.post('/operations/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    assert client.get('/operations/export/vehicles?format=csv').status_code==200
    r=client.get('/operations/export/incidents?format=xlsx');assert r.status_code==200 and load_workbook(BytesIO(r.content)).active.max_row==2
    remove('incident.read');assert 'source_incident_ids' not in client.get('/operations/statistics?year=2026').json()
    assert client.get('/search?q=Synthetic&modules=operations').status_code==403
    assert client.get('/operations/export/summary?year=2026').status_code==200

def test_atomic_import_validation(client):
    r=client.post('/operations/import/vehicles',files={'file':('synthetic.csv',b'code,name,odometer\nOK,Valid,0\nBAD,Invalid,-1\n','text/csv')});assert r.status_code==422
    assert client.get('/operations/vehicles').json()==[]
    assert client.get('/operations/statistics?year=2026&month=13').status_code==422

def test_same_shell_frontend():
    from pathlib import Path
    import subprocess
    root=Path(__file__).resolve().parents[2];html=(root/'frontend/index.html').read_text();assert 'operationsBtn' in html and 'operations.js' in html
    js=root/'frontend/operations.js';r=subprocess.run(['node','--check',str(js)],capture_output=True,text=True);assert r.returncode==0,r.stderr
    s=js.read_text();assert '/review' in s and '/approve' in s and '/import/' in s and '/export/' in s and '/history' in s and '??' in s

def test_import_preview_binding_explicit_confirm(client):
    raw=b'code,name,odometer\nPREVIEW,Synthetic preview,0\n'
    p=client.post('/operations/import/vehicles',files={'file':('synthetic.csv',raw,'text/csv')});assert p.status_code==200,p.text
    preview=p.json();assert preview['status']=='preview' and preview['rows']==1
    assert client.get('/operations/vehicles').json()==[]
    path='/operations/import-previews/'+preview['preview_id']+'/confirm'
    assert client.post(path,json={'expected_version':1,'file_sha256':'0'*64}).status_code==409
    applied=client.post(path,json={'expected_version':1,'file_sha256':preview['file_sha256']});assert applied.status_code==200,applied.text
    assert applied.json()['inserted']==1
    assert client.post(path,json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==409

def test_reviewed_dispatch_cannot_change_crew_and_stale_source_rejected(client):
    i=incident(client);v=vehicle(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U','vehicle_id':v['vehicle_id']})
    post(client,'/dispatches/'+d['dispatch_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    with SessionLocal() as db:eid=db.scalar(select(Employee.employee_id))
    assert client.post('/operations/dispatches/'+d['dispatch_id']+'/crew',json={'expected_version':2,'employee_id':eid,'role':'crew'}).status_code==409
    assert client.patch('/operations/incidents/'+i['incident_id'],json={'expected_version':2,'title':'Changed source'}).status_code==200
    assert client.post('/operations/dispatches/'+d['dispatch_id']+'/approve',json={'expected_version':2,'note':'Approve stale'}).status_code==409

def test_service_approval_checks_expected_vehicle_version(client):
    v=vehicle(client);s=post(client,'/vehicles/'+v['vehicle_id']+'/services',{'expected_version':1,'kind':'inspection','performed_on':'2026-10-01','description':'Synthetic'})
    post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    assert client.patch('/operations/vehicles/'+v['vehicle_id'],json={'expected_version':2,'notes':'Vehicle changed'}).status_code==200
    assert client.post('/operations/services/'+s['service_id']+'/approve',json={'expected_version':2,'expected_vehicle_version':2,'note':'Stale'}).status_code==409

def test_document_permission_formula_export_and_crew_privacy(client):
    from app.models import Document
    with SessionLocal() as db:
        doc=Document(original_filename='Synthetic.pdf',mime_type='application/pdf',storage_path='synthetic',sha256='a'*64,size_bytes=1)
        db.add(doc);db.commit();did=doc.document_id
    i=incident(client);v=vehicle(client)
    remove('document.read')
    assert client.post('/operations/incidents/'+i['incident_id']+'/dispatches',json={'expected_version':1,'unit':'U','document_id':did}).status_code==403
    assert client.post('/operations/vehicles/'+v['vehicle_id']+'/fuel',json={'expected_version':1,'kind':'receipt','liters':'1','amount':'0','occurred_at':'2026-10-01T10:00:00Z','document_id':did}).status_code==403
    assert client.patch('/operations/vehicles/'+v['vehicle_id'],json={'expected_version':1,'name':'=SYNTHETIC()'}).status_code==200
    assert "'=SYNTHETIC()" in client.get('/operations/export/vehicles?format=csv').text

def test_cancel_retains_typed_approved_decision_excludes_totals(client):
    i=incident(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U'})
    rate=post(client,'/allowance-rates',{'rounding':'half_up','code':'cancel-rate','label':'Synthetic','basis':'per_dispatch','amount':'5.55','approval_reference':'Synthetic'})
    post(client,'/allowance-rates/'+rate['rate_id']+'/approve',{'expected_version':1,'note':'Checked'},200)
    post(client,'/dispatches/'+d['dispatch_id']+'/calculate',{'expected_version':1,'rate_id':rate['rate_id']},200)
    post(client,'/dispatches/'+d['dispatch_id']+'/review',{'expected_version':2,'note':'Checked'},200)
    post(client,'/dispatches/'+d['dispatch_id']+'/approve',{'expected_version':3,'note':'Approved'},200)
    cancelled=post(client,'/dispatches/'+d['dispatch_id']+'/cancel',{'expected_version':4,'note':'Synthetic cancelled'},200)
    assert cancelled['official_amount']=='5.55' and cancelled['approved_by']
    assert client.get('/operations/statistics?year=2026').json()['official_allowance']=='0.00'

def test_crew_import_confirm_and_sensitive_exports(client):
    i=incident(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U'})
    with SessionLocal() as db:eid=db.scalar(select(Employee.employee_id))
    raw=f"dispatch_id,expected_version,employee_id,role\n{d['dispatch_id']},1,{eid},crew\n".encode()
    p=client.post('/operations/import/crew',files={'file':('synthetic.csv',raw,'text/csv')});assert p.status_code==200,p.text
    preview=p.json();assert client.post('/operations/import-previews/'+preview['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':preview['file_sha256']}).status_code==200
    assert eid in client.get('/operations/export/crew').text
    remove('incident.crew.read');assert client.get('/operations/export/crew').status_code==403
    assert eid not in client.get('/operations/export/dispatches').text

def test_snapshot_locks_real_linked_source_model(client,monkeypatch):
    with SessionLocal() as db:
        s=EmergencyCase(source_case_key='lock-synthetic',call_date=date(2026,10,1));db.add(s);db.commit();eid=s.emergency_case_id
    i=post(client,'/incidents',{'kind':'emergency_support','title':'Synthetic lock','emergency_case_id':eid});d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U'})
    from sqlalchemy.orm import Session
    actual=Session.scalar;locked=[]
    def record(self,stmt,*args,**kwargs):
        if getattr(stmt,'_for_update_arg',None) is not None:locked.append(str(stmt))
        return actual(self,stmt,*args,**kwargs)
    monkeypatch.setattr(Session,'scalar',record)
    post(client,'/dispatches/'+d['dispatch_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    assert any('FROM emergency_cases' in stmt for stmt in locked)

def test_database_constraints_and_shared_foreign_keys(client):
    from app.operations_models import Incident,Vehicle,VehicleTrip,DispatchCrew,VehicleService
    from sqlalchemy.exc import IntegrityError
    with SessionLocal() as db:
        db.add(Vehicle(code='NEG',name='Synthetic',odometer=-1))
        with pytest.raises(IntegrityError):db.flush()
        db.rollback()
        db.add(Incident(kind='other',title='Copied source',emergency_case_id='00000000-0000-0000-0000-000000000001',address='Copy'))
        with pytest.raises(IntegrityError):db.flush()
        db.rollback()
    assert {f.target_fullname for f in DispatchCrew.__table__.foreign_keys}>={'employees.employee_id','operation_dispatches.dispatch_id'}
    assert {f.target_fullname for f in VehicleTrip.__table__.foreign_keys}>={'operation_vehicles.vehicle_id','operation_dispatches.vehicle_id','documents.document_id','employees.employee_id'}
    assert any(c.name=='ck_service_cost' for c in VehicleService.__table__.constraints)

def test_import_confirmation_rechecks_and_rolls_back_all_rows(client):
    raw=b'code,name,odometer\nFIRST,Synthetic one,0\nCONFLICT,Synthetic two,0\n'
    p=client.post('/operations/import/vehicles',files={'file':('synthetic.csv',raw,'text/csv')}).json()
    post(client,'/vehicles',{'code':'CONFLICT','name':'Concurrent registry'})
    r=client.post('/operations/import-previews/'+p['preview_id']+'/confirm',json={'expected_version':1,'file_sha256':p['file_sha256']});assert r.status_code==409
    assert [v['code'] for v in client.get('/operations/vehicles').json()]==['CONFLICT']

def test_human_approval_rbac_and_rate_rounding_is_explicit(client):
    i=incident(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U'})
    assert client.post('/operations/allowance-rates',json={'code':'NO-POLICY','label':'Synthetic','amount':'1','basis':'per_hour','approval_reference':'Synthetic'}).status_code==422
    remove('incident.approve');assert client.post('/operations/dispatches/'+d['dispatch_id']+'/approve',json={'expected_version':1,'note':'Denied'}).status_code==403
    remove('fleet.create')
    assert client.post('/operations/vehicles',json={'code':'DENIED','name':'Denied'}).status_code==403

def test_frontend_real_form_values_zero_escape_and_human_paths():
    from pathlib import Path
    import subprocess
    root=Path(__file__).resolve().parents[2]
    script=r'''const fs=require('fs'),vm=require('vm'),assert=require('assert');
const nodes={operationsContent:{innerHTML:''},operationsForm:{},operationsFormBack:{},operationsField_amount:{value:'0'},operationsField_note:{value:'Synthetic'},operationsMessage:{textContent:''}};
const requests=[];let returned=false;
const ctx=vm.createContext({console,Date,URLSearchParams,$:id=>nodes[id],api:async(url,options)=>{requests.push([url,JSON.parse(options.body)]);return{}},document:{},fetch:()=>{}});
vm.runInContext(fs.readFileSync(process.argv[2],'utf8').match(/^const esc=.*$/m)[0],ctx);
vm.runInContext(fs.readFileSync(process.argv[1],'utf8'),ctx);
assert(vm.runInContext(`operationsField(['amount','Amount','number'],{amount:0})`,ctx).includes('value="0"'));
assert(vm.runInContext(`operationsField(['name','Name'],{name:'<script>'})`,ctx).includes('&lt;script&gt;'));
assert.strictEqual(vm.runInContext(`operationsValues([['amount','Amount','number']]).amount`,ctx),'0');
(async()=>{await vm.runInContext(`operationsHumanAction('/dispatches/synthetic/review',{version:7},async()=>{},'Human review')`,ctx);await nodes.operationsForm.onsubmit({preventDefault(){}});await new Promise(r=>setImmediate(r));assert.strictEqual(requests[0][0],'/operations/dispatches/synthetic/review');assert.strictEqual(requests[0][1].expected_version,7);assert.strictEqual(requests[0][1].note,'Synthetic');})().catch(e=>{console.error(e);process.exitCode=1});'''
    r=subprocess.run(['node','-e',script,str(root/'frontend/operations.js'),str(root/'frontend/index.html')],capture_output=True,text=True);assert r.returncode==0,r.stderr

def test_repair_explicit_fault_resolution_human_version_gate(client):
    v=vehicle(client);key=v['vehicle_id']
    fault=post(client,'/vehicles/'+key+'/services',{'expected_version':1,'kind':'fault','performed_on':'2026-10-01','description':'Synthetic specific fault'})
    unrelated=post(client,'/vehicles/'+key+'/services',{'expected_version':2,'kind':'repair','performed_on':'2026-10-02','description':'Unrelated repair'})
    post(client,'/services/'+unrelated['service_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    post(client,'/services/'+unrelated['service_id']+'/approve',{'expected_version':2,'expected_vehicle_version':3,'note':'Approved unrelated'},200)
    assert any(x.get('service_id')==fault['service_id'] for x in client.get('/operations/alerts').json())
    repair=post(client,'/vehicles/'+key+'/services',{'expected_version':4,'kind':'repair','performed_on':'2026-10-03','description':'Specific fault repaired','resolves_fault_id':fault['service_id']})
    post(client,'/services/'+repair['service_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    assert client.post('/operations/services/'+repair['service_id']+'/approve',json={'expected_version':2,'expected_vehicle_version':5,'expected_fault_version':2,'note':'Stale fault'}).status_code==409
    post(client,'/services/'+repair['service_id']+'/approve',{'expected_version':2,'expected_vehicle_version':5,'expected_fault_version':1,'note':'Human confirms fault repaired'},200)
    assert not any(x.get('service_id')==fault['service_id'] for x in client.get('/operations/alerts').json())
    history=client.get('/operations/vehicles/'+key+'/history').json()['services'];original=next(s for s in history if s['service_id']==fault['service_id'])
    assert original['description']=='Synthetic specific fault' and original['resolved_by_service_id']==repair['service_id'] and original['version']==2

def test_dispatch_timezone_normalization_and_configured_hourly_rounding(client):
    i=incident(client);d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'U','departed_at':'2026-10-01T10:00:00+09:00','returned_at':'2026-10-01T02:00:00Z'})
    rate=post(client,'/allowance-rates',{'code':'HOUR','label':'Synthetic hourly','basis':'per_hour','amount':'12.34','rounding':'half_up','approval_reference':'Synthetic'})
    post(client,'/allowance-rates/'+rate['rate_id']+'/approve',{'expected_version':1,'note':'Human checked'},200)
    calc=post(client,'/dispatches/'+d['dispatch_id']+'/calculate',{'expected_version':1,'rate_id':rate['rate_id']},200)
    assert calc['candidate_amount']=='12.34' and calc['calculation']['units']=='1.0'


def synthetic_document():
    from app.models import Document
    with SessionLocal() as db:
        doc=Document(original_filename='Synthetic-review.pdf',mime_type='application/pdf',storage_path='synthetic-review',sha256='b'*64,size_bytes=1)
        db.add(doc);db.commit();return doc.document_id


@pytest.mark.parametrize('action',['calculate','review','approve','cancel'])
def test_dispatch_mutation_responses_filter_revoked_document_permission(client,action):
    import json
    from app.operations_models import Dispatch
    did=synthetic_document();i=incident(client)
    d=post(client,'/incidents/'+i['incident_id']+'/dispatches',{'expected_version':1,'unit':'Synthetic projection','document_id':did})
    assert d['document_id']==did
    rate=post(client,'/allowance-rates',{'code':'PROJECTION','label':'Synthetic','amount':'1','basis':'per_dispatch','rounding':'half_up','approval_reference':'Synthetic'})
    post(client,'/allowance-rates/'+rate['rate_id']+'/approve',{'expected_version':1,'note':'Checked'},200)
    transitions=['calculate','review','approve','cancel']
    for step in transitions[:transitions.index(action)]:
        data={'expected_version':d['version'],**({'rate_id':rate['rate_id']} if step=='calculate' else {'note':'Checked'})}
        d=post(client,'/dispatches/'+d['dispatch_id']+'/'+step,data,200)
    remove('document.read')
    data={'expected_version':d['version'],**({'rate_id':rate['rate_id']} if action=='calculate' else {'note':'Checked'})}
    output=post(client,'/dispatches/'+d['dispatch_id']+'/'+action,data,200)
    assert 'document_id' not in output and did not in json.dumps(output)
    for field in ['calculation','review_snapshot']:
        assert 'document_id' not in output[field].get('dispatch_fields',{})
    with SessionLocal() as db:
        stored=db.get(Dispatch,d['dispatch_id']);assert stored.document_id==did
        if stored.calculation:assert stored.calculation['dispatch_fields']['document_id']==did


@pytest.mark.parametrize('action',['review','approve','cancel'])
def test_service_mutation_responses_filter_revoked_document_permission(client,action):
    from app.operations_models import VehicleService
    did=synthetic_document();v=vehicle(client)
    s=post(client,'/vehicles/'+v['vehicle_id']+'/services',{'expected_version':1,'kind':'inspection','performed_on':'2026-10-01','description':'Synthetic projection','document_id':did})
    if action=='approve':s=post(client,'/services/'+s['service_id']+'/review',{'expected_version':1,'note':'Checked'},200)
    remove('document.read')
    output=post(client,'/services/'+s['service_id']+'/'+action,{'expected_version':s['version'],'note':'Checked',**({'expected_vehicle_version':2} if action=='approve' else {})},200)
    assert 'document_id' not in output and did not in str(output)
    with SessionLocal() as db:assert db.get(VehicleService,s['service_id']).document_id==did


@pytest.mark.parametrize('format',['csv','xlsx'])
def test_mixed_incident_exports_preserve_later_source_projection(client,format):
    import csv,json
    from io import StringIO
    from openpyxl import load_workbook
    local=incident(client)
    with SessionLocal() as db:
        e=EmergencyCase(source_case_key='mixed-export',incident_address='Synthetic emergency source',call_date=date(2026,10,1),dispatch_number='MIXED-E')
        f=FireInvestigationCase(title='Synthetic fire source',case_number='MIXED-F',location_text='Synthetic fire source')
        db.add_all([e,f]);db.commit();eid=e.emergency_case_id;fid=f.fire_investigation_case_id
    linked_e=post(client,'/incidents',{'kind':'emergency_support','title':'Synthetic linked emergency','emergency_case_id':eid})
    linked_f=post(client,'/incidents',{'kind':'fire','title':'Synthetic linked fire','fire_investigation_case_id':fid})
    def table():
        r=client.get('/operations/export/incidents?format='+format);assert r.status_code==200,r.text
        if format=='csv':return list(csv.DictReader(StringIO(r.content.decode('utf-8-sig'))))
        sheet=load_workbook(BytesIO(r.content)).active;values=list(sheet.values);return [dict(zip(values[0],row)) for row in values[1:]]
    rows=table();by_id={r['incident_id']:r for r in rows}
    assert not by_id[local['incident_id']].get('source')
    assert json.loads(by_id[linked_e['incident_id']]['source'])['source_id']==eid
    assert json.loads(by_id[linked_f['incident_id']]['source'])['source_id']==fid
    remove('emergency.case.read');rows=table();by_id={r['incident_id']:r for r in rows}
    assert str(by_id[linked_e['incident_id']]['source_restricted']).lower()=='true'
    assert not by_id[linked_e['incident_id']].get('source') and eid not in str(rows)
    assert json.loads(by_id[linked_f['incident_id']]['source'])['source_id']==fid


def test_fuel_purchase_expense_excludes_issue_valuation(client):
    import csv
    from io import StringIO
    v=vehicle(client)
    entries=[('receipt','10','100.01'),('issue','5','50.50'),('refuel','3','30.02')]
    for version,(kind,liters,amount) in enumerate(entries,1):
        post(client,'/vehicles/'+v['vehicle_id']+'/fuel',{'expected_version':version,'kind':kind,'liters':liters,'amount':amount,'occurred_at':'2026-10-01T10:00:00Z'})
    monthly=client.get('/operations/fleet-statistics?year=2026&month=10').json()
    assert monthly['fuel_purchase_expense']=='130.03'
    assert monthly['fuel_issue_valuation']=='50.50'
    assert 'fuel_cost' not in monthly
    yearly=client.get('/operations/fleet-statistics?year=2026').json();assert yearly['fuel_purchase_expense']=='130.03'
    rows=list(csv.DictReader(StringIO(client.get('/operations/export/fleet-summary?year=2026').content.decode('utf-8-sig'))))
    assert rows[0]['fuel_purchase_expense']=='130.03' and rows[0]['fuel_issue_valuation']=='50.50'

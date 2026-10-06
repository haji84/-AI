"""Synthetic inquiry evidence, privacy and Human authority contracts."""
import os
os.environ['FIRE_AI_DATABASE_URL']='sqlite+pysqlite:///:memory:'
from io import BytesIO
from pathlib import Path
import subprocess,sys
import pytest
from sqlalchemy import select
from fastapi.testclient import TestClient
from app.main import app
from app.db import Base,engine,SessionLocal
from app.models import User,Employee,Role,Permission,UserRole,RolePermission,Document,AuditLog
from app.security import hash_password
CODES=['inquiry.'+x for x in ['read','create','update','review','approve','admin','import','export']]+['document.read','document.create','personnel.read','emergency.case.read','search.use','template.read']
@pytest.fixture
def client():
 with engine.begin() as connection:
  for table in reversed(list(Base.metadata.tables.values())):table.drop(connection,checkfirst=True)
 Base.metadata.create_all(engine)
 with SessionLocal() as db:
  u=User(username='inquiry',password_hash=hash_password('synthetic-password'));r=Role(code='inquiry-test',name='Synthetic');db.add_all([u,r]);db.flush();db.add(UserRole(user_id=u.user_id,role_id=r.role_id))
  for code in CODES:
   p=Permission(code=code);db.add(p);db.flush();db.add(RolePermission(role_id=r.role_id,permission_id=p.permission_id))
  db.commit()
 with TestClient(app) as c:
  assert c.post('/auth/login',json={'username':'inquiry','password':'synthetic-password'}).status_code==200
  yield c

def remove(code):
 with SessionLocal() as db:
  p=db.scalar(select(Permission).where(Permission.code==code))
  for r in db.scalars(select(RolePermission).where(RolePermission.permission_id==p.permission_id)):db.delete(r)
  db.commit()
def create(c):
 r=c.post('/inquiries',json={'year':2026,'question':'Synthetic council expenditure question'});assert r.status_code==201,r.text;return r.json()
def evidence(c,r,text=b'Expenditure was 9999999999999999.99 JPY.'):
 d=c.post('/documents/upload',files={'file':('synthetic.txt',text,'text/plain')}).json()['document_id']
 p=c.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'document','source_id':d,'query_parameters':{},'excerpt':text.decode()});assert p.status_code==201,p.text
 return p.json(),d

def draft(c,r,e,text='Expenditure was 9999999999999999.99 JPY.',value='9999999999999999.99'):
 return c.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':text,'claims':[{'text':value,'value':value,'unit':'JPY','evidence_id':e['evidence_id']}]})
def action(c,r,name,status=200):
 p=c.post('/inquiries/'+r['inquiry_id']+'/'+name,json={'expected_version':r['version'],'reason':'Synthetic explicit Human source confirmation'});assert p.status_code==status,p.text;return p.json()
def test_exact_evidence_human_review_immutable_revision(client):
 r=create(client);e,d=evidence(client,r);r=client.get('/inquiries/'+r['inquiry_id']).json();p=draft(client,r,e);assert p.status_code==200,p.text;r=p.json()
 assert r['claims'][0]['value']=='9999999999999999.99'
 action(client,r,'approve',409);r=action(client,r,'review');r=action(client,r,'approve');assert r['status']=='approved' and r['reviewed_at'] and r['approved_at']
 assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'changed','claims':[]}).status_code==409
 n=client.post('/inquiries/'+r['inquiry_id']+'/revisions',json={'expected_version':r['version'],'reason':'Synthetic amendment'});assert n.status_code==201,n.text;assert n.json()['revision_of']==r['inquiry_id'] and n.json()['status']=='draft'
 assert client.get('/inquiries/'+r['inquiry_id']).json()['draft']==r['draft']
 with SessionLocal() as db:assert db.scalar(select(AuditLog).where(AuditLog.action=='inquiry.approve'))
def test_forgery_unsupported_number_stale_sha_and_review_invalidation(client):
 r=create(client);e,d=evidence(client,r,b'Synthetic count 12 people.');r=client.get('/inquiries/'+r['inquiry_id']).json()
 assert draft(client,r,e,'Count 13 people.','13').status_code==422
 p=draft(client,r,e,'Count 12 people.','12');assert p.status_code==422 # unit must match source, not invented JPY
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Count 12 people.','claims':[{'text':'12','value':'12','unit':'people','evidence_id':e['evidence_id']}]});assert p.status_code==200,p.text;r=action(client,p.json(),'review')
 assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':1,'draft':'stale','claims':[]}).status_code==409
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Human corrected prose: count 12 people.','claims':r['claims']});assert p.status_code==200,p.text;assert p.json()['status']=='draft' and p.json()['reviewed_by'] is None
 r=action(client,p.json(),'review')
 with SessionLocal() as db:db.get(Document,d).sha256='0'*64;db.commit()
 action(client,r,'approve',409)
def test_source_prose_privacy_list_search_similarity_export(client):
 from app.models import EmergencyCase
 for kind in ['personnel','emergency']:
  with SessionLocal() as db:
   source=Employee(display_name='RestrictedSynthetic',title='Chief 73') if kind=='personnel' else EmergencyCase(source_case_key='restricted',dispatch_number='73',incident_address='RestrictedSynthetic')
   db.add(source);db.commit();key=source.employee_id if kind=='personnel' else source.emergency_case_id
  r=create(client);p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':kind,'source_id':key,'excerpt':'73','query_parameters':{}});assert p.status_code==201,p.text
  r=client.get('/inquiries/'+r['inquiry_id']).json();p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'RestrictedSynthetic count 73','claims':[{'text':'73','value':'73','unit':'','evidence_id':p.json()['evidence_id']}]});assert p.status_code==200,p.text
  if kind=='emergency':r=action(client,p.json(),'review');action(client,r,'approve')
  remove('personnel.read' if kind=='personnel' else 'emergency.case.read')
  assert client.get('/inquiries/'+r['inquiry_id']).status_code==403
  assert client.get('/inquiries?q=RestrictedSynthetic').json()==[]
  assert client.get('/inquiries/similar?q=RestrictedSynthetic').json()['items']==[]
  assert 'RestrictedSynthetic' not in client.get('/inquiries/export?format=csv').text
  assert client.get('/search?q=RestrictedSynthetic&modules=inquiries').json()['hits']==[]
  assert client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':kind,'source_id':key,'excerpt':'73','query_parameters':{}}).status_code==403

def test_atomic_csv_xlsx_exchange_and_provenance(client):
 import csv,json
 from openpyxl import load_workbook
 r=create(client);e,d=evidence(client,r);r=client.get('/inquiries/'+r['inquiry_id']).json();assert draft(client,r,e).status_code==200
 for fmt in ['csv','xlsx']:
  exported=client.get('/inquiries/export?format='+fmt);assert exported.status_code==200
  before=len(client.get('/inquiries').json());p=client.post('/inquiries/import',files={'file':('synthetic.'+fmt,exported.content)});assert p.status_code==201,p.text
  assert len(client.get('/inquiries').json())==before+p.json()['inserted']
  assert p.json()['file_sha256'] and p.json()['source_document_id']
 raw=client.get('/inquiries/export').content.decode('utf-8-sig');rows=list(csv.DictReader(raw.splitlines()));rows[-1]['year']='not-year';out=__import__('io').StringIO();w=csv.DictWriter(out,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
 before=len(client.get('/inquiries').json());assert client.post('/inquiries/import',files={'file':('bad.csv',out.getvalue().encode())}).status_code==422;assert len(client.get('/inquiries').json())==before
 remove('inquiry.review');r=client.get('/inquiries').json()[0];action(client,r,'review',403)
def test_fresh_cli_bootstrap_registration(tmp_path):
 env={**os.environ,'PYTHONPATH':'backend','FIRE_AI_DATABASE_URL':'sqlite+pysqlite:///'+str(tmp_path/'fresh.db'),'FIRE_AI_PRODUCTION_MODE':'false'};env.pop('FIRE_AI_TENANT_ID',None)
 p=subprocess.run([sys.executable,'-m','app.bootstrap','--username','synthetic','--display-name','Synthetic','--password','synthetic-password'],env=env,capture_output=True,text=True);assert p.returncode==0,p.stderr
 import sqlite3
 with sqlite3.connect(tmp_path/'fresh.db') as db:
  for table in ('inquiries','violation_cases','violation_measures','corrective_actions','correction_events'):
   assert db.execute("select count(*) from sqlite_master where name=?",(table,)).fetchone()[0]==1,table

def test_candidate_never_official_local_adapter_outage_and_provenance(client,monkeypatch):
 from app import inquiries_service as svc
 r=create(client);e,d=evidence(client,r,b'Synthetic count 12 people.');r=client.get('/inquiries/'+r['inquiry_id']).json()
 p=client.post('/inquiries/'+r['inquiry_id']+'/ai-draft',json={'expected_version':r['version'],'reason':'Synthetic local model request'});assert p.status_code==201,p.text;c=p.json()
 assert c['model']=='deterministic-evidence-extract' and c['input_provenance']['sources'][e['evidence_id']]['documents'][0]['sha256']
 unchanged=client.get('/inquiries/'+r['inquiry_id']).json();assert unchanged['draft']=='' and unchanged['status']=='draft'
 p=client.post('/inquiries/'+r['inquiry_id']+'/adopt-candidate',json={'expected_version':r['version'],'candidate_id':c['candidate_id'],'reason':'Synthetic verified adoption'});assert p.status_code==200,p.text;r=p.json();assert r['status']=='draft' and r['provenance']['generated_at']
 monkeypatch.setattr(svc,'local_model_adapter',lambda inputs:{'model':'synthetic-local','model_version':'fixture-v1','draft':'Candidate count 13 people.','confidence':'0.8','claims':[{'text':'13','value':'13','unit':'people','evidence_id':e['evidence_id']}]})
 p=client.post('/inquiries/'+r['inquiry_id']+'/ai-draft',json={'expected_version':r['version'],'reason':'Synthetic model provenance'});assert p.status_code==201,p.text;c=p.json();assert c['confidence']=='0.8' and c['model_version']=='fixture-v1'
 assert client.post('/inquiries/'+r['inquiry_id']+'/adopt-candidate',json={'expected_version':r['version'],'candidate_id':c['candidate_id'],'reason':'Forged adoption'}).status_code==422
 def unavailable(inputs):raise RuntimeError('synthetic outage')
 monkeypatch.setattr(svc,'local_model_adapter',unavailable)
 assert client.post('/inquiries/'+r['inquiry_id']+'/ai-draft',json={'expected_version':r['version'],'reason':'outage'}).status_code==503
 assert client.get('/inquiries').status_code==200
 assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Manual count 12 people.','claims':r['claims']}).status_code==200

def test_removed_evidence_retains_source_prose_lineage(client):
 with SessionLocal() as db:
  employee=Employee(display_name='RestrictedProseSynthetic');db.add(employee);db.commit();key=employee.employee_id
 r=create(client);p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'personnel','source_id':key,'excerpt':'RestrictedProseSynthetic','query_parameters':{}});assert p.status_code==201,p.text;e=p.json();r=client.get('/inquiries/'+r['inquiry_id']).json()
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'RestrictedProseSynthetic holds office.','claims':[]});assert p.status_code==200,p.text;r=p.json()
 assert client.request('DELETE','/inquiries/'+r['inquiry_id']+'/evidence/'+e['evidence_id'],json={'expected_version':r['version'],'reason':'Synthetic source removed'}).status_code==200
 remove('personnel.read');assert client.get('/inquiries/'+r['inquiry_id']).status_code==403;assert client.get('/inquiries?q=RestrictedProseSynthetic').json()==[]

def test_declared_formula_exact_decimal_and_unsupported_punctuation(client):
 r=create(client);e,d=evidence(client,r,b'Exact decimal sources: 0.1 JPY and 0.2 JPY.');r=client.get('/inquiries/'+r['inquiry_id']).json()
 claim={'text':'0.3','value':'0.3','unit':'JPY','evidence_id':e['evidence_id'],'formula':'sum','operands':[{'evidence_id':e['evidence_id'],'value':v,'unit':'JPY'} for v in ['0.1','0.2']]}
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Declared sum 0.3 JPY.','claims':[claim]});assert p.status_code==200,p.text;r=p.json();assert action(client,r,'review')['status']=='reviewed'
 for text in ['Unsupported 13.','Unsupported １３.','Unsupported 1.3e1.','Unsupported 1,300 people.','Unsupported thirteen people.','Unsupported 十三人。']:
  assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version']+1,'draft':text,'claims':[]}).status_code==422
 bad={**claim,'value':'0.30000000000000001','text':'0.30000000000000001'}
 assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version']+1,'draft':'Declared sum 0.30000000000000001 JPY.','claims':[bad]}).status_code==422

def test_original_template_and_derived_original_guard(client):
 from openpyxl import Workbook,load_workbook
 r=create(client);e,d=evidence(client,r);r=client.get('/inquiries/'+r['inquiry_id']).json();r=draft(client,r,e).json();r=action(client,action(client,r,'review'),'approve')
 book=Workbook();book.active['A1']='Synthetic original official layout';raw=BytesIO();book.save(raw)
 template_doc=client.post('/documents/upload',files={'file':('official-synthetic.xlsx',raw.getvalue())}).json()
 from app.models import FormTemplate
 with SessionLocal() as db:
  t=FormTemplate(template_code='SYN-INQUIRY',name='Synthetic original',module_code='inquiries',document_id=template_doc['document_id'],version_label='1',modification_policy='fill_only',field_mapping={'answer':{'sheet':'Sheet','cell':'A2'}});db.add(t);db.commit();tid=t.form_template_id
 p=client.post('/inquiries/'+r['inquiry_id']+'/render',json={'form_template_id':tid});assert p.status_code==201,p.text;x=p.json();result=client.get('/documents/'+x['document_id']+'/download');assert result.status_code==200
 sheet=load_workbook(BytesIO(result.content)).active;assert sheet['A1'].value=='Synthetic original official layout' and sheet['A2'].value==r['draft'];assert x['manifest']['template_sha256']==template_doc['sha256'];assert client.get('/documents/'+template_doc['document_id']+'/download').content==raw.getvalue()
 remove('inquiry.read');assert client.get('/documents/'+x['document_id']).status_code==403;assert client.get('/documents/'+x['document_id']+'/download').status_code==403
 assert client.get('/search?q=inquiry&modules=documents').json()['hits']==[]

def test_restricted_import_original_guard_and_formula_neutralization(client):
 with SessionLocal() as db:
  employee=Employee(display_name='RestrictedExportSynthetic');db.add(employee);db.commit();key=employee.employee_id
 r=create(client);e=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'personnel','source_id':key,'excerpt':'RestrictedExportSynthetic'});assert e.status_code==201,e.text;r=client.get('/inquiries/'+r['inquiry_id']).json()
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'=RestrictedExportSynthetic','question':'+Synthetic formula question','claims':[]});assert p.status_code==200,p.text
 csvdata=client.get('/inquiries/export').content;assert b"'=RestrictedExportSynthetic" in csvdata and b"'+Synthetic formula question" in csvdata
 imported=client.post('/inquiries/import',files={'file':('restricted-synthetic.csv',csvdata)});assert imported.status_code==201,imported.text;did=imported.json()['source_document_id'];assert client.get('/documents/'+did+'/download').status_code==200
 remove('personnel.read');assert client.get('/documents/'+did+'/download').status_code==403;assert client.get('/documents/'+did).status_code==403
 assert client.get('/search?q=restricted-synthetic&modules=documents').json()['hits']==[]

def test_actual_postgresql_inquiry_migration_source_locks_and_concurrent_review(tmp_path,monkeypatch):
 url=os.environ.get('FIRE_AI_TEST_POSTGRES_URL')
 if not url:pytest.skip('actual PostgreSQL inquiry migration/source locks/concurrency execute in CI')
 from uuid import uuid4
 from hashlib import sha256
 from concurrent.futures import ThreadPoolExecutor
 from threading import Barrier
 from sqlalchemy import create_engine,event,text,update
 from sqlalchemy.orm import sessionmaker
 from sqlalchemy.exc import IntegrityError,DBAPIError
 from fastapi import HTTPException
 from app.rbac_seed import seed_rbac
 from app.inquiries_models import Inquiry,InquiryEvidence
 from app.inquiries_schemas import InquiryInput,EvidenceInput,InquiryPatch,Action
 from app import inquiries_service as svc
 from app.settings import settings
 from app.migrations import split_sql
 schema='synthetic_inquiry_'+uuid4().hex+'-q';eng=create_engine(url)
 @event.listens_for(eng,'connect')
 def scope(connection,record):
  with connection.cursor() as cursor:cursor.execute('SET search_path TO "'+schema+'",public')
 scoped=eng.execution_options(schema_translate_map={None:schema})
 monkeypatch.setattr(settings,'storage_root',str(tmp_path));proof=b'Exact source amount 9999999999999999.99 JPY.';(tmp_path/'proof.txt').write_bytes(proof)
 try:
  with eng.begin() as db:db.execute(text('CREATE SCHEMA "'+schema+'"'))
  Base.metadata.create_all(scoped,tables=[t for t in Base.metadata.sorted_tables if t.name!='inquiries' and not t.name.startswith('inquiry_')])
  migration=Path(__file__).resolve().parents[2]/'db/migrations/050_run_b_inquiries.sql'
  with eng.begin() as db:
   for statement in split_sql(migration.read_text()):db.exec_driver_sql(statement,execution_options={'no_parameters':True})
  factory=sessionmaker(bind=scoped,expire_on_commit=False,autoflush=False)
  with factory() as db:
   roles=seed_rbac(db);u=User(username='synthetic-inquiry-pg',password_hash=hash_password('synthetic-password'));db.add(u);db.flush();db.add(UserRole(user_id=u.user_id,role_id=roles['system_admin'].role_id));db.flush()
   d=Document(original_filename='Synthetic original.txt',storage_path='proof.txt',sha256=sha256(proof).hexdigest(),size_bytes=len(proof),mime_type='text/plain');db.add(d);db.flush()
   r=svc.create(db,u,InquiryInput(year=2026,question='Synthetic PG concurrency'));e=svc.add_evidence(db,u,r.inquiry_id,EvidenceInput(expected_version=r.version,source_type='document',source_id=d.document_id,excerpt=proof.decode()))
   r=svc.patch(db,u,r.inquiry_id,InquiryPatch(expected_version=r.version,draft=proof.decode(),claims=[{'text':'9999999999999999.99','value':'9999999999999999.99','unit':'JPY','evidence_id':e.evidence_id}]))
   uid=u.user_id;key=r.inquiry_id;version=r.version;did=d.document_id;eid=e.evidence_id;db.commit()
  barrier=Barrier(2)
  def race(index):
   with factory() as db:
    user=db.get(User,uid);barrier.wait(timeout=10)
    try:svc.action(db,user,key,Action(expected_version=version,reason='Synthetic concurrent review'),'review');db.commit();return 200
    except HTTPException as exc:db.rollback();return exc.status_code
  with ThreadPoolExecutor(max_workers=2) as pool:assert sorted(pool.map(race,[0,1]))==[200,409]
  # Review holds a native lock on source originals; mutation in another connection must time out.
  with factory() as first,factory() as second:
   u=first.get(User,uid);r=first.get(Inquiry,key);svc.action(first,u,key,Action(expected_version=r.version,reason='Synthetic held source review'),'review')
   second.execute(text("SET LOCAL lock_timeout='200ms'"))
   with pytest.raises(DBAPIError):second.execute(update(Document).where(Document.document_id==did).values(sha256='0'*64))
   second.rollback();first.commit()
  with factory() as db:
   r=svc.action(db,db.get(User,uid),key,Action(expected_version=db.get(Inquiry,key).version,reason='Synthetic explicit official confirmation'),'approve');db.commit();assert r.claims[0]['value']=='9999999999999999.99'
  # Trigger resolves the qualified source schema even when caller search_path omits it.
  for model,pk,identity,values in [(Inquiry,'inquiry_id',key,{'draft':'forged'}),(InquiryEvidence,'evidence_id',eid,{'excerpt':'forged'})]:
   with factory() as db:
    db.execute(text('SET LOCAL search_path TO pg_catalog'))
    with pytest.raises(IntegrityError):db.execute(update(model).where(getattr(model,pk)==identity).values(**values))
    db.rollback()
  with factory() as db:
   assert db.get(Inquiry,key).draft==proof.decode()
 finally:
  with eng.begin() as db:db.execute(text('DROP SCHEMA "'+schema+'" CASCADE'))
  eng.dispose()

def test_other_common_document_consumers_honor_inquiry_source_guard(client):
 # Preserve unchanged standalone originals while closing two real alternate extraction routes.
 with SessionLocal() as db:
  role=db.scalar(select(Role))
  for code in ['finance.read','intake.read','intake.analyze']:
   p=Permission(code=code);db.add(p);db.flush();db.add(RolePermission(role_id=role.role_id,permission_id=p.permission_id))
  employee=Employee(display_name='RestrictedExtractorSynthetic');db.add(employee);db.commit();key=employee.employee_id
 standalone=client.post('/documents/upload',files={'file':('standalone.txt',b'Synthetic unrelated original','text/plain')}).json()['document_id']
 r=create(client);p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'personnel','source_id':key,'excerpt':'RestrictedExtractorSynthetic'});assert p.status_code==201,p.text;r=client.get('/inquiries/'+r['inquiry_id']).json();assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'RestrictedExtractorSynthetic','claims':[]}).status_code==200
 exported=client.get('/inquiries/export').content;p=client.post('/inquiries/import',files={'file':('RestrictedExtractorSynthetic.csv',exported)});assert p.status_code==201,p.text;did=p.json()['source_document_id']
 analysis=client.post('/document-analyses',json={'document_id':did});assert analysis.status_code==201,analysis.text;aid=analysis.json()['document_analysis_id']
 remove('personnel.read')
 assert client.get('/finance/documents/'+did+'/extract').status_code==403
 assert all(v['document_id']!=did for v in client.get('/finance/documents').json())
 assert client.get('/document-analyses/'+aid).status_code==403
 assert all(v['document_analysis_id']!=aid for v in client.get('/document-analyses').json())
 assert client.post('/document-analyses',json={'document_id':did}).status_code==403
 assert client.get('/finance/documents/'+standalone+'/extract').status_code==200
 assert client.get('/documents/'+standalone+'/download').content==b'Synthetic unrelated original'
 assert client.post('/document-analyses',json={'document_id':standalone}).status_code==201

def test_forged_quotes_source_versions_missing_sources_and_all_action_gates(client):
 with SessionLocal() as db:
  employee=Employee(display_name='SyntheticSource',title='Chief');db.add(employee);db.commit();key=employee.employee_id
 r=create(client);p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'personnel','source_id':key,'excerpt':'Client invented prose'});assert p.status_code==422
 p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'personnel','source_id':key,'excerpt':'Chief'});assert p.status_code==201,p.text;r=client.get('/inquiries/'+r['inquiry_id']).json();r=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'SyntheticSource is Chief.','claims':[]}).json();r=action(client,r,'review')
 with SessionLocal() as db:source=db.get(Employee,key);source.version+=1;db.commit()
 action(client,r,'approve',409)
 with SessionLocal() as db:db.delete(db.get(Employee,key));db.commit()
 action(client,r,'approve',409)
 remove('inquiry.update');assert client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Unauthorized','claims':[]}).status_code==403
 remove('inquiry.approve');action(client,r,'approve',403)
 remove('inquiry.create');assert client.post('/inquiries',json={'year':2026,'question':'Denied'}).status_code==403
 remove('inquiry.admin');assert client.request('DELETE','/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'reason':'Denied'}).status_code==403
 remove('inquiry.import');assert client.post('/inquiries/import',files={'file':('denied.csv',b'denied')}).status_code==403
 remove('inquiry.export');assert client.get('/inquiries/export').status_code==403

def test_new_canonical_workforce_source_and_document_permissions(client):
 from datetime import date,datetime,time,timezone
 from app.personnel import OrganizationUnit
 from app.workforce_models import WorkforceShiftType,WorkforceRosterEntry
 with SessionLocal() as db:
  role=db.scalar(select(Role));p=Permission(code='workforce.read');db.add(p);db.flush();db.add(RolePermission(role_id=role.role_id,permission_id=p.permission_id))
  employee=Employee(display_name='Synthetic workforce source');org=OrganizationUnit(code='SYN-INQUIRY',name='Synthetic organization');db.add_all([employee,org]);db.flush()
  shift=WorkforceShiftType(code='SYN-INQUIRY',name='Synthetic shift',start_time=time(8),end_time=time(17),payable_minutes=480);db.add(shift);db.flush()
  actor=db.scalar(select(User));roster=WorkforceRosterEntry(employee_id=employee.employee_id,organization_id=org.organization_id,shift_type_id=shift.shift_type_id,work_date=date(2026,10,6),starts_at=datetime(2026,10,6,8,tzinfo=timezone.utc),ends_at=datetime(2026,10,6,17,tzinfo=timezone.utc),payable_minutes=480,note='RestrictedWorkforceSynthetic',created_by=actor.user_id);db.add(roster);db.commit();key=roster.roster_entry_id
 r=create(client);p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'workforce','source_id':key,'excerpt':'RestrictedWorkforceSynthetic'});assert p.status_code==201,p.text;e=p.json();r=client.get('/inquiries/'+r['inquiry_id']).json()
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'RestrictedWorkforceSynthetic worked 480 minutes.','claims':[{'text':'480','value':'480','unit':'minutes','evidence_id':e['evidence_id']}]});assert p.status_code==200,p.text;action(client,p.json(),'review')
 remove('workforce.read');assert client.get('/inquiries/'+r['inquiry_id']).status_code==403;assert client.get('/inquiries?q=RestrictedWorkforceSynthetic').json()==[]
 assert client.get('/inquiries/sources?source_type=workforce').status_code==403

def test_model_work_releases_write_locks_then_revalidates_source_and_session(client,monkeypatch):
 from app import inquiries_service as svc,authz
 from app.models import UserSession,now_utc
 r=create(client);e,d=evidence(client,r,b'Synthetic count 12 people.');r=client.get('/inquiries/'+r['inquiry_id']).json()
 locked=[];original=authz.account_change_lock
 def gate(db):locked.append(True);return original(db)
 monkeypatch.setattr(authz,'account_change_lock',gate)
 def changed_source(inputs):
  assert not locked,'external model must run outside canonical write gate'
  with SessionLocal() as db:db.get(Document,d).sha256='0'*64;db.commit()
  return {'model':'synthetic-local','model_version':'v1','draft':'Synthetic count 12 people.','claims':[],'confidence':None}
 monkeypatch.setattr(svc,'local_model_adapter',changed_source)
 p=client.post('/inquiries/'+r['inquiry_id']+'/ai-draft',json={'expected_version':r['version'],'reason':'Synthetic source race'});assert p.status_code==409,p.text;assert client.get('/inquiries/'+r['inquiry_id']+'/candidates').json()==[];assert locked
 # Restore the source and prove a logout during external model work rejects the candidate write.
 from hashlib import sha256
 with SessionLocal() as db:db.get(Document,d).sha256=sha256(b'Synthetic count 12 people.').hexdigest();db.commit()
 locked.clear()
 def revoked_session(inputs):
  assert not locked
  with SessionLocal() as db:
   for session in db.scalars(select(UserSession)):session.revoked_at=now_utc()
   db.commit()
  return {'model':'synthetic-local','model_version':'v1','draft':'Synthetic count 12 people.','claims':[],'confidence':None}
 monkeypatch.setattr(svc,'local_model_adapter',revoked_session)
 p=client.post('/inquiries/'+r['inquiry_id']+'/ai-draft',json={'expected_version':r['version'],'reason':'Synthetic logout race'});assert p.status_code==401,p.text
 with SessionLocal() as db:
  from app.inquiries_models import InquiryCandidate
  assert not list(db.scalars(select(InquiryCandidate)))

def test_large_import_decode_precedes_canonical_write_gate(client,monkeypatch):
 from app import inquiries_service as svc,authz
 r=create(client);raw=client.get('/inquiries/export').content;locked=[];original=authz.account_change_lock;decoder=svc.read_tabular
 def gate(db):locked.append(True);return original(db)
 def decode(data,filename):assert not locked;return decoder(data,filename)
 monkeypatch.setattr(authz,'account_change_lock',gate);monkeypatch.setattr(svc,'read_tabular',decode)
 p=client.post('/inquiries/import',files={'file':('synthetic.csv',raw)});assert p.status_code==201,p.text;assert locked and p.json()['inserted']==1

@pytest.mark.parametrize('literal',['1e61','1e'+'2'*62,'1'*64,'1e'+'2'*61+'f'])
def test_source_number_limits_reject_before_unbounded_decimal_formatting(client,literal):
 doc=client.post('/documents/upload',files={'file':('synthetic-exponent.txt',('Synthetic count '+literal+' people.').encode(),'text/plain')}).json()['document_id']
 p=client.get('/inquiries/source/document/'+doc);assert p.status_code==422,p.text

def test_nonfinite_and_deep_json_import_rejected_atomically(client):
 import csv,json
 from io import StringIO
 create(client);raw=client.get('/inquiries/export').content.decode('utf-8-sig');base=list(csv.DictReader(raw.splitlines()))[0]
 for history in ['{"provenance":'+'{"nested":'*1100+'{}'+'}'*1100+'}','{"provenance":{"invalid":NaN}}','{"provenance":{"invalid":Infinity}}',json.dumps({'provenance':{'nested':__import__('functools').reduce(lambda value,_:{'next':value},range(70),{})}})]:
  row={**base,'history':history};out=StringIO();w=csv.DictWriter(out,fieldnames=list(row));w.writeheader();w.writerow(row);before=len(client.get('/inquiries').json());p=client.post('/inquiries/import',files={'file':('synthetic-bounded.csv',out.getvalue().encode())});assert p.status_code==422,p.text;assert len(client.get('/inquiries').json())==before

def test_official_spreadsheet_output_neutralizes_formula_text(client):
 from openpyxl import Workbook,load_workbook
 from app.models import FormTemplate
 r=create(client);e,d=evidence(client,r,b'Synthetic literal answer source.');r=client.get('/inquiries/'+r['inquiry_id']).json();p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'=SyntheticLiteralAnswer','claims':[]});assert p.status_code==200,p.text;r=action(client,action(client,p.json(),'review'),'approve')
 book=Workbook();book.active['A1']='Synthetic original layout';book.active['B1']='=1+1';raw=BytesIO();book.save(raw);doc=client.post('/documents/upload',files={'file':('synthetic-literal.xlsx',raw.getvalue())}).json()
 with SessionLocal() as db:
  t=FormTemplate(template_code='SYN-LITERAL',name='Synthetic literal output',module_code='inquiries',document_id=doc['document_id'],version_label='1',modification_policy='fill_only',field_mapping={'answer':{'sheet':'Sheet','cell':'A2'}});db.add(t);db.commit();tid=t.form_template_id
 p=client.post('/inquiries/'+r['inquiry_id']+'/render',json={'form_template_id':tid});assert p.status_code==201,p.text;sheet=load_workbook(BytesIO(client.get('/documents/'+p.json()['document_id']+'/download').content),data_only=False).active
 assert sheet['A2'].data_type!='f' and sheet['A2'].value==r['draft']
 assert sheet['B1'].data_type=='f' and sheet['B1'].value=='=1+1'
 assert client.get('/documents/'+doc['document_id']+'/download').content==raw.getvalue()

def test_audit_only_role_cannot_read_inquiry_source_prose_numbers_or_candidate_inputs(client,monkeypatch):
 from app.rbac_seed import seed_rbac
 from app import inquiries_service as svc
 marker='PrivateSyntheticCouncilMarker';number='783491823.73'
 r=client.post('/inquiries',json={'year':2026,'question':marker}).json();e,d=evidence(client,r,(marker+' '+number+' JPY.').encode());r=client.get('/inquiries/'+r['inquiry_id']).json()
 p=draft(client,r,e,marker+' '+number+' JPY.',number);assert p.status_code==200,p.text;r=p.json()
 monkeypatch.setattr(svc,'local_model_adapter',lambda inputs:{'model':marker,'model_version':marker,'draft':marker+' '+number+' JPY.','claims':r['claims'],'confidence':'0.7'})
 p=client.post('/inquiries/'+r['inquiry_id']+'/ai-draft',json={'expected_version':r['version'],'reason':marker});assert p.status_code==201,p.text
 p=client.post('/inquiries/'+r['inquiry_id']+'/adopt-candidate',json={'expected_version':r['version'],'candidate_id':p.json()['candidate_id'],'reason':marker});assert p.status_code==200,p.text
 r=action(client,p.json(),'review');action(client,r,'approve')
 with SessionLocal() as db:
  roles=seed_rbac(db);u=User(username='audit-only',password_hash=hash_password('synthetic-password'));db.add(u);db.flush();db.add(UserRole(user_id=u.user_id,role_id=roles['auditor'].role_id));db.commit()
 client.post('/auth/logout');assert client.post('/auth/login',json={'username':'audit-only','password':'synthetic-password'}).status_code==200
 assert client.get('/inquiries/'+r['inquiry_id']).status_code==403
 p=client.get('/administration/audit?limit=500');assert p.status_code==200,p.text
 rows=[row for row in p.json() if row['action'].startswith('inquiry.')];assert rows
 raw=p.text;assert marker not in raw and number not in raw
 assert {'inquiry.create','inquiry.update','inquiry.evidence.add','inquiry.candidate.generate','inquiry.candidate.adopt','inquiry.review','inquiry.approve'}<={row['action'] for row in rows}
 for row in rows:
  assert row['entity_id'] and row['after_data']['change_sha256']
  if row['before_data']:assert row['before_data']['change_sha256']

@pytest.mark.parametrize('answer,claims',[
 ('Amount .5 JPY.',[]),
 ('Amount −5 JPY.',[{'text':'5','value':'5','unit':'JPY'}]),
 ('Ratio 1/2.',[{'text':'1','value':'1','unit':''},{'text':'2','value':'2','unit':''}]),
])
def test_review_round1_numeric_forms_cannot_bypass_patch_review_or_approval(client,answer,claims):
 from app.inquiries_models import Inquiry
 r=create(client);e,d=evidence(client,r,b'Synthetic values: 5 JPY; 1; 2.');r=client.get('/inquiries/'+r['inquiry_id']).json();claims=[{**c,'evidence_id':e['evidence_id']} for c in claims]
 p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':answer,'claims':claims});assert p.status_code==422,p.text
 # Exercise the actual Human gates against historical/imported unsafe content too.
 with SessionLocal() as db:
  row=db.get(Inquiry,r['inquiry_id']);row.draft=answer;row.claims=claims;db.commit()
 action(client,r,'review',422);action(client,r,'approve',422)

def test_review_round1_source_and_answer_numeric_literals_exact_or_fail_closed(client):
 r=create(client);e,d=evidence(client,r,b'Synthetic value .5 JPY and negative \xe2\x88\x925 JPY.');values=e['snapshot']['values'];assert any(v['value']=='0.5' for v in values) and any(v['value']=='-5' for v in values)
 r=client.get('/inquiries/'+r['inquiry_id']).json();p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Exact .5 JPY and −5 JPY.','claims':[{'text':'.5','value':'0.5','unit':'JPY','evidence_id':e['evidence_id']},{'text':'−5','value':'-5','unit':'JPY','evidence_id':e['evidence_id']}]});assert p.status_code==200,p.text;action(client,action(client,p.json(),'review'),'approve')
 for text in ['Ratio 1/2.','Value 1,5 JPY.','Value 1.2.3 JPY.','Value 1٫5 JPY.','Value –5 JPY.']:
  r=create(client);d=client.post('/documents/upload',files={'file':('unsupported.txt',text.encode(),'text/plain')}).json()['document_id'];p=client.post('/inquiries/'+r['inquiry_id']+'/evidence',json={'expected_version':r['version'],'source_type':'document','source_id':d,'excerpt':text});assert p.status_code==422,(text,p.text)

def _grant(codes):
 with SessionLocal() as db:
  role=db.scalar(select(Role).where(Role.code=='inquiry-test'))
  for code in codes:
   p=db.scalar(select(Permission).where(Permission.code==code))
   if not p:p=Permission(code=code);db.add(p);db.flush()
   if not db.scalar(select(RolePermission).where(RolePermission.role_id==role.role_id,RolePermission.permission_id==p.permission_id)):db.add(RolePermission(role_id=role.role_id,permission_id=p.permission_id))
  db.commit()

def test_review_round1_rendered_template_snapshot_and_current_rights_all_consumers(client):
 from app.models import Facility,FormTemplate
 from openpyxl import Workbook
 _grant(['facility.read','finance.read','intake.read','intake.analyze'])
 r=create(client);e,d=evidence(client,r,b'Synthetic supported prose.');r=client.get('/inquiries/'+r['inquiry_id']).json();p=client.patch('/inquiries/'+r['inquiry_id'],json={'expected_version':r['version'],'draft':'Synthetic supported prose.','claims':[]});r=action(client,action(client,p.json(),'review'),'approve')
 outputs=[];originals=[];building_ids=[]
 for at_render in [True,False]:
  book=Workbook();book.active['A1']='RestrictedUnmappedTemplateMarker';raw=BytesIO();book.save(raw);doc=client.post('/documents/upload',files={'file':('synthetic-template.xlsx',raw.getvalue())}).json();originals.append(doc['document_id'])
  with SessionLocal() as db:
   f=Facility(name='Synthetic restricted template');db.add(f);db.flush();building_ids.append(f.building_id)
   if at_render:db.get(Document,doc['document_id']).building_id=f.building_id
   t=FormTemplate(template_code='SYN-RIGHTS-'+str(at_render),name='Synthetic original',module_code='inquiries',document_id=doc['document_id'],version_label='1',modification_policy='fill_only',field_mapping={'answer':{'sheet':'Sheet','cell':'A2'}});db.add(t);db.commit();tid=t.form_template_id
  p=client.post('/inquiries/'+r['inquiry_id']+'/render',json={'form_template_id':tid});assert p.status_code==201,p.text;outputs.append(p.json()['document_id'])
  analysis=client.post('/document-analyses',json={'document_id':outputs[-1]});assert analysis.status_code==201,analysis.text
  outputs[-1]=(outputs[-1],analysis.json()['document_analysis_id'])
 # First restriction is removed after rendering: snapshot must still require it.
 # Second restriction is added after rendering: current source must require it.
 with SessionLocal() as db:
  db.get(Document,originals[0]).building_id=None;db.get(Document,originals[1]).building_id=building_ids[1];db.commit()
 remove('facility.read')
 for did,aid in outputs:
  assert client.get('/documents/'+did).status_code==403
  assert client.get('/documents/'+did+'/download').status_code==403
  assert client.get('/finance/documents/'+did+'/extract').status_code==403
  assert all(x['document_id']!=did for x in client.get('/finance/documents').json())
  assert client.post('/document-analyses',json={'document_id':did}).status_code==403
  assert client.get('/document-analyses/'+aid).status_code==403
  assert all(x['document_analysis_id']!=aid for x in client.get('/document-analyses').json())
  assert client.get('/search?q='+did+'&modules=documents').json()['hits']==[]
 assert client.get('/inquiries/'+r['inquiry_id']).status_code==200

def test_review_round1_import_original_self_and_multirecord_cycles_keep_rights_and_crud(client):
 # Empty drafts make their import original a source without unrelated numeric assertions.
 rows=[create(client),create(client)];evidence(client,rows[0],b'Synthetic cycle source 0.');raw=client.get('/inquiries/export').content
 p=client.post('/inquiries/import',files={'file':('Synthetic shared.csv',raw)});assert p.status_code==201,p.text;keys=p.json()['inquiry_ids'];did=p.json()['source_document_id']
 for key in keys:
  r=client.get('/inquiries/'+key).json();p=client.post('/inquiries/'+key+'/evidence',json={'expected_version':r['version'],'source_type':'document','source_id':did,'excerpt':'Synthetic council expenditure question'});assert p.status_code==201,p.text
  r=client.get('/inquiries/'+key);assert r.status_code==200,r.text
 assert client.get('/inquiries').status_code==200 and client.get('/search?q=Synthetic&modules=inquiries').status_code==200
 # A private source added to one member must propagate through the shared original
 # into the other member, even though traversal revisits the original/owners.
 with SessionLocal() as db:
  employee=Employee(display_name='PrivateCycleMarker');db.add(employee);db.commit();eid=employee.employee_id
 r=client.get('/inquiries/'+keys[0]).json();p=client.post('/inquiries/'+keys[0]+'/evidence',json={'expected_version':r['version'],'source_type':'personnel','source_id':eid,'excerpt':'PrivateCycleMarker'});assert p.status_code==201,p.text
 r=client.get('/inquiries/'+keys[1]).json();own_evidence=next(e['evidence_id'] for e in r['evidence'] if e['source_id']==did);assert client.request('DELETE','/inquiries/'+keys[1]+'/evidence/'+own_evidence,json={'expected_version':r['version'],'reason':'Synthetic remove own original'}).status_code==200
 remove('personnel.read')
 for key in keys:assert client.get('/inquiries/'+key).status_code==403
 assert client.get('/documents/'+did+'/download').status_code==403
 _grant(['personnel.read'])
 r=client.get('/inquiries/'+keys[0]).json();own=next(e for e in r['evidence'] if e['source_id']==did);assert client.request('DELETE','/inquiries/'+keys[0]+'/evidence/'+own['evidence_id'],json={'expected_version':r['version'],'reason':'Replace source after propagated permission change'}).status_code==200
 r=client.get('/inquiries/'+keys[0]).json();assert client.patch('/inquiries/'+keys[0],json={'expected_version':r['version'],'draft':'Synthetic safe edit','claims':[]}).status_code==200
